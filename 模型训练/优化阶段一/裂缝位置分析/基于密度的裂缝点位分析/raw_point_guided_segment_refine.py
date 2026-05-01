from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import joblib
import numpy as np
import xml.etree.ElementTree as ET
import pandas as pd
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from sklearn.cluster import KMeans
from sklearn.ensemble import HistGradientBoostingRegressor
from sklearn.linear_model import LogisticRegression, PoissonRegressor, TweedieRegressor
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

import build_segment_count_dataset as dataset_builder

try:
    from xgboost import XGBRegressor
except ImportError:
    XGBRegressor = None


ROOT = Path(__file__).resolve().parent
DEFAULT_EXIST_EXP_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/裂缝存在性预测/LSTM/单井验证/成像测井裂缝预测/cnn+lstm/exp40_missing_drop_ac_gr_3x3_seq5_AC_GR"
)
DEFAULT_SAMPLE_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/成像测井/裂缝样本"
)
DEFAULT_RAW_LABEL_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/研究内容一/成像测井/裂缝标注"
)
DEFAULT_EXTRA_RAW_LABEL_DIRS = [
    Path(
        r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/成像测井/裂缝提取"
    )
]
BASE_SAVE_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/裂缝位置预测/基于密度的裂缝点位分析/裂缝点位精细化"
)
FLOW_RESULT_ROOT = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/两阶段裂缝预测流程结果"
)
BASE_SAVE_DIR = FLOW_RESULT_ROOT / "manual_runs" / "raw_point_guided_segment_refine"
DEFAULT_DOCX_PATH = Path(r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/实验记录/实验记录20260323.docx")
DOC_TITLE = "实验记录 20260323"

WELL_TO_RAW_FILE = {
    "车151HF": "车151HF_fractures.csv",
    "车660-1": "车660_1.xlsx",
    "车660-2": "车660_2.xlsx",
    "车662": "车662.xlsx",
    "车663": "车663.xlsx",
    "车页1HF": "车页1导眼_fractures.csv",
    "车页1导眼": "车页1导眼_fractures.csv",
}
WELL_TO_SAMPLE_FILE = {
    "车151HF": "车151HF_sample.csv",
    "车660-1": "车660-1_sample.csv",
    "车660-2": "车660-2_sample.csv",
    "车662": "车662_sample.csv",
    "车663": "车663_sample.csv",
    "车页1HF": "车页1导眼_sample.csv",
    "车页1导眼": "车页1导眼_sample.csv",
}
ANALYSIS_ONLY_FEATURES = {
    "GTPointCountInPredSegment",
    "RawGTCountInPredSegment",
    "HasRawPoint",
    "HasOrientationPoint",
    "GTDevOverlapLen",
    "GTDevOverlapRatio",
    "GTDevOverlapCount",
    "GTDevOverlapFlag",
    "TargetWeight",
    "TargetLabelType",
    "RawP10MeanInPredSegment",
    "RawP10MaxInPredSegment",
    "RawP10MassInPredSegment",
    "RawP10MassPerLength",
    "RawP10PositiveFracInPredSegment",
    "RawOrientationPointCountInPredSegment",
    "RawAzimuthMeanInPredSegment",
    "RawDipMeanInPredSegment",
    "RawOrientationAzimuthStdInPredSegment",
    "RawOrientationDipStdInPredSegment",
}
DEFAULT_RESULT_HEADERS = [
    "Well",
    "GT points",
    "Pred points",
    "Count diff",
    "Count err %",
    "pred->gt mean",
    "pred->gt p90",
    "gt->pred mean",
    "gt->pred p90",
    "SegMAE",
    "CoveredRaw",
    "MissedRaw",
]
DEFAULT_RESULT_KEYS = [
    "well",
    "N_gt_points",
    "N_pred_points",
    "count_diff",
    "count_error_pct",
    "pred_to_gt_mean_dist",
    "pred_to_gt_p90_dist",
    "gt_to_pred_mean_dist",
    "gt_to_pred_p90_dist",
    "segment_count_mae",
    "raw_points_covered_by_pred_segments",
    "raw_points_missed_outside_pred_segments",
]
MODEL_ARTIFACT_FILENAME = "segment_refine_model.joblib"
MODEL_META_FILENAME = "segment_refine_model_meta.json"
BOUNDARY_EXPAND_EFFECTIVE_MASK_COL = "__PRED_LABEL_BOUNDARY_EXPANDED"
COUNT_BIN_LABEL_ORDER = ["zero", "one", "two", "three_four", "ge_five"]
COUNT_BIN_DEFAULT_VALUE_MAP = {
    "zero": 0.0,
    "one": 1.0,
    "two": 2.0,
    "three_four": 3.5,
    "ge_five": 6.0,
}
RAW_REFINE_CONFIG_PROFILES = {
    "normal_strata_v1": {
        "fallback_count_mode": "global_length_prob_mass_scale",
        "count_train_mode": "all_segments",
        "count_fusion_mode": "blend",
        "fusion_learned_weight": 0.7,
        "global_post_scale_mode": "train_sum_match",
        "selective_post_mode": "train_zero_dual_quantile_downscale",
        "selective_post_feature": "ProbMassPerLength",
        "selective_post_feature_2": "ProbMean",
        "selective_post_quantile": 0.6,
        "selective_post_quantile_2": 0.5,
        "selective_post_max_pred_float": 4.0,
        "boundary_expand_mode": "prob_shoulder",
        "boundary_expand_prob_min": 0.35,
        "boundary_expand_max_steps": 2,
        "gate_mode": "logistic_soft",
        "gate_prob_threshold": 0.5,
        "model_name": "tweedie",
        "type_calibration_mode": "rule_v1",
        "count_bin_mode": "logistic_expected_blend",
        "count_bin_blend_weight": 0.65,
        "density_strength_mode": "segment_p10_mass_per_length",
        "density_strength_basis_col": "ProbMassPerLength",
    },
    "sparse_strata_v1": {
        "fallback_count_mode": "global_length_prob_mass_scale",
        "count_train_mode": "positive_only",
        "count_fusion_mode": "uplift",
        "fusion_learned_weight": 0.8,
        "global_post_scale_mode": "train_sum_match_up_only",
        "selective_post_mode": "train_zero_dual_quantile_downscale",
        "selective_post_feature": "ProbMassPerLength",
        "selective_post_feature_2": "ProbMean",
        "selective_post_quantile": 0.65,
        "selective_post_quantile_2": 0.55,
        "selective_post_max_pred_float": 3.0,
        "boundary_expand_mode": "prob_shoulder",
        "boundary_expand_prob_min": 0.4,
        "boundary_expand_max_steps": 2,
        "gate_mode": "logistic_hard",
        "gate_prob_threshold": 0.55,
        "model_name": "tweedie",
        "type_calibration_mode": "rule_v1",
        "count_bin_mode": "logistic_expected_blend",
        "count_bin_blend_weight": 0.55,
        "density_strength_mode": "segment_p10_mass_per_length",
        "density_strength_basis_col": "ProbMassPerLength",
    },
    "cross_strata_safe_v1": {
        "fallback_count_mode": "global_length_prob_mass_scale",
        "count_train_mode": "positive_only",
        "global_post_scale_mode": "none",
        "selective_post_mode": "train_zero_dual_quantile_downscale",
        "selective_post_feature": "ProbMassPerLength",
        "selective_post_feature_2": "ProbMean",
        "selective_post_quantile": 0.65,
        "selective_post_quantile_2": 0.55,
        "selective_post_max_pred_float": 2.5,
        "boundary_expand_mode": "prob_shoulder",
        "boundary_expand_prob_min": 0.4,
        "boundary_expand_max_steps": 1,
        "gate_mode": "logistic_hard",
        "gate_prob_threshold": 0.55,
        "model_name": "rule_only",
        "type_calibration_mode": "rule_v1",
        "count_bin_mode": "logistic_expected_blend",
        "count_bin_blend_weight": 0.45,
        "density_strength_mode": "segment_p10_mass_per_length",
        "density_strength_basis_col": "ProbMassPerLength",
    },
    "sand4_recall_v1": {
        "fallback_count_mode": "global_length_prob_mass_scale",
        "count_train_mode": "positive_only",
        "count_fusion_mode": "uplift",
        "fusion_learned_weight": 0.8,
        "global_post_scale_mode": "train_sum_match_up_only",
        "selective_post_mode": "train_zero_dual_quantile_downscale",
        "selective_post_feature": "ProbMassPerLength",
        "selective_post_feature_2": "ProbMean",
        "selective_post_quantile": 0.65,
        "selective_post_quantile_2": 0.55,
        "selective_post_max_pred_float": 6.0,
        "boundary_expand_mode": "prob_shoulder",
        "boundary_expand_prob_min": 0.35,
        "boundary_expand_max_steps": 2,
        "gate_mode": "none",
        "count_bin_mode": "none",
        "density_strength_mode": "segment_p10_mass_per_length",
        "density_strength_basis_col": "ProbMassPerLength",
        "type_calibration_mode": "rule_v1",
        "model_name": "tweedie",
    },
    "sand4_balanced_v2": {
        "fallback_count_mode": "global_length_prob_mass_scale",
        "count_train_mode": "positive_only",
        "count_fusion_mode": "blend",
        "fusion_learned_weight": 0.55,
        "global_post_scale_mode": "none",
        "selective_post_mode": "train_zero_dual_quantile_downscale",
        "selective_post_feature": "ProbMassPerLength",
        "selective_post_feature_2": "ProbMean",
        "selective_post_quantile": 0.60,
        "selective_post_quantile_2": 0.50,
        "selective_post_max_pred_float": 2.0,
        "boundary_expand_mode": "prob_shoulder",
        "boundary_expand_prob_min": 0.45,
        "boundary_expand_max_steps": 1,
        "gate_mode": "logistic_hard",
        "gate_prob_threshold": 0.55,
        "count_bin_mode": "logistic_expected_blend",
        "count_bin_blend_weight": 0.45,
        "density_strength_mode": "segment_p10_mass_per_length",
        "density_strength_basis_col": "ProbMassPerLength",
        "type_calibration_mode": "rule_v1",
        "model_name": "tweedie",
    },
    "sand3_probmass_rule_v1": {
        "fallback_count_mode": "global_prob_mass_scale",
        "count_train_mode": "all_segments",
        "count_fusion_mode": "learned_only",
        "fusion_learned_weight": 0.7,
        "global_post_scale_mode": "train_sum_match",
        "selective_post_mode": "train_zero_dual_quantile_downscale",
        "selective_post_feature": "ProbMassPerLength",
        "selective_post_feature_2": "ProbMean",
        "selective_post_quantile": 0.6,
        "selective_post_quantile_2": 0.5,
        "selective_post_max_pred_float": 2.0,
        "boundary_expand_mode": "prob_shoulder",
        "boundary_expand_prob_min": 0.35,
        "boundary_expand_max_steps": 2,
        "gate_mode": "logistic_hard",
        "gate_prob_threshold": 0.6,
        "gate_target_wells": "车662",
        "count_bin_mode": "none",
        "density_strength_mode": "segment_p10_mass_per_length",
        "density_strength_basis_col": "ProbMassPerLength",
        "type_calibration_mode": "rule_v1",
        "model_name": "rule_only",
    },
}


def sanitize(text: str) -> str:
    return "".join(ch if ch.isalnum() or ch in {"_", "-"} else "_" for ch in text).strip("_")


def parse_list_arg(raw: str) -> list[str]:
    return [item.strip() for item in str(raw).split(",") if item.strip()]


def parse_float_list_arg(raw: str) -> list[float]:
    values: list[float] = []
    for item in parse_list_arg(raw):
        try:
            values.append(float(item))
        except ValueError as exc:
            raise ValueError(f"Expected float list, got invalid value: {item}") from exc
    return values


def apply_named_profile(args, profile_name: str, default_args=None) -> str:
    normalized = str(profile_name or "").strip().lower()
    if not normalized:
        return ""
    if normalized not in RAW_REFINE_CONFIG_PROFILES:
        raise ValueError(
            f"Unsupported config_profile: {profile_name}, "
            f"available={sorted(RAW_REFINE_CONFIG_PROFILES.keys())}"
        )
    default_values = vars(default_args) if default_args is not None else {}
    for key, value in RAW_REFINE_CONFIG_PROFILES[normalized].items():
        if key in default_values and getattr(args, key, None) != default_values.get(key):
            continue
        setattr(args, key, value)
    return normalized


def format_verify_well_dir_name(well_name: str) -> str:
    well_name = str(well_name).strip()
    return well_name if well_name.startswith("verify_") else f"verify_{well_name}"


def resolve_raw_label_path(raw_label_dir: Path, raw_file_name: str) -> Path:
    raw_path = Path(raw_file_name)
    candidate_paths: list[Path] = []
    if raw_path.is_absolute():
        candidate_paths.append(raw_path)

    search_dirs = [Path(raw_label_dir)]
    for extra_dir in DEFAULT_EXTRA_RAW_LABEL_DIRS:
        extra_dir = Path(extra_dir)
        if extra_dir not in search_dirs:
            search_dirs.append(extra_dir)

    for search_dir in search_dirs:
        candidate_paths.append(search_dir / raw_file_name)

    for candidate_path in candidate_paths:
        if candidate_path.exists():
            return candidate_path

    searched_paths = ", ".join(str(path) for path in candidate_paths)
    raise FileNotFoundError(f"Raw label file not found: {raw_file_name}; searched: {searched_paths}")


def load_strata_range_df(strata_range_csv: str) -> pd.DataFrame:
    csv_path = str(strata_range_csv or "").strip()
    if not csv_path:
        return pd.DataFrame()
    path = Path(csv_path)
    if not path.exists():
        raise FileNotFoundError(f"strata_range_csv not found: {path}")
    df = pd.read_csv(path, encoding="utf-8-sig").copy()
    required_cols = {"WellName", "StrataName", "DepthMin", "DepthMax"}
    missing_cols = sorted(required_cols - set(df.columns))
    if missing_cols:
        raise ValueError(f"strata_range_csv missing required cols: {missing_cols}")
    df["WellName"] = df["WellName"].astype(str).str.strip()
    df["StrataName"] = df["StrataName"].astype(str).str.strip()
    df["DepthMin"] = dataset_builder.POINT.safe_to_numeric(df["DepthMin"])
    df["DepthMax"] = dataset_builder.POINT.safe_to_numeric(df["DepthMax"])
    if "StrataTop" in df.columns:
        df["StrataTop"] = dataset_builder.POINT.safe_to_numeric(df["StrataTop"])
    if "StrataBase" in df.columns:
        df["StrataBase"] = dataset_builder.POINT.safe_to_numeric(df["StrataBase"])
    df = df.dropna(subset=["DepthMin", "DepthMax"]).copy()
    if df.empty:
        raise ValueError(f"No valid rows found in strata_range_csv: {path}")
    return df


def resolve_well_strata_range(
    well_name: str,
    strata_name: str,
    strata_range_df: pd.DataFrame | None,
) -> dict | None:
    normalized_strata = str(strata_name or "").strip()
    if not normalized_strata:
        return None
    if strata_range_df is None or strata_range_df.empty:
        raise ValueError("strata_name is provided but strata_range_csv is missing or empty")

    subset = strata_range_df[
        (strata_range_df["WellName"].astype(str).str.strip() == str(well_name).strip())
        & (strata_range_df["StrataName"].astype(str).str.strip() == normalized_strata)
    ].copy()
    if subset.empty:
        raise KeyError(f"Missing strata range for well={well_name}, strata={normalized_strata}")
    row = subset.iloc[0]
    depth_min = float(min(float(row["DepthMin"]), float(row["DepthMax"])))
    depth_max = float(max(float(row["DepthMin"]), float(row["DepthMax"])))
    return {
        "WellName": str(well_name).strip(),
        "StrataName": normalized_strata,
        "DepthMin": depth_min,
        "DepthMax": depth_max,
        "StrataTop": normalize_optional_positive_float(row.get("StrataTop", np.nan))
        if "StrataTop" in subset.columns
        else np.nan,
        "StrataBase": normalize_optional_positive_float(row.get("StrataBase", np.nan))
        if "StrataBase" in subset.columns
        else np.nan,
    }


def filter_df_to_depth_range(
    df: pd.DataFrame,
    depth_col: str,
    depth_min: float,
    depth_max: float,
) -> pd.DataFrame:
    out = df.copy()
    if depth_col not in out.columns:
        raise ValueError(f"Missing depth col {depth_col}")
    out[depth_col] = dataset_builder.POINT.safe_to_numeric(out[depth_col])
    mask = out[depth_col].ge(float(depth_min) - 1e-8) & out[depth_col].le(float(depth_max) + 1e-8)
    return out.loc[mask].copy()


def validate_train_distance_weight_mode(mode: str) -> str:
    normalized = str(mode or "none").strip().lower()
    valid_modes = {"none", "inverse"}
    if normalized not in valid_modes:
        raise ValueError(f"Unsupported train_distance_weight_mode: {mode}, available={sorted(valid_modes)}")
    return normalized


def infer_train_distance_csv_from_exist_exp_dir(exist_exp_dir: Path) -> Path | None:
    exist_exp_dir = Path(exist_exp_dir)
    if not exist_exp_dir.exists():
        return None
    for verify_dir in sorted(exist_exp_dir.glob("verify_*")):
        config_path = verify_dir / "config.json"
        if not config_path.exists():
            continue
        try:
            with config_path.open("r", encoding="utf-8-sig") as file_obj:
                config = json.load(file_obj)
        except Exception:
            continue
        distance_path = str(config.get("DIST_MATRIX_PATH") or "").strip()
        if distance_path:
            return Path(distance_path)
    return None


def load_train_distance_matrix(distance_csv: Path) -> pd.DataFrame:
    distance_csv = Path(distance_csv)
    if not distance_csv.exists():
        raise FileNotFoundError(f"Train distance CSV not found: {distance_csv}")
    distance_df = pd.read_csv(distance_csv, encoding="utf-8-sig", index_col=0)
    if distance_df.empty:
        raise ValueError(f"Train distance CSV is empty: {distance_csv}")
    distance_df.index = [str(idx).strip() for idx in distance_df.index]
    distance_df.columns = [str(col).strip() for col in distance_df.columns]
    distance_df = distance_df.apply(pd.to_numeric, errors="coerce")
    return distance_df


def resolve_train_distance_resources(args) -> tuple[Path | None, pd.DataFrame | None]:
    mode = validate_train_distance_weight_mode(str(args.train_distance_weight_mode))
    if mode == "none":
        return None, None
    train_distance_csv = str(args.train_distance_csv or "").strip()
    resolved_path = Path(train_distance_csv) if train_distance_csv else infer_train_distance_csv_from_exist_exp_dir(Path(args.exist_exp_dir))
    if resolved_path is None:
        raise ValueError(
            "train_distance_weight_mode is enabled but no distance CSV was provided or inferred from exist_exp_dir"
        )
    return resolved_path, load_train_distance_matrix(resolved_path)


def build_train_distance_weight_info(
    target_well: str,
    train_wells: list[str],
    train_distance_weight_mode: str,
    distance_df: pd.DataFrame | None,
    weight_power: float,
    weight_min: float,
    weight_max: float,
) -> dict:
    well_distances = {str(well): np.nan for well in train_wells}
    well_weights = {str(well): 1.0 for well in train_wells}
    info = {
        "mode": train_distance_weight_mode,
        "target_well": str(target_well),
        "well_distances": well_distances,
        "well_weights": well_weights,
        "weight_power": float(weight_power),
        "weight_min": float(weight_min),
        "weight_max": float(weight_max),
        "raw_similarity_mean": np.nan,
    }
    if train_distance_weight_mode == "none" or not train_wells:
        return info
    if distance_df is None:
        raise ValueError("distance_df is required when train_distance_weight_mode is enabled")
    if target_well not in distance_df.index:
        raise ValueError(f"Target well {target_well} not found in train distance matrix")
    missing_train_wells = [well for well in train_wells if well not in distance_df.columns]
    if missing_train_wells:
        raise ValueError(f"Train wells missing in train distance matrix: {missing_train_wells}")

    distances = distance_df.loc[target_well, train_wells]
    distances = pd.to_numeric(distances, errors="coerce")
    if distances.isna().any():
        missing_values = distances[distances.isna()].index.tolist()
        raise ValueError(f"Invalid train distance values for target well {target_well}: {missing_values}")

    clipped_power = max(float(weight_power), 1e-6)
    raw_similarity = 1.0 / np.power(np.maximum(distances.to_numpy(dtype=np.float64), 1e-6), clipped_power)
    similarity_mean = float(np.mean(raw_similarity)) if raw_similarity.size else 1.0
    normalized_weights = raw_similarity / max(similarity_mean, 1e-6)
    clipped_weights = np.clip(normalized_weights, float(weight_min), float(weight_max))

    info["well_distances"] = {str(well): float(distances.loc[well]) for well in train_wells}
    info["well_weights"] = {str(well): float(clipped_weights[idx]) for idx, well in enumerate(train_wells)}
    info["raw_similarity_mean"] = similarity_mean
    return info


def apply_train_distance_weight_info(train_df: pd.DataFrame, distance_weight_info: dict) -> pd.DataFrame:
    out = train_df.copy()
    well_distances = dict(distance_weight_info.get("well_distances", {}))
    well_weights = dict(distance_weight_info.get("well_weights", {}))
    out["TrainDistanceTargetWell"] = str(distance_weight_info.get("target_well", ""))
    out["TrainDistanceToTargetWell"] = out["WellName"].map(well_distances)
    out["TrainWellDistanceWeight"] = out["WellName"].map(well_weights).fillna(1.0).astype(np.float64)
    out["TrainSampleWeightUsed"] = (
        out["TargetWeight"].fillna(1.0).astype(np.float64) * out["TrainWellDistanceWeight"].astype(np.float64)
    )
    return out


def to_jsonable(value):
    if isinstance(value, dict):
        return {str(key): to_jsonable(val) for key, val in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [to_jsonable(item) for item in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.integer, np.floating, np.bool_)):
        return value.item()
    return value


def validate_optional_count_mode(mode: str) -> str:
    normalized = str(mode or "").strip()
    if not normalized:
        return ""
    if normalized not in dataset_builder.POINT.COUNT_MODE_TO_BASIS_COL:
        raise ValueError(
            f"Unsupported fallback_count_mode: {normalized}, "
            f"available={sorted(dataset_builder.POINT.COUNT_MODE_TO_BASIS_COL.keys())}"
        )
    return normalized


def validate_count_train_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "all_segments"
    if normalized not in {"all_segments", "positive_only"}:
        raise ValueError(
            f"Unsupported count_train_mode: {normalized}, "
            "available=['all_segments', 'positive_only']"
        )
    return normalized


def validate_type_calibration_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "rule_v1"}:
        raise ValueError(
            f"Unsupported type_calibration_mode: {normalized}, "
            "available=['none', 'rule_v1']"
        )
    return normalized


def resolve_positive_floor_allowed_labels(raw: str) -> list[str]:
    valid_labels = {"low_conf", "short_strong", "long_weak", "other"}
    labels = parse_list_arg(raw)
    if not labels:
        labels = ["other", "short_strong"]
    normalized = []
    for label in labels:
        normalized_label = str(label).strip().lower()
        if normalized_label not in valid_labels:
            raise ValueError(
                f"Unsupported positive floor label: {label}, available={sorted(valid_labels)}"
            )
        if normalized_label not in normalized:
            normalized.append(normalized_label)
    return normalized


def validate_global_post_scale_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "train_sum_match", "train_sum_match_up_only"}:
        raise ValueError(
            f"Unsupported global_post_scale_mode: {normalized}, "
            "available=['none', 'train_sum_match', 'train_sum_match_up_only']"
        )
    return normalized


def validate_selective_post_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "train_zero_quantile_downscale", "train_zero_dual_quantile_downscale"}:
        raise ValueError(
            f"Unsupported selective_post_mode: {normalized}, "
            "available=['none', 'train_zero_quantile_downscale', 'train_zero_dual_quantile_downscale']"
        )
    return normalized


def validate_boundary_expand_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "prob_shoulder"}:
        raise ValueError(
            f"Unsupported boundary_expand_mode: {normalized}, "
            "available=['none', 'prob_shoulder']"
        )
    return normalized


def validate_gate_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "logistic_hard", "logistic_soft"}:
        raise ValueError(
            f"Unsupported gate_mode: {normalized}, "
            "available=['none', 'logistic_hard', 'logistic_soft']"
        )
    return normalized


def validate_count_bin_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "logistic_expected_blend"}:
        raise ValueError(
            f"Unsupported count_bin_mode: {normalized}, "
            "available=['none', 'logistic_expected_blend']"
        )
    return normalized


def validate_density_strength_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {
        "none",
        "segment_p10_mean",
        "segment_p10_mass_per_length",
        "segment_p10_mean_regression",
        "segment_p10_mass_per_length_regression",
    }:
        raise ValueError(
            f"Unsupported density_strength_mode: {normalized}, "
            "available=['none', 'segment_p10_mean', 'segment_p10_mass_per_length', "
            "'segment_p10_mean_regression', 'segment_p10_mass_per_length_regression']"
        )
    return normalized


def density_strength_uses_regression(mode: str) -> bool:
    normalized = validate_density_strength_mode(mode)
    return normalized.endswith("_regression")


def validate_density_strength_regression_train_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "all_segments"
    if normalized not in {"all_segments", "positive_only"}:
        raise ValueError(
            f"Unsupported density_strength_regression_train_mode: {normalized}, "
            "available=['all_segments', 'positive_only']"
        )
    return normalized


def validate_density_strength_regression_target_transform(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "log1p"}:
        raise ValueError(
            f"Unsupported density_strength_regression_target_transform: {normalized}, "
            "available=['none', 'log1p']"
        )
    return normalized


def validate_density_strength_regression_calibration_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "quantile_map"}:
        raise ValueError(
            f"Unsupported density_strength_regression_calibration_mode: {normalized}, "
            "available=['none', 'quantile_map']"
        )
    return normalized


def normalize_density_strength_regression_calibration_quantiles(raw: str) -> list[float]:
    quantiles = parse_float_list_arg(raw)
    if not quantiles:
        quantiles = [0.10, 0.25, 0.50, 0.75, 0.90, 0.95]
    normalized: list[float] = []
    for value in sorted(set(float(item) for item in quantiles)):
        if not (0.0 < value < 1.0):
            raise ValueError(
                f"density strength regression calibration quantiles must be inside (0, 1), got {value}"
            )
        normalized.append(value)
    return normalized


def validate_orientation_mode(mode: str) -> str:
    normalized = str(mode or "").strip().lower()
    if not normalized:
        return "none"
    if normalized not in {"none", "family_classifier"}:
        raise ValueError(
            f"Unsupported orientation_mode: {normalized}, "
            "available=['none', 'family_classifier']"
        )
    return normalized


def resolve_feature_cols(feature_preset: str, feature_cols_raw: str) -> list[str]:
    if feature_preset:
        if feature_preset not in dataset_builder.FEATURE_PRESETS:
            raise ValueError(
                f"Unsupported feature_preset: {feature_preset}, "
                f"available={sorted(dataset_builder.FEATURE_PRESETS.keys())}"
            )
        return list(dataset_builder.FEATURE_PRESETS[feature_preset])
    return parse_list_arg(feature_cols_raw)


def validate_feature_cols(feature_cols: list[str], dataset_df: pd.DataFrame) -> None:
    missing_features = [feature for feature in feature_cols if feature not in dataset_df.columns]
    if missing_features:
        raise ValueError(f"Missing feature cols: {missing_features}")
    leakage_features = [feature for feature in feature_cols if feature in ANALYSIS_ONLY_FEATURES]
    if leakage_features:
        raise ValueError(f"Feature cols contain analysis-only GT fields: {leakage_features}")


def dedupe_preserve_order(values: list[str] | tuple[str, ...]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        name = str(value or "").strip()
        if not name or name in seen:
            continue
        seen.add(name)
        out.append(name)
    return out


def resolve_strategy_feature_cols(
    dataset_df: pd.DataFrame,
    base_feature_cols: list[str],
    extra_feature_cols: list[str] | tuple[str, ...],
    required_feature_cols: list[str] | tuple[str, ...] = (),
) -> list[str]:
    candidate_cols = dedupe_preserve_order(
        [*required_feature_cols, *base_feature_cols, *extra_feature_cols]
    )
    resolved = []
    for feature_name in candidate_cols:
        if feature_name in ANALYSIS_ONLY_FEATURES:
            continue
        if feature_name in dataset_df.columns:
            resolved.append(feature_name)
    return resolved


def build_continuous_regressor(
    learning_rate: float,
    max_depth: int,
    max_iter: int,
) -> Pipeline:
    _ = (learning_rate, max_depth)
    model = TweedieRegressor(
        power=1.5,
        alpha=0.03,
        max_iter=int(np.clip(int(max_iter), 200, 5000)),
        link="log",
    )
    return Pipeline([("scaler", StandardScaler()), ("model", model)])


def apply_density_strength_target_transform(values: np.ndarray, transform_mode: str) -> np.ndarray:
    normalized_mode = validate_density_strength_regression_target_transform(transform_mode)
    arr = np.clip(np.asarray(values, dtype=np.float64), a_min=0.0, a_max=None)
    if normalized_mode == "log1p":
        return np.log1p(arr)
    return arr


def invert_density_strength_target_transform(values: np.ndarray, transform_mode: str) -> np.ndarray:
    normalized_mode = validate_density_strength_regression_target_transform(transform_mode)
    arr = np.asarray(values, dtype=np.float64)
    if normalized_mode == "log1p":
        return np.expm1(arr)
    return arr


def build_density_strength_quantile_map(
    pred_raw: np.ndarray,
    target_raw: np.ndarray,
    quantiles: list[float],
) -> dict:
    pred = np.clip(np.asarray(pred_raw, dtype=np.float64), a_min=0.0, a_max=None)
    target = np.clip(np.asarray(target_raw, dtype=np.float64), a_min=0.0, a_max=None)
    mask = np.isfinite(pred) & np.isfinite(target) & (target > 0.0)
    if int(np.sum(mask)) < 4:
        return {"mode": "none", "quantiles": [], "pred_knots": [], "target_knots": []}

    pred_pos = pred[mask]
    target_pos = target[mask]
    pred_knots = [0.0]
    target_knots = [0.0]
    for quantile in quantiles:
        pred_knots.append(float(np.quantile(pred_pos, quantile)))
        target_knots.append(float(np.quantile(target_pos, quantile)))

    pred_arr = np.maximum.accumulate(np.asarray(pred_knots, dtype=np.float64))
    target_arr = np.maximum.accumulate(np.asarray(target_knots, dtype=np.float64))
    dedup_pred: list[float] = []
    dedup_target: list[float] = []
    for pred_value, target_value in zip(pred_arr.tolist(), target_arr.tolist()):
        if dedup_pred and abs(pred_value - dedup_pred[-1]) <= 1e-10:
            dedup_target[-1] = max(dedup_target[-1], float(target_value))
            continue
        dedup_pred.append(float(pred_value))
        dedup_target.append(float(target_value))

    if len(dedup_pred) < 2 or dedup_pred[-1] <= 1e-10:
        return {"mode": "none", "quantiles": [], "pred_knots": [], "target_knots": []}
    return {
        "mode": "quantile_map",
        "quantiles": [float(value) for value in quantiles],
        "pred_knots": dedup_pred,
        "target_knots": dedup_target,
    }


def apply_density_strength_quantile_map(pred_raw: np.ndarray, calibration_info: dict) -> np.ndarray:
    pred = np.clip(np.asarray(pred_raw, dtype=np.float64), a_min=0.0, a_max=None)
    if validate_density_strength_regression_calibration_mode(str(calibration_info.get("mode", "none"))) == "none":
        return pred

    pred_knots = np.asarray(calibration_info.get("pred_knots") or [], dtype=np.float64)
    target_knots = np.asarray(calibration_info.get("target_knots") or [], dtype=np.float64)
    if pred_knots.size < 2 or pred_knots.size != target_knots.size:
        return pred
    clipped_pred = np.clip(pred, a_min=0.0, a_max=float(pred_knots[-1]))
    mapped = np.interp(clipped_pred, pred_knots, target_knots)
    return np.clip(mapped, a_min=0.0, a_max=float(target_knots[-1]))


def open_or_create_doc(docx_path: Path) -> Document:
    if docx_path.exists():
        try:
            return Document(str(docx_path))
        except (PackageNotFoundError, BadZipFile):
            pass
    doc = Document()
    doc.add_heading(DOC_TITLE, level=1)
    return doc


def format_float(value: object) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return str(value)
    if not np.isfinite(numeric):
        return "nan"
    return f"{numeric:.4f}"


def mean_of_numeric(rows: list[dict], key: str) -> float | None:
    values: list[float] = []
    for row in rows:
        raw_value = row.get(key, "")
        if raw_value in {"", None}:
            continue
        try:
            numeric_value = float(raw_value)
        except (TypeError, ValueError):
            continue
        if not np.isfinite(numeric_value):
            continue
        values.append(numeric_value)
    if not values:
        return None
    return float(sum(values) / len(values))


def normalize_optional_positive_float(value: object) -> float:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return np.nan
    if not np.isfinite(numeric) or numeric <= 0.0:
        return np.nan
    return float(numeric)


def compute_count_error_pct(count_diff: object, gt_count: object) -> float:
    try:
        gt_numeric = float(gt_count)
        diff_numeric = abs(float(count_diff))
    except (TypeError, ValueError):
        return np.nan
    if not np.isfinite(gt_numeric) or gt_numeric <= 0.0 or not np.isfinite(diff_numeric):
        return np.nan
    return float((diff_numeric / gt_numeric) * 100.0)


def compute_overall_count_error_pct(rows: list[dict]) -> float | None:
    total_gt = 0.0
    total_abs_diff = 0.0
    for row in rows:
        try:
            gt_numeric = float(row.get("N_gt_points", np.nan))
            diff_numeric = abs(float(row.get("count_diff", np.nan)))
        except (TypeError, ValueError):
            continue
        if not np.isfinite(gt_numeric) or gt_numeric <= 0.0 or not np.isfinite(diff_numeric):
            continue
        total_gt += gt_numeric
        total_abs_diff += diff_numeric
    if total_gt <= 0.0:
        return None
    return float((total_abs_diff / total_gt) * 100.0)


def append_result_to_docx(docx_path: Path, title: str, config: dict, results: list[dict]) -> None:
    doc = open_or_create_doc(docx_path)
    doc.add_heading(title, level=2)
    doc.add_paragraph(f"Record time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    config_lines = [
        "task: raw-point-guided segment refine",
        f"script: {ROOT / 'raw_point_guided_segment_refine.py'}",
        f"config_profile: {config.get('config_profile', '')}",
        f"exist_exp_dir: {config['exist_exp_dir']}",
        f"sample_dir: {config['sample_dir']}",
        f"raw_label_dir: {config['raw_label_dir']}",
        f"strata_name: {config['strata_name']}",
        f"strata_range_csv: {config['strata_range_csv']}",
        f"gt_dev_rule: {config['gt_dev_rule']}",
        f"model_name: {config['model_name']}",
        f"model_alpha: {config['model_alpha']}",
        f"model_max_iter: {config['model_max_iter']}",
        f"tweedie_power: {config['tweedie_power']}",
        f"hgb_learning_rate: {config['hgb_learning_rate']}",
        f"hgb_max_depth: {config['hgb_max_depth']}",
        f"feature_preset: {config['feature_preset']}",
        f"feature_cols: {config['feature_cols']}",
        f"prob_weight_gamma: {config['prob_weight_gamma']}",
        f"fallback_count_mode: {config['fallback_count_mode']}",
        f"positive_floor_target_wells: {config['positive_floor_target_wells']}",
        f"count_fusion_mode: {config['count_fusion_mode']}",
        f"fusion_learned_weight: {config['fusion_learned_weight']}",
        f"global_post_scale_mode: {config['global_post_scale_mode']}",
        f"global_post_scale_min: {config['global_post_scale_min']}",
        f"global_post_scale_max: {config['global_post_scale_max']}",
        f"selective_post_mode: {config['selective_post_mode']}",
        f"selective_post_target_wells: {config['selective_post_target_wells']}",
        f"selective_post_feature: {config['selective_post_feature']}",
        f"selective_post_feature_2: {config['selective_post_feature_2']}",
        f"selective_post_quantile: {config['selective_post_quantile']}",
        f"selective_post_quantile_2: {config['selective_post_quantile_2']}",
        f"selective_post_max_pred_float: {config['selective_post_max_pred_float']}",
        f"boundary_expand_mode: {config['boundary_expand_mode']}",
        f"boundary_expand_prob_min: {config['boundary_expand_prob_min']}",
        f"boundary_expand_max_steps: {config['boundary_expand_max_steps']}",
        f"boundary_expand_max_depth: {config['boundary_expand_max_depth']}",
        f"count_bin_mode: {config.get('count_bin_mode', 'none')}",
        f"count_bin_blend_weight: {config.get('count_bin_blend_weight', np.nan)}",
        f"density_strength_mode: {config.get('density_strength_mode', 'none')}",
        f"density_strength_basis_col: {config.get('density_strength_basis_col', '')}",
        f"soft_negative_weight: {config['soft_negative_weight']}",
        f"pred_min_points_per_segment: {config['pred_min_points_per_segment']}",
        f"rounding_mode: {config['rounding_mode']}",
        f"save_dir: {config['save_dir']}",
        f"summary_csv: {config['summary_csv']}",
    ]
    if config["model_name"] == "xgboost":
        config_lines.extend(
            [
                f"xgb_n_estimators: {config['xgb_n_estimators']}",
                f"xgb_max_depth: {config['xgb_max_depth']}",
                f"xgb_learning_rate: {config['xgb_learning_rate']}",
                f"xgb_subsample: {config['xgb_subsample']}",
                f"xgb_colsample_bytree: {config['xgb_colsample_bytree']}",
            ]
        )
    doc.add_paragraph("\n".join(config_lines))

    table = doc.add_table(rows=1, cols=len(DEFAULT_RESULT_HEADERS))
    for idx, header in enumerate(DEFAULT_RESULT_HEADERS):
        table.rows[0].cells[idx].text = header
    for row in results:
        cells = table.add_row().cells
        for idx, key in enumerate(DEFAULT_RESULT_KEYS):
            cells[idx].text = str(row[key]) if key == "well" else format_float(row[key])

    summary_parts = []
    for key, label in (
        ("count_diff", "AVG_count_diff"),
        ("count_error_pct", "AVG_count_error_pct"),
        ("pred_to_gt_mean_dist", "AVG_pred_to_gt_mean"),
        ("gt_to_pred_mean_dist", "AVG_gt_to_pred_mean"),
        ("segment_count_mae", "AVG_segment_count_mae"),
    ):
        avg_value = mean_of_numeric(results, key)
        if avg_value is not None:
            summary_parts.append(f"{label}={avg_value:.4f}")
    overall_count_error_pct = compute_overall_count_error_pct(results)
    if overall_count_error_pct is not None:
        summary_parts.append(f"TOTAL_count_error_pct={overall_count_error_pct:.4f}")
    doc.add_paragraph("Summary: " + ", ".join(summary_parts) if summary_parts else "Summary: recorded")
    doc.add_paragraph("")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def interval_overlap_length(start_a: float, end_a: float, start_b: float, end_b: float) -> float:
    left = max(float(start_a), float(start_b))
    right = min(float(end_a), float(end_b))
    return float(max(right - left, 0.0))


def compute_overlap_with_gt_dev(start_depth: float, end_depth: float, gt_dev_segments: pd.DataFrame) -> tuple[float, int]:
    if gt_dev_segments.empty:
        return 0.0, 0
    overlap_len = 0.0
    overlap_count = 0
    for row in gt_dev_segments.itertuples(index=False):
        current_len = interval_overlap_length(start_depth, end_depth, row.SegStartDepth, row.SegEndDepth)
        if current_len > 0:
            overlap_len += current_len
            overlap_count += 1
    return float(overlap_len), int(overlap_count)


def count_points_in_any_segments(point_depths: np.ndarray, segment_df: pd.DataFrame) -> int:
    if point_depths.size == 0 or segment_df.empty:
        return 0
    count = 0
    for depth_value in point_depths:
        in_any = (
            (segment_df["SegStartDepth"].to_numpy(dtype=np.float64) <= depth_value)
            & (segment_df["SegEndDepth"].to_numpy(dtype=np.float64) >= depth_value)
        )
        if bool(np.any(in_any)):
            count += 1
    return int(count)


def apply_boundary_expand_to_pred_mask(
    df: pd.DataFrame,
    config,
    mode: str,
    prob_min: float,
    max_steps: int,
    max_depth: float,
) -> tuple[pd.DataFrame, str, dict]:
    if config.depth_col not in df.columns:
        raise ValueError(f"Missing depth col {config.depth_col}")
    if config.pred_label_col not in df.columns:
        raise ValueError(f"Missing pred label col {config.pred_label_col}")

    normalized_mode = validate_boundary_expand_mode(mode)
    out = df.copy()
    out[config.depth_col] = dataset_builder.POINT.safe_to_numeric(out[config.depth_col])
    out = out.sort_values(config.depth_col).reset_index(drop=True)

    base_mask = (
        dataset_builder.POINT.safe_to_numeric(out[config.pred_label_col])
        .fillna(0.0)
        .astype(int)
        .to_numpy(dtype=np.int8)
        == 1
    )
    effective_mask = base_mask.copy()
    base_segments = dataset_builder.POINT.find_binary_segments(base_mask)

    clipped_prob_min = float(np.clip(prob_min, 0.0, 1.0))
    clipped_max_steps = max(int(max_steps), 0)
    clipped_max_depth = float(max_depth)
    limit_depth = np.isfinite(clipped_max_depth) and clipped_max_depth > 0.0
    added_left = 0
    added_right = 0

    if normalized_mode == "prob_shoulder" and base_segments:
        if config.pred_prob_col not in out.columns:
            raise ValueError(f"Missing pred prob col {config.pred_prob_col}")
        depth = out[config.depth_col].to_numpy(dtype=np.float64)
        prob = (
            dataset_builder.POINT.safe_to_numeric(out[config.pred_prob_col])
            .fillna(0.0)
            .clip(lower=0.0, upper=1.0)
            .to_numpy(dtype=np.float64)
        )

        for start_idx, end_idx in base_segments:
            start_depth = float(depth[start_idx])
            end_depth = float(depth[end_idx])

            current_idx = start_idx - 1
            steps_used = 0
            while current_idx >= 0 and steps_used < clipped_max_steps:
                if effective_mask[current_idx]:
                    current_idx -= 1
                    continue
                current_prob = float(prob[current_idx])
                current_depth = float(depth[current_idx])
                if not np.isfinite(current_prob) or current_prob < clipped_prob_min:
                    break
                if limit_depth and np.isfinite(current_depth) and (start_depth - current_depth) > clipped_max_depth:
                    break
                effective_mask[current_idx] = True
                added_left += 1
                steps_used += 1
                current_idx -= 1

            current_idx = end_idx + 1
            steps_used = 0
            while current_idx < len(effective_mask) and steps_used < clipped_max_steps:
                if effective_mask[current_idx]:
                    current_idx += 1
                    continue
                current_prob = float(prob[current_idx])
                current_depth = float(depth[current_idx])
                if not np.isfinite(current_prob) or current_prob < clipped_prob_min:
                    break
                if limit_depth and np.isfinite(current_depth) and (current_depth - end_depth) > clipped_max_depth:
                    break
                effective_mask[current_idx] = True
                added_right += 1
                steps_used += 1
                current_idx += 1

    out[BOUNDARY_EXPAND_EFFECTIVE_MASK_COL] = effective_mask.astype(np.int8)
    effective_segments = dataset_builder.POINT.find_binary_segments(effective_mask)
    stats = {
        "boundary_expand_mode": normalized_mode,
        "boundary_expand_prob_min": clipped_prob_min,
        "boundary_expand_max_steps": clipped_max_steps,
        "boundary_expand_max_depth": clipped_max_depth,
        "boundary_base_positive_samples": int(np.sum(base_mask.astype(np.int64))),
        "boundary_effective_positive_samples": int(np.sum(effective_mask.astype(np.int64))),
        "boundary_added_positive_samples": int(np.sum(effective_mask.astype(np.int64)) - np.sum(base_mask.astype(np.int64))),
        "boundary_base_segment_count": int(len(base_segments)),
        "boundary_effective_segment_count": int(len(effective_segments)),
        "boundary_added_left_samples": int(added_left),
        "boundary_added_right_samples": int(added_right),
    }
    return out, BOUNDARY_EXPAND_EFFECTIVE_MASK_COL, stats




def excel_ref_to_col_idx(cell_ref: str) -> int:
    col = 0
    for ch in str(cell_ref):
        if not ch.isalpha():
            break
        col = (col * 26) + (ord(ch.upper()) - ord("A") + 1)
    return max(col - 1, 0)


def load_shared_strings(xlsx_file: ZipFile) -> list[str]:
    try:
        xml_bytes = xlsx_file.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    root = ET.fromstring(xml_bytes)
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    values = []
    for item in root.findall(f"{ns}si"):
        text_parts = []
        for text_node in item.iter(f"{ns}t"):
            text_parts.append(text_node.text or "")
        values.append("".join(text_parts))
    return values


def read_xlsx_cell_value(cell, shared_strings: list[str]) -> str:
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    cell_type = cell.attrib.get("t", "")
    if cell_type == "inlineStr":
        text_node = cell.find(f"{ns}is/{ns}t")
        return "" if text_node is None else str(text_node.text or "")
    value_node = cell.find(f"{ns}v")
    if value_node is None:
        return ""
    raw_text = "" if value_node.text is None else str(value_node.text)
    if cell_type == "s":
        try:
            idx = int(float(raw_text))
        except ValueError:
            return raw_text
        return shared_strings[idx] if 0 <= idx < len(shared_strings) else ""
    return raw_text


def read_simple_xlsx(path: Path, names: list[str]) -> pd.DataFrame:
    with ZipFile(path) as xlsx_file:
        shared_strings = load_shared_strings(xlsx_file)
        sheet_xml = xlsx_file.read("xl/worksheets/sheet1.xml")
    root = ET.fromstring(sheet_xml)
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rows = []
    for row_node in root.iter(f"{ns}row"):
        current_row = {}
        for cell in row_node.findall(f"{ns}c"):
            col_idx = excel_ref_to_col_idx(cell.attrib.get("r", ""))
            current_row[col_idx] = read_xlsx_cell_value(cell, shared_strings)
        if current_row:
            width = max(current_row.keys()) + 1
            rows.append([current_row.get(idx, "") for idx in range(width)])
    if not rows:
        return pd.DataFrame(columns=names)
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    df = pd.DataFrame(normalized)
    if df.shape[1] < len(names):
        for _ in range(len(names) - df.shape[1]):
            df[df.shape[1]] = ""
    df = df.iloc[:, : len(names)].copy()
    df.columns = names
    if not df.empty:
        first_value = pd.to_numeric(df[names[0]], errors="coerce").iloc[0]
        if not np.isfinite(first_value):
            df = df.iloc[1:].reset_index(drop=True)
    return df


def load_raw_fracture_point_table(raw_path: Path) -> pd.DataFrame:
    suffix = raw_path.suffix.lower()
    if suffix == ".csv":
        raw_df = pd.read_csv(raw_path, encoding="utf-8-sig").copy()
    elif suffix == ".xlsx":
        raw_df = read_simple_xlsx(raw_path, ["MD", "Angle(0~90)", "Azimuth(0~360)"]).copy()
    else:
        raise ValueError(f"Unsupported raw fracture file type: {raw_path}")

    raw_df.columns = [str(col).strip() for col in raw_df.columns]
    required_cols = ["MD", "Angle(0~90)", "Azimuth(0~360)"]
    missing_cols = [col for col in required_cols if col not in raw_df.columns]
    if missing_cols:
        raise ValueError(f"Raw fracture file missing required cols {missing_cols}: {raw_path}")
    return raw_df[required_cols].copy()

def load_raw_fracture_points(
    raw_label_dir: Path,
    well_name: str,
    depth_col: str = "TVD",
    strata_name: str = "",
    strata_range_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    raw_file_name = WELL_TO_RAW_FILE.get(well_name)
    if not raw_file_name:
        raise KeyError(f"No raw label file mapping for well: {well_name}")
    raw_path = resolve_raw_label_path(Path(raw_label_dir), raw_file_name)

    raw_df = load_raw_fracture_point_table(raw_path)
    raw_df = raw_df.rename(
        columns={
            "Angle(0~90)": "Frac_Dip",
            "Azimuth(0~360)": "Frac_Azimuth",
        }
    )
    raw_df["MD"] = dataset_builder.POINT.safe_to_numeric(raw_df["MD"])
    raw_df["Frac_Dip"] = dataset_builder.POINT.safe_to_numeric(raw_df["Frac_Dip"])
    raw_df["Frac_Azimuth"] = dataset_builder.POINT.safe_to_numeric(raw_df["Frac_Azimuth"])
    raw_df = raw_df.dropna(subset=["MD"]).copy()
    raw_df = raw_df.sort_values("MD").drop_duplicates(subset=["MD"], keep="first").reset_index(drop=True)
    raw_df[depth_col] = raw_df["MD"].astype(np.float64)
    raw_df["RawPointID"] = np.arange(1, len(raw_df) + 1, dtype=np.int64)
    raw_df["WellName"] = well_name
    raw_df["StrataName"] = str(strata_name or "").strip()
    raw_df["SourceFile"] = str(raw_path)
    strata_range = resolve_well_strata_range(well_name, strata_name, strata_range_df)
    if strata_range is not None:
        raw_df = filter_df_to_depth_range(
            raw_df,
            depth_col=depth_col,
            depth_min=float(strata_range["DepthMin"]),
            depth_max=float(strata_range["DepthMax"]),
        ).copy()
    return raw_df[["WellName", "StrataName", "RawPointID", "MD", depth_col, "Frac_Dip", "Frac_Azimuth", "SourceFile"]]


def build_gt_dev_segments(
    sample_dir: Path,
    well_name: str,
    gt_dev_rule: str,
    depth_col: str = "TVD",
    strata_name: str = "",
    strata_range_df: pd.DataFrame | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    sample_file_name = WELL_TO_SAMPLE_FILE.get(well_name)
    if not sample_file_name:
        raise KeyError(f"No sample file mapping for well: {well_name}")
    sample_path = sample_dir / sample_file_name
    if not sample_path.exists():
        raise FileNotFoundError(f"Sample file not found for {well_name}: {sample_path}")

    sample_df = pd.read_csv(sample_path, encoding="utf-8-sig")
    if depth_col not in sample_df.columns:
        raise ValueError(f"Missing depth col {depth_col} in sample file: {sample_path}")

    sample_df = sample_df.copy()
    sample_df[depth_col] = dataset_builder.POINT.safe_to_numeric(sample_df[depth_col])
    for col in ["P10", "P21", "P33"]:
        if col in sample_df.columns:
            sample_df[col] = dataset_builder.POINT.safe_to_numeric(sample_df[col]).fillna(0.0)
    if "Frac_Azimuth" in sample_df.columns:
        sample_df["Frac_Azimuth"] = dataset_builder.POINT.safe_to_numeric(sample_df["Frac_Azimuth"])
    if "GT_LABEL" in sample_df.columns:
        sample_df["GT_LABEL"] = dataset_builder.POINT.safe_to_numeric(sample_df["GT_LABEL"]).fillna(0.0)

    sample_df = sample_df.dropna(subset=[depth_col]).sort_values(depth_col).reset_index(drop=True)
    strata_range = resolve_well_strata_range(well_name, strata_name, strata_range_df)
    if strata_range is not None:
        sample_df = filter_df_to_depth_range(
            sample_df,
            depth_col=depth_col,
            depth_min=float(strata_range["DepthMin"]),
            depth_max=float(strata_range["DepthMax"]),
        ).copy()
        sample_df = sample_df.sort_values(depth_col).reset_index(drop=True)

    if sample_df.empty:
        return sample_df, pd.DataFrame(
            columns=[
                "WellName",
                "GTDevSegment_ID",
                "SegStartDepth",
                "SegEndDepth",
                "SegLength",
                "SampleCount",
                "SourceFile",
                "GTDevRule",
                "StrataName",
            ]
        )

    if gt_dev_rule == "any_density":
        mask = (
            sample_df[[col for col in ["P10", "P21", "P33"] if col in sample_df.columns]]
            .fillna(0.0)
            .max(axis=1)
            .gt(0.0)
        )
    elif gt_dev_rule == "p10_only":
        if "P10" not in sample_df.columns:
            raise ValueError(f"P10 missing in sample file: {sample_path}")
        mask = sample_df["P10"].fillna(0.0).gt(0.0)
    elif gt_dev_rule == "frac_azimuth_notna":
        if "Frac_Azimuth" not in sample_df.columns:
            raise ValueError(f"Frac_Azimuth missing in sample file: {sample_path}")
        mask = sample_df["Frac_Azimuth"].notna()
    elif gt_dev_rule == "gt_label":
        if "GT_LABEL" not in sample_df.columns:
            raise ValueError(f"GT_LABEL missing in sample file: {sample_path}")
        mask = sample_df["GT_LABEL"].fillna(0.0).ge(1.0)
    else:
        raise ValueError(f"Unsupported gt_dev_rule: {gt_dev_rule}")

    depth = sample_df[depth_col].to_numpy(dtype=np.float64)
    segments = dataset_builder.POINT.find_binary_segments(mask.to_numpy(dtype=bool))
    gt_rows = []
    for seg_id, (start_idx, end_idx) in enumerate(segments, start=1):
        depth_seg = depth[start_idx : end_idx + 1]
        seg_length = float(max(depth_seg[-1] - depth_seg[0], 0.0)) if depth_seg.size >= 2 else 0.0
        gt_rows.append(
            {
                "WellName": well_name,
                "GTDevSegment_ID": seg_id,
                "SegStartDepth": float(depth_seg[0]),
                "SegEndDepth": float(depth_seg[-1]),
                "SegLength": seg_length,
                "SampleCount": int(len(depth_seg)),
                "SourceFile": str(sample_path),
                "GTDevRule": gt_dev_rule,
                "StrataName": str(strata_name or "").strip(),
            }
        )
    gt_dev_df = pd.DataFrame(gt_rows)
    return sample_df, gt_dev_df


def summarize_segment_p10(seg_df: pd.DataFrame, depth_col: str, seg_length: float) -> dict[str, float]:
    if "P10" not in seg_df.columns:
        return {
            "RawP10MeanInPredSegment": np.nan,
            "RawP10MaxInPredSegment": np.nan,
            "RawP10MassInPredSegment": np.nan,
            "RawP10MassPerLength": np.nan,
            "RawP10PositiveFracInPredSegment": np.nan,
        }

    p10 = pd.to_numeric(seg_df["P10"], errors="coerce").fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
    if p10.size == 0:
        return {
            "RawP10MeanInPredSegment": 0.0,
            "RawP10MaxInPredSegment": 0.0,
            "RawP10MassInPredSegment": 0.0,
            "RawP10MassPerLength": 0.0,
            "RawP10PositiveFracInPredSegment": 0.0,
        }

    depth = pd.to_numeric(seg_df.get(depth_col), errors="coerce").to_numpy(dtype=np.float64)
    valid_mask = np.isfinite(depth) & np.isfinite(p10)
    depth_valid = depth[valid_mask]
    p10_valid = p10[valid_mask]
    if depth_valid.size >= 2:
        if hasattr(np, "trapezoid"):
            density_mass = float(np.trapezoid(p10_valid, depth_valid))
        else:
            density_mass = float(np.trapz(p10_valid, depth_valid))
    else:
        density_mass = float(np.mean(p10_valid) * max(float(seg_length), 0.0))
    density_mean = float(np.mean(p10_valid)) if p10_valid.size > 0 else 0.0
    density_max = float(np.max(p10_valid)) if p10_valid.size > 0 else 0.0
    density_pos_frac = float(np.mean(p10_valid > 0.0)) if p10_valid.size > 0 else 0.0
    density_mass = max(density_mass, 0.0)
    return {
        "RawP10MeanInPredSegment": density_mean,
        "RawP10MaxInPredSegment": density_max,
        "RawP10MassInPredSegment": density_mass,
        "RawP10MassPerLength": float(density_mass / seg_length) if seg_length > 1e-8 else 0.0,
        "RawP10PositiveFracInPredSegment": density_pos_frac,
    }


def summarize_segment_log_stats(seg_df: pd.DataFrame, log_cols: tuple[str, ...] = ("AC", "GR")) -> dict[str, float]:
    stats: dict[str, float] = {}
    for col in log_cols:
        values = pd.to_numeric(seg_df.get(col), errors="coerce").to_numpy(dtype=np.float64)
        finite_values = values[np.isfinite(values)]
        if finite_values.size == 0:
            stats[f"{col}_MeanInPredSegment"] = np.nan
            stats[f"{col}_StdInPredSegment"] = np.nan
            continue
        stats[f"{col}_MeanInPredSegment"] = float(np.mean(finite_values))
        stats[f"{col}_StdInPredSegment"] = float(np.std(finite_values))
    return stats


def wrap_azimuth_deg(values: np.ndarray | pd.Series | list[float]) -> np.ndarray:
    array = np.asarray(values, dtype=np.float64)
    if array.size == 0:
        return array.astype(np.float64)
    return np.mod(array, 360.0)


def circular_mean_deg(values: np.ndarray | pd.Series | list[float]) -> float:
    array = wrap_azimuth_deg(values)
    finite = array[np.isfinite(array)]
    if finite.size == 0:
        return np.nan
    radians = np.deg2rad(finite)
    sin_mean = float(np.mean(np.sin(radians)))
    cos_mean = float(np.mean(np.cos(radians)))
    if abs(sin_mean) <= 1e-12 and abs(cos_mean) <= 1e-12:
        return np.nan
    angle = np.degrees(np.arctan2(sin_mean, cos_mean))
    return float(np.mod(angle, 360.0))


def angular_abs_diff_deg(a_deg: object, b_deg: object) -> float:
    if not np.isfinite(pd.to_numeric(pd.Series([a_deg]), errors="coerce").iloc[0]):
        return np.nan
    if not np.isfinite(pd.to_numeric(pd.Series([b_deg]), errors="coerce").iloc[0]):
        return np.nan
    a = float(a_deg) % 360.0
    b = float(b_deg) % 360.0
    diff = abs(a - b)
    return float(min(diff, 360.0 - diff))


def summarize_segment_orientation(segment_raw_points: pd.DataFrame) -> dict[str, object]:
    if segment_raw_points.empty:
        return {
            "RawOrientationPointCountInPredSegment": 0,
            "RawAzimuthMeanInPredSegment": np.nan,
            "RawDipMeanInPredSegment": np.nan,
            "RawOrientationAzimuthStdInPredSegment": np.nan,
            "RawOrientationDipStdInPredSegment": np.nan,
            "HasOrientationPoint": 0,
        }

    azimuth = pd.to_numeric(segment_raw_points.get("Frac_Azimuth"), errors="coerce").to_numpy(dtype=np.float64)
    dip = pd.to_numeric(segment_raw_points.get("Frac_Dip"), errors="coerce").to_numpy(dtype=np.float64)
    valid_mask = np.isfinite(azimuth) & np.isfinite(dip)
    azimuth = azimuth[valid_mask]
    dip = dip[valid_mask]
    if azimuth.size == 0:
        return {
            "RawOrientationPointCountInPredSegment": 0,
            "RawAzimuthMeanInPredSegment": np.nan,
            "RawDipMeanInPredSegment": np.nan,
            "RawOrientationAzimuthStdInPredSegment": np.nan,
            "RawOrientationDipStdInPredSegment": np.nan,
            "HasOrientationPoint": 0,
        }

    azimuth_wrapped = wrap_azimuth_deg(azimuth)
    azimuth_mean = circular_mean_deg(azimuth_wrapped)
    azimuth_diff = np.asarray([angular_abs_diff_deg(value, azimuth_mean) for value in azimuth_wrapped], dtype=np.float64)
    return {
        "RawOrientationPointCountInPredSegment": int(azimuth_wrapped.size),
        "RawAzimuthMeanInPredSegment": azimuth_mean,
        "RawDipMeanInPredSegment": float(np.mean(dip)),
        "RawOrientationAzimuthStdInPredSegment": float(np.nanstd(azimuth_diff)),
        "RawOrientationDipStdInPredSegment": float(np.nanstd(dip)),
        "HasOrientationPoint": 1,
    }


def build_orientation_cluster_features(azimuth_deg: np.ndarray, dip_deg: np.ndarray) -> np.ndarray:
    azimuth = wrap_azimuth_deg(azimuth_deg)
    dip = np.clip(np.asarray(dip_deg, dtype=np.float64), a_min=0.0, a_max=90.0)
    radians = np.deg2rad(azimuth)
    return np.column_stack(
        [
            np.cos(radians),
            np.sin(radians),
            dip / 90.0,
        ]
    )


def extract_segment_raw_points_by_depth(
    raw_gt_points: pd.DataFrame,
    depth_col: str,
    start_depth: float,
    end_depth: float,
) -> pd.DataFrame:
    if raw_gt_points.empty:
        return raw_gt_points.iloc[0:0].copy()
    depth = pd.to_numeric(raw_gt_points.get(depth_col), errors="coerce")
    mask = depth.ge(float(start_depth) - 1e-6) & depth.le(float(end_depth) + 1e-6)
    return raw_gt_points.loc[mask].copy()


def resolve_density_prediction_outputs(
    pred_density_strength: np.ndarray,
    segment_df: pd.DataFrame,
    strategy: dict,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    pred = np.clip(np.asarray(pred_density_strength, dtype=np.float64), a_min=0.0, a_max=None)
    pred_p10_mean = pred.copy()
    pred_p10_mass_per_length = pred.copy()
    target_col = str(strategy.get("target_col") or "")
    mean_over_mass_per_length_ratio = float(strategy.get("mean_over_mass_per_length_ratio", 1.0))
    if not np.isfinite(mean_over_mass_per_length_ratio) or mean_over_mass_per_length_ratio <= 1e-8:
        mean_over_mass_per_length_ratio = 1.0
    if target_col == "RawP10MeanInPredSegment":
        pred_p10_mass_per_length = pred_p10_mean / mean_over_mass_per_length_ratio
    elif target_col == "RawP10MassPerLength":
        pred_p10_mean = pred_p10_mass_per_length * mean_over_mass_per_length_ratio
    seg_length = (
        segment_df["SegLength"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
        if "SegLength" in segment_df.columns
        else np.zeros(len(pred), dtype=np.float64)
    )
    pred_p10_mass = pred_p10_mass_per_length * seg_length
    return (
        np.clip(pred_p10_mean, a_min=0.0, a_max=None),
        np.clip(pred_p10_mass_per_length, a_min=0.0, a_max=None),
        np.clip(pred_p10_mass, a_min=0.0, a_max=None),
    )


def compute_orientation_error_stats(segment_df: pd.DataFrame) -> dict[str, float]:
    if segment_df.empty:
        return {
            "orientation_num_eval_segments": 0,
            "orientation_azimuth_mae_deg": np.nan,
            "orientation_dip_mae_deg": np.nan,
        }
    azimuth_true = pd.to_numeric(segment_df.get("RawAzimuthMeanInPredSegment"), errors="coerce").to_numpy(dtype=np.float64)
    dip_true = pd.to_numeric(segment_df.get("RawDipMeanInPredSegment"), errors="coerce").to_numpy(dtype=np.float64)
    azimuth_pred = pd.to_numeric(segment_df.get("PredAzimuth"), errors="coerce").to_numpy(dtype=np.float64)
    dip_pred = pd.to_numeric(segment_df.get("PredDip"), errors="coerce").to_numpy(dtype=np.float64)
    valid_mask = (
        np.isfinite(azimuth_true)
        & np.isfinite(dip_true)
        & np.isfinite(azimuth_pred)
        & np.isfinite(dip_pred)
    )
    if not np.any(valid_mask):
        return {
            "orientation_num_eval_segments": 0,
            "orientation_azimuth_mae_deg": np.nan,
            "orientation_dip_mae_deg": np.nan,
        }
    azimuth_errors = np.asarray(
        [
            angular_abs_diff_deg(pred_value, true_value)
            for pred_value, true_value in zip(azimuth_pred[valid_mask], azimuth_true[valid_mask])
        ],
        dtype=np.float64,
    )
    dip_errors = np.abs(dip_pred[valid_mask] - dip_true[valid_mask])
    return {
        "orientation_num_eval_segments": int(np.sum(valid_mask.astype(np.int64))),
        "orientation_azimuth_mae_deg": float(np.nanmean(azimuth_errors)) if azimuth_errors.size > 0 else np.nan,
        "orientation_dip_mae_deg": float(np.nanmean(dip_errors)) if dip_errors.size > 0 else np.nan,
    }


def summarize_orientation_spread_values(values: np.ndarray | pd.Series | list[float]) -> dict[str, float]:
    array = np.asarray(values, dtype=np.float64)
    finite = array[np.isfinite(array)]
    if finite.size == 0:
        return {"median": np.nan, "q80": np.nan}
    clipped = np.clip(finite, a_min=0.0, a_max=None)
    return {
        "median": float(np.quantile(clipped, 0.50)),
        "q80": float(np.quantile(clipped, 0.80)),
    }


def first_finite_float(*values: object, default: float = np.nan) -> float:
    for value in values:
        try:
            numeric = float(value)
        except (TypeError, ValueError):
            continue
        if np.isfinite(numeric):
            return numeric
    return float(default)


def build_point_orientation_rank_template(picked_depths: np.ndarray | list[float]) -> np.ndarray:
    depth_array = np.asarray(picked_depths, dtype=np.float64)
    num_points = int(depth_array.size)
    if num_points <= 1:
        return np.zeros(num_points, dtype=np.float64)
    template_sorted = np.linspace(-1.0, 1.0, num_points, dtype=np.float64)
    template_sorted -= float(np.mean(template_sorted))
    template_std = float(np.std(template_sorted))
    if template_std <= 1e-8:
        return np.zeros(num_points, dtype=np.float64)
    template_sorted = template_sorted / template_std
    sort_order = np.argsort(depth_array, kind="mergesort")
    template = np.zeros(num_points, dtype=np.float64)
    template[sort_order] = template_sorted
    return template


def expand_segment_orientation_to_points(
    pred_info: dict,
    orientation_strategy: dict | None,
    picked_depths: np.ndarray | list[float],
) -> dict[str, object]:
    num_points = int(len(picked_depths))
    seg_azimuth = first_finite_float(pred_info.get("PredAzimuth"))
    seg_dip = first_finite_float(pred_info.get("PredDip"))
    point_azimuth = np.full(num_points, seg_azimuth, dtype=np.float64)
    point_dip = np.full(num_points, seg_dip, dtype=np.float64)
    template = np.zeros(num_points, dtype=np.float64)
    azimuth_offset = np.zeros(num_points, dtype=np.float64)
    dip_offset = np.zeros(num_points, dtype=np.float64)
    spread_mode = "segment_constant"
    confidence_scale = 1.0
    azimuth_std_ref = np.nan
    dip_std_ref = np.nan

    strategy = dict(orientation_strategy or {})
    if (
        num_points <= 1
        or validate_orientation_mode(str(strategy.get("mode", "none"))) == "none"
        or not np.isfinite(seg_azimuth)
        or not np.isfinite(seg_dip)
    ):
        return {
            "mode": spread_mode,
            "template": template,
            "confidence_scale": confidence_scale,
            "azimuth_std_ref_deg": azimuth_std_ref,
            "dip_std_ref_deg": dip_std_ref,
            "azimuth_offset_deg": azimuth_offset,
            "dip_offset_deg": dip_offset,
            "point_azimuth": point_azimuth,
            "point_dip": point_dip,
        }

    family_name = str(pred_info.get("PredOrientationFamily") or "").strip()
    family_spreads = {
        str(name): dict(value)
        for name, value in dict(strategy.get("family_spreads") or {}).items()
    }
    family_spread = dict(family_spreads.get(family_name) or {})
    global_spread = dict(strategy.get("global_spread") or {})
    azimuth_std_ref = first_finite_float(
        family_spread.get("azimuth_std_deg"),
        global_spread.get("azimuth_std_deg"),
        default=0.0,
    )
    dip_std_ref = first_finite_float(
        family_spread.get("dip_std_deg"),
        global_spread.get("dip_std_deg"),
        default=0.0,
    )
    if azimuth_std_ref <= 1e-8 and dip_std_ref <= 1e-8:
        return {
            "mode": spread_mode,
            "template": template,
            "confidence_scale": confidence_scale,
            "azimuth_std_ref_deg": azimuth_std_ref,
            "dip_std_ref_deg": dip_std_ref,
            "azimuth_offset_deg": azimuth_offset,
            "dip_offset_deg": dip_offset,
            "point_azimuth": point_azimuth,
            "point_dip": point_dip,
        }

    pred_confidence = first_finite_float(pred_info.get("PredOrientationConfidence"), default=np.nan)
    conf_scale_min = float(np.clip(first_finite_float(strategy.get("point_spread_conf_scale_min"), default=0.45), 0.05, 1.0))
    conf_scale_max = float(max(conf_scale_min, first_finite_float(strategy.get("point_spread_conf_scale_max"), default=1.0)))
    if np.isfinite(pred_confidence):
        clipped_confidence = float(np.clip(pred_confidence, 0.0, 1.0))
        confidence_scale = float(np.clip(conf_scale_max - 0.60 * clipped_confidence, conf_scale_min, conf_scale_max))
    else:
        confidence_scale = conf_scale_max

    template = build_point_orientation_rank_template(picked_depths)
    azimuth_cap = first_finite_float(
        family_spread.get("azimuth_q80_deg"),
        global_spread.get("azimuth_q80_deg"),
        azimuth_std_ref * 2.5,
        default=max(azimuth_std_ref * 2.5, 0.0),
    )
    dip_cap = first_finite_float(
        family_spread.get("dip_q80_deg"),
        global_spread.get("dip_q80_deg"),
        dip_std_ref * 2.5,
        default=max(dip_std_ref * 2.5, 0.0),
    )
    azimuth_offset = template * azimuth_std_ref * confidence_scale
    dip_offset = template * dip_std_ref * confidence_scale
    if np.isfinite(azimuth_cap) and azimuth_cap >= 0.0:
        azimuth_offset = np.clip(azimuth_offset, a_min=-azimuth_cap, a_max=azimuth_cap)
    if np.isfinite(dip_cap) and dip_cap >= 0.0:
        dip_offset = np.clip(dip_offset, a_min=-dip_cap, a_max=dip_cap)
    point_azimuth = wrap_azimuth_deg(seg_azimuth + azimuth_offset)
    point_dip = np.clip(seg_dip + dip_offset, a_min=0.0, a_max=90.0)
    spread_mode = "family_confidence_rank_template_v1"
    return {
        "mode": spread_mode,
        "template": template,
        "confidence_scale": confidence_scale,
        "azimuth_std_ref_deg": azimuth_std_ref,
        "dip_std_ref_deg": dip_std_ref,
        "azimuth_offset_deg": azimuth_offset,
        "dip_offset_deg": dip_offset,
        "point_azimuth": point_azimuth,
        "point_dip": point_dip,
    }


def build_density_point_floor_info(density_strength_strategy: dict) -> dict:
    strategy = dict(density_strength_strategy or {})
    quantiles = dict(strategy.get("positive_quantiles") or {})
    q50 = float(quantiles.get("q50", np.nan))
    q80 = float(quantiles.get("q80", np.nan))
    q95 = float(quantiles.get("q95", np.nan))
    threshold_mass_per_length = np.nan
    enabled = False
    source_quantile = ""
    if density_strength_uses_regression(str(strategy.get("mode", "none"))):
        if np.isfinite(q95) and q95 > 1e-8:
            threshold_mass_per_length = float(q95)
            enabled = True
            source_quantile = "q95"
        elif np.isfinite(q80) and q80 > 1e-8:
            threshold_mass_per_length = float(q80)
            enabled = True
            source_quantile = "q80"
        elif np.isfinite(q50) and q50 > 1e-8:
            threshold_mass_per_length = float(q50)
            enabled = True
            source_quantile = "q50"
    elif np.isfinite(q50) and q50 > 1e-8:
        threshold_mass_per_length = float(q50 * 0.5)
        enabled = True
        source_quantile = "q50_x_0.5"
    elif validate_density_strength_mode(str(strategy.get("mode", "none"))) != "none":
        threshold_mass_per_length = 0.0
        enabled = True
        source_quantile = "zero_fallback"
    return {
        "enabled": bool(enabled),
        "threshold_mass_per_length": threshold_mass_per_length,
        "source_quantile": source_quantile if enabled else "",
    }


def apply_density_point_floor(
    pred_point_count: np.ndarray,
    pred_p10_mass_per_length: np.ndarray,
    floor_info: dict,
) -> tuple[np.ndarray, np.ndarray]:
    counts = np.asarray(pred_point_count, dtype=np.int64).copy()
    applied_mask = np.zeros(counts.shape[0], dtype=bool)
    if counts.size == 0:
        return counts, applied_mask
    if not bool(floor_info.get("enabled", False)):
        return counts, applied_mask
    threshold = float(floor_info.get("threshold_mass_per_length", np.nan))
    density = np.asarray(pred_p10_mass_per_length, dtype=np.float64)
    if not np.isfinite(threshold):
        return counts, applied_mask
    applied_mask = (counts <= 0) & np.isfinite(density) & (density >= threshold)
    counts[applied_mask] = 1
    return counts, applied_mask


def build_well_segment_dataset(
    exist_csv: str,
    well_name: str,
    config,
    sample_dir: Path,
    raw_label_dir: Path,
    gt_dev_rule: str,
    strata_name: str = "",
    strata_range_df: pd.DataFrame | None = None,
    boundary_expand_mode: str = "none",
    boundary_expand_prob_min: float = 0.0,
    boundary_expand_max_steps: int = 0,
    boundary_expand_max_depth: float = 0.0,
) -> dict:
    df = pd.read_csv(exist_csv, encoding="utf-8-sig")
    strata_range = resolve_well_strata_range(well_name, strata_name, strata_range_df)
    if strata_range is not None:
        df = filter_df_to_depth_range(
            df,
            depth_col=config.depth_col,
            depth_min=float(strata_range["DepthMin"]),
            depth_max=float(strata_range["DepthMax"]),
        ).copy()
    raw_gt_points = load_raw_fracture_points(
        raw_label_dir=raw_label_dir,
        well_name=well_name,
        depth_col=config.depth_col,
        strata_name=strata_name,
        strata_range_df=strata_range_df,
    )
    raw_gt_depths = (
        raw_gt_points[config.depth_col].to_numpy(dtype=np.float64)
        if not raw_gt_points.empty
        else np.array([], dtype=np.float64)
    )
    _, gt_dev_segments = build_gt_dev_segments(
        sample_dir=sample_dir,
        well_name=well_name,
        gt_dev_rule=gt_dev_rule,
        depth_col=config.depth_col,
        strata_name=strata_name,
        strata_range_df=strata_range_df,
    )
    df_for_segment, effective_mask_col, boundary_stats = apply_boundary_expand_to_pred_mask(
        df=df,
        config=config,
        mode=boundary_expand_mode,
        prob_min=boundary_expand_prob_min,
        max_steps=boundary_expand_max_steps,
        max_depth=boundary_expand_max_depth,
    )
    df_sorted, payloads = dataset_builder.POINT.extract_segment_payloads(
        df_for_segment,
        mask_col=effective_mask_col,
        density_col=None,
        config=config,
        prob_col=config.pred_prob_col,
        shape_mode="prob_only",
    )

    rows = []
    for payload in payloads:
        seg_df = df_sorted.iloc[payload["StartIdx"] : payload["EndIdx"] + 1].copy()
        segment_raw_points = extract_segment_raw_points_by_depth(
            raw_gt_points=raw_gt_points,
            depth_col=config.depth_col,
            start_depth=float(payload["SegStartDepth"]),
            end_depth=float(payload["SegEndDepth"]),
        )
        row = dataset_builder.build_segment_row(
            well_name=well_name,
            payload=payload,
            df_sorted=df_sorted,
            gt_depths=raw_gt_depths,
            config=config,
        )
        raw_gt_count = int(row.get("GTPointCountInPredSegment", 0))
        overlap_len, overlap_count = compute_overlap_with_gt_dev(
            start_depth=float(payload["SegStartDepth"]),
            end_depth=float(payload["SegEndDepth"]),
            gt_dev_segments=gt_dev_segments,
        )
        seg_length = float(max(row.get("SegLength", 0.0), 0.0))
        row.update(
            {
                "GTPointCountInPredSegment": raw_gt_count,
                "RawGTCountInPredSegment": raw_gt_count,
                "HasRawPoint": int(raw_gt_count > 0),
                "GTDevOverlapLen": float(overlap_len),
                "GTDevOverlapRatio": float(overlap_len / seg_length) if seg_length > 1e-8 else 0.0,
                "GTDevOverlapCount": int(overlap_count),
                "GTDevOverlapFlag": int(overlap_len > 0),
            }
        )
        row.update(
            summarize_segment_p10(
                seg_df=seg_df,
                depth_col=config.depth_col,
                seg_length=float(max(row.get("SegLength", 0.0), 0.0)),
            )
        )
        row.update(summarize_segment_log_stats(seg_df=seg_df))
        row.update(summarize_segment_orientation(segment_raw_points=segment_raw_points))
        rows.append(row)

    segment_df = pd.DataFrame(rows)
    if not segment_df.empty:
        segment_df["StrataName"] = str(strata_name or "").strip()
        if strata_range is not None:
            segment_df["StrataDepthMin"] = float(strata_range["DepthMin"])
            segment_df["StrataDepthMax"] = float(strata_range["DepthMax"])
    raw_points_covered = count_points_in_any_segments(raw_gt_depths, segment_df)
    return {
        "well_name": well_name,
        "strata_name": str(strata_name or "").strip(),
        "strata_range": strata_range,
        "exist_csv": exist_csv,
        "df": df,
        "df_effective": df_for_segment,
        "df_sorted": df_sorted,
        "payloads": payloads,
        "raw_gt_points": raw_gt_points,
        "raw_gt_depths": raw_gt_depths,
        "gt_dev_segments": gt_dev_segments,
        "segment_df": segment_df,
        "boundary_expand_stats": boundary_stats,
        "raw_points_covered_by_pred_segments": int(raw_points_covered),
        "raw_points_missed_outside_pred_segments": int(max(len(raw_gt_depths) - raw_points_covered, 0)),
    }


def build_predict_segment_dataset(
    exist_csv: str,
    well_name: str,
    config,
    strata_name: str = "",
    strata_range_df: pd.DataFrame | None = None,
    boundary_expand_mode: str = "none",
    boundary_expand_prob_min: float = 0.0,
    boundary_expand_max_steps: int = 0,
    boundary_expand_max_depth: float = 0.0,
) -> dict:
    df = pd.read_csv(exist_csv, encoding="utf-8-sig")
    strata_range = resolve_well_strata_range(well_name, strata_name, strata_range_df)
    if strata_range is not None:
        df = filter_df_to_depth_range(
            df,
            depth_col=config.depth_col,
            depth_min=float(strata_range["DepthMin"]),
            depth_max=float(strata_range["DepthMax"]),
        ).copy()
    df_for_segment, effective_mask_col, boundary_stats = apply_boundary_expand_to_pred_mask(
        df=df,
        config=config,
        mode=boundary_expand_mode,
        prob_min=boundary_expand_prob_min,
        max_steps=boundary_expand_max_steps,
        max_depth=boundary_expand_max_depth,
    )
    df_sorted, payloads = dataset_builder.POINT.extract_segment_payloads(
        df_for_segment,
        mask_col=effective_mask_col,
        density_col=None,
        config=config,
        prob_col=config.pred_prob_col,
        shape_mode="prob_only",
    )

    empty_gt_depths = np.array([], dtype=np.float64)
    rows = []
    for payload in payloads:
        seg_df = df_sorted.iloc[payload["StartIdx"] : payload["EndIdx"] + 1].copy()
        row = dataset_builder.build_segment_row(
            well_name=well_name,
            payload=payload,
            df_sorted=df_sorted,
            gt_depths=empty_gt_depths,
            config=config,
        )
        row.update(summarize_segment_log_stats(seg_df=seg_df))
        row.update(
            {
                "StrataName": str(strata_name or "").strip(),
                "RawGTCountInPredSegment": 0,
                "HasRawPoint": 0,
                "GTDevOverlapLen": 0.0,
                "GTDevOverlapRatio": 0.0,
                "GTDevOverlapCount": 0,
                "GTDevOverlapFlag": 0,
                "TargetWeight": np.nan,
                "TargetLabelType": "",
                "RawP10MeanInPredSegment": np.nan,
                "RawP10MaxInPredSegment": np.nan,
                "RawP10MassInPredSegment": np.nan,
                "RawP10MassPerLength": np.nan,
                "RawP10PositiveFracInPredSegment": np.nan,
                "RawOrientationPointCountInPredSegment": 0,
                "RawAzimuthMeanInPredSegment": np.nan,
                "RawDipMeanInPredSegment": np.nan,
                "RawOrientationAzimuthStdInPredSegment": np.nan,
                "RawOrientationDipStdInPredSegment": np.nan,
                "HasOrientationPoint": 0,
            }
        )
        rows.append(row)

    return {
        "well_name": well_name,
        "strata_name": str(strata_name or "").strip(),
        "strata_range": strata_range,
        "exist_csv": exist_csv,
        "df": df,
        "df_effective": df_for_segment,
        "df_sorted": df_sorted,
        "payloads": payloads,
        "boundary_expand_stats": boundary_stats,
        "segment_df": pd.DataFrame(rows),
    }


def compute_target_weights(segment_df: pd.DataFrame, soft_negative_weight: float) -> pd.DataFrame:
    if segment_df.empty:
        out = segment_df.copy()
        out["TargetWeight"] = []
        out["TargetLabelType"] = []
        return out

    clipped_soft_weight = float(np.clip(soft_negative_weight, 0.0, 1.0))
    target_weights = []
    target_types = []
    for row in segment_df.itertuples(index=False):
        raw_count = int(getattr(row, "RawGTCountInPredSegment", 0))
        overlap_len = float(getattr(row, "GTDevOverlapLen", 0.0))
        if raw_count > 0:
            target_weights.append(1.0)
            target_types.append("hard_positive")
        elif overlap_len > 0:
            target_weights.append(clipped_soft_weight)
            target_types.append("soft_zero_overlap_gt_dev")
        else:
            target_weights.append(1.0)
            target_types.append("hard_zero")

    out = segment_df.copy()
    out["TargetWeight"] = np.asarray(target_weights, dtype=np.float64)
    out["TargetLabelType"] = target_types
    return out


def fit_basis_scale(train_df: pd.DataFrame, count_mode: str) -> tuple[str, float]:
    basis_col = dataset_builder.POINT.COUNT_MODE_TO_BASIS_COL[count_mode]
    basis_sum = float(train_df[basis_col].fillna(0.0).clip(lower=0.0).sum())
    target_sum = float(train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0).sum())
    scale = (target_sum / basis_sum) if basis_sum > 1e-8 else 1.0
    return basis_col, float(max(scale, 0.0))


def fuse_count_predictions(
    learned_pred: np.ndarray,
    fallback_pred: np.ndarray | None,
    fusion_mode: str,
    learned_weight: float,
) -> np.ndarray:
    learned_pred = np.clip(np.asarray(learned_pred, dtype=np.float64), a_min=0.0, a_max=None)
    if fallback_pred is None or not fusion_mode or fusion_mode == "learned_only":
        return learned_pred

    fallback_pred = np.clip(np.asarray(fallback_pred, dtype=np.float64), a_min=0.0, a_max=None)
    if fusion_mode == "max":
        return np.maximum(learned_pred, fallback_pred)
    if fusion_mode == "blend":
        weight = float(np.clip(learned_weight, 0.0, 1.0))
        return (weight * learned_pred) + ((1.0 - weight) * fallback_pred)
    if fusion_mode == "uplift":
        weight = float(np.clip(learned_weight, 0.0, 1.0))
        positive_gap = np.clip(fallback_pred - learned_pred, a_min=0.0, a_max=None)
        return learned_pred + ((1.0 - weight) * positive_gap)
    raise ValueError(f"Unsupported count_fusion_mode: {fusion_mode}")


def fit_global_post_scale(
    train_target: np.ndarray,
    train_pred: np.ndarray,
    mode: str,
    scale_min: float,
    scale_max: float,
) -> dict:
    normalized_mode = validate_global_post_scale_mode(mode)
    lower = float(scale_min)
    upper = float(scale_max)
    if upper < lower:
        raise ValueError(f"global_post_scale_max must be >= global_post_scale_min, got {upper} < {lower}")

    target = np.clip(np.asarray(train_target, dtype=np.float64), a_min=0.0, a_max=None)
    pred = np.clip(np.asarray(train_pred, dtype=np.float64), a_min=0.0, a_max=None)
    train_target_sum = float(np.sum(target))
    train_pred_sum = float(np.sum(pred))
    raw_scale = (train_target_sum / train_pred_sum) if train_pred_sum > 1e-8 else 1.0

    if normalized_mode == "none":
        applied_scale = 1.0
    elif normalized_mode == "train_sum_match":
        applied_scale = raw_scale
    elif normalized_mode == "train_sum_match_up_only":
        applied_scale = max(raw_scale, 1.0)
    else:
        raise ValueError(f"Unsupported global_post_scale_mode: {normalized_mode}")

    applied_scale = float(np.clip(applied_scale, lower, upper))
    return {
        "mode": normalized_mode,
        "raw_scale": float(raw_scale),
        "scale": applied_scale,
        "train_target_sum": train_target_sum,
        "train_pred_sum": train_pred_sum,
        "scale_min": lower,
        "scale_max": upper,
    }


def apply_global_post_scale(pred_count_float: np.ndarray, post_scale_info: dict) -> np.ndarray:
    scale = float(post_scale_info.get("scale", 1.0))
    pred = np.asarray(pred_count_float, dtype=np.float64)
    return np.clip(pred * scale, a_min=0.0, a_max=None)


def fit_selective_post_strategy(
    train_df: pd.DataFrame,
    train_pred_after_global_scale: np.ndarray,
    mode: str,
    feature_name: str,
    feature_name_2: str,
    quantile: float,
    quantile_2: float,
    fallback_scale: float,
    max_pred_float: float,
) -> dict:
    normalized_mode = validate_selective_post_mode(mode)
    clipped_quantile = float(np.clip(float(quantile), 0.0, 1.0))
    clipped_quantile_2 = float(np.clip(float(quantile_2), 0.0, 1.0))
    clipped_max_pred_float = normalize_optional_positive_float(max_pred_float)
    feature_values = train_df.get(feature_name)
    feature_values_2 = train_df.get(feature_name_2) if feature_name_2 else None
    if normalized_mode == "none":
        return {
            "mode": normalized_mode,
            "feature_name": feature_name,
            "feature_name_2": feature_name_2,
            "quantile": clipped_quantile,
            "quantile_2": clipped_quantile_2,
            "threshold": np.nan,
            "threshold_2": np.nan,
            "lowconf_scale": 1.0,
            "max_pred_float": clipped_max_pred_float,
            "num_train_zero_segments": 0,
        }

    if feature_values is None:
        raise ValueError(f"Missing selective post feature: {feature_name}")
    if normalized_mode == "train_zero_dual_quantile_downscale" and feature_values_2 is None:
        raise ValueError(f"Missing selective post feature 2: {feature_name_2}")

    zero_mask = train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64) <= 0.0
    zero_feature_values = (
        pd.to_numeric(feature_values, errors="coerce")
        .to_numpy(dtype=np.float64)
    )
    zero_feature_values = zero_feature_values[zero_mask]
    zero_feature_values = zero_feature_values[np.isfinite(zero_feature_values)]
    if zero_feature_values.size == 0:
        threshold = np.nan
    else:
        threshold = float(np.quantile(zero_feature_values, clipped_quantile))

    threshold_2 = np.nan
    if normalized_mode == "train_zero_dual_quantile_downscale":
        zero_feature_values_2 = (
            pd.to_numeric(feature_values_2, errors="coerce")
            .to_numpy(dtype=np.float64)
        )
        zero_feature_values_2 = zero_feature_values_2[zero_mask]
        zero_feature_values_2 = zero_feature_values_2[np.isfinite(zero_feature_values_2)]
        if zero_feature_values_2.size > 0:
            threshold_2 = float(np.quantile(zero_feature_values_2, clipped_quantile_2))

    lowconf_scale = float(np.clip(fallback_scale, 0.0, 1.0))
    return {
        "mode": normalized_mode,
        "feature_name": feature_name,
        "feature_name_2": feature_name_2,
        "quantile": clipped_quantile,
        "quantile_2": clipped_quantile_2,
        "threshold": threshold,
        "threshold_2": threshold_2,
        "lowconf_scale": lowconf_scale,
        "max_pred_float": clipped_max_pred_float,
        "num_train_zero_segments": int(zero_feature_values.size),
    }


def apply_selective_post_strategy(
    pred_count_float: np.ndarray,
    val_df: pd.DataFrame,
    strategy: dict,
    pred_count_float_for_protection: np.ndarray | None = None,
) -> tuple[np.ndarray, np.ndarray]:
    pred = np.clip(np.asarray(pred_count_float, dtype=np.float64), a_min=0.0, a_max=None)
    lowconf_mask = np.zeros(pred.shape[0], dtype=bool)
    if strategy.get("mode") not in {"train_zero_quantile_downscale", "train_zero_dual_quantile_downscale"}:
        return pred, lowconf_mask

    threshold = strategy.get("threshold")
    feature_name = str(strategy.get("feature_name"))
    if feature_name not in val_df.columns or not np.isfinite(threshold):
        return pred, lowconf_mask

    feature_values = pd.to_numeric(val_df[feature_name], errors="coerce").to_numpy(dtype=np.float64)
    lowconf_mask = np.isfinite(feature_values) & (feature_values <= float(threshold))
    if strategy.get("mode") == "train_zero_dual_quantile_downscale":
        threshold_2 = strategy.get("threshold_2")
        feature_name_2 = str(strategy.get("feature_name_2"))
        if feature_name_2 not in val_df.columns or not np.isfinite(threshold_2):
            return pred, np.zeros(pred.shape[0], dtype=bool)
        feature_values_2 = pd.to_numeric(val_df[feature_name_2], errors="coerce").to_numpy(dtype=np.float64)
        lowconf_mask = lowconf_mask & np.isfinite(feature_values_2) & (feature_values_2 <= float(threshold_2))
    max_pred_float = normalize_optional_positive_float(strategy.get("max_pred_float", np.nan))
    if np.isfinite(max_pred_float):
        protect_values = pred if pred_count_float_for_protection is None else np.asarray(
            pred_count_float_for_protection,
            dtype=np.float64,
        )
        lowconf_mask = lowconf_mask & np.isfinite(protect_values) & (protect_values <= max_pred_float)
    lowconf_scale = float(strategy.get("lowconf_scale", 1.0))
    if lowconf_scale >= 0.999999:
        return pred, lowconf_mask

    out = pred.copy()
    out[lowconf_mask] = out[lowconf_mask] * lowconf_scale
    return np.clip(out, a_min=0.0, a_max=None), lowconf_mask


def assign_segment_type_labels(
    segment_df: pd.DataFrame,
    strategy: dict,
) -> np.ndarray:
    labels = np.full(len(segment_df), "other", dtype=object)
    mode = validate_type_calibration_mode(str(strategy.get("mode", "none")))
    if mode == "none" or len(segment_df) == 0:
        return labels

    sample_count = pd.to_numeric(segment_df.get("SampleCount"), errors="coerce").to_numpy(dtype=np.float64)
    prob_mean = pd.to_numeric(segment_df.get("ProbMean"), errors="coerce").to_numpy(dtype=np.float64)
    prob_mass_per_length = pd.to_numeric(segment_df.get("ProbMassPerLength"), errors="coerce").to_numpy(dtype=np.float64)

    sample_low = strategy.get("sample_count_low", np.nan)
    sample_high = strategy.get("sample_count_high", np.nan)
    prob_low = strategy.get("prob_mean_low", np.nan)
    mass_low = strategy.get("prob_mass_per_length_low", np.nan)
    mass_high = strategy.get("prob_mass_per_length_high", np.nan)

    low_conf_mask = (
        np.isfinite(prob_mean)
        & np.isfinite(prob_mass_per_length)
        & np.isfinite(prob_low)
        & np.isfinite(mass_low)
        & (prob_mean <= float(prob_low))
        & (prob_mass_per_length <= float(mass_low))
    )
    labels[low_conf_mask] = "low_conf"

    short_strong_mask = (
        (labels == "other")
        & np.isfinite(sample_count)
        & np.isfinite(prob_mass_per_length)
        & np.isfinite(sample_low)
        & np.isfinite(mass_high)
        & (sample_count <= float(sample_low))
        & (prob_mass_per_length >= float(mass_high))
    )
    labels[short_strong_mask] = "short_strong"

    long_weak_mask = (
        (labels == "other")
        & np.isfinite(sample_count)
        & np.isfinite(prob_mass_per_length)
        & np.isfinite(sample_high)
        & np.isfinite(mass_low)
        & (sample_count >= float(sample_high))
        & (prob_mass_per_length <= float(mass_low))
    )
    labels[long_weak_mask] = "long_weak"
    return labels


def fit_type_calibration_strategy(
    train_df: pd.DataFrame,
    train_pred_count_float: np.ndarray,
    mode: str,
    scale_min: float,
    scale_max: float,
) -> dict:
    normalized_mode = validate_type_calibration_mode(mode)
    clipped_scale_min = float(np.clip(float(scale_min), 0.05, 1.0))
    clipped_scale_max = float(max(float(scale_max), clipped_scale_min))
    strategy = {
        "mode": normalized_mode,
        "sample_count_low": np.nan,
        "sample_count_high": np.nan,
        "prob_mean_low": np.nan,
        "prob_mass_per_length_low": np.nan,
        "prob_mass_per_length_high": np.nan,
        "scale_min": clipped_scale_min,
        "scale_max": clipped_scale_max,
        "scales": {
            "low_conf": 1.0,
            "short_strong": 1.0,
            "long_weak": 1.0,
            "other": 1.0,
        },
        "stats": {},
    }
    if normalized_mode == "none" or train_df.empty:
        return strategy

    sample_count = pd.to_numeric(train_df.get("SampleCount"), errors="coerce").to_numpy(dtype=np.float64)
    prob_mean = pd.to_numeric(train_df.get("ProbMean"), errors="coerce").to_numpy(dtype=np.float64)
    prob_mass_per_length = pd.to_numeric(train_df.get("ProbMassPerLength"), errors="coerce").to_numpy(dtype=np.float64)

    finite_sample = sample_count[np.isfinite(sample_count)]
    finite_prob = prob_mean[np.isfinite(prob_mean)]
    finite_mass = prob_mass_per_length[np.isfinite(prob_mass_per_length)]
    if finite_sample.size == 0 or finite_prob.size == 0 or finite_mass.size == 0:
        return strategy

    strategy["sample_count_low"] = float(np.quantile(finite_sample, 0.35))
    strategy["sample_count_high"] = float(np.quantile(finite_sample, 0.65))
    strategy["prob_mean_low"] = float(np.quantile(finite_prob, 0.30))
    strategy["prob_mass_per_length_low"] = float(np.quantile(finite_mass, 0.35))
    strategy["prob_mass_per_length_high"] = float(np.quantile(finite_mass, 0.65))

    labels = assign_segment_type_labels(train_df, strategy)
    pred = np.clip(np.asarray(train_pred_count_float, dtype=np.float64), a_min=0.0, a_max=None)
    target = train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
    for label_name in ["low_conf", "short_strong", "long_weak", "other"]:
        mask = labels == label_name
        n_segments = int(np.sum(mask))
        pred_sum = float(pred[mask].sum()) if n_segments > 0 else 0.0
        target_sum = float(target[mask].sum()) if n_segments > 0 else 0.0
        raw_scale = target_sum / pred_sum if pred_sum > 1e-8 else 1.0
        clipped_scale = float(np.clip(raw_scale, clipped_scale_min, clipped_scale_max))
        shrink = float(n_segments / (n_segments + 10.0)) if n_segments > 0 else 0.0
        blended_scale = float(1.0 + (clipped_scale - 1.0) * shrink)
        strategy["scales"][label_name] = blended_scale
        strategy["stats"][label_name] = {
            "n_segments": n_segments,
            "pred_sum": pred_sum,
            "target_sum": target_sum,
            "raw_scale": raw_scale,
            "clipped_scale": clipped_scale,
            "blended_scale": blended_scale,
        }
    return strategy


def apply_type_calibration(
    pred_count_float: np.ndarray,
    segment_df: pd.DataFrame,
    strategy: dict,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    pred = np.clip(np.asarray(pred_count_float, dtype=np.float64), a_min=0.0, a_max=None)
    labels = assign_segment_type_labels(segment_df, strategy)
    scales = np.ones(pred.shape[0], dtype=np.float64)
    if validate_type_calibration_mode(str(strategy.get("mode", "none"))) == "none":
        return pred, labels, scales

    out = pred.copy()
    scale_map = dict(strategy.get("scales") or {})
    for label_name in ["low_conf", "short_strong", "long_weak", "other"]:
        mask = labels == label_name
        if not np.any(mask):
            continue
        scale_value = float(scale_map.get(label_name, 1.0))
        out[mask] = out[mask] * scale_value
        scales[mask] = scale_value
    return np.clip(out, a_min=0.0, a_max=None), labels, scales


def assign_count_bin_labels(raw_counts: np.ndarray) -> np.ndarray:
    counts = np.clip(np.asarray(raw_counts, dtype=np.float64), a_min=0.0, a_max=None)
    labels = np.full(counts.shape[0], "zero", dtype=object)
    labels[counts >= 1.0] = "one"
    labels[counts >= 2.0] = "two"
    labels[counts >= 3.0] = "three_four"
    labels[counts >= 5.0] = "ge_five"
    return labels


def fit_count_bin_strategy(
    train_df: pd.DataFrame,
    feature_cols: list[str],
    mode: str,
    blend_weight: float,
) -> dict:
    normalized_mode = validate_count_bin_mode(mode)
    clipped_blend_weight = float(np.clip(float(blend_weight), 0.0, 1.0))
    strategy = {
        "mode": normalized_mode,
        "feature_cols": list(feature_cols),
        "blend_weight": clipped_blend_weight,
        "model": None,
        "classes": [],
        "constant_label": "",
        "expected_value_map": dict(COUNT_BIN_DEFAULT_VALUE_MAP),
    }
    if normalized_mode == "none" or train_df.empty:
        return strategy

    raw_counts = train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
    bin_labels = assign_count_bin_labels(raw_counts)
    for label_name in COUNT_BIN_LABEL_ORDER:
        label_mask = bin_labels == label_name
        if np.any(label_mask):
            strategy["expected_value_map"][label_name] = float(np.mean(raw_counts[label_mask]))

    unique_labels = [label for label in COUNT_BIN_LABEL_ORDER if np.any(bin_labels == label)]
    if len(unique_labels) <= 1:
        strategy["constant_label"] = unique_labels[0] if unique_labels else "zero"
        strategy["classes"] = unique_labels
        return strategy

    x_train = train_df[feature_cols].fillna(0.0)
    sample_weight = train_df["TrainSampleWeightUsed"].fillna(1.0).to_numpy(dtype=np.float64)
    classifier = Pipeline(
        [
            ("scaler", StandardScaler()),
            ("model", LogisticRegression(max_iter=2000)),
        ]
    )
    classifier.fit(x_train, bin_labels, model__sample_weight=sample_weight)
    strategy["model"] = classifier
    strategy["classes"] = list(classifier.named_steps["model"].classes_)
    return strategy


def apply_count_bin_strategy(
    pred_count_float: np.ndarray,
    segment_df: pd.DataFrame,
    strategy: dict,
    gate_label: np.ndarray | None = None,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    pred = np.clip(np.asarray(pred_count_float, dtype=np.float64), a_min=0.0, a_max=None)
    expected = pred.copy()
    pred_labels = np.full(pred.shape[0], "", dtype=object)
    pred_conf = np.full(pred.shape[0], np.nan, dtype=np.float64)
    if validate_count_bin_mode(str(strategy.get("mode", "none"))) == "none" or len(segment_df) == 0:
        return pred, expected, pred_labels, pred_conf

    expected_value_map = {
        str(key): float(value)
        for key, value in dict(strategy.get("expected_value_map") or COUNT_BIN_DEFAULT_VALUE_MAP).items()
    }
    classifier = strategy.get("model")
    if classifier is None:
        constant_label = str(strategy.get("constant_label") or "zero")
        expected_value = float(expected_value_map.get(constant_label, COUNT_BIN_DEFAULT_VALUE_MAP["zero"]))
        expected = np.full(pred.shape[0], expected_value, dtype=np.float64)
        pred_labels[:] = constant_label
        pred_conf[:] = 1.0
    else:
        x_val = segment_df[list(strategy.get("feature_cols", []))].fillna(0.0)
        prob = classifier.predict_proba(x_val)
        classes = [str(item) for item in classifier.named_steps["model"].classes_]
        expected = np.zeros(pred.shape[0], dtype=np.float64)
        for idx, class_name in enumerate(classes):
            expected += prob[:, idx] * float(expected_value_map.get(class_name, COUNT_BIN_DEFAULT_VALUE_MAP["zero"]))
        best_idx = np.argmax(prob, axis=1)
        pred_labels = np.asarray([classes[idx] for idx in best_idx], dtype=object)
        pred_conf = np.max(prob, axis=1).astype(np.float64)

    # If the dominant count bin is still zero, do not let the blended expected value
    # re-inflate the segment count purely from soft probability mass.
    zero_label_mask = pred_labels == "zero"
    if np.any(zero_label_mask):
        expected = expected.copy()
        expected[zero_label_mask] = 0.0

    blended = (
        float(np.clip(float(strategy.get("blend_weight", 0.7)), 0.0, 1.0)) * pred
        + (1.0 - float(np.clip(float(strategy.get("blend_weight", 0.7)), 0.0, 1.0))) * expected
    )
    if gate_label is not None:
        gate_mask = np.asarray(gate_label, dtype=np.int64) > 0
        blended = np.where(gate_mask, blended, 0.0)
        expected = np.where(gate_mask, expected, 0.0)
    return np.clip(blended, a_min=0.0, a_max=None), expected, pred_labels, pred_conf


def resolve_density_strength_target_col(mode: str) -> str:
    normalized_mode = validate_density_strength_mode(mode)
    if normalized_mode in {"segment_p10_mean", "segment_p10_mean_regression"}:
        return "RawP10MeanInPredSegment"
    if normalized_mode in {"segment_p10_mass_per_length", "segment_p10_mass_per_length_regression"}:
        return "RawP10MassPerLength"
    return ""


def fit_density_strength_strategy(
    train_df: pd.DataFrame,
    mode: str,
    basis_col: str,
    scale_min: float,
    scale_max: float,
    feature_cols: list[str],
    hgb_learning_rate: float,
    hgb_max_depth: int,
    model_max_iter: int,
    regression_train_mode: str,
    regression_target_transform: str,
    regression_clip_quantile: float,
    regression_calibration_mode: str,
    regression_calibration_quantiles: str,
) -> dict:
    normalized_mode = validate_density_strength_mode(mode)
    normalized_regression_train_mode = validate_density_strength_regression_train_mode(regression_train_mode)
    normalized_regression_target_transform = validate_density_strength_regression_target_transform(
        regression_target_transform
    )
    normalized_regression_calibration_mode = validate_density_strength_regression_calibration_mode(
        regression_calibration_mode
    )
    normalized_regression_calibration_quantiles = (
        normalize_density_strength_regression_calibration_quantiles(regression_calibration_quantiles)
    )
    target_col = resolve_density_strength_target_col(normalized_mode)
    clipped_scale_min = float(np.clip(float(scale_min), 0.05, 10.0))
    clipped_scale_max = float(max(float(scale_max), clipped_scale_min))
    strategy = {
        "mode": normalized_mode,
        "target_col": target_col,
        "basis_col": str(basis_col or ""),
        "raw_scale": np.nan,
        "applied_scale": 1.0,
        "type_strategy": {"mode": "none"},
        "type_scale_map": {
            "low_conf": 1.0,
            "short_strong": 1.0,
            "long_weak": 1.0,
            "other": 1.0,
        },
        "positive_quantiles": {
            "q50": np.nan,
            "q80": np.nan,
            "q95": np.nan,
        },
        "feature_cols": [],
        "model": None,
        "constant_value": np.nan,
        "regression_model_name": "",
        "regression_train_mode_requested": normalized_regression_train_mode,
        "regression_train_mode_effective": "all_segments",
        "regression_target_transform": normalized_regression_target_transform,
        "regression_clip_quantile": np.nan,
        "regression_clip_value": np.nan,
        "regression_calibration_mode": normalized_regression_calibration_mode,
        "regression_calibration_quantiles": normalized_regression_calibration_quantiles,
        "regression_calibration_info": {"mode": "none", "quantiles": [], "pred_knots": [], "target_knots": []},
        "mean_over_mass_per_length_ratio": 1.0,
        "prediction_cap": np.nan,
    }
    if normalized_mode == "none" or train_df.empty:
        return strategy
    if not basis_col:
        raise ValueError("density_strength_basis_col is required when density_strength_mode is enabled")
    if basis_col not in train_df.columns:
        raise ValueError(f"Missing density_strength_basis_col in train_df: {basis_col}")
    if target_col not in train_df.columns:
        raise ValueError(f"Missing density strength target col in train_df: {target_col}")

    basis = train_df[basis_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
    target = train_df[target_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
    if {"RawP10MeanInPredSegment", "RawP10MassPerLength"}.issubset(train_df.columns):
        p10_mean = train_df["RawP10MeanInPredSegment"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
        p10_mass_per_length = (
            train_df["RawP10MassPerLength"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
        )
        ratio_mask = np.isfinite(p10_mean) & np.isfinite(p10_mass_per_length) & (p10_mass_per_length > 1e-8)
        if np.any(ratio_mask):
            strategy["mean_over_mass_per_length_ratio"] = float(
                np.median(p10_mean[ratio_mask] / p10_mass_per_length[ratio_mask])
            )

    target_for_fit = target.copy()
    positive_target_all = target[target > 0.0]
    regression_clip_quantile_value = float(regression_clip_quantile)
    if (
        density_strength_uses_regression(normalized_mode)
        and np.isfinite(regression_clip_quantile_value)
        and 0.0 < regression_clip_quantile_value < 1.0
        and positive_target_all.size > 0
    ):
        clip_value = float(np.quantile(positive_target_all, regression_clip_quantile_value))
        target_for_fit = np.clip(target_for_fit, a_min=0.0, a_max=clip_value)
        strategy["regression_clip_quantile"] = regression_clip_quantile_value
        strategy["regression_clip_value"] = clip_value

    if density_strength_uses_regression(normalized_mode):
        regression_feature_cols = resolve_strategy_feature_cols(
            dataset_df=train_df,
            base_feature_cols=feature_cols,
            extra_feature_cols=[
                basis_col,
                "SegLength",
                "SampleCount",
                "ProbMean",
                "ProbMax",
                "ProbStd",
                "ProbMass",
                "ProbMassPerLength",
                "HighProbLen_05",
                "HighProbLen_06",
                "HighProbLen_08",
                "HighProbFrac_05",
                "ProbExcessMass_06",
                "ProbExcessMass_08",
                "LengthXProbMean",
                "AC_MeanInPredSegment",
                "AC_StdInPredSegment",
                "GR_MeanInPredSegment",
                "GR_StdInPredSegment",
            ],
            required_feature_cols=[basis_col],
        )
        if not regression_feature_cols:
            raise ValueError("No usable feature cols resolved for density regression")
        strategy["feature_cols"] = regression_feature_cols
        strategy["regression_model_name"] = "tweedie_linear"
        fit_mask = np.isfinite(target_for_fit)
        if normalized_regression_train_mode == "positive_only":
            positive_mask = fit_mask & (target_for_fit > 0.0)
            if int(np.sum(positive_mask)) >= 8 and np.unique(np.round(target_for_fit[positive_mask], 8)).size > 1:
                fit_mask = positive_mask
                strategy["regression_train_mode_effective"] = "positive_only"
        fit_target_raw = target_for_fit[fit_mask]
        fit_x = train_df.loc[fit_mask, regression_feature_cols].fillna(0.0)
        fit_sample_weight = train_df.loc[fit_mask, "TrainSampleWeightUsed"].fillna(1.0).to_numpy(dtype=np.float64)
        unique_target = np.unique(np.round(fit_target_raw, 8))
        if unique_target.size <= 1:
            strategy["constant_value"] = float(np.mean(fit_target_raw)) if fit_target_raw.size > 0 else 0.0
        else:
            regressor = build_continuous_regressor(
                learning_rate=hgb_learning_rate,
                max_depth=hgb_max_depth,
                max_iter=model_max_iter,
            )
            fit_target_model = apply_density_strength_target_transform(
                fit_target_raw,
                normalized_regression_target_transform,
            )
            regressor.fit(
                fit_x,
                fit_target_model,
                model__sample_weight=fit_sample_weight,
            )
            strategy["model"] = regressor
            if normalized_regression_calibration_mode != "none":
                fit_pred_raw = invert_density_strength_target_transform(
                    regressor.predict(fit_x),
                    normalized_regression_target_transform,
                )
                calibration_info = build_density_strength_quantile_map(
                    pred_raw=fit_pred_raw,
                    target_raw=fit_target_raw,
                    quantiles=normalized_regression_calibration_quantiles,
                )
                strategy["regression_calibration_info"] = calibration_info
        strategy["raw_scale"] = np.nan
        strategy["applied_scale"] = np.nan
    else:
        basis_sum = float(np.sum(basis))
        target_sum = float(np.sum(target))
        raw_scale = (target_sum / basis_sum) if basis_sum > 1e-8 else 1.0
        applied_scale = float(np.clip(raw_scale, clipped_scale_min, clipped_scale_max))
        strategy["raw_scale"] = raw_scale
        strategy["applied_scale"] = applied_scale

        base_pred = np.clip(basis * applied_scale, a_min=0.0, a_max=None)
        type_strategy = fit_type_calibration_strategy(
            train_df=train_df,
            train_pred_count_float=base_pred,
            mode="rule_v1",
            scale_min=0.5,
            scale_max=1.8,
        )
        type_labels = assign_segment_type_labels(train_df, type_strategy)
        type_scale_map = dict(strategy["type_scale_map"])
        for label_name in ["low_conf", "short_strong", "long_weak", "other"]:
            mask = type_labels == label_name
            pred_sum = float(base_pred[mask].sum()) if np.any(mask) else 0.0
            target_sum_label = float(target[mask].sum()) if np.any(mask) else 0.0
            raw_label_scale = (target_sum_label / pred_sum) if pred_sum > 1e-8 else 1.0
            type_scale_map[label_name] = float(np.clip(raw_label_scale, 0.5, 1.8))
        strategy["type_strategy"] = type_strategy
        strategy["type_scale_map"] = type_scale_map

    positive_target = target_for_fit[target_for_fit > 0.0]
    if positive_target.size > 0:
        strategy["positive_quantiles"] = {
            "q50": float(np.quantile(positive_target, 0.50)),
            "q80": float(np.quantile(positive_target, 0.80)),
            "q95": float(np.quantile(positive_target, 0.95)),
        }
        strategy["prediction_cap"] = float(
            max(
                np.quantile(positive_target, 0.50),
                np.quantile(positive_target, 0.95) * 2.0,
            )
        )
    elif target.size > 0:
        strategy["prediction_cap"] = float(np.max(target))
    return strategy


def assign_density_strength_level(pred_density_strength: np.ndarray, quantiles: dict) -> np.ndarray:
    pred = np.clip(np.asarray(pred_density_strength, dtype=np.float64), a_min=0.0, a_max=None)
    levels = np.full(pred.shape[0], "zero", dtype=object)
    q50 = float(quantiles.get("q50", np.nan))
    q80 = float(quantiles.get("q80", np.nan))
    q95 = float(quantiles.get("q95", np.nan))
    positive_mask = pred > 0.0
    levels[positive_mask] = "low"
    if np.isfinite(q50):
        levels[pred >= q50] = "medium"
    if np.isfinite(q80):
        levels[pred >= q80] = "high"
    if np.isfinite(q95):
        levels[pred >= q95] = "very_high"
    return levels


def apply_density_strength_strategy(segment_df: pd.DataFrame, strategy: dict) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    normalized_mode = validate_density_strength_mode(str(strategy.get("mode", "none")))
    pred = np.full(len(segment_df), np.nan, dtype=np.float64)
    size_scale = np.full(len(segment_df), np.nan, dtype=np.float64)
    labels = np.full(len(segment_df), "", dtype=object)
    type_scales = np.ones(len(segment_df), dtype=np.float64)
    if normalized_mode == "none" or len(segment_df) == 0:
        return pred, size_scale, labels, type_scales

    if density_strength_uses_regression(normalized_mode):
        regression_feature_cols = list(strategy.get("feature_cols") or [])
        validate_feature_cols(regression_feature_cols, segment_df)
        regressor = strategy.get("model")
        if regressor is None:
            pred = np.full(len(segment_df), float(strategy.get("constant_value", 0.0)), dtype=np.float64)
        else:
            pred = invert_density_strength_target_transform(
                regressor.predict(segment_df[regression_feature_cols].fillna(0.0)),
                str(strategy.get("regression_target_transform", "none")),
            )
        pred = np.clip(pred, a_min=0.0, a_max=None)
        pred = apply_density_strength_quantile_map(
            pred,
            dict(strategy.get("regression_calibration_info") or {"mode": "none"}),
        )
    else:
        basis_col = str(strategy.get("basis_col") or "")
        if basis_col not in segment_df.columns:
            raise ValueError(f"Missing density strength basis col in segment_df: {basis_col}")
        basis = segment_df[basis_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64)
        pred = np.clip(basis * float(strategy.get("applied_scale", 1.0)), a_min=0.0, a_max=None)
        type_strategy = dict(strategy.get("type_strategy") or {"mode": "none"})
        labels = assign_segment_type_labels(segment_df, type_strategy)
        scale_map = dict(strategy.get("type_scale_map") or {})
        for label_name in ["low_conf", "short_strong", "long_weak", "other"]:
            mask = labels == label_name
            if not np.any(mask):
                continue
            scale_value = float(scale_map.get(label_name, 1.0))
            pred[mask] = pred[mask] * scale_value
            type_scales[mask] = scale_value

    prediction_cap = float(strategy.get("prediction_cap", np.nan))
    if np.isfinite(prediction_cap) and prediction_cap >= 0.0:
        pred = np.clip(pred, a_min=0.0, a_max=prediction_cap)
    q50 = float(dict(strategy.get("positive_quantiles") or {}).get("q50", np.nan))
    if np.isfinite(q50) and q50 > 1e-8:
        size_scale = np.clip(pred / q50, a_min=0.0, a_max=10.0)
    labels = assign_density_strength_level(pred, dict(strategy.get("positive_quantiles") or {}))
    return pred, size_scale, labels, type_scales


def fit_orientation_strategy(
    train_df: pd.DataFrame,
    feature_cols: list[str],
    mode: str,
    num_families: int,
    min_points_per_segment: int,
) -> dict:
    normalized_mode = validate_orientation_mode(mode)
    clipped_num_families = int(np.clip(int(num_families), 1, 12))
    clipped_min_points = int(max(int(min_points_per_segment), 1))
    strategy = {
        "mode": normalized_mode,
        "feature_cols": [],
        "classifier": None,
        "constant_family": "",
        "family_centers": {},
        "global_spread": {
            "azimuth_std_deg": np.nan,
            "azimuth_q80_deg": np.nan,
            "dip_std_deg": np.nan,
            "dip_q80_deg": np.nan,
        },
        "family_spreads": {},
        "point_spread_mode": "segment_constant" if normalized_mode == "none" else "family_confidence_rank_template_v1",
        "point_spread_conf_scale_min": 0.45,
        "point_spread_conf_scale_max": 1.0,
        "num_families_requested": clipped_num_families,
        "num_families_fitted": 0,
        "min_points_per_segment": clipped_min_points,
        "num_train_segments": 0,
    }
    if normalized_mode == "none" or train_df.empty:
        return strategy

    orientation_df = train_df[
        train_df["RawOrientationPointCountInPredSegment"].fillna(0.0).ge(clipped_min_points)
        & train_df["RawAzimuthMeanInPredSegment"].notna()
        & train_df["RawDipMeanInPredSegment"].notna()
    ].copy()
    if orientation_df.empty:
        return strategy
    strategy["num_train_segments"] = int(len(orientation_df))
    azimuth_spread_summary = summarize_orientation_spread_values(
        orientation_df["RawOrientationAzimuthStdInPredSegment"].to_numpy(dtype=np.float64)
    )
    dip_spread_summary = summarize_orientation_spread_values(
        orientation_df.get("RawOrientationDipStdInPredSegment", pd.Series(dtype=np.float64)).to_numpy(dtype=np.float64)
    )
    strategy["global_spread"] = {
        "azimuth_std_deg": azimuth_spread_summary["median"],
        "azimuth_q80_deg": azimuth_spread_summary["q80"],
        "dip_std_deg": dip_spread_summary["median"],
        "dip_q80_deg": dip_spread_summary["q80"],
    }

    azimuth = orientation_df["RawAzimuthMeanInPredSegment"].to_numpy(dtype=np.float64)
    dip = orientation_df["RawDipMeanInPredSegment"].to_numpy(dtype=np.float64)
    cluster_features = build_orientation_cluster_features(azimuth, dip)
    num_clusters = int(min(clipped_num_families, len(orientation_df)))
    if num_clusters <= 1:
        raw_cluster_labels = np.zeros(len(orientation_df), dtype=np.int64)
    else:
        kmeans = KMeans(n_clusters=num_clusters, n_init=10, random_state=42)
        raw_cluster_labels = kmeans.fit_predict(cluster_features)

    center_rows = []
    for raw_label in sorted(pd.unique(raw_cluster_labels)):
        label_mask = raw_cluster_labels == raw_label
        center_rows.append(
            {
                "raw_label": int(raw_label),
                "azimuth_deg": circular_mean_deg(azimuth[label_mask]),
                "dip_deg": float(np.mean(dip[label_mask])),
                "sample_count": int(np.sum(label_mask.astype(np.int64))),
            }
        )
    center_rows = sorted(
        center_rows,
        key=lambda item: (
            float(item.get("azimuth_deg", np.nan))
            if np.isfinite(item.get("azimuth_deg", np.nan))
            else 1e9,
            float(item.get("dip_deg", np.nan))
            if np.isfinite(item.get("dip_deg", np.nan))
            else 1e9,
        ),
    )

    label_map = {}
    family_centers = {}
    for family_idx, center in enumerate(center_rows, start=1):
        family_name = f"family_{family_idx}"
        label_map[int(center["raw_label"])] = family_name
        family_centers[family_name] = {
            "azimuth_deg": float(center["azimuth_deg"]) if np.isfinite(center["azimuth_deg"]) else np.nan,
            "dip_deg": float(center["dip_deg"]) if np.isfinite(center["dip_deg"]) else np.nan,
            "sample_count": int(center["sample_count"]),
        }
    strategy["family_centers"] = family_centers
    strategy["num_families_fitted"] = int(len(family_centers))

    orientation_labels = np.asarray([label_map[int(label)] for label in raw_cluster_labels], dtype=object)
    orientation_df = orientation_df.reset_index(drop=True)
    orientation_df["OrientationFamilyLabel"] = orientation_labels
    family_spreads = {}
    global_azimuth_std_values = []
    global_azimuth_q80_values = []
    global_dip_std_values = []
    global_dip_q80_values = []
    for family_name, family_df in orientation_df.groupby("OrientationFamilyLabel", dropna=False):
        family_azimuth_within_spread = summarize_orientation_spread_values(
            family_df["RawOrientationAzimuthStdInPredSegment"].to_numpy(dtype=np.float64)
        )
        family_dip_within_spread = summarize_orientation_spread_values(
            family_df.get("RawOrientationDipStdInPredSegment", pd.Series(dtype=np.float64)).to_numpy(dtype=np.float64)
        )
        family_center = dict(family_centers.get(str(family_name)) or {})
        family_center_azimuth = first_finite_float(family_center.get("azimuth_deg"), default=np.nan)
        family_center_dip = first_finite_float(family_center.get("dip_deg"), default=np.nan)
        if np.isfinite(family_center_azimuth):
            family_segment_azimuth_diff = np.asarray(
                [
                    angular_abs_diff_deg(value, family_center_azimuth)
                    for value in family_df["RawAzimuthMeanInPredSegment"].to_numpy(dtype=np.float64)
                ],
                dtype=np.float64,
            )
        else:
            family_segment_azimuth_diff = np.asarray([], dtype=np.float64)
        if np.isfinite(family_center_dip):
            family_segment_dip_diff = np.abs(
                family_df["RawDipMeanInPredSegment"].to_numpy(dtype=np.float64) - family_center_dip
            )
        else:
            family_segment_dip_diff = np.asarray([], dtype=np.float64)
        family_azimuth_between_spread = summarize_orientation_spread_values(family_segment_azimuth_diff)
        family_dip_between_spread = summarize_orientation_spread_values(family_segment_dip_diff)
        family_azimuth_std = max(
            first_finite_float(family_azimuth_within_spread["median"], default=0.0),
            first_finite_float(family_azimuth_between_spread["median"], default=0.0),
        )
        family_azimuth_q80 = max(
            first_finite_float(family_azimuth_within_spread["q80"], default=0.0),
            first_finite_float(family_azimuth_between_spread["q80"], default=0.0),
        )
        family_dip_std = max(
            first_finite_float(family_dip_within_spread["median"], default=0.0),
            first_finite_float(family_dip_between_spread["median"], default=0.0),
        )
        family_dip_q80 = max(
            first_finite_float(family_dip_within_spread["q80"], default=0.0),
            first_finite_float(family_dip_between_spread["q80"], default=0.0),
        )
        family_spreads[str(family_name)] = {
            "azimuth_std_deg": family_azimuth_std,
            "azimuth_q80_deg": family_azimuth_q80,
            "dip_std_deg": family_dip_std,
            "dip_q80_deg": family_dip_q80,
            "segment_count": int(len(family_df)),
        }
        global_azimuth_std_values.append(family_azimuth_std)
        global_azimuth_q80_values.append(family_azimuth_q80)
        global_dip_std_values.append(family_dip_std)
        global_dip_q80_values.append(family_dip_q80)
    strategy["family_spreads"] = family_spreads
    strategy["global_spread"] = {
        "azimuth_std_deg": max(
            first_finite_float(strategy["global_spread"].get("azimuth_std_deg"), default=0.0),
            first_finite_float(np.nanmedian(np.asarray(global_azimuth_std_values, dtype=np.float64)), default=0.0),
        ),
        "azimuth_q80_deg": max(
            first_finite_float(strategy["global_spread"].get("azimuth_q80_deg"), default=0.0),
            first_finite_float(np.nanmedian(np.asarray(global_azimuth_q80_values, dtype=np.float64)), default=0.0),
        ),
        "dip_std_deg": max(
            first_finite_float(strategy["global_spread"].get("dip_std_deg"), default=0.0),
            first_finite_float(np.nanmedian(np.asarray(global_dip_std_values, dtype=np.float64)), default=0.0),
        ),
        "dip_q80_deg": max(
            first_finite_float(strategy["global_spread"].get("dip_q80_deg"), default=0.0),
            first_finite_float(np.nanmedian(np.asarray(global_dip_q80_values, dtype=np.float64)), default=0.0),
        ),
    }
    unique_labels = dedupe_preserve_order(list(orientation_labels))
    if len(unique_labels) <= 1:
        strategy["constant_family"] = unique_labels[0] if unique_labels else ""
        return strategy

    orientation_feature_cols = resolve_strategy_feature_cols(
        dataset_df=train_df,
        base_feature_cols=feature_cols,
        extra_feature_cols=[
            "SegLength",
            "SampleCount",
            "ProbMean",
            "ProbMax",
            "ProbStd",
            "ProbMass",
            "ProbMassPerLength",
            "HighProbLen_05",
            "HighProbLen_06",
            "HighProbLen_08",
            "HighProbFrac_05",
            "ProbExcessMass_06",
            "ProbExcessMass_08",
            "LengthXProbMean",
            "AC_MeanInPredSegment",
            "AC_StdInPredSegment",
            "GR_MeanInPredSegment",
            "GR_StdInPredSegment",
        ],
    )
    if not orientation_feature_cols:
        strategy["constant_family"] = unique_labels[0]
        return strategy

    classifier = Pipeline(
        [
            ("scaler", StandardScaler()),
            (
                "model",
                LogisticRegression(
                    max_iter=2000,
                    class_weight="balanced",
                    random_state=42,
                ),
            ),
        ]
    )
    sample_weight = (
        orientation_df["TrainSampleWeightUsed"].fillna(1.0).to_numpy(dtype=np.float64)
        * np.clip(
            orientation_df["RawOrientationPointCountInPredSegment"].fillna(1.0).to_numpy(dtype=np.float64),
            a_min=1.0,
            a_max=None,
        )
    )
    classifier.fit(
        orientation_df[orientation_feature_cols].fillna(0.0),
        orientation_labels,
        model__sample_weight=sample_weight,
    )
    strategy["feature_cols"] = orientation_feature_cols
    strategy["classifier"] = classifier
    return strategy


def apply_orientation_strategy(
    segment_df: pd.DataFrame,
    strategy: dict,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    normalized_mode = validate_orientation_mode(str(strategy.get("mode", "none")))
    family_labels = np.full(len(segment_df), "", dtype=object)
    confidence = np.full(len(segment_df), np.nan, dtype=np.float64)
    pred_azimuth = np.full(len(segment_df), np.nan, dtype=np.float64)
    pred_dip = np.full(len(segment_df), np.nan, dtype=np.float64)
    if normalized_mode == "none" or len(segment_df) == 0:
        return family_labels, confidence, pred_azimuth, pred_dip

    family_centers = {
        str(name): dict(center)
        for name, center in dict(strategy.get("family_centers") or {}).items()
    }
    if not family_centers:
        return family_labels, confidence, pred_azimuth, pred_dip

    classifier = strategy.get("classifier")
    if classifier is None:
        constant_family = str(strategy.get("constant_family") or next(iter(family_centers.keys()), ""))
        family_labels[:] = constant_family
        confidence[:] = 1.0 if constant_family else np.nan
    else:
        orientation_feature_cols = list(strategy.get("feature_cols") or [])
        validate_feature_cols(orientation_feature_cols, segment_df)
        prob = classifier.predict_proba(segment_df[orientation_feature_cols].fillna(0.0))
        classes = [str(item) for item in classifier.named_steps["model"].classes_]
        best_idx = np.argmax(prob, axis=1)
        family_labels = np.asarray([classes[idx] for idx in best_idx], dtype=object)
        confidence = np.max(prob, axis=1).astype(np.float64)

    for idx, family_name in enumerate(family_labels):
        center = family_centers.get(str(family_name))
        if not center:
            continue
        pred_azimuth[idx] = float(center.get("azimuth_deg", np.nan))
        pred_dip[idx] = float(center.get("dip_deg", np.nan))
    return family_labels, confidence, pred_azimuth, pred_dip


def build_regressor(
    model_name: str,
    alpha: float,
    max_iter: int,
    tweedie_power: float,
    xgb_n_estimators: int,
    xgb_max_depth: int,
    xgb_learning_rate: float,
    xgb_subsample: float,
    xgb_colsample_bytree: float,
    hgb_learning_rate: float,
    hgb_max_depth: int,
):
    if model_name == "poisson":
        model = PoissonRegressor(alpha=alpha, max_iter=max_iter)
        return Pipeline([("scaler", StandardScaler()), ("model", model)])
    if model_name == "tweedie":
        model = TweedieRegressor(power=tweedie_power, alpha=alpha, max_iter=max_iter, link="log")
        return Pipeline([("scaler", StandardScaler()), ("model", model)])
    if model_name == "hgbt_poisson":
        model = HistGradientBoostingRegressor(
            loss="poisson",
            learning_rate=hgb_learning_rate,
            max_iter=max_iter,
            max_depth=hgb_max_depth,
            min_samples_leaf=5,
            l2_regularization=max(float(alpha), 0.0),
            random_state=42,
        )
        return Pipeline([("model", model)])
    if model_name == "xgboost":
        if XGBRegressor is None:
            raise ImportError("xgboost is not installed in current environment")
        model = XGBRegressor(
            objective="count:poisson",
            n_estimators=xgb_n_estimators,
            max_depth=xgb_max_depth,
            learning_rate=xgb_learning_rate,
            subsample=xgb_subsample,
            colsample_bytree=xgb_colsample_bytree,
            reg_lambda=max(float(alpha), 0.0),
            random_state=42,
        )
        return Pipeline([("model", model)])
    raise ValueError(f"Unsupported model_name: {model_name}")


def fit_regressor(regressor: Pipeline, x_train: pd.DataFrame, y_train: pd.Series, sample_weight: np.ndarray | None):
    if sample_weight is None:
        regressor.fit(x_train, y_train)
    else:
        regressor.fit(x_train, y_train, model__sample_weight=sample_weight)
    return regressor


def select_count_train_df(train_df: pd.DataFrame, count_train_mode: str) -> tuple[pd.DataFrame, str]:
    normalized_mode = validate_count_train_mode(count_train_mode)
    if normalized_mode == "positive_only":
        positive_df = train_df[
            train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0) > 0.0
        ].copy()
        if not positive_df.empty:
            return positive_df, "positive_only"
    return train_df.copy(), "all_segments"


def fit_gate_classifier(
    train_df: pd.DataFrame,
    feature_cols: list[str],
    gate_mode: str,
    gate_prob_threshold: float,
) -> tuple[dict, list[dict]]:
    normalized_gate_mode = validate_gate_mode(gate_mode)
    clipped_gate_prob_threshold = float(np.clip(float(gate_prob_threshold), 0.0, 1.0))
    gate_info = {
        "mode": normalized_gate_mode,
        "prob_threshold": clipped_gate_prob_threshold,
        "feature_cols": list(feature_cols),
        "model": None,
        "constant_prob": np.nan,
        "positive_class": 1,
    }
    coef_rows: list[dict] = []
    if normalized_gate_mode == "none":
        return gate_info, coef_rows

    x_train = train_df[feature_cols].fillna(0.0)
    y_gate = train_df["HasRawPoint"].fillna(0.0).astype(np.int64).clip(lower=0, upper=1)
    sample_weight = train_df["TrainSampleWeightUsed"].fillna(1.0).to_numpy(dtype=np.float64)
    unique_gate_values = sorted(pd.unique(y_gate))
    if len(unique_gate_values) <= 1:
        gate_info["constant_prob"] = float(unique_gate_values[0]) if unique_gate_values else 0.0
        return gate_info, coef_rows

    gate_model = Pipeline(
        [
            ("scaler", StandardScaler()),
            (
                "model",
                LogisticRegression(
                    max_iter=1000,
                    class_weight="balanced",
                    random_state=42,
                ),
            ),
        ]
    )
    gate_model.fit(x_train, y_gate, model__sample_weight=sample_weight)
    gate_info["model"] = gate_model
    gate_info["positive_class"] = 1

    model = gate_model.named_steps["model"]
    if hasattr(model, "coef_"):
        for feature_name, coef in zip(feature_cols, model.coef_[0]):
            coef_rows.append({"val_well": "__gate__", "feature": feature_name, "coef": float(coef)})
    if hasattr(model, "intercept_"):
        coef_rows.append({"val_well": "__gate__", "feature": "__intercept__", "coef": float(model.intercept_[0])})
    return gate_info, coef_rows


def apply_gate_to_count_predictions(
    pred_count_float: np.ndarray,
    segment_df: pd.DataFrame,
    gate_info: dict,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    pred = np.clip(np.asarray(pred_count_float, dtype=np.float64), a_min=0.0, a_max=None)
    gate_prob = np.ones(pred.shape[0], dtype=np.float64)
    gate_label = np.ones(pred.shape[0], dtype=np.int64)
    if validate_gate_mode(str(gate_info.get("mode", "none"))) == "none":
        return pred, gate_prob, gate_label

    clipped_threshold = float(np.clip(float(gate_info.get("prob_threshold", 0.5)), 0.0, 1.0))
    constant_prob = gate_info.get("constant_prob", np.nan)
    if np.isfinite(constant_prob):
        gate_prob = np.full(pred.shape[0], float(constant_prob), dtype=np.float64)
    else:
        gate_model = gate_info.get("model")
        if gate_model is None:
            raise ValueError("gate_mode is enabled but gate model is missing")
        feature_cols = list(gate_info.get("feature_cols") or [])
        validate_feature_cols(feature_cols, segment_df)
        x_val = segment_df[feature_cols].fillna(0.0)
        positive_class = int(gate_info.get("positive_class", 1))
        positive_idx = int(np.where(gate_model.named_steps["model"].classes_ == positive_class)[0][0])
        gate_prob = gate_model.predict_proba(x_val)[:, positive_idx]

    gate_label = (gate_prob >= clipped_threshold).astype(np.int64)
    gated_pred = pred.copy()
    gate_mode = validate_gate_mode(str(gate_info.get("mode", "none")))
    if gate_mode == "logistic_hard":
        gated_pred[gate_label <= 0] = 0.0
    elif gate_mode == "logistic_soft":
        gated_pred = gated_pred * np.clip(gate_prob, a_min=0.0, a_max=1.0)
    return gated_pred, gate_prob, gate_label


def compute_pred_count(raw_value: float, pred_min_points_per_segment: int, rounding_mode: str) -> int:
    if not np.isfinite(raw_value):
        return 0
    max_safe_value = float(np.iinfo(np.int64).max // 4)
    safe_value = float(min(max(raw_value, 0.0), max_safe_value))
    return int(
        dataset_builder.POINT.compute_n_points(
            safe_value,
            int(pred_min_points_per_segment),
            rounding_mode,
        )
    )


def apply_pred_point_count_sample_cap(pred_segment_df: pd.DataFrame) -> pd.DataFrame:
    if pred_segment_df.empty or "PredPointCount" not in pred_segment_df.columns:
        return pred_segment_df

    out = pred_segment_df.copy()
    pred_count = pd.to_numeric(out["PredPointCount"], errors="coerce").fillna(0).astype(np.int64)
    if "SampleCount" in out.columns:
        sample_cap_series = pd.to_numeric(out["SampleCount"], errors="coerce")
    else:
        sample_cap_series = pd.Series(np.nan, index=out.index, dtype=np.float64)

    sample_cap_float = sample_cap_series.to_numpy(dtype=np.float64, copy=False)
    finite_cap_mask = np.isfinite(sample_cap_float)
    sample_cap_int = np.full(len(out), np.iinfo(np.int64).max, dtype=np.int64)
    if finite_cap_mask.any():
        sample_cap_int[finite_cap_mask] = np.maximum(sample_cap_float[finite_cap_mask], 0.0).astype(np.int64)

    pred_count_arr = pred_count.to_numpy(dtype=np.int64, copy=False)
    capped_count = np.minimum(pred_count_arr, sample_cap_int)
    cap_applied = capped_count != pred_count_arr

    out["PredPointCountBeforeSampleCap"] = pred_count_arr
    out["PredPointCountSampleCap"] = sample_cap_series
    out["PredPointCountSampleCapApplied"] = cap_applied.astype(int)
    out["PredPointCount"] = capped_count

    if "PredPointCountFloat" in out.columns:
        pred_count_float = pd.to_numeric(out["PredPointCountFloat"], errors="coerce")
        pred_count_float_arr = pred_count_float.to_numpy(dtype=np.float64, copy=False)
        capped_float = pred_count_float_arr.copy()
        if finite_cap_mask.any():
            capped_float[finite_cap_mask] = np.minimum(
                pred_count_float_arr[finite_cap_mask],
                sample_cap_float[finite_cap_mask],
            )
        out["PredPointCountFloatBeforeSampleCap"] = pred_count_float
        out["PredPointCountFloat"] = capped_float

    return out


def compute_point_support_intervals(
    seg_start_depth: float,
    seg_end_depth: float,
    picked_depths: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    picked = np.asarray(picked_depths, dtype=np.float64)
    if picked.size == 0:
        empty = np.array([], dtype=np.float64)
        return empty, empty, empty

    seg_start = float(min(seg_start_depth, seg_end_depth))
    seg_end = float(max(seg_start_depth, seg_end_depth))
    clipped = np.clip(picked, a_min=seg_start, a_max=seg_end)
    left_bounds = np.empty(clipped.size, dtype=np.float64)
    right_bounds = np.empty(clipped.size, dtype=np.float64)
    if clipped.size == 1:
        left_bounds[0] = seg_start
        right_bounds[0] = seg_end
    else:
        midpoints = 0.5 * (clipped[:-1] + clipped[1:])
        left_bounds[0] = seg_start
        left_bounds[1:] = midpoints
        right_bounds[:-1] = midpoints
        right_bounds[-1] = seg_end
    support_lengths = np.clip(right_bounds - left_bounds, a_min=0.0, a_max=None)
    return left_bounds, right_bounds, support_lengths


def allocate_segment_density_to_points(
    pred_info: dict,
    payload: dict,
    picked_depths: np.ndarray,
    picked_rel_idx: np.ndarray,
) -> dict[str, np.ndarray | float | str]:
    n_points = len(picked_depths)
    empty_float = np.array([], dtype=np.float64)
    if n_points == 0:
        return {
            "allocation_mode": "shape_intensity_x_support",
            "segment_density_mass_budget": np.nan,
            "segment_density_mass_per_length": np.nan,
            "point_support_left_depth": empty_float,
            "point_support_right_depth": empty_float,
            "point_support_length": empty_float,
            "point_weight_raw": empty_float,
            "point_weight_norm": empty_float,
            "point_density_mass_allocated": empty_float,
            "point_density_mass_per_length_allocated": empty_float,
        }

    seg_start = float(pred_info.get("SegStartDepth", payload.get("SegStartDepth", np.nan)))
    seg_end = float(pred_info.get("SegEndDepth", payload.get("SegEndDepth", np.nan)))
    seg_length = float(max(pred_info.get("SegLength", payload.get("SegLength", 0.0)), 0.0))
    segment_density_mass_per_length = pd.to_numeric(
        pd.Series([pred_info.get("PredP10MassPerLength", pred_info.get("PredDensityStrength", np.nan))]),
        errors="coerce",
    ).iloc[0]
    segment_density_mass_budget = pd.to_numeric(
        pd.Series([pred_info.get("PredP10Mass", np.nan)]),
        errors="coerce",
    ).iloc[0]
    if not np.isfinite(segment_density_mass_budget):
        if np.isfinite(segment_density_mass_per_length):
            segment_density_mass_budget = float(max(segment_density_mass_per_length, 0.0) * seg_length)
        else:
            segment_density_mass_budget = np.nan

    support_left_depth, support_right_depth, support_lengths = compute_point_support_intervals(
        seg_start_depth=seg_start,
        seg_end_depth=seg_end,
        picked_depths=picked_depths,
    )

    shape_intensity = np.asarray(payload.get("shape_intensity"), dtype=np.float64)
    rel_idx = np.asarray(picked_rel_idx, dtype=np.int64)
    if shape_intensity.size == 0 or rel_idx.size != n_points:
        local_intensity = np.ones(n_points, dtype=np.float64)
    else:
        clipped_rel_idx = np.clip(rel_idx, a_min=0, a_max=max(shape_intensity.size - 1, 0))
        local_intensity = shape_intensity[clipped_rel_idx]
    local_intensity = np.nan_to_num(local_intensity, nan=0.0, posinf=0.0, neginf=0.0)
    local_intensity = np.clip(local_intensity, a_min=0.0, a_max=None)

    weight_raw = local_intensity * np.where(support_lengths > 1e-8, support_lengths, 0.0)
    if not np.any(weight_raw > 0.0):
        weight_raw = np.ones(n_points, dtype=np.float64)
    weight_norm = weight_raw / max(float(np.sum(weight_raw)), 1e-8)

    if np.isfinite(segment_density_mass_budget):
        point_density_mass_allocated = weight_norm * float(max(segment_density_mass_budget, 0.0))
    else:
        point_density_mass_allocated = np.full(n_points, np.nan, dtype=np.float64)

    point_density_mass_per_length_allocated = np.full(n_points, np.nan, dtype=np.float64)
    valid_support_mask = support_lengths > 1e-8
    if np.any(valid_support_mask):
        point_density_mass_per_length_allocated[valid_support_mask] = (
            point_density_mass_allocated[valid_support_mask] / support_lengths[valid_support_mask]
        )
    zero_support_mask = ~valid_support_mask
    if np.any(zero_support_mask) and np.isfinite(segment_density_mass_per_length):
        point_density_mass_per_length_allocated[zero_support_mask] = float(
            max(segment_density_mass_per_length, 0.0)
        )

    return {
        "allocation_mode": "shape_intensity_x_support",
        "segment_density_mass_budget": float(segment_density_mass_budget) if np.isfinite(segment_density_mass_budget) else np.nan,
        "segment_density_mass_per_length": float(segment_density_mass_per_length) if np.isfinite(segment_density_mass_per_length) else np.nan,
        "point_support_left_depth": support_left_depth,
        "point_support_right_depth": support_right_depth,
        "point_support_length": support_lengths,
        "point_weight_raw": weight_raw,
        "point_weight_norm": weight_norm,
        "point_density_mass_allocated": point_density_mass_allocated,
        "point_density_mass_per_length_allocated": point_density_mass_per_length_allocated,
    }


def build_pred_points_from_segment_df(
    df_sorted: pd.DataFrame,
    payloads: list[dict],
    pred_segment_df: pd.DataFrame,
    config,
    orientation_strategy: dict | None = None,
) -> pd.DataFrame:
    pred_map = pred_segment_df.set_index("Segment_ID").to_dict("index")
    segment_passthrough_cols = [
        "DensityStrengthMode",
        "DensityStrengthBasisCol",
        "DensityStrengthTargetCol",
        "PredDensityStrength",
        "PredP10Mean",
        "PredP10MassPerLength",
        "PredP10Mass",
        "PredDensitySizeScale",
        "PredDensityStrengthLevel",
        "DensityPointFloorEnabled",
        "DensityPointFloorThresholdMassPerLength",
        "DensityPointFloorSourceQuantile",
        "DensityPointFloorApplied",
        "PredOrientationFamily",
        "PredOrientationConfidence",
        "PredAzimuth",
        "PredDip",
    ]
    rows = []
    for payload in payloads:
        segment_id = int(payload["Segment_ID"])
        pred_info = pred_map.get(segment_id)
        if pred_info is None:
            continue
        max_points = int(max(len(payload["depth_seg"]), 0))
        n_points = int(min(int(pred_info["PredPointCount"]), max_points))
        if n_points <= 0:
            continue
        picked_depths, picked_rel_idx = dataset_builder.POINT.pick_points_by_equal_intensity(
            payload["depth_seg"],
            payload["shape_intensity"],
            n_points,
        )
        density_allocation = allocate_segment_density_to_points(
            pred_info=pred_info,
            payload=payload,
            picked_depths=picked_depths,
            picked_rel_idx=picked_rel_idx,
        )
        orientation_point_info = expand_segment_orientation_to_points(
            pred_info=pred_info,
            orientation_strategy=orientation_strategy,
            picked_depths=picked_depths,
        )
        for point_idx, (depth_value, rel_idx) in enumerate(zip(picked_depths, picked_rel_idx), start=1):
            abs_idx = int(payload["StartIdx"] + rel_idx)
            row = df_sorted.iloc[abs_idx]
            point_row = {
                "Segment_ID": segment_id,
                "Point_ID_In_Segment": point_idx,
                "WellName": pred_info.get("WellName", row.get("WellName", "")),
                "StrataName": pred_info.get("StrataName", ""),
                config.depth_col: float(depth_value),
                "NearestSampleDepth": float(row[config.depth_col]),
                config.x_col: float(row[config.x_col]) if config.x_col in df_sorted.columns else np.nan,
                config.y_col: float(row[config.y_col]) if config.y_col in df_sorted.columns else np.nan,
                config.time_col: float(row[config.time_col]) if config.time_col in df_sorted.columns else np.nan,
                "SegStartDepth": float(pred_info["SegStartDepth"]),
                "SegEndDepth": float(pred_info["SegEndDepth"]),
                "SegLength": float(pred_info["SegLength"]),
                "PredPointCountFloat": float(pred_info["PredPointCountFloat"]),
                "PredPointCount": n_points,
                "RawGTCountInPredSegment": int(pred_info["RawGTCountInPredSegment"]),
                "PredProbMean": float(pred_info["ProbMean"]),
                "PredProbMax": float(pred_info["ProbMax"]),
                "StrataDepthMin": pred_info.get("StrataDepthMin", np.nan),
                "StrataDepthMax": pred_info.get("StrataDepthMax", np.nan),
                "PointDensityAllocationMode": density_allocation["allocation_mode"],
                "SegmentDensityMassBudget": density_allocation["segment_density_mass_budget"],
                "SegmentDensityMassPerLength": density_allocation["segment_density_mass_per_length"],
                "PointSupportLeftDepth": float(density_allocation["point_support_left_depth"][point_idx - 1]),
                "PointSupportRightDepth": float(density_allocation["point_support_right_depth"][point_idx - 1]),
                "PointSupportLength": float(density_allocation["point_support_length"][point_idx - 1]),
                "PointDensityWeightRaw": float(density_allocation["point_weight_raw"][point_idx - 1]),
                "PointDensityWeightNorm": float(density_allocation["point_weight_norm"][point_idx - 1]),
                "PointDensityMassAllocated": float(density_allocation["point_density_mass_allocated"][point_idx - 1]),
                "PointDensityMassPerLengthAllocated": float(
                    density_allocation["point_density_mass_per_length_allocated"][point_idx - 1]
                ),
                "PointOrientationMode": orientation_point_info["mode"],
                "PointOrientationTemplateValue": float(orientation_point_info["template"][point_idx - 1]),
                "PointOrientationConfidenceScale": float(orientation_point_info["confidence_scale"]),
                "PointAzimuthStdRefDeg": float(orientation_point_info["azimuth_std_ref_deg"]),
                "PointDipStdRefDeg": float(orientation_point_info["dip_std_ref_deg"]),
                "PointAzimuthOffsetDeg": float(orientation_point_info["azimuth_offset_deg"][point_idx - 1]),
                "PointDipOffsetDeg": float(orientation_point_info["dip_offset_deg"][point_idx - 1]),
                "PointAzimuth": float(orientation_point_info["point_azimuth"][point_idx - 1]),
                "PointDip": float(orientation_point_info["point_dip"][point_idx - 1]),
            }
            for col_name in segment_passthrough_cols:
                point_row[col_name] = pred_info.get(col_name, np.nan)
            rows.append(point_row)
    return pd.DataFrame(rows)


def build_coef_rows(well_name: str, regressor: Pipeline, feature_cols: list[str]) -> list[dict]:
    model = regressor.named_steps["model"]
    rows = []
    if hasattr(model, "coef_"):
        for feature_name, coef in zip(feature_cols, model.coef_):
            rows.append({"val_well": well_name, "feature": feature_name, "coef": float(coef)})
    if hasattr(model, "intercept_"):
        rows.append({"val_well": well_name, "feature": "__intercept__", "coef": float(model.intercept_)})
    elif hasattr(model, "feature_importances_"):
        for feature_name, importance in zip(feature_cols, model.feature_importances_):
            rows.append({"val_well": well_name, "feature": feature_name, "coef": float(importance)})
    return rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", required=True)
    parser.add_argument("--exist-exp-dir", default=str(DEFAULT_EXIST_EXP_DIR))
    parser.add_argument("--sample-dir", default=str(DEFAULT_SAMPLE_DIR))
    parser.add_argument("--raw-label-dir", default=str(DEFAULT_RAW_LABEL_DIR))
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument("--result-dir")
    parser.add_argument("--well-names", default="")
    parser.add_argument("--gt-dev-rule", default="any_density")
    parser.add_argument("--prob-weight-gamma", type=float, default=2.0)
    parser.add_argument("--pred-min-points-per-segment", type=int, default=0)
    parser.add_argument("--rounding-mode", default="round")
    parser.add_argument("--fallback-count-mode", default="")
    parser.add_argument("--count-fusion-mode", default="learned_only")
    parser.add_argument("--fusion-learned-weight", type=float, default=0.7)
    parser.add_argument("--global-post-scale-mode", default="none")
    parser.add_argument("--global-post-scale-min", type=float, default=0.5)
    parser.add_argument("--global-post-scale-max", type=float, default=1.5)
    parser.add_argument("--selective-post-mode", default="none")
    parser.add_argument("--selective-post-target-wells", default="")
    parser.add_argument("--selective-post-feature", default="ProbMassPerLength")
    parser.add_argument("--selective-post-feature-2", default="ProbMean")
    parser.add_argument("--selective-post-quantile", type=float, default=0.5)
    parser.add_argument("--selective-post-quantile-2", type=float, default=0.5)
    parser.add_argument("--soft-negative-weight", type=float, default=0.5)
    parser.add_argument("--model-name", default="tweedie")
    parser.add_argument("--model-alpha", type=float, default=0.03)
    parser.add_argument("--model-max-iter", type=int, default=3000)
    parser.add_argument("--tweedie-power", type=float, default=1.5)
    parser.add_argument("--feature-preset", default="high_density_compact_v1")
    parser.add_argument("--feature-cols", default=",".join(dataset_builder.DEFAULT_FEATURE_COLS))
    parser.add_argument("--xgb-n-estimators", type=int, default=300)
    parser.add_argument("--xgb-max-depth", type=int, default=3)
    parser.add_argument("--xgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--xgb-subsample", type=float, default=0.8)
    parser.add_argument("--xgb-colsample-bytree", type=float, default=0.8)
    parser.add_argument("--hgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--hgb-max-depth", type=int, default=3)
    args = parser.parse_args()

    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / sanitize(args.exp_id))
    save_dir.mkdir(parents=True, exist_ok=True)
    summary_csv = save_dir / "raw_point_guided_refine_summary.csv"

    verify_logs = dataset_builder.POINT.list_verify_logs(str(args.exist_exp_dir))
    selected_wells = parse_list_arg(args.well_names) if args.well_names else sorted(verify_logs.keys())
    missing_wells = [well for well in selected_wells if well not in verify_logs]
    if missing_wells:
        raise ValueError(f"Wells not found in verify logs: {missing_wells}")

    feature_cols = resolve_feature_cols(str(args.feature_preset), str(args.feature_cols))
    fallback_count_mode = validate_optional_count_mode(str(args.fallback_count_mode))
    global_post_scale_mode = validate_global_post_scale_mode(str(args.global_post_scale_mode))
    selective_post_mode = validate_selective_post_mode(str(args.selective_post_mode))
    selective_post_target_wells = set(parse_list_arg(args.selective_post_target_wells))
    config = dataset_builder.make_config(
        exist_exp_dir=str(args.exist_exp_dir),
        prob_weight_gamma=float(args.prob_weight_gamma),
        gt_min_points_per_segment=1,
        gt_rounding_mode="round",
        density_gt_col="P10",
    )

    per_well = {}
    all_rows = []
    dataset_save_dir = save_dir / "segment_dataset"
    dataset_save_dir.mkdir(parents=True, exist_ok=True)
    for well_name in selected_wells:
        well_data = build_well_segment_dataset(
            exist_csv=verify_logs[well_name],
            well_name=well_name,
            config=config,
            sample_dir=Path(args.sample_dir),
            raw_label_dir=Path(args.raw_label_dir),
            gt_dev_rule=str(args.gt_dev_rule),
        )
        well_data["segment_df"] = compute_target_weights(
            well_data["segment_df"],
            soft_negative_weight=float(args.soft_negative_weight),
        )
        per_well[well_name] = well_data
        well_dir = dataset_save_dir / well_name
        well_dir.mkdir(parents=True, exist_ok=True)
        well_data["segment_df"].to_csv(well_dir / "segment_dataset.csv", index=False, encoding="utf-8-sig")
        well_data["raw_gt_points"].to_csv(well_dir / "raw_gt_points.csv", index=False, encoding="utf-8-sig")
        well_data["gt_dev_segments"].to_csv(well_dir / "gt_dev_segments.csv", index=False, encoding="utf-8-sig")
        if not well_data["segment_df"].empty:
            all_rows.append(well_data["segment_df"])

    dataset_df = pd.concat(all_rows, ignore_index=True) if all_rows else pd.DataFrame()
    if dataset_df.empty:
        raise ValueError("Segment dataset is empty")
    dataset_df.to_csv(dataset_save_dir / "all_segment_dataset.csv", index=False, encoding="utf-8-sig")
    validate_feature_cols(feature_cols, dataset_df)

    summary_rows = []
    coef_rows = []
    for well_name in selected_wells:
        train_df = dataset_df[dataset_df["WellName"] != well_name].copy()
        val_df = dataset_df[dataset_df["WellName"] == well_name].copy()
        if train_df.empty or val_df.empty:
            continue

        x_train = train_df[feature_cols].fillna(0.0)
        y_train = train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0)
        x_val = val_df[feature_cols].fillna(0.0)
        sample_weight = train_df["TargetWeight"].fillna(1.0).to_numpy(dtype=np.float64)

        learned_pred_count_float = np.zeros(len(val_df), dtype=np.float64)
        learned_train_pred_count_float = np.zeros(len(train_df), dtype=np.float64)
        regressor = None
        if str(args.model_name) != "rule_only":
            regressor = build_regressor(
                model_name=str(args.model_name),
                alpha=float(args.model_alpha),
                max_iter=int(args.model_max_iter),
                tweedie_power=float(args.tweedie_power),
                xgb_n_estimators=int(args.xgb_n_estimators),
                xgb_max_depth=int(args.xgb_max_depth),
                xgb_learning_rate=float(args.xgb_learning_rate),
                xgb_subsample=float(args.xgb_subsample),
                xgb_colsample_bytree=float(args.xgb_colsample_bytree),
                hgb_learning_rate=float(args.hgb_learning_rate),
                hgb_max_depth=int(args.hgb_max_depth),
            )
            regressor = fit_regressor(regressor, x_train, y_train, sample_weight)
            learned_train_pred_count_float = np.clip(regressor.predict(x_train), a_min=0.0, a_max=None)
            learned_pred_count_float = np.clip(regressor.predict(x_val), a_min=0.0, a_max=None)
            coef_rows.extend(build_coef_rows(well_name, regressor, feature_cols))

        fallback_basis_col = ""
        fallback_scale = np.nan
        fallback_pred_count_float = None
        fallback_train_pred_count_float = None
        if fallback_count_mode:
            fallback_basis_col, fallback_scale = fit_basis_scale(train_df, fallback_count_mode)
            fallback_train_pred_count_float = np.clip(
                train_df[fallback_basis_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64) * fallback_scale,
                a_min=0.0,
                a_max=None,
            )
            fallback_pred_count_float = np.clip(
                val_df[fallback_basis_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64) * fallback_scale,
                a_min=0.0,
                a_max=None,
            )

        if str(args.model_name) == "rule_only":
            if fallback_pred_count_float is None:
                raise ValueError("model_name=rule_only requires fallback_count_mode")
            train_pred_count_float_before_post_scale = fallback_train_pred_count_float
            pred_count_float_before_post_scale = fallback_pred_count_float
        else:
            train_pred_count_float_before_post_scale = fuse_count_predictions(
                learned_pred=learned_train_pred_count_float,
                fallback_pred=fallback_train_pred_count_float,
                fusion_mode=str(args.count_fusion_mode),
                learned_weight=float(args.fusion_learned_weight),
            )
            pred_count_float_before_post_scale = fuse_count_predictions(
                learned_pred=learned_pred_count_float,
                fallback_pred=fallback_pred_count_float,
                fusion_mode=str(args.count_fusion_mode),
                learned_weight=float(args.fusion_learned_weight),
            )

        post_scale_info = fit_global_post_scale(
            train_target=train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64),
            train_pred=train_pred_count_float_before_post_scale,
            mode=global_post_scale_mode,
            scale_min=float(args.global_post_scale_min),
            scale_max=float(args.global_post_scale_max),
        )
        train_pred_count_float_after_global_scale = apply_global_post_scale(
            train_pred_count_float_before_post_scale,
            post_scale_info,
        )
        pred_count_float_after_global_scale = apply_global_post_scale(
            pred_count_float_before_post_scale,
            post_scale_info,
        )
        selective_post_mode_this_well = (
            selective_post_mode
            if (not selective_post_target_wells or well_name in selective_post_target_wells)
            else "none"
        )
        selective_post_strategy = fit_selective_post_strategy(
            train_df=train_df,
            train_pred_after_global_scale=train_pred_count_float_after_global_scale,
            mode=selective_post_mode_this_well,
            feature_name=str(args.selective_post_feature),
            feature_name_2=str(args.selective_post_feature_2),
            quantile=float(args.selective_post_quantile),
            quantile_2=float(args.selective_post_quantile_2),
            fallback_scale=float(post_scale_info["raw_scale"]),
        )
        pred_count_float, lowconf_mask = apply_selective_post_strategy(
            pred_count_float_after_global_scale,
            val_df,
            selective_post_strategy,
        )

        pred_segment_df = val_df.copy()
        pred_segment_df["LearnedPredPointCountFloat"] = learned_pred_count_float
        pred_segment_df["FallbackCountBasisCol"] = fallback_basis_col
        pred_segment_df["FallbackCountScale"] = fallback_scale
        pred_segment_df["FallbackPredPointCountFloat"] = (
            fallback_pred_count_float if fallback_pred_count_float is not None else np.nan
        )
        pred_segment_df["PredPointCountFloatBeforePostScale"] = pred_count_float_before_post_scale
        pred_segment_df["GlobalPostScaleMode"] = post_scale_info["mode"]
        pred_segment_df["GlobalPostScaleRawScale"] = post_scale_info["raw_scale"]
        pred_segment_df["GlobalPostScaleAppliedScale"] = post_scale_info["scale"]
        pred_segment_df["GlobalPostScaleTrainTargetSum"] = post_scale_info["train_target_sum"]
        pred_segment_df["GlobalPostScaleTrainPredSum"] = post_scale_info["train_pred_sum"]
        pred_segment_df["SelectivePostMode"] = selective_post_strategy["mode"]
        pred_segment_df["SelectivePostFeature"] = selective_post_strategy["feature_name"]
        pred_segment_df["SelectivePostFeature2"] = selective_post_strategy["feature_name_2"]
        pred_segment_df["SelectivePostThreshold"] = selective_post_strategy["threshold"]
        pred_segment_df["SelectivePostThreshold2"] = selective_post_strategy["threshold_2"]
        pred_segment_df["SelectivePostQuantile"] = selective_post_strategy["quantile"]
        pred_segment_df["SelectivePostQuantile2"] = selective_post_strategy["quantile_2"]
        pred_segment_df["SelectivePostLowconfScale"] = selective_post_strategy["lowconf_scale"]
        pred_segment_df["SelectivePostIsLowConf"] = lowconf_mask.astype(int)
        pred_segment_df["PredPointCountFloat"] = pred_count_float
        pred_segment_df["PredPointCount"] = pred_segment_df["PredPointCountFloat"].apply(
            lambda value: compute_pred_count(
                raw_value=float(value),
                pred_min_points_per_segment=int(args.pred_min_points_per_segment),
                rounding_mode=str(args.rounding_mode),
            )
        )
        pred_segment_df = apply_pred_point_count_sample_cap(pred_segment_df)
        pred_segment_df["SegmentCountAbsError"] = (
            pred_segment_df["PredPointCount"] - pred_segment_df["RawGTCountInPredSegment"]
        ).abs()

        well_data = per_well[well_name]
        pred_points = build_pred_points_from_segment_df(
            df_sorted=well_data["df_sorted"],
            payloads=well_data["payloads"],
            pred_segment_df=pred_segment_df,
            config=config,
            orientation_strategy=dict(artifact.get("orientation_strategy") or {}),
        )
        raw_gt_points = well_data["raw_gt_points"]
        raw_gt_depths = well_data["raw_gt_depths"]
        pred_depths = (
            pred_points[config.depth_col].to_numpy(dtype=np.float64)
            if not pred_points.empty
            else np.array([], dtype=np.float64)
        )
        pred_to_gt = dataset_builder.POINT.nearest_distance_stats(pred_depths, raw_gt_depths)
        gt_to_pred = dataset_builder.POINT.nearest_distance_stats(raw_gt_depths, pred_depths)

        well_save_dir = save_dir / well_name
        well_save_dir.mkdir(parents=True, exist_ok=True)
        raw_gt_points.to_csv(well_save_dir / "raw_gt_points.csv", index=False, encoding="utf-8-sig")
        well_data["gt_dev_segments"].to_csv(well_save_dir / "gt_dev_segments.csv", index=False, encoding="utf-8-sig")
        pred_points.to_csv(well_save_dir / "pred_fracture_points.csv", index=False, encoding="utf-8-sig")
        pred_segment_df.to_csv(well_save_dir / "pred_segment_summary.csv", index=False, encoding="utf-8-sig")
        train_df.to_csv(well_save_dir / "train_segment_dataset.csv", index=False, encoding="utf-8-sig")

        metrics = {
            "well": well_name,
            "exist_csv": well_data["exist_csv"],
            "model_name": args.model_name,
            "model_alpha": args.model_alpha,
            "model_max_iter": args.model_max_iter,
            "tweedie_power": args.tweedie_power,
            "hgb_learning_rate": args.hgb_learning_rate,
            "hgb_max_depth": args.hgb_max_depth,
            "feature_preset": args.feature_preset,
            "feature_cols": feature_cols,
            "prob_weight_gamma": args.prob_weight_gamma,
            "fallback_count_mode": fallback_count_mode,
            "count_fusion_mode": args.count_fusion_mode,
            "fusion_learned_weight": args.fusion_learned_weight,
            "global_post_scale_mode": global_post_scale_mode,
            "global_post_scale_raw_scale": post_scale_info["raw_scale"],
            "global_post_scale_applied_scale": post_scale_info["scale"],
            "global_post_scale_train_target_sum": post_scale_info["train_target_sum"],
            "global_post_scale_train_pred_sum": post_scale_info["train_pred_sum"],
            "selective_post_mode": selective_post_mode_this_well,
            "selective_post_target_wells": sorted(selective_post_target_wells),
            "selective_post_feature": args.selective_post_feature,
            "selective_post_feature_2": args.selective_post_feature_2,
            "selective_post_quantile": args.selective_post_quantile,
            "selective_post_quantile_2": args.selective_post_quantile_2,
            "selective_post_threshold": selective_post_strategy["threshold"],
            "selective_post_threshold_2": selective_post_strategy["threshold_2"],
            "selective_post_lowconf_scale": selective_post_strategy["lowconf_scale"],
            "selective_post_num_lowconf_segments": int(np.sum(lowconf_mask.astype(np.int64))),
            "soft_negative_weight": args.soft_negative_weight,
            "pred_min_points_per_segment": args.pred_min_points_per_segment,
            "rounding_mode": args.rounding_mode,
            "num_train_segments": int(len(train_df)),
            "num_val_segments": int(len(val_df)),
            "segment_count_mae": float(pred_segment_df["SegmentCountAbsError"].mean()),
            "segment_count_rmse": float(
                np.sqrt(
                    np.mean(
                        np.square(
                            pred_segment_df["PredPointCount"].to_numpy(dtype=np.float64)
                            - pred_segment_df["RawGTCountInPredSegment"].to_numpy(dtype=np.float64)
                        )
                    )
                )
            ),
            "N_gt_points": int(len(raw_gt_depths)),
            "N_pred_points": int(len(pred_depths)),
            "count_diff": int(len(pred_depths) - len(raw_gt_depths)),
            "pred_to_gt_mean_dist": pred_to_gt["mean"],
            "pred_to_gt_median_dist": pred_to_gt["median"],
            "pred_to_gt_p90_dist": pred_to_gt["p90"],
            "gt_to_pred_mean_dist": gt_to_pred["mean"],
            "gt_to_pred_median_dist": gt_to_pred["median"],
            "gt_to_pred_p90_dist": gt_to_pred["p90"],
            "raw_points_covered_by_pred_segments": int(well_data["raw_points_covered_by_pred_segments"]),
            "raw_points_missed_outside_pred_segments": int(well_data["raw_points_missed_outside_pred_segments"]),
            "pred_segments_with_raw_gt_0": int((pred_segment_df["RawGTCountInPredSegment"] <= 0).sum()),
            "pred_segments_overlap_gt_dev_and_raw_gt_0": int(
                ((pred_segment_df["RawGTCountInPredSegment"] <= 0) & (pred_segment_df["GTDevOverlapFlag"] == 1)).sum()
            ),
        }
        with open(well_save_dir / "metrics.json", "w", encoding="utf-8") as file_obj:
            json.dump(metrics, file_obj, ensure_ascii=False, indent=2)
        summary_rows.append(metrics)

    summary_df = pd.DataFrame(summary_rows)
    if summary_df.empty:
        raise ValueError("No validation summary rows were produced")
    summary_df.to_csv(summary_csv, index=False, encoding="utf-8-sig")
    pd.DataFrame(coef_rows).to_csv(save_dir / "loo_model_coefficients.csv", index=False, encoding="utf-8-sig")

    with open(save_dir / "config.json", "w", encoding="utf-8") as file_obj:
        json.dump(
            {
                "exp_id": args.exp_id,
                "exist_exp_dir": str(args.exist_exp_dir),
                "sample_dir": str(args.sample_dir),
                "raw_label_dir": str(args.raw_label_dir),
                "well_names": selected_wells,
                "gt_dev_rule": args.gt_dev_rule,
                "feature_cols": feature_cols,
                "feature_preset": args.feature_preset,
                "prob_weight_gamma": args.prob_weight_gamma,
                "fallback_count_mode": fallback_count_mode,
                "count_fusion_mode": args.count_fusion_mode,
                "fusion_learned_weight": args.fusion_learned_weight,
                "global_post_scale_mode": global_post_scale_mode,
                "global_post_scale_min": args.global_post_scale_min,
                "global_post_scale_max": args.global_post_scale_max,
                "selective_post_mode": selective_post_mode,
                "selective_post_target_wells": sorted(selective_post_target_wells),
                "selective_post_feature": args.selective_post_feature,
                "selective_post_feature_2": args.selective_post_feature_2,
                "selective_post_quantile": args.selective_post_quantile,
                "selective_post_quantile_2": args.selective_post_quantile_2,
                "soft_negative_weight": args.soft_negative_weight,
                "pred_min_points_per_segment": args.pred_min_points_per_segment,
                "rounding_mode": args.rounding_mode,
                "model_name": args.model_name,
                "model_alpha": args.model_alpha,
                "model_max_iter": args.model_max_iter,
                "tweedie_power": args.tweedie_power,
                "hgb_learning_rate": args.hgb_learning_rate,
                "hgb_max_depth": args.hgb_max_depth,
                "xgb_n_estimators": args.xgb_n_estimators,
                "xgb_max_depth": args.xgb_max_depth,
                "xgb_learning_rate": args.xgb_learning_rate,
                "xgb_subsample": args.xgb_subsample,
                "xgb_colsample_bytree": args.xgb_colsample_bytree,
                "save_dir": str(save_dir),
                "summary_csv": str(summary_csv),
            },
            file_obj,
            ensure_ascii=False,
            indent=2,
        )

    with summary_csv.open("r", encoding="utf-8-sig", newline="") as file_obj:
        results = list(csv.DictReader(file_obj))
    append_result_to_docx(
        Path(args.docx_path),
        f"实验 {args.exp_id}",
        {
            "exist_exp_dir": str(args.exist_exp_dir),
            "sample_dir": str(args.sample_dir),
            "raw_label_dir": str(args.raw_label_dir),
            "gt_dev_rule": args.gt_dev_rule,
            "model_name": args.model_name,
            "model_alpha": args.model_alpha,
            "model_max_iter": args.model_max_iter,
            "tweedie_power": args.tweedie_power,
            "hgb_learning_rate": args.hgb_learning_rate,
            "hgb_max_depth": args.hgb_max_depth,
            "feature_preset": args.feature_preset,
            "feature_cols": feature_cols,
            "prob_weight_gamma": args.prob_weight_gamma,
            "fallback_count_mode": fallback_count_mode,
            "count_fusion_mode": args.count_fusion_mode,
            "fusion_learned_weight": args.fusion_learned_weight,
            "global_post_scale_mode": global_post_scale_mode,
            "global_post_scale_min": args.global_post_scale_min,
            "global_post_scale_max": args.global_post_scale_max,
            "selective_post_mode": selective_post_mode,
            "selective_post_target_wells": sorted(selective_post_target_wells),
            "selective_post_feature": args.selective_post_feature,
            "selective_post_feature_2": args.selective_post_feature_2,
            "selective_post_quantile": args.selective_post_quantile,
            "selective_post_quantile_2": args.selective_post_quantile_2,
            "soft_negative_weight": args.soft_negative_weight,
            "pred_min_points_per_segment": args.pred_min_points_per_segment,
            "rounding_mode": args.rounding_mode,
            "save_dir": str(save_dir),
            "summary_csv": str(summary_csv),
            "xgb_n_estimators": args.xgb_n_estimators,
            "xgb_max_depth": args.xgb_max_depth,
            "xgb_learning_rate": args.xgb_learning_rate,
            "xgb_subsample": args.xgb_subsample,
            "xgb_colsample_bytree": args.xgb_colsample_bytree,
        },
        results,
    )

    print(
        json.dumps(
            {
                "exp_id": args.exp_id,
                "save_dir": str(save_dir),
                "summary_csv": str(summary_csv),
                "feature_cols": feature_cols,
                "results": results,
            },
            ensure_ascii=False,
        )
    )
    return 0


def build_run_config_dict(
    args,
    selected_wells: list[str],
    feature_cols: list[str],
    fallback_count_mode: str,
    global_post_scale_mode: str,
    selective_post_mode: str,
    selective_post_target_wells: set[str],
    gate_target_wells: set[str],
    train_distance_csv: Path | None,
    save_dir: Path,
    summary_csv: Path,
    saved_model_root: Path | None = None,
) -> dict:
    config_dict = {
        "exp_id": args.exp_id,
        "run_mode": args.run_mode,
        "config_profile": str(getattr(args, "config_profile", "")),
        "exist_exp_dir": str(args.exist_exp_dir),
        "sample_dir": str(args.sample_dir),
        "raw_label_dir": str(args.raw_label_dir),
        "strata_name": str(args.strata_name or "").strip(),
        "strata_range_csv": str(args.strata_range_csv or "").strip(),
        "well_names": selected_wells,
        "gt_dev_rule": args.gt_dev_rule,
        "feature_cols": feature_cols,
        "feature_preset": args.feature_preset,
        "prob_weight_gamma": args.prob_weight_gamma,
        "count_train_mode": validate_count_train_mode(str(args.count_train_mode)),
        "positive_floor_prob_min": float(args.positive_floor_prob_min),
        "positive_floor_allowed_labels": resolve_positive_floor_allowed_labels(
            str(args.positive_floor_allowed_labels)
        ),
        "positive_floor_target_wells": sorted(parse_list_arg(str(args.positive_floor_target_wells))),
        "type_calibration_mode": validate_type_calibration_mode(str(args.type_calibration_mode)),
        "type_calibration_scale_min": float(args.type_calibration_scale_min),
        "type_calibration_scale_max": float(args.type_calibration_scale_max),
        "fallback_count_mode": fallback_count_mode,
        "count_fusion_mode": args.count_fusion_mode,
        "fusion_learned_weight": args.fusion_learned_weight,
        "global_post_scale_mode": global_post_scale_mode,
        "global_post_scale_min": args.global_post_scale_min,
        "global_post_scale_max": args.global_post_scale_max,
        "selective_post_mode": selective_post_mode,
        "selective_post_target_wells": sorted(selective_post_target_wells),
        "selective_post_feature": args.selective_post_feature,
        "selective_post_feature_2": args.selective_post_feature_2,
        "selective_post_quantile": args.selective_post_quantile,
        "selective_post_quantile_2": args.selective_post_quantile_2,
        "selective_post_max_pred_float": args.selective_post_max_pred_float,
        "train_distance_weight_mode": validate_train_distance_weight_mode(str(args.train_distance_weight_mode)),
        "train_distance_weight_power": args.train_distance_weight_power,
        "train_distance_weight_min": args.train_distance_weight_min,
        "train_distance_weight_max": args.train_distance_weight_max,
        "train_distance_csv": str(train_distance_csv) if train_distance_csv is not None else "",
        "boundary_expand_mode": validate_boundary_expand_mode(str(args.boundary_expand_mode)),
        "boundary_expand_prob_min": args.boundary_expand_prob_min,
        "boundary_expand_max_steps": args.boundary_expand_max_steps,
        "boundary_expand_max_depth": args.boundary_expand_max_depth,
        "gate_mode": validate_gate_mode(str(args.gate_mode)),
        "gate_prob_threshold": args.gate_prob_threshold,
        "gate_target_wells": sorted(gate_target_wells),
        "count_bin_mode": validate_count_bin_mode(str(getattr(args, "count_bin_mode", "none"))),
        "count_bin_blend_weight": float(getattr(args, "count_bin_blend_weight", 0.7)),
        "density_strength_mode": validate_density_strength_mode(str(getattr(args, "density_strength_mode", "none"))),
        "density_strength_basis_col": str(getattr(args, "density_strength_basis_col", "ProbMassPerLength")),
        "density_strength_scale_min": float(getattr(args, "density_strength_scale_min", 0.5)),
        "density_strength_scale_max": float(getattr(args, "density_strength_scale_max", 3.0)),
        "density_strength_regression_train_mode": validate_density_strength_regression_train_mode(
            str(getattr(args, "density_strength_regression_train_mode", "all_segments"))
        ),
        "density_strength_regression_target_transform": validate_density_strength_regression_target_transform(
            str(getattr(args, "density_strength_regression_target_transform", "none"))
        ),
        "density_strength_regression_clip_quantile": float(
            getattr(args, "density_strength_regression_clip_quantile", np.nan)
        ),
        "density_strength_regression_calibration_mode": validate_density_strength_regression_calibration_mode(
            str(getattr(args, "density_strength_regression_calibration_mode", "none"))
        ),
        "density_strength_regression_calibration_quantiles": (
            normalize_density_strength_regression_calibration_quantiles(
                str(
                    getattr(
                        args,
                        "density_strength_regression_calibration_quantiles",
                        "0.1,0.25,0.5,0.75,0.9,0.95",
                    )
                )
            )
        ),
        "orientation_mode": validate_orientation_mode(str(getattr(args, "orientation_mode", "none"))),
        "orientation_num_families": int(getattr(args, "orientation_num_families", 3)),
        "orientation_min_points_per_segment": int(getattr(args, "orientation_min_points_per_segment", 1)),
        "density_point_floor_enabled": validate_density_strength_mode(
            str(getattr(args, "density_strength_mode", "none"))
        ) != "none",
        "soft_negative_weight": args.soft_negative_weight,
        "pred_min_points_per_segment": args.pred_min_points_per_segment,
        "rounding_mode": args.rounding_mode,
        "model_name": args.model_name,
        "model_alpha": args.model_alpha,
        "model_max_iter": args.model_max_iter,
        "tweedie_power": args.tweedie_power,
        "hgb_learning_rate": args.hgb_learning_rate,
        "hgb_max_depth": args.hgb_max_depth,
        "xgb_n_estimators": args.xgb_n_estimators,
        "xgb_max_depth": args.xgb_max_depth,
        "xgb_learning_rate": args.xgb_learning_rate,
        "xgb_subsample": args.xgb_subsample,
        "xgb_colsample_bytree": args.xgb_colsample_bytree,
        "save_model_artifacts": int(args.save_model_artifacts),
        "save_dir": str(save_dir),
        "summary_csv": str(summary_csv),
    }
    if saved_model_root is not None:
        config_dict["saved_model_root"] = str(saved_model_root)
    return config_dict


def fit_segment_refine_artifact(
    train_df: pd.DataFrame,
    feature_cols: list[str],
    args,
    fallback_count_mode: str,
    global_post_scale_mode: str,
    selective_post_mode_this_well: str,
    gate_mode_this_well: str,
    artifact_type: str,
    artifact_target: str,
    train_wells: list[str],
    distance_weight_info: dict,
    train_distance_csv: Path | None,
) -> tuple[dict, list[dict]]:
    count_train_mode = validate_count_train_mode(str(getattr(args, "count_train_mode", "all_segments")))
    count_train_df, effective_count_train_mode = select_count_train_df(train_df, count_train_mode)
    x_train = count_train_df[feature_cols].fillna(0.0)
    y_train = count_train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0)
    sample_weight = count_train_df["TrainSampleWeightUsed"].fillna(1.0).to_numpy(dtype=np.float64)

    learned_train_pred_count_float = np.zeros(len(train_df), dtype=np.float64)
    regressor = None
    coef_rows: list[dict] = []
    if str(args.model_name) != "rule_only":
        regressor = build_regressor(
            model_name=str(args.model_name),
            alpha=float(args.model_alpha),
            max_iter=int(args.model_max_iter),
            tweedie_power=float(args.tweedie_power),
            xgb_n_estimators=int(args.xgb_n_estimators),
            xgb_max_depth=int(args.xgb_max_depth),
            xgb_learning_rate=float(args.xgb_learning_rate),
            xgb_subsample=float(args.xgb_subsample),
            xgb_colsample_bytree=float(args.xgb_colsample_bytree),
            hgb_learning_rate=float(args.hgb_learning_rate),
            hgb_max_depth=int(args.hgb_max_depth),
        )
        regressor = fit_regressor(regressor, x_train, y_train, sample_weight)
        learned_train_pred_count_float = np.clip(
            regressor.predict(train_df[feature_cols].fillna(0.0)),
            a_min=0.0,
            a_max=None,
        )
        coef_rows.extend(build_coef_rows(artifact_target, regressor, feature_cols))

    fallback_basis_col = ""
    fallback_scale = np.nan
    fallback_train_pred_count_float = None
    if fallback_count_mode:
        fallback_basis_col, fallback_scale = fit_basis_scale(train_df, fallback_count_mode)
        fallback_train_pred_count_float = np.clip(
            train_df[fallback_basis_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64) * fallback_scale,
            a_min=0.0,
            a_max=None,
        )

    if str(args.model_name) == "rule_only":
        if fallback_train_pred_count_float is None:
            raise ValueError("model_name=rule_only requires fallback_count_mode")
        train_pred_count_float_before_post_scale = fallback_train_pred_count_float
    else:
        train_pred_count_float_before_post_scale = fuse_count_predictions(
            learned_pred=learned_train_pred_count_float,
            fallback_pred=fallback_train_pred_count_float,
            fusion_mode=str(args.count_fusion_mode),
            learned_weight=float(args.fusion_learned_weight),
        )

    post_scale_info = fit_global_post_scale(
        train_target=train_df["RawGTCountInPredSegment"].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64),
        train_pred=train_pred_count_float_before_post_scale,
        mode=global_post_scale_mode,
        scale_min=float(args.global_post_scale_min),
        scale_max=float(args.global_post_scale_max),
    )
    train_pred_count_float_after_global_scale = apply_global_post_scale(
        train_pred_count_float_before_post_scale,
        post_scale_info,
    )
    selective_post_strategy = fit_selective_post_strategy(
        train_df=train_df,
        train_pred_after_global_scale=train_pred_count_float_after_global_scale,
        mode=selective_post_mode_this_well,
        feature_name=str(args.selective_post_feature),
        feature_name_2=str(args.selective_post_feature_2),
        quantile=float(args.selective_post_quantile),
        quantile_2=float(args.selective_post_quantile_2),
        fallback_scale=float(post_scale_info["raw_scale"]),
        max_pred_float=float(args.selective_post_max_pred_float),
    )
    train_pred_count_float_after_selective_post, _ = apply_selective_post_strategy(
        train_pred_count_float_after_global_scale,
        train_df,
        selective_post_strategy,
        pred_count_float_for_protection=train_pred_count_float_before_post_scale,
    )
    gate_info, gate_coef_rows = fit_gate_classifier(
        train_df=train_df,
        feature_cols=feature_cols,
        gate_mode=str(gate_mode_this_well),
        gate_prob_threshold=float(getattr(args, "gate_prob_threshold", 0.5)),
    )
    train_pred_count_float_after_gate, train_gate_prob, train_gate_label = apply_gate_to_count_predictions(
        train_pred_count_float_after_selective_post,
        train_df,
        gate_info,
    )
    positive_floor_prob_min = float(np.clip(float(getattr(args, "positive_floor_prob_min", 0.7)), 0.0, 1.0))
    positive_floor_allowed_labels = resolve_positive_floor_allowed_labels(
        str(getattr(args, "positive_floor_allowed_labels", ""))
    )
    positive_floor_target_wells = set(parse_list_arg(str(getattr(args, "positive_floor_target_wells", ""))))
    positive_floor_enabled = (not positive_floor_target_wells) or (str(artifact_target) in positive_floor_target_wells)
    pre_type_calibration_strategy = fit_type_calibration_strategy(
        train_df=train_df,
        train_pred_count_float=train_pred_count_float_after_gate,
        mode=str(getattr(args, "type_calibration_mode", "none")),
        scale_min=float(getattr(args, "type_calibration_scale_min", 0.4)),
        scale_max=float(getattr(args, "type_calibration_scale_max", 1.3)),
    )
    if (
        effective_count_train_mode == "positive_only"
        and validate_gate_mode(str(gate_info.get("mode", "none"))) != "none"
        and positive_floor_enabled
    ):
        train_type_labels = assign_segment_type_labels(train_df, pre_type_calibration_strategy)
        train_positive_floor_mask = (
            (train_gate_label > 0)
            & (train_gate_prob >= positive_floor_prob_min)
            & np.isin(train_type_labels, np.asarray(positive_floor_allowed_labels, dtype=object))
        )
        train_pred_count_float_after_gate = train_pred_count_float_after_gate.copy()
        train_pred_count_float_after_gate[train_positive_floor_mask] = np.maximum(
            train_pred_count_float_after_gate[train_positive_floor_mask],
            1.0,
        )
    count_bin_strategy = fit_count_bin_strategy(
        train_df=train_df,
        feature_cols=feature_cols,
        mode=str(getattr(args, "count_bin_mode", "none")),
        blend_weight=float(getattr(args, "count_bin_blend_weight", 0.7)),
    )
    train_pred_count_float_after_bin, _, _, _ = apply_count_bin_strategy(
        train_pred_count_float_after_gate,
        train_df,
        count_bin_strategy,
        gate_label=train_gate_label,
    )
    type_calibration_strategy = fit_type_calibration_strategy(
        train_df=train_df,
        train_pred_count_float=train_pred_count_float_after_bin,
        mode=str(getattr(args, "type_calibration_mode", "none")),
        scale_min=float(getattr(args, "type_calibration_scale_min", 0.4)),
        scale_max=float(getattr(args, "type_calibration_scale_max", 1.3)),
    )
    density_strength_strategy = fit_density_strength_strategy(
        train_df=train_df,
        mode=str(getattr(args, "density_strength_mode", "none")),
        basis_col=str(getattr(args, "density_strength_basis_col", "ProbMassPerLength")),
        scale_min=float(getattr(args, "density_strength_scale_min", 0.5)),
        scale_max=float(getattr(args, "density_strength_scale_max", 3.0)),
        feature_cols=feature_cols,
        hgb_learning_rate=float(args.hgb_learning_rate),
        hgb_max_depth=int(args.hgb_max_depth),
        model_max_iter=int(args.model_max_iter),
        regression_train_mode=str(getattr(args, "density_strength_regression_train_mode", "all_segments")),
        regression_target_transform=str(
            getattr(args, "density_strength_regression_target_transform", "none")
        ),
        regression_clip_quantile=float(
            getattr(args, "density_strength_regression_clip_quantile", np.nan)
        ),
        regression_calibration_mode=str(
            getattr(args, "density_strength_regression_calibration_mode", "none")
        ),
        regression_calibration_quantiles=str(
            getattr(
                args,
                "density_strength_regression_calibration_quantiles",
                "0.1,0.25,0.5,0.75,0.9,0.95",
            )
        ),
    )
    orientation_strategy = fit_orientation_strategy(
        train_df=train_df,
        feature_cols=feature_cols,
        mode=str(getattr(args, "orientation_mode", "none")),
        num_families=int(getattr(args, "orientation_num_families", 3)),
        min_points_per_segment=int(getattr(args, "orientation_min_points_per_segment", 1)),
    )
    density_point_floor_info = build_density_point_floor_info(density_strength_strategy)
    for row in gate_coef_rows:
        gate_row = dict(row)
        gate_row["val_well"] = f"{artifact_target}__gate"
        coef_rows.append(gate_row)

    artifact = {
        "artifact_version": 2,
        "artifact_type": artifact_type,
        "artifact_target": artifact_target,
        "train_wells": list(train_wells),
        "feature_cols": list(feature_cols),
        "model_name": str(args.model_name),
        "count_train_mode": count_train_mode,
        "count_train_mode_effective": effective_count_train_mode,
        "positive_floor_prob_min": positive_floor_prob_min,
        "positive_floor_allowed_labels": list(positive_floor_allowed_labels),
        "positive_floor_target_wells": sorted(positive_floor_target_wells),
        "positive_floor_enabled": bool(positive_floor_enabled),
        "num_count_train_segments": int(len(count_train_df)),
        "num_count_train_positive_segments": int(
            count_train_df["HasRawPoint"].fillna(0.0).astype(np.int64).clip(lower=0, upper=1).sum()
        ),
        "count_bin_strategy": count_bin_strategy,
        "type_calibration_strategy": type_calibration_strategy,
        "density_strength_strategy": density_strength_strategy,
        "orientation_strategy": orientation_strategy,
        "density_point_floor_info": density_point_floor_info,
        "model_alpha": float(args.model_alpha),
        "model_max_iter": int(args.model_max_iter),
        "tweedie_power": float(args.tweedie_power),
        "hgb_learning_rate": float(args.hgb_learning_rate),
        "hgb_max_depth": int(args.hgb_max_depth),
        "xgb_n_estimators": int(args.xgb_n_estimators),
        "xgb_max_depth": int(args.xgb_max_depth),
        "xgb_learning_rate": float(args.xgb_learning_rate),
        "xgb_subsample": float(args.xgb_subsample),
        "xgb_colsample_bytree": float(args.xgb_colsample_bytree),
        "fallback_count_mode": fallback_count_mode,
        "fallback_basis_col": fallback_basis_col,
        "fallback_scale": float(fallback_scale) if np.isfinite(fallback_scale) else np.nan,
        "count_fusion_mode": str(args.count_fusion_mode),
        "fusion_learned_weight": float(args.fusion_learned_weight),
        "global_post_scale_info": post_scale_info,
        "selective_post_strategy": selective_post_strategy,
        "gate_mode": str(gate_info.get("mode", "none")),
        "gate_prob_threshold": float(gate_info.get("prob_threshold", getattr(args, "gate_prob_threshold", 0.5))),
        "gate_constant_prob": float(gate_info.get("constant_prob", np.nan))
        if np.isfinite(gate_info.get("constant_prob", np.nan))
        else np.nan,
        "gate_positive_class": int(gate_info.get("positive_class", 1)),
        "pred_min_points_per_segment": int(args.pred_min_points_per_segment),
        "rounding_mode": str(args.rounding_mode),
        "prob_weight_gamma": float(args.prob_weight_gamma),
        "boundary_expand_mode": validate_boundary_expand_mode(str(args.boundary_expand_mode)),
        "boundary_expand_prob_min": float(args.boundary_expand_prob_min),
        "boundary_expand_max_steps": int(args.boundary_expand_max_steps),
        "boundary_expand_max_depth": float(args.boundary_expand_max_depth),
        "gt_dev_rule": str(args.gt_dev_rule),
        "strata_name": str(args.strata_name or "").strip(),
        "strata_range_csv": str(args.strata_range_csv or "").strip(),
        "soft_negative_weight": float(args.soft_negative_weight),
        "train_distance_weight_mode": str(distance_weight_info.get("mode", "none")),
        "train_distance_weight_power": float(distance_weight_info.get("weight_power", getattr(args, "train_distance_weight_power", 1.0))),
        "train_distance_weight_min": float(distance_weight_info.get("weight_min", getattr(args, "train_distance_weight_min", 1.0))),
        "train_distance_weight_max": float(distance_weight_info.get("weight_max", getattr(args, "train_distance_weight_max", 1.0))),
        "train_distance_csv": str(train_distance_csv) if train_distance_csv is not None else "",
        "train_distance_weight_map": dict(distance_weight_info.get("well_weights", {})),
        "train_distance_map": dict(distance_weight_info.get("well_distances", {})),
        "exist_exp_dir": str(args.exist_exp_dir),
        "sample_dir": str(args.sample_dir),
        "raw_label_dir": str(args.raw_label_dir),
        "strata_name": str(args.strata_name or "").strip(),
        "strata_range_csv": str(args.strata_range_csv or "").strip(),
        "gate_info": gate_info,
        "regressor": regressor,
    }
    return artifact, coef_rows


def predict_segments_with_artifact(
    segment_df: pd.DataFrame,
    artifact: dict,
    include_gt_columns: bool,
) -> tuple[pd.DataFrame, np.ndarray]:
    feature_cols = list(artifact["feature_cols"])
    validate_feature_cols(feature_cols, segment_df)
    x_val = segment_df[feature_cols].fillna(0.0)

    learned_pred_count_float = np.zeros(len(segment_df), dtype=np.float64)
    regressor = artifact.get("regressor")
    if str(artifact.get("model_name")) != "rule_only":
        if regressor is None:
            raise ValueError("Artifact does not contain a fitted regressor")
        learned_pred_count_float = np.clip(regressor.predict(x_val), a_min=0.0, a_max=None)

    fallback_basis_col = str(artifact.get("fallback_basis_col") or "")
    fallback_scale = artifact.get("fallback_scale", np.nan)
    fallback_pred_count_float = None
    if fallback_basis_col:
        if fallback_basis_col not in segment_df.columns:
            raise ValueError(f"Missing fallback basis col in inference dataset: {fallback_basis_col}")
        scale_value = float(fallback_scale) if np.isfinite(fallback_scale) else 1.0
        fallback_pred_count_float = np.clip(
            segment_df[fallback_basis_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64) * scale_value,
            a_min=0.0,
            a_max=None,
        )

    if str(artifact.get("model_name")) == "rule_only":
        if fallback_pred_count_float is None:
            raise ValueError("Saved rule_only artifact requires fallback basis data for inference")
        pred_count_float_before_post_scale = fallback_pred_count_float
    else:
        pred_count_float_before_post_scale = fuse_count_predictions(
            learned_pred=learned_pred_count_float,
            fallback_pred=fallback_pred_count_float,
            fusion_mode=str(artifact.get("count_fusion_mode", "learned_only")),
            learned_weight=float(artifact.get("fusion_learned_weight", 0.7)),
        )

    post_scale_info = artifact.get("global_post_scale_info", {"mode": "none", "scale": 1.0})
    pred_count_float_after_global_scale = apply_global_post_scale(
        pred_count_float_before_post_scale,
        post_scale_info,
    )
    selective_post_strategy = artifact.get(
        "selective_post_strategy",
        {
            "mode": "none",
            "feature_name": "",
            "feature_name_2": "",
            "quantile": np.nan,
            "quantile_2": np.nan,
            "threshold": np.nan,
            "threshold_2": np.nan,
            "lowconf_scale": 1.0,
        },
    )
    pred_count_float, lowconf_mask = apply_selective_post_strategy(
        pred_count_float_after_global_scale,
        segment_df,
        selective_post_strategy,
        pred_count_float_for_protection=pred_count_float_before_post_scale,
    )
    gate_info = dict(artifact.get("gate_info") or {})
    if not gate_info:
        gate_info = {
            "mode": artifact.get("gate_mode", "none"),
            "prob_threshold": artifact.get("gate_prob_threshold", 0.5),
            "constant_prob": artifact.get("gate_constant_prob", np.nan),
            "positive_class": artifact.get("gate_positive_class", 1),
            "feature_cols": feature_cols,
            "model": artifact.get("gate_model"),
        }
    gated_pred_count_float, gate_prob, gate_label = apply_gate_to_count_predictions(
        pred_count_float,
        segment_df,
        gate_info,
    )
    type_calibration_strategy = dict(artifact.get("type_calibration_strategy") or {"mode": "none"})
    type_labels_before_floor = assign_segment_type_labels(segment_df, type_calibration_strategy)
    count_floor_mask = np.zeros(len(segment_df), dtype=bool)
    if (
        validate_count_train_mode(
            str(artifact.get("count_train_mode_effective", artifact.get("count_train_mode", "all_segments")))
        )
        == "positive_only"
        and validate_gate_mode(str(gate_info.get("mode", "none"))) != "none"
        and bool(artifact.get("positive_floor_enabled", True))
    ):
        positive_floor_prob_min = float(np.clip(float(artifact.get("positive_floor_prob_min", 0.7)), 0.0, 1.0))
        positive_floor_allowed_labels = resolve_positive_floor_allowed_labels(
            ",".join(artifact.get("positive_floor_allowed_labels", []))
        )
        count_floor_mask = (
            (gate_label > 0)
            & (gate_prob >= positive_floor_prob_min)
            & np.isin(type_labels_before_floor, np.asarray(positive_floor_allowed_labels, dtype=object))
        )
        gated_pred_count_float = gated_pred_count_float.copy()
        gated_pred_count_float[count_floor_mask] = np.maximum(gated_pred_count_float[count_floor_mask], 1.0)
    count_bin_strategy = dict(artifact.get("count_bin_strategy") or {"mode": "none"})
    bin_blended_pred_count_float, bin_expected_count_float, count_bin_labels, count_bin_conf = apply_count_bin_strategy(
        gated_pred_count_float,
        segment_df,
        count_bin_strategy,
        gate_label=gate_label,
    )
    final_pred_count_float, type_labels, type_scales = apply_type_calibration(
        bin_blended_pred_count_float,
        segment_df,
        type_calibration_strategy,
    )
    density_strength_strategy = dict(artifact.get("density_strength_strategy") or {"mode": "none"})
    pred_density_strength, pred_density_size_scale, density_strength_labels, density_strength_type_scales = apply_density_strength_strategy(
        segment_df,
        density_strength_strategy,
    )
    pred_p10_mean, pred_p10_mass_per_length, pred_p10_mass = resolve_density_prediction_outputs(
        pred_density_strength=pred_density_strength,
        segment_df=segment_df,
        strategy=density_strength_strategy,
    )
    orientation_strategy = dict(artifact.get("orientation_strategy") or {"mode": "none"})
    pred_orientation_family, pred_orientation_confidence, pred_azimuth, pred_dip = apply_orientation_strategy(
        segment_df,
        orientation_strategy,
    )

    pred_segment_df = segment_df.copy()
    pred_segment_df["LearnedPredPointCountFloat"] = learned_pred_count_float
    pred_segment_df["FallbackCountBasisCol"] = fallback_basis_col
    pred_segment_df["FallbackCountScale"] = fallback_scale
    pred_segment_df["FallbackPredPointCountFloat"] = (
        fallback_pred_count_float if fallback_pred_count_float is not None else np.nan
    )
    pred_segment_df["PredPointCountFloatBeforePostScale"] = pred_count_float_before_post_scale
    pred_segment_df["GlobalPostScaleMode"] = post_scale_info.get("mode", "none")
    pred_segment_df["GlobalPostScaleRawScale"] = post_scale_info.get("raw_scale", np.nan)
    pred_segment_df["GlobalPostScaleAppliedScale"] = post_scale_info.get("scale", 1.0)
    pred_segment_df["GlobalPostScaleTrainTargetSum"] = post_scale_info.get("train_target_sum", np.nan)
    pred_segment_df["GlobalPostScaleTrainPredSum"] = post_scale_info.get("train_pred_sum", np.nan)
    pred_segment_df["SelectivePostMode"] = selective_post_strategy.get("mode", "none")
    pred_segment_df["SelectivePostFeature"] = selective_post_strategy.get("feature_name", "")
    pred_segment_df["SelectivePostFeature2"] = selective_post_strategy.get("feature_name_2", "")
    pred_segment_df["SelectivePostThreshold"] = selective_post_strategy.get("threshold", np.nan)
    pred_segment_df["SelectivePostThreshold2"] = selective_post_strategy.get("threshold_2", np.nan)
    pred_segment_df["SelectivePostQuantile"] = selective_post_strategy.get("quantile", np.nan)
    pred_segment_df["SelectivePostQuantile2"] = selective_post_strategy.get("quantile_2", np.nan)
    pred_segment_df["SelectivePostLowconfScale"] = selective_post_strategy.get("lowconf_scale", 1.0)
    pred_segment_df["SelectivePostMaxPredFloat"] = selective_post_strategy.get("max_pred_float", np.nan)
    pred_segment_df["SelectivePostIsLowConf"] = lowconf_mask.astype(int)
    pred_segment_df["PredPointCountFloatBeforeSelectivePost"] = pred_count_float_after_global_scale
    pred_segment_df["PredPointCountFloatBeforeGate"] = pred_count_float
    pred_segment_df["GateMode"] = gate_info.get("mode", "none")
    pred_segment_df["GateProbThreshold"] = gate_info.get("prob_threshold", np.nan)
    pred_segment_df["GatePredProb"] = gate_prob
    pred_segment_df["GatePredLabel"] = gate_label
    pred_segment_df["PositiveOnlyCountFloorApplied"] = count_floor_mask.astype(int)
    pred_segment_df["PositiveFloorProbMin"] = artifact.get("positive_floor_prob_min", 0.7)
    pred_segment_df["PositiveFloorAllowedLabels"] = ",".join(artifact.get("positive_floor_allowed_labels", []))
    pred_segment_df["PositiveFloorTargetWells"] = ",".join(artifact.get("positive_floor_target_wells", []))
    pred_segment_df["PositiveFloorEnabled"] = int(bool(artifact.get("positive_floor_enabled", True)))
    pred_segment_df["TypeCalibrationLabelBeforeFloor"] = type_labels_before_floor
    pred_segment_df["PredPointCountFloatBeforeCountBin"] = gated_pred_count_float
    pred_segment_df["CountBinMode"] = count_bin_strategy.get("mode", "none")
    pred_segment_df["CountBinBlendWeight"] = float(count_bin_strategy.get("blend_weight", np.nan))
    pred_segment_df["CountBinExpectedFloat"] = bin_expected_count_float
    pred_segment_df["CountBinPredLabel"] = count_bin_labels
    pred_segment_df["CountBinConfidence"] = count_bin_conf
    pred_segment_df["PredPointCountFloatBeforeTypeCalibration"] = bin_blended_pred_count_float
    pred_segment_df["TypeCalibrationMode"] = type_calibration_strategy.get("mode", "none")
    pred_segment_df["TypeCalibrationLabel"] = type_labels
    pred_segment_df["TypeCalibrationScale"] = type_scales
    pred_segment_df["PredPointCountFloat"] = final_pred_count_float
    pred_segment_df["DensityStrengthMode"] = density_strength_strategy.get("mode", "none")
    pred_segment_df["DensityStrengthBasisCol"] = density_strength_strategy.get("basis_col", "")
    pred_segment_df["DensityStrengthTargetCol"] = density_strength_strategy.get("target_col", "")
    pred_segment_df["DensityStrengthGlobalScale"] = density_strength_strategy.get("applied_scale", np.nan)
    pred_segment_df["DensityStrengthTypeScale"] = density_strength_type_scales
    pred_segment_df["PredDensityStrength"] = pred_density_strength
    pred_segment_df["PredP10Mean"] = pred_p10_mean
    pred_segment_df["PredP10MassPerLength"] = pred_p10_mass_per_length
    pred_segment_df["PredP10Mass"] = pred_p10_mass
    pred_segment_df["PredDensitySizeScale"] = pred_density_size_scale
    pred_segment_df["PredDensityStrengthLevel"] = density_strength_labels
    pred_segment_df["OrientationMode"] = orientation_strategy.get("mode", "none")
    pred_segment_df["PredOrientationFamily"] = pred_orientation_family
    pred_segment_df["PredOrientationConfidence"] = pred_orientation_confidence
    pred_segment_df["PredAzimuth"] = pred_azimuth
    pred_segment_df["PredDip"] = pred_dip
    pred_point_count_base = pred_segment_df["PredPointCountFloat"].apply(
        lambda value: compute_pred_count(
            raw_value=float(value),
            pred_min_points_per_segment=int(artifact.get("pred_min_points_per_segment", 0)),
            rounding_mode=str(artifact.get("rounding_mode", "round")),
        )
    ).to_numpy(dtype=np.int64)
    density_point_floor_info = dict(artifact.get("density_point_floor_info") or {})
    pred_point_count, density_point_floor_mask = apply_density_point_floor(
        pred_point_count=pred_point_count_base,
        pred_p10_mass_per_length=pred_p10_mass_per_length,
        floor_info=density_point_floor_info,
    )
    pred_segment_df["PredPointCountBeforeDensityPointFloor"] = pred_point_count_base
    pred_segment_df["DensityPointFloorEnabled"] = int(bool(density_point_floor_info.get("enabled", False)))
    pred_segment_df["DensityPointFloorThresholdMassPerLength"] = density_point_floor_info.get(
        "threshold_mass_per_length",
        np.nan,
    )
    pred_segment_df["DensityPointFloorSourceQuantile"] = density_point_floor_info.get("source_quantile", "")
    pred_segment_df["DensityPointFloorApplied"] = density_point_floor_mask.astype(int)
    pred_segment_df["PredPointCount"] = pred_point_count
    pred_segment_df = apply_pred_point_count_sample_cap(pred_segment_df)
    if include_gt_columns and "RawGTCountInPredSegment" in pred_segment_df.columns:
        pred_segment_df["SegmentCountAbsError"] = (
            pred_segment_df["PredPointCount"] - pred_segment_df["RawGTCountInPredSegment"]
        ).abs()
    else:
        pred_segment_df["SegmentCountAbsError"] = np.nan
    return pred_segment_df, lowconf_mask


def save_model_artifact(model_dir: Path, artifact: dict) -> None:
    model_dir.mkdir(parents=True, exist_ok=True)
    joblib.dump(artifact, model_dir / MODEL_ARTIFACT_FILENAME)

    meta = dict(artifact)
    regressor = meta.pop("regressor", None)
    meta["regressor_class"] = None if regressor is None else regressor.__class__.__name__
    gate_info = dict(meta.get("gate_info") or {})
    gate_model = gate_info.pop("model", None)
    if gate_info:
        meta["gate_info"] = gate_info
    meta["gate_model_class"] = None if gate_model is None else gate_model.__class__.__name__
    count_bin_strategy = dict(meta.get("count_bin_strategy") or {})
    count_bin_model = count_bin_strategy.pop("model", None)
    if count_bin_strategy:
        meta["count_bin_strategy"] = count_bin_strategy
    meta["count_bin_model_class"] = None if count_bin_model is None else count_bin_model.__class__.__name__
    density_strength_strategy = dict(meta.get("density_strength_strategy") or {})
    density_strength_model = density_strength_strategy.pop("model", None)
    if density_strength_strategy:
        meta["density_strength_strategy"] = density_strength_strategy
    meta["density_strength_model_class"] = (
        None if density_strength_model is None else density_strength_model.__class__.__name__
    )
    orientation_strategy = dict(meta.get("orientation_strategy") or {})
    orientation_classifier = orientation_strategy.pop("classifier", None)
    if orientation_strategy:
        meta["orientation_strategy"] = orientation_strategy
    meta["orientation_classifier_class"] = (
        None if orientation_classifier is None else orientation_classifier.__class__.__name__
    )
    with (model_dir / MODEL_META_FILENAME).open("w", encoding="utf-8") as file_obj:
        json.dump(to_jsonable(meta), file_obj, ensure_ascii=False, indent=2)


def load_model_artifact(model_dir: Path) -> dict:
    artifact_path = model_dir / MODEL_ARTIFACT_FILENAME
    if not artifact_path.exists():
        raise FileNotFoundError(f"Saved model artifact not found: {artifact_path}")
    artifact = joblib.load(artifact_path)
    if not isinstance(artifact, dict):
        raise ValueError(f"Unexpected artifact payload in {artifact_path}")
    return artifact


def resolve_saved_model_root(args, save_dir: Path) -> Path:
    return Path(args.saved_model_root) if args.saved_model_root else save_dir


def run_loo_eval(args) -> int:
    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / sanitize(args.exp_id))
    save_dir.mkdir(parents=True, exist_ok=True)
    summary_csv = save_dir / "raw_point_guided_refine_summary.csv"

    verify_logs = dataset_builder.POINT.list_verify_logs(str(args.exist_exp_dir))
    selected_wells = parse_list_arg(args.well_names) if args.well_names else sorted(verify_logs.keys())
    missing_wells = [well for well in selected_wells if well not in verify_logs]
    if missing_wells:
        raise ValueError(f"Wells not found in verify logs: {missing_wells}")

    feature_cols = resolve_feature_cols(str(args.feature_preset), str(args.feature_cols))
    fallback_count_mode = validate_optional_count_mode(str(args.fallback_count_mode))
    global_post_scale_mode = validate_global_post_scale_mode(str(args.global_post_scale_mode))
    selective_post_mode = validate_selective_post_mode(str(args.selective_post_mode))
    selective_post_target_wells = set(parse_list_arg(args.selective_post_target_wells))
    gate_mode = validate_gate_mode(str(args.gate_mode))
    gate_target_wells = set(parse_list_arg(getattr(args, "gate_target_wells", "")))
    boundary_expand_mode = validate_boundary_expand_mode(str(args.boundary_expand_mode))
    train_distance_weight_mode = validate_train_distance_weight_mode(str(args.train_distance_weight_mode))
    train_distance_csv, train_distance_df = resolve_train_distance_resources(args)
    strata_name = str(args.strata_name or "").strip()
    strata_range_df = load_strata_range_df(str(args.strata_range_csv or ""))
    config = dataset_builder.make_config(
        exist_exp_dir=str(args.exist_exp_dir),
        prob_weight_gamma=float(args.prob_weight_gamma),
        gt_min_points_per_segment=1,
        gt_rounding_mode="round",
        density_gt_col="P10",
    )

    per_well = {}
    all_rows = []
    dataset_save_dir = save_dir / "segment_dataset"
    dataset_save_dir.mkdir(parents=True, exist_ok=True)
    for well_name in selected_wells:
        well_data = build_well_segment_dataset(
            exist_csv=verify_logs[well_name],
            well_name=well_name,
            config=config,
            sample_dir=Path(args.sample_dir),
            raw_label_dir=Path(args.raw_label_dir),
            gt_dev_rule=str(args.gt_dev_rule),
            strata_name=strata_name,
            strata_range_df=strata_range_df,
            boundary_expand_mode=boundary_expand_mode,
            boundary_expand_prob_min=float(args.boundary_expand_prob_min),
            boundary_expand_max_steps=int(args.boundary_expand_max_steps),
            boundary_expand_max_depth=float(args.boundary_expand_max_depth),
        )
        well_data["segment_df"] = compute_target_weights(
            well_data["segment_df"],
            soft_negative_weight=float(args.soft_negative_weight),
        )
        per_well[well_name] = well_data
        well_dir = dataset_save_dir / well_name
        well_dir.mkdir(parents=True, exist_ok=True)
        well_data["segment_df"].to_csv(well_dir / "segment_dataset.csv", index=False, encoding="utf-8-sig")
        well_data["raw_gt_points"].to_csv(well_dir / "raw_gt_points.csv", index=False, encoding="utf-8-sig")
        well_data["gt_dev_segments"].to_csv(well_dir / "gt_dev_segments.csv", index=False, encoding="utf-8-sig")
        if not well_data["segment_df"].empty:
            all_rows.append(well_data["segment_df"])

    dataset_df = pd.concat(all_rows, ignore_index=True) if all_rows else pd.DataFrame()
    if dataset_df.empty:
        raise ValueError("Segment dataset is empty")
    dataset_df.to_csv(dataset_save_dir / "all_segment_dataset.csv", index=False, encoding="utf-8-sig")
    validate_feature_cols(feature_cols, dataset_df)

    summary_rows = []
    coef_rows = []
    saved_model_root = resolve_saved_model_root(args, save_dir)
    if int(args.save_model_artifacts):
        saved_model_root.mkdir(parents=True, exist_ok=True)

    for well_name in selected_wells:
        val_df = per_well[well_name]["segment_df"].copy()
        if val_df.empty:
            continue
        train_parts = [per_well[other]["segment_df"] for other in selected_wells if other != well_name]
        train_parts = [part for part in train_parts if not part.empty]
        if not train_parts:
            raise ValueError(f"No training segments available for {well_name}")
        train_df = pd.concat(train_parts, ignore_index=True)
        train_wells = [other for other in selected_wells if other != well_name]
        distance_weight_info = build_train_distance_weight_info(
            target_well=well_name,
            train_wells=train_wells,
            train_distance_weight_mode=train_distance_weight_mode,
            distance_df=train_distance_df,
            weight_power=float(args.train_distance_weight_power),
            weight_min=float(args.train_distance_weight_min),
            weight_max=float(args.train_distance_weight_max),
        )
        train_df = apply_train_distance_weight_info(train_df, distance_weight_info)

        selective_post_mode_this_well = (
            selective_post_mode
            if (not selective_post_target_wells or well_name in selective_post_target_wells)
            else "none"
        )
        gate_mode_this_well = (
            gate_mode
            if (not gate_target_wells or well_name in gate_target_wells)
            else "none"
        )
        artifact, artifact_coef_rows = fit_segment_refine_artifact(
            train_df=train_df,
            feature_cols=feature_cols,
            args=args,
            fallback_count_mode=fallback_count_mode,
            global_post_scale_mode=global_post_scale_mode,
            selective_post_mode_this_well=selective_post_mode_this_well,
            gate_mode_this_well=gate_mode_this_well,
            artifact_type="loo_split_model",
            artifact_target=well_name,
            train_wells=train_wells,
            distance_weight_info=distance_weight_info,
            train_distance_csv=train_distance_csv,
        )
        coef_rows.extend(artifact_coef_rows)

        pred_segment_df, lowconf_mask = predict_segments_with_artifact(
            segment_df=val_df,
            artifact=artifact,
            include_gt_columns=True,
        )

        well_data = per_well[well_name]
        pred_points = build_pred_points_from_segment_df(
            df_sorted=well_data["df_sorted"],
            payloads=well_data["payloads"],
            pred_segment_df=pred_segment_df,
            config=config,
            orientation_strategy=dict(artifact.get("orientation_strategy") or {}),
        )
        raw_gt_points = well_data["raw_gt_points"]
        raw_gt_depths = well_data["raw_gt_depths"]
        pred_depths = (
            pred_points[config.depth_col].to_numpy(dtype=np.float64)
            if not pred_points.empty
            else np.array([], dtype=np.float64)
        )
        pred_to_gt = dataset_builder.POINT.nearest_distance_stats(pred_depths, raw_gt_depths)
        gt_to_pred = dataset_builder.POINT.nearest_distance_stats(raw_gt_depths, pred_depths)
        orientation_error_stats = compute_orientation_error_stats(pred_segment_df)
        density_eval_mask = (
            pred_segment_df["PredP10MassPerLength"].notna()
            & pred_segment_df["RawP10MassPerLength"].notna()
        ) if {"PredP10MassPerLength", "RawP10MassPerLength"}.issubset(pred_segment_df.columns) else pd.Series([], dtype=bool)
        density_mass_per_length_mae = (
            float(
                (
                    pred_segment_df.loc[density_eval_mask, "PredP10MassPerLength"]
                    - pred_segment_df.loc[density_eval_mask, "RawP10MassPerLength"]
                ).abs().mean()
            )
            if len(density_eval_mask) > 0 and bool(np.any(density_eval_mask.to_numpy(dtype=bool)))
            else np.nan
        )

        well_save_dir = save_dir / well_name
        well_save_dir.mkdir(parents=True, exist_ok=True)
        raw_gt_points.to_csv(well_save_dir / "raw_gt_points.csv", index=False, encoding="utf-8-sig")
        well_data["gt_dev_segments"].to_csv(well_save_dir / "gt_dev_segments.csv", index=False, encoding="utf-8-sig")
        pred_points.to_csv(well_save_dir / "pred_fracture_points.csv", index=False, encoding="utf-8-sig")
        pred_segment_df.to_csv(well_save_dir / "pred_segment_summary.csv", index=False, encoding="utf-8-sig")
        train_df.to_csv(well_save_dir / "train_segment_dataset.csv", index=False, encoding="utf-8-sig")

        if int(args.save_model_artifacts):
            save_model_artifact(saved_model_root / format_verify_well_dir_name(well_name), artifact)

        metrics = {
            "well": well_name,
            "strata_name": strata_name,
            "config_profile": str(getattr(args, "config_profile", "")),
            "exist_csv": well_data["exist_csv"],
            "model_name": args.model_name,
            "model_alpha": args.model_alpha,
            "model_max_iter": args.model_max_iter,
            "tweedie_power": args.tweedie_power,
            "hgb_learning_rate": args.hgb_learning_rate,
            "hgb_max_depth": args.hgb_max_depth,
            "feature_preset": args.feature_preset,
            "feature_cols": feature_cols,
            "prob_weight_gamma": args.prob_weight_gamma,
            "count_train_mode": validate_count_train_mode(str(args.count_train_mode)),
            "count_train_mode_effective": artifact.get(
                "count_train_mode_effective",
                validate_count_train_mode(str(args.count_train_mode)),
            ),
            "positive_floor_prob_min": float(artifact.get("positive_floor_prob_min", getattr(args, "positive_floor_prob_min", 0.7))),
            "positive_floor_allowed_labels": artifact.get(
                "positive_floor_allowed_labels",
                resolve_positive_floor_allowed_labels(str(getattr(args, "positive_floor_allowed_labels", ""))),
            ),
            "positive_floor_target_wells": artifact.get(
                "positive_floor_target_wells",
                sorted(parse_list_arg(str(getattr(args, "positive_floor_target_wells", "")))),
            ),
            "positive_floor_enabled": bool(artifact.get("positive_floor_enabled", True)),
            "type_calibration_mode": artifact.get("type_calibration_strategy", {}).get("mode", "none"),
            "fallback_count_mode": fallback_count_mode,
            "count_fusion_mode": args.count_fusion_mode,
            "fusion_learned_weight": args.fusion_learned_weight,
            "global_post_scale_mode": global_post_scale_mode,
            "global_post_scale_raw_scale": pred_segment_df["GlobalPostScaleRawScale"].iloc[0] if not pred_segment_df.empty else np.nan,
            "global_post_scale_applied_scale": pred_segment_df["GlobalPostScaleAppliedScale"].iloc[0] if not pred_segment_df.empty else np.nan,
            "global_post_scale_train_target_sum": pred_segment_df["GlobalPostScaleTrainTargetSum"].iloc[0] if not pred_segment_df.empty else np.nan,
            "global_post_scale_train_pred_sum": pred_segment_df["GlobalPostScaleTrainPredSum"].iloc[0] if not pred_segment_df.empty else np.nan,
            "selective_post_mode": selective_post_mode_this_well,
            "selective_post_target_wells": sorted(selective_post_target_wells),
            "selective_post_feature": args.selective_post_feature,
            "selective_post_feature_2": args.selective_post_feature_2,
            "selective_post_quantile": args.selective_post_quantile,
            "selective_post_quantile_2": args.selective_post_quantile_2,
            "selective_post_threshold": pred_segment_df["SelectivePostThreshold"].iloc[0] if not pred_segment_df.empty else np.nan,
            "selective_post_threshold_2": pred_segment_df["SelectivePostThreshold2"].iloc[0] if not pred_segment_df.empty else np.nan,
            "selective_post_lowconf_scale": pred_segment_df["SelectivePostLowconfScale"].iloc[0] if not pred_segment_df.empty else np.nan,
            "selective_post_max_pred_float": pred_segment_df["SelectivePostMaxPredFloat"].iloc[0] if not pred_segment_df.empty else np.nan,
            "selective_post_num_lowconf_segments": int(np.sum(lowconf_mask.astype(np.int64))),
            "gate_mode": str(gate_mode_this_well),
            "gate_prob_threshold": float(getattr(args, "gate_prob_threshold", 0.5)),
            "gate_target_wells": sorted(gate_target_wells),
            "gate_num_positive_segments": int(pred_segment_df["GatePredLabel"].fillna(0).astype(np.int64).sum()) if "GatePredLabel" in pred_segment_df.columns else int(len(pred_segment_df)),
            "gate_num_zeroed_segments": int((pred_segment_df["GatePredLabel"].fillna(1).astype(np.int64) <= 0).sum()) if "GatePredLabel" in pred_segment_df.columns else 0,
            "gate_mean_prob": float(pred_segment_df["GatePredProb"].fillna(1.0).mean()) if "GatePredProb" in pred_segment_df.columns else 1.0,
            "count_bin_mode": str(artifact.get("count_bin_strategy", {}).get("mode", "none")),
            "count_bin_blend_weight": float(artifact.get("count_bin_strategy", {}).get("blend_weight", np.nan)),
            "count_bin_mean_confidence": float(pred_segment_df["CountBinConfidence"].fillna(0.0).mean()) if "CountBinConfidence" in pred_segment_df.columns else np.nan,
            "density_strength_mode": str(artifact.get("density_strength_strategy", {}).get("mode", "none")),
            "density_strength_basis_col": str(artifact.get("density_strength_strategy", {}).get("basis_col", "")),
            "density_strength_global_scale": float(artifact.get("density_strength_strategy", {}).get("applied_scale", np.nan)),
            "pred_density_strength_mean": float(pred_segment_df["PredDensityStrength"].fillna(0.0).mean()) if "PredDensityStrength" in pred_segment_df.columns else np.nan,
            "pred_density_strength_max": float(pred_segment_df["PredDensityStrength"].fillna(0.0).max()) if "PredDensityStrength" in pred_segment_df.columns else np.nan,
            "gt_density_strength_mean": float(pred_segment_df["RawP10MassPerLength"].fillna(0.0).mean()) if "RawP10MassPerLength" in pred_segment_df.columns else np.nan,
            "density_mass_per_length_mae": density_mass_per_length_mae,
            "density_point_floor_enabled": bool(artifact.get("density_point_floor_info", {}).get("enabled", False)),
            "density_point_floor_threshold_mass_per_length": float(
                artifact.get("density_point_floor_info", {}).get("threshold_mass_per_length", np.nan)
            ),
            "density_point_floor_source_quantile": str(
                artifact.get("density_point_floor_info", {}).get("source_quantile", "")
            ),
            "density_point_floor_segments": int(pred_segment_df["DensityPointFloorApplied"].sum())
            if "DensityPointFloorApplied" in pred_segment_df.columns
            else 0,
            "orientation_mode": str(artifact.get("orientation_strategy", {}).get("mode", "none")),
            "orientation_num_families": int(artifact.get("orientation_strategy", {}).get("num_families_fitted", 0)),
            "orientation_min_points_per_segment": int(artifact.get("orientation_strategy", {}).get("min_points_per_segment", 1)),
            "orientation_num_train_segments": int(artifact.get("orientation_strategy", {}).get("num_train_segments", 0)),
            "pred_orientation_conf_mean": float(pred_segment_df["PredOrientationConfidence"].fillna(0.0).mean()) if "PredOrientationConfidence" in pred_segment_df.columns else np.nan,
            "orientation_num_eval_segments": int(orientation_error_stats["orientation_num_eval_segments"]),
            "orientation_azimuth_mae_deg": orientation_error_stats["orientation_azimuth_mae_deg"],
            "orientation_dip_mae_deg": orientation_error_stats["orientation_dip_mae_deg"],
            "train_distance_weight_mode": train_distance_weight_mode,
            "train_distance_weight_power": float(args.train_distance_weight_power),
            "train_distance_weight_min": float(args.train_distance_weight_min),
            "train_distance_weight_max": float(args.train_distance_weight_max),
            "train_distance_csv": str(train_distance_csv) if train_distance_csv is not None else "",
            "train_distance_mean_weight": float(train_df["TrainWellDistanceWeight"].mean()),
            "train_distance_min_weight": float(train_df["TrainWellDistanceWeight"].min()),
            "train_distance_max_weight": float(train_df["TrainWellDistanceWeight"].max()),
            "train_distance_weight_map": distance_weight_info.get("well_weights", {}),
            "train_distance_map": distance_weight_info.get("well_distances", {}),
            "boundary_expand_mode": boundary_expand_mode,
            "boundary_expand_prob_min": float(args.boundary_expand_prob_min),
            "boundary_expand_max_steps": int(args.boundary_expand_max_steps),
            "boundary_expand_max_depth": float(args.boundary_expand_max_depth),
            "boundary_base_positive_samples": int(well_data["boundary_expand_stats"]["boundary_base_positive_samples"]),
            "boundary_effective_positive_samples": int(well_data["boundary_expand_stats"]["boundary_effective_positive_samples"]),
            "boundary_added_positive_samples": int(well_data["boundary_expand_stats"]["boundary_added_positive_samples"]),
            "boundary_base_segment_count": int(well_data["boundary_expand_stats"]["boundary_base_segment_count"]),
            "boundary_effective_segment_count": int(well_data["boundary_expand_stats"]["boundary_effective_segment_count"]),
            "boundary_added_left_samples": int(well_data["boundary_expand_stats"]["boundary_added_left_samples"]),
            "boundary_added_right_samples": int(well_data["boundary_expand_stats"]["boundary_added_right_samples"]),
            "soft_negative_weight": args.soft_negative_weight,
            "pred_min_points_per_segment": args.pred_min_points_per_segment,
            "rounding_mode": args.rounding_mode,
            "num_train_segments": int(len(train_df)),
            "num_train_segments_for_count": int(artifact.get("num_count_train_segments", len(train_df))),
            "num_train_positive_segments_for_count": int(
                artifact.get(
                    "num_count_train_positive_segments",
                    train_df["HasRawPoint"].fillna(0.0).astype(np.int64).clip(lower=0, upper=1).sum(),
                )
            ),
            "num_val_segments": int(len(val_df)),
            "positive_only_count_floor_segments": int(pred_segment_df["PositiveOnlyCountFloorApplied"].sum())
            if "PositiveOnlyCountFloorApplied" in pred_segment_df.columns
            else 0,
            "type_calibration_adjusted_segments": int(
                np.sum(np.abs(pred_segment_df["TypeCalibrationScale"].to_numpy(dtype=np.float64) - 1.0) > 1e-8)
            )
            if "TypeCalibrationScale" in pred_segment_df.columns
            else 0,
            "segment_count_mae": float(pred_segment_df["SegmentCountAbsError"].mean()),
            "segment_count_rmse": float(
                np.sqrt(
                    np.mean(
                        np.square(
                            pred_segment_df["PredPointCount"].to_numpy(dtype=np.float64)
                            - pred_segment_df["RawGTCountInPredSegment"].to_numpy(dtype=np.float64)
                        )
                    )
                )
            ),
            "N_gt_points": int(len(raw_gt_depths)),
            "N_pred_points": int(len(pred_depths)),
            "count_diff": int(len(pred_depths) - len(raw_gt_depths)),
            "count_error_pct": compute_count_error_pct(int(len(pred_depths) - len(raw_gt_depths)), int(len(raw_gt_depths))),
            "pred_to_gt_mean_dist": pred_to_gt["mean"],
            "pred_to_gt_median_dist": pred_to_gt["median"],
            "pred_to_gt_p90_dist": pred_to_gt["p90"],
            "gt_to_pred_mean_dist": gt_to_pred["mean"],
            "gt_to_pred_median_dist": gt_to_pred["median"],
            "gt_to_pred_p90_dist": gt_to_pred["p90"],
            "raw_points_covered_by_pred_segments": int(well_data["raw_points_covered_by_pred_segments"]),
            "raw_points_missed_outside_pred_segments": int(well_data["raw_points_missed_outside_pred_segments"]),
            "pred_segments_with_raw_gt_0": int((pred_segment_df["RawGTCountInPredSegment"] <= 0).sum()),
            "pred_segments_overlap_gt_dev_and_raw_gt_0": int(
                ((pred_segment_df["RawGTCountInPredSegment"] <= 0) & (pred_segment_df["GTDevOverlapFlag"] == 1)).sum()
            ),
            "saved_model_dir": str(
                (saved_model_root / format_verify_well_dir_name(well_name)) if int(args.save_model_artifacts) else ""
            ),
        }
        with open(well_save_dir / "metrics.json", "w", encoding="utf-8") as file_obj:
            json.dump(metrics, file_obj, ensure_ascii=False, indent=2)
        summary_rows.append(metrics)

    summary_df = pd.DataFrame(summary_rows)
    if summary_df.empty:
        raise ValueError("No validation summary rows were produced")
    summary_df.to_csv(summary_csv, index=False, encoding="utf-8-sig")
    pd.DataFrame(coef_rows).to_csv(save_dir / "loo_model_coefficients.csv", index=False, encoding="utf-8-sig")
    overall_metrics = {
        "num_wells": int(len(summary_df)),
        "total_gt_points": int(summary_df["N_gt_points"].fillna(0).sum()),
        "total_pred_points": int(summary_df["N_pred_points"].fillna(0).sum()),
        "total_abs_count_error": int(summary_df["count_diff"].abs().fillna(0).sum()),
        "overall_count_error_pct": compute_count_error_pct(
            float(summary_df["count_diff"].abs().fillna(0).sum()),
            float(summary_df["N_gt_points"].fillna(0).sum()),
        ),
        "avg_count_error_pct": float(summary_df["count_error_pct"].dropna().mean()) if "count_error_pct" in summary_df.columns else np.nan,
    }
    with open(save_dir / "overall_metrics.json", "w", encoding="utf-8") as file_obj:
        json.dump(to_jsonable(overall_metrics), file_obj, ensure_ascii=False, indent=2)

    run_config = build_run_config_dict(
        args=args,
        selected_wells=selected_wells,
        feature_cols=feature_cols,
        fallback_count_mode=fallback_count_mode,
        global_post_scale_mode=global_post_scale_mode,
        selective_post_mode=selective_post_mode,
        selective_post_target_wells=selective_post_target_wells,
        gate_target_wells=gate_target_wells,
        train_distance_csv=train_distance_csv,
        save_dir=save_dir,
        summary_csv=summary_csv,
        saved_model_root=saved_model_root if int(args.save_model_artifacts) else None,
    )
    with open(save_dir / "config.json", "w", encoding="utf-8") as file_obj:
        json.dump(to_jsonable(run_config), file_obj, ensure_ascii=False, indent=2)

    with summary_csv.open("r", encoding="utf-8-sig", newline="") as file_obj:
        results = list(csv.DictReader(file_obj))
    append_result_to_docx(
        Path(args.docx_path),
        f"实验 {args.exp_id}",
        run_config,
        results,
    )

    print(
        json.dumps(
            {
                "exp_id": args.exp_id,
                "run_mode": args.run_mode,
                "save_dir": str(save_dir),
                "summary_csv": str(summary_csv),
                "saved_model_root": str(saved_model_root) if int(args.save_model_artifacts) else "",
                "overall_metrics": overall_metrics,
                "feature_cols": feature_cols,
                "results": results,
            },
            ensure_ascii=False,
        )
    )
    return 0


def run_predict(args) -> int:
    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / sanitize(args.exp_id))
    save_dir.mkdir(parents=True, exist_ok=True)
    saved_model_root = resolve_saved_model_root(args, save_dir)

    if not args.predict_exist_csv:
        raise ValueError("run_mode=predict requires --predict-exist-csv")
    if not args.predict_well_name:
        raise ValueError("run_mode=predict requires --predict-well-name")

    if args.load_model_dir:
        model_dir = Path(args.load_model_dir)
    else:
        if not args.load_model_well:
            raise ValueError("run_mode=predict requires --load-model-dir or --load-model-well")
        model_dir = saved_model_root / format_verify_well_dir_name(args.load_model_well)

    artifact = load_model_artifact(model_dir)
    strata_name = str(args.strata_name or artifact.get("strata_name", "") or "").strip()
    strata_range_csv = str(args.strata_range_csv or artifact.get("strata_range_csv", "") or "").strip()
    strata_range_df = load_strata_range_df(strata_range_csv)
    config = dataset_builder.make_config(
        exist_exp_dir=str(artifact.get("exist_exp_dir") or args.exist_exp_dir),
        prob_weight_gamma=float(artifact.get("prob_weight_gamma", args.prob_weight_gamma)),
        gt_min_points_per_segment=1,
        gt_rounding_mode="round",
        density_gt_col="P10",
    )
    boundary_expand_mode = validate_boundary_expand_mode(
        str(artifact.get("boundary_expand_mode", args.boundary_expand_mode))
    )
    boundary_expand_prob_min = float(artifact.get("boundary_expand_prob_min", args.boundary_expand_prob_min))
    boundary_expand_max_steps = int(artifact.get("boundary_expand_max_steps", args.boundary_expand_max_steps))
    boundary_expand_max_depth = float(artifact.get("boundary_expand_max_depth", args.boundary_expand_max_depth))
    well_data = build_predict_segment_dataset(
        exist_csv=str(args.predict_exist_csv),
        well_name=str(args.predict_well_name),
        config=config,
        strata_name=strata_name,
        strata_range_df=strata_range_df,
        boundary_expand_mode=boundary_expand_mode,
        boundary_expand_prob_min=boundary_expand_prob_min,
        boundary_expand_max_steps=boundary_expand_max_steps,
        boundary_expand_max_depth=boundary_expand_max_depth,
    )
    if well_data["segment_df"].empty:
        raise ValueError(f"No predicted fracture segments found in: {args.predict_exist_csv}")

    pred_segment_df, lowconf_mask = predict_segments_with_artifact(
        segment_df=well_data["segment_df"].copy(),
        artifact=artifact,
        include_gt_columns=False,
    )
    pred_points = build_pred_points_from_segment_df(
        df_sorted=well_data["df_sorted"],
        payloads=well_data["payloads"],
        pred_segment_df=pred_segment_df,
        config=config,
        orientation_strategy=dict(artifact.get("orientation_strategy") or {}),
    )

    inference_dir = save_dir / f"{args.predict_well_name}__with__{artifact.get('artifact_target', 'selected_model')}"
    inference_dir.mkdir(parents=True, exist_ok=True)
    pred_segment_df.to_csv(inference_dir / "pred_segment_summary.csv", index=False, encoding="utf-8-sig")
    pred_points.to_csv(inference_dir / "pred_fracture_points.csv", index=False, encoding="utf-8-sig")
    well_data["df"].to_csv(inference_dir / "input_predict_log.csv", index=False, encoding="utf-8-sig")

    inference_config = {
        "exp_id": args.exp_id,
        "run_mode": args.run_mode,
        "config_profile": str(getattr(args, "config_profile", "")),
        "predict_exist_csv": str(args.predict_exist_csv),
        "predict_well_name": str(args.predict_well_name),
        "strata_name": strata_name,
        "strata_range_csv": strata_range_csv,
        "load_model_dir": str(model_dir),
        "artifact_target": artifact.get("artifact_target"),
        "artifact_type": artifact.get("artifact_type"),
        "train_wells": artifact.get("train_wells", []),
        "train_distance_weight_mode": artifact.get("train_distance_weight_mode", "none"),
        "train_distance_weight_power": artifact.get("train_distance_weight_power", 1.0),
        "train_distance_weight_min": artifact.get("train_distance_weight_min", 1.0),
        "train_distance_weight_max": artifact.get("train_distance_weight_max", 1.0),
        "train_distance_csv": artifact.get("train_distance_csv", ""),
        "train_distance_weight_map": artifact.get("train_distance_weight_map", {}),
        "train_distance_map": artifact.get("train_distance_map", {}),
        "feature_cols": artifact.get("feature_cols", []),
        "count_train_mode": artifact.get("count_train_mode", "all_segments"),
        "count_train_mode_effective": artifact.get(
            "count_train_mode_effective",
            artifact.get("count_train_mode", "all_segments"),
        ),
        "positive_floor_prob_min": artifact.get("positive_floor_prob_min", 0.7),
        "positive_floor_allowed_labels": artifact.get("positive_floor_allowed_labels", []),
        "positive_floor_target_wells": artifact.get("positive_floor_target_wells", []),
        "positive_floor_enabled": bool(artifact.get("positive_floor_enabled", True)),
        "type_calibration_mode": artifact.get("type_calibration_strategy", {}).get("mode", "none"),
        "fallback_count_mode": artifact.get("fallback_count_mode", ""),
        "count_fusion_mode": artifact.get("count_fusion_mode", "learned_only"),
        "count_bin_mode": artifact.get("count_bin_strategy", {}).get("mode", "none"),
        "count_bin_blend_weight": artifact.get("count_bin_strategy", {}).get("blend_weight", np.nan),
        "density_strength_mode": artifact.get("density_strength_strategy", {}).get("mode", "none"),
        "density_strength_basis_col": artifact.get("density_strength_strategy", {}).get("basis_col", ""),
        "density_strength_global_scale": artifact.get("density_strength_strategy", {}).get("applied_scale", np.nan),
        "density_point_floor_enabled": artifact.get("density_point_floor_info", {}).get("enabled", False),
        "density_point_floor_threshold_mass_per_length": artifact.get("density_point_floor_info", {}).get("threshold_mass_per_length", np.nan),
        "density_point_floor_source_quantile": artifact.get("density_point_floor_info", {}).get("source_quantile", ""),
        "orientation_mode": artifact.get("orientation_strategy", {}).get("mode", "none"),
        "orientation_num_families": artifact.get("orientation_strategy", {}).get("num_families_fitted", 0),
        "orientation_min_points_per_segment": artifact.get("orientation_strategy", {}).get("min_points_per_segment", 1),
        "boundary_expand_mode": boundary_expand_mode,
        "boundary_expand_prob_min": boundary_expand_prob_min,
        "boundary_expand_max_steps": boundary_expand_max_steps,
        "boundary_expand_max_depth": boundary_expand_max_depth,
        "boundary_base_positive_samples": int(well_data["boundary_expand_stats"]["boundary_base_positive_samples"]),
        "boundary_effective_positive_samples": int(well_data["boundary_expand_stats"]["boundary_effective_positive_samples"]),
        "boundary_added_positive_samples": int(well_data["boundary_expand_stats"]["boundary_added_positive_samples"]),
        "boundary_base_segment_count": int(well_data["boundary_expand_stats"]["boundary_base_segment_count"]),
        "boundary_effective_segment_count": int(well_data["boundary_expand_stats"]["boundary_effective_segment_count"]),
        "selective_post_max_pred_float": artifact.get("selective_post_strategy", {}).get("max_pred_float", np.nan),
        "pred_min_points_per_segment": artifact.get("pred_min_points_per_segment", 0),
        "rounding_mode": artifact.get("rounding_mode", "round"),
        "num_segments": int(len(pred_segment_df)),
        "num_lowconf_segments": int(np.sum(lowconf_mask.astype(np.int64))),
        "num_pred_points": int(len(pred_points)),
        "pred_density_strength_mean": float(pred_segment_df["PredDensityStrength"].fillna(0.0).mean()) if "PredDensityStrength" in pred_segment_df.columns else np.nan,
        "density_point_floor_segments": int(pred_segment_df["DensityPointFloorApplied"].sum()) if "DensityPointFloorApplied" in pred_segment_df.columns else 0,
        "pred_orientation_conf_mean": float(pred_segment_df["PredOrientationConfidence"].fillna(0.0).mean()) if "PredOrientationConfidence" in pred_segment_df.columns else np.nan,
        "save_dir": str(inference_dir),
    }
    with (inference_dir / "inference_config.json").open("w", encoding="utf-8") as file_obj:
        json.dump(to_jsonable(inference_config), file_obj, ensure_ascii=False, indent=2)

    print(
        json.dumps(
            {
                "exp_id": args.exp_id,
                "run_mode": args.run_mode,
                "model_dir": str(model_dir),
                "predict_well_name": str(args.predict_well_name),
                "save_dir": str(inference_dir),
                "num_segments": int(len(pred_segment_df)),
                "num_pred_points": int(len(pred_points)),
                "pred_density_strength_mean": float(pred_segment_df["PredDensityStrength"].fillna(0.0).mean()) if "PredDensityStrength" in pred_segment_df.columns else np.nan,
                "density_point_floor_segments": int(pred_segment_df["DensityPointFloorApplied"].sum()) if "DensityPointFloorApplied" in pred_segment_df.columns else 0,
                "pred_orientation_conf_mean": float(pred_segment_df["PredOrientationConfidence"].fillna(0.0).mean()) if "PredOrientationConfidence" in pred_segment_df.columns else np.nan,
            },
            ensure_ascii=False,
        )
    )
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", required=True)
    parser.add_argument("--run-mode", default="loo_eval", choices=["loo_eval", "predict"])
    parser.add_argument("--config-profile", default="")
    parser.add_argument("--exist-exp-dir", default=str(DEFAULT_EXIST_EXP_DIR))
    parser.add_argument("--sample-dir", default=str(DEFAULT_SAMPLE_DIR))
    parser.add_argument("--raw-label-dir", default=str(DEFAULT_RAW_LABEL_DIR))
    parser.add_argument("--strata-name", default="")
    parser.add_argument("--strata-range-csv", default="")
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument("--result-dir")
    parser.add_argument("--well-names", default="")
    parser.add_argument("--gt-dev-rule", default="any_density")
    parser.add_argument("--prob-weight-gamma", type=float, default=2.0)
    parser.add_argument("--pred-min-points-per-segment", type=int, default=0)
    parser.add_argument("--rounding-mode", default="round")
    parser.add_argument("--fallback-count-mode", default="")
    parser.add_argument("--count-train-mode", default="all_segments")
    parser.add_argument("--positive-floor-prob-min", type=float, default=0.7)
    parser.add_argument("--positive-floor-allowed-labels", default="other,short_strong")
    parser.add_argument("--positive-floor-target-wells", default="")
    parser.add_argument("--type-calibration-mode", default="none")
    parser.add_argument("--type-calibration-scale-min", type=float, default=0.4)
    parser.add_argument("--type-calibration-scale-max", type=float, default=1.3)
    parser.add_argument("--count-fusion-mode", default="learned_only")
    parser.add_argument("--fusion-learned-weight", type=float, default=0.7)
    parser.add_argument("--global-post-scale-mode", default="none")
    parser.add_argument("--global-post-scale-min", type=float, default=0.5)
    parser.add_argument("--global-post-scale-max", type=float, default=1.5)
    parser.add_argument("--selective-post-mode", default="none")
    parser.add_argument("--selective-post-target-wells", default="")
    parser.add_argument("--selective-post-feature", default="ProbMassPerLength")
    parser.add_argument("--selective-post-feature-2", default="ProbMean")
    parser.add_argument("--selective-post-quantile", type=float, default=0.5)
    parser.add_argument("--selective-post-quantile-2", type=float, default=0.5)
    parser.add_argument("--selective-post-max-pred-float", type=float, default=-1.0)
    parser.add_argument("--soft-negative-weight", type=float, default=0.5)
    parser.add_argument("--train-distance-weight-mode", default="none")
    parser.add_argument("--train-distance-csv", default="")
    parser.add_argument("--train-distance-weight-power", type=float, default=1.0)
    parser.add_argument("--train-distance-weight-min", type=float, default=0.7)
    parser.add_argument("--train-distance-weight-max", type=float, default=1.3)
    parser.add_argument("--boundary-expand-mode", default="none")
    parser.add_argument("--boundary-expand-prob-min", type=float, default=0.35)
    parser.add_argument("--boundary-expand-max-steps", type=int, default=2)
    parser.add_argument("--boundary-expand-max-depth", type=float, default=0.0)
    parser.add_argument("--gate-mode", default="none")
    parser.add_argument("--gate-prob-threshold", type=float, default=0.5)
    parser.add_argument("--gate-target-wells", default="")
    parser.add_argument("--count-bin-mode", default="none")
    parser.add_argument("--count-bin-blend-weight", type=float, default=0.7)
    parser.add_argument("--density-strength-mode", default="none")
    parser.add_argument("--density-strength-basis-col", default="ProbMassPerLength")
    parser.add_argument("--density-strength-scale-min", type=float, default=0.5)
    parser.add_argument("--density-strength-scale-max", type=float, default=3.0)
    parser.add_argument("--density-strength-regression-train-mode", default="all_segments")
    parser.add_argument("--density-strength-regression-target-transform", default="none")
    parser.add_argument("--density-strength-regression-clip-quantile", type=float, default=np.nan)
    parser.add_argument("--density-strength-regression-calibration-mode", default="none")
    parser.add_argument(
        "--density-strength-regression-calibration-quantiles",
        default="0.1,0.25,0.5,0.75,0.9,0.95",
    )
    parser.add_argument("--orientation-mode", default="none")
    parser.add_argument("--orientation-num-families", type=int, default=3)
    parser.add_argument("--orientation-min-points-per-segment", type=int, default=1)
    parser.add_argument("--model-name", default="tweedie")
    parser.add_argument("--model-alpha", type=float, default=0.03)
    parser.add_argument("--model-max-iter", type=int, default=3000)
    parser.add_argument("--tweedie-power", type=float, default=1.5)
    parser.add_argument("--feature-preset", default="high_density_compact_v1")
    parser.add_argument("--feature-cols", default=",".join(dataset_builder.DEFAULT_FEATURE_COLS))
    parser.add_argument("--xgb-n-estimators", type=int, default=300)
    parser.add_argument("--xgb-max-depth", type=int, default=3)
    parser.add_argument("--xgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--xgb-subsample", type=float, default=0.8)
    parser.add_argument("--xgb-colsample-bytree", type=float, default=0.8)
    parser.add_argument("--hgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--hgb-max-depth", type=int, default=3)
    parser.add_argument("--save-model-artifacts", type=int, default=1)
    parser.add_argument("--saved-model-root", default="")
    parser.add_argument("--load-model-dir", default="")
    parser.add_argument("--load-model-well", default="")
    parser.add_argument("--predict-exist-csv", default="")
    parser.add_argument("--predict-well-name", default="")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    default_args = parser.parse_args(["--exp-id", "__profile_defaults__"])
    args.config_profile = apply_named_profile(args, getattr(args, "config_profile", ""), default_args=default_args)
    if args.run_mode == "predict":
        return run_predict(args)
    return run_loo_eval(args)


if __name__ == "__main__":
    raise SystemExit(main())
