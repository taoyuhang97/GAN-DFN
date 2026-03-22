import argparse
import csv
import json
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

import numpy as np
import pandas as pd
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from sklearn.ensemble import HistGradientBoostingRegressor
from sklearn.linear_model import PoissonRegressor, TweedieRegressor
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

import build_segment_count_dataset as dataset_builder

try:
    from xgboost import XGBRegressor
except ImportError:
    XGBRegressor = None


ROOT = Path(__file__).resolve().parent
BASE_SAVE_DIR = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容一\裂缝位置预测\基于密度的裂缝点位分析\裂缝点位精细化"
)
DEFAULT_DOCX_PATH = Path(r"D:\项目\石油开采\断缝储实验\实验记录20260319.docx")


def sanitize(text: str) -> str:
    return "".join(ch if ch.isalnum() or ch in {"_", "-", "中", "文"} else "_" for ch in text).strip("_")


def parse_list_arg(raw: str) -> list[str]:
    return [item.strip() for item in raw.split(",") if item.strip()]


def validate_optional_count_mode(mode: str) -> str:
    normalized = str(mode or "").strip()
    if not normalized:
        return ""
    if normalized not in dataset_builder.POINT.COUNT_MODE_TO_BASIS_COL:
        raise ValueError(
            f"Unsupported fallback_count_mode: {normalized}, "
            f"available={sorted(dataset_builder.POINT.COUNT_MODE_TO_BASIS_COL.keys())}"
        )
    return normalized


def fit_basis_scale(train_df: pd.DataFrame, count_mode: str) -> tuple[str, float]:
    basis_col = dataset_builder.POINT.COUNT_MODE_TO_BASIS_COL[count_mode]
    basis_sum = float(train_df[basis_col].fillna(0.0).clip(lower=0.0).sum())
    target_sum = float(train_df["GTPointCountInPredSegment"].fillna(0.0).clip(lower=0.0).sum())
    scale = (target_sum / basis_sum) if basis_sum > 1e-8 else 1.0
    return basis_col, float(max(scale, 0.0))


def fuse_count_predictions(
    learned_pred: np.ndarray,
    fallback_pred: np.ndarray | None,
    fusion_mode: str,
    learned_weight: float,
) -> np.ndarray:
    learned_pred = np.clip(np.asarray(learned_pred, dtype=np.float64), a_min=0.0, a_max=None)
    if fallback_pred is None or not fusion_mode or fusion_mode == "learned_only":
        return learned_pred

    fallback_pred = np.clip(np.asarray(fallback_pred, dtype=np.float64), a_min=0.0, a_max=None)
    if fusion_mode == "max":
        return np.maximum(learned_pred, fallback_pred)
    if fusion_mode == "blend":
        weight = float(np.clip(learned_weight, 0.0, 1.0))
        return (weight * learned_pred) + ((1.0 - weight) * fallback_pred)
    if fusion_mode == "uplift":
        weight = float(np.clip(learned_weight, 0.0, 1.0))
        positive_gap = np.clip(fallback_pred - learned_pred, a_min=0.0, a_max=None)
        return learned_pred + ((1.0 - weight) * positive_gap)
    raise ValueError(f"Unsupported count_fusion_mode: {fusion_mode}")


def resolve_feature_cols(feature_preset: str, feature_cols_raw: str) -> list[str]:
    if feature_preset:
        if feature_preset not in dataset_builder.FEATURE_PRESETS:
            raise ValueError(
                f"Unsupported feature_preset: {feature_preset}, "
                f"available={sorted(dataset_builder.FEATURE_PRESETS.keys())}"
            )
        return list(dataset_builder.FEATURE_PRESETS[feature_preset])
    return parse_list_arg(feature_cols_raw)


def open_or_create_doc(docx_path: Path) -> Document:
    if docx_path.exists():
        try:
            return Document(str(docx_path))
        except (PackageNotFoundError, BadZipFile):
            pass
    doc = Document()
    doc.add_heading("实验记录 20260319", level=1)
    return doc


def format_float(value: object) -> str:
    try:
        return f"{float(value):.4f}"
    except (TypeError, ValueError):
        return str(value)


def mean_of_numeric(rows: list[dict], key: str) -> float | None:
    values = []
    for row in rows:
        raw_value = row.get(key, "")
        if raw_value in {"", None}:
            continue
        try:
            numeric_value = float(raw_value)
        except (TypeError, ValueError):
            continue
        if numeric_value != numeric_value:
            continue
        values.append(numeric_value)
    if not values:
        return None
    return sum(values) / len(values)


def append_result_to_docx(docx_path: Path, title: str, config: dict, results: list[dict]) -> None:
    doc = open_or_create_doc(docx_path)
    doc.add_heading(title, level=2)
    doc.add_paragraph(f"记录时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    config_lines = [
        "task: 预测段级点数 + 段内概率点位精细化",
        f"script: {ROOT / 'segment_count_predict_then_refine.py'}",
        f"exist_exp_dir: {config['exist_exp_dir']}",
        f"model_name: {config['model_name']}",
        f"model_alpha: {config['model_alpha']}",
        f"model_max_iter: {config['model_max_iter']}",
        f"tweedie_power: {config['tweedie_power']}",
        f"hgb_learning_rate: {config['hgb_learning_rate']}",
        f"hgb_max_depth: {config['hgb_max_depth']}",
        f"feature_preset: {config['feature_preset']}",
        f"feature_cols: {config['feature_cols']}",
        f"prob_weight_gamma: {config['prob_weight_gamma']}",
        f"gt_min_points_per_segment: {config['gt_min_points_per_segment']}",
        f"gt_rounding_mode: {config['gt_rounding_mode']}",
        f"fallback_count_mode: {config['fallback_count_mode']}",
        f"count_fusion_mode: {config['count_fusion_mode']}",
        f"fusion_learned_weight: {config['fusion_learned_weight']}",
        f"pred_min_points_per_segment: {config['pred_min_points_per_segment']}",
        f"pred_rounding_mode: {config['rounding_mode']}",
        f"save_dir: {config['save_dir']}",
        f"summary_csv: {config['summary_csv']}",
    ]
    if config["model_name"] == "xgboost":
        config_lines.extend(
            [
                f"xgb_n_estimators: {config['xgb_n_estimators']}",
                f"xgb_max_depth: {config['xgb_max_depth']}",
                f"xgb_learning_rate: {config['xgb_learning_rate']}",
                f"xgb_subsample: {config['xgb_subsample']}",
                f"xgb_colsample_bytree: {config['xgb_colsample_bytree']}",
            ]
        )
    doc.add_paragraph("\n".join(config_lines))

    headers = [
        "井名",
        "GT点数",
        "Pred点数",
        "点数差",
        "pred->gt_mean",
        "pred->gt_p90",
        "gt->pred_mean",
        "gt->pred_p90",
        "SegMAE",
    ]
    key_order = [
        "well",
        "N_gt_points",
        "N_pred_points",
        "count_diff",
        "pred_to_gt_mean_dist",
        "pred_to_gt_p90_dist",
        "gt_to_pred_mean_dist",
        "gt_to_pred_p90_dist",
        "segment_count_mae",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in results:
        cells = table.add_row().cells
        for idx, key in enumerate(key_order):
            cells[idx].text = str(row[key]) if key == "well" else format_float(row[key])

    summary_parts = []
    for key, label in (
        ("count_diff", "AVG_count_diff"),
        ("pred_to_gt_mean_dist", "AVG_pred_to_gt_mean"),
        ("gt_to_pred_mean_dist", "AVG_gt_to_pred_mean"),
        ("segment_count_mae", "AVG_segment_count_mae"),
    ):
        avg_value = mean_of_numeric(results, key)
        if avg_value is not None:
            summary_parts.append(f"{label}={avg_value:.4f}")
    doc.add_paragraph("总结: " + ", ".join(summary_parts) if summary_parts else "总结: 已记录本次实验结果。")
    doc.add_paragraph("")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def build_regressor(
    model_name: str,
    alpha: float,
    max_iter: int,
    tweedie_power: float,
    xgb_n_estimators: int,
    xgb_max_depth: int,
    xgb_learning_rate: float,
    xgb_subsample: float,
    xgb_colsample_bytree: float,
    hgb_learning_rate: float,
    hgb_max_depth: int,
):
    if model_name == "poisson":
        model = PoissonRegressor(alpha=alpha, max_iter=max_iter)
        return Pipeline(
            [
                ("scaler", StandardScaler()),
                ("model", model),
            ]
        )
    elif model_name == "tweedie":
        model = TweedieRegressor(power=tweedie_power, alpha=alpha, max_iter=max_iter, link="log")
        return Pipeline(
            [
                ("scaler", StandardScaler()),
                ("model", model),
            ]
        )
    if model_name == "hgbt_poisson":
        model = HistGradientBoostingRegressor(
            loss="poisson",
            learning_rate=hgb_learning_rate,
            max_iter=max_iter,
            max_depth=hgb_max_depth,
            min_samples_leaf=5,
            l2_regularization=max(float(alpha), 0.0),
            random_state=42,
        )
        return Pipeline([("model", model)])
    if model_name == "xgboost":
        if XGBRegressor is None:
            raise ImportError("xgboost is not installed in current environment")
        model = XGBRegressor(
            objective="count:poisson",
            n_estimators=xgb_n_estimators,
            max_depth=xgb_max_depth,
            learning_rate=xgb_learning_rate,
            subsample=xgb_subsample,
            colsample_bytree=xgb_colsample_bytree,
            reg_lambda=max(float(alpha), 0.0),
            random_state=42,
        )
        return Pipeline([("model", model)])
    raise ValueError(f"Unsupported model_name: {model_name}")


def compute_pred_count(raw_value: float, pred_min_points_per_segment: int, rounding_mode: str) -> int:
    return int(
        dataset_builder.POINT.compute_n_points(
            float(max(raw_value, 0.0)),
            int(pred_min_points_per_segment),
            rounding_mode,
        )
    )


def build_pred_points_from_segment_df(
    df_sorted: pd.DataFrame,
    payloads: list[dict],
    pred_segment_df: pd.DataFrame,
    config,
) -> pd.DataFrame:
    pred_map = pred_segment_df.set_index("Segment_ID").to_dict("index")
    rows = []
    for payload in payloads:
        segment_id = int(payload["Segment_ID"])
        pred_info = pred_map.get(segment_id)
        if pred_info is None:
            continue
        n_points = int(pred_info["PredPointCount"])
        if n_points <= 0:
            continue
        picked_depths, picked_rel_idx = dataset_builder.POINT.pick_points_by_equal_intensity(
            payload["depth_seg"],
            payload["shape_intensity"],
            n_points,
        )
        for point_idx, (depth_value, rel_idx) in enumerate(zip(picked_depths, picked_rel_idx), start=1):
            abs_idx = int(payload["StartIdx"] + rel_idx)
            row = df_sorted.iloc[abs_idx]
            rows.append(
                {
                    "Segment_ID": segment_id,
                    "Point_ID_In_Segment": point_idx,
                    config.depth_col: float(depth_value),
                    "NearestSampleDepth": float(row[config.depth_col]),
                    config.x_col: float(row[config.x_col]) if config.x_col in df_sorted.columns else np.nan,
                    config.y_col: float(row[config.y_col]) if config.y_col in df_sorted.columns else np.nan,
                    config.time_col: float(row[config.time_col]) if config.time_col in df_sorted.columns else np.nan,
                    "SegStartDepth": float(pred_info["SegStartDepth"]),
                    "SegEndDepth": float(pred_info["SegEndDepth"]),
                    "SegLength": float(pred_info["SegLength"]),
                    "PredPointCountFloat": float(pred_info["PredPointCountFloat"]),
                    "PredPointCount": n_points,
                    "GTPointCountInPredSegment": int(pred_info["GTPointCountInPredSegment"]),
                }
            )
    return pd.DataFrame(rows)


def build_coef_rows(well_name: str, regressor: Pipeline, feature_cols: list[str]) -> list[dict]:
    model = regressor.named_steps["model"]
    rows = []
    if hasattr(model, "coef_"):
        for feature_name, coef in zip(feature_cols, model.coef_):
            rows.append(
                {
                    "val_well": well_name,
                    "feature": feature_name,
                    "coef": float(coef),
                }
            )
    if hasattr(model, "intercept_"):
        rows.append(
            {
                "val_well": well_name,
                "feature": "__intercept__",
                "coef": float(model.intercept_),
            }
        )
    elif hasattr(model, "feature_importances_"):
        for feature_name, importance in zip(feature_cols, model.feature_importances_):
            rows.append(
                {
                    "val_well": well_name,
                    "feature": feature_name,
                    "coef": float(importance),
                }
            )
    return rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", required=True)
    parser.add_argument("--exist-exp-dir", default=str(dataset_builder.DEFAULT_EXIST_EXP_DIR))
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument("--result-dir")
    parser.add_argument("--well-names", default="")
    parser.add_argument("--prob-weight-gamma", type=float, default=2.0)
    parser.add_argument("--gt-min-points-per-segment", type=int, default=1)
    parser.add_argument("--gt-rounding-mode", default="round")
    parser.add_argument("--pred-min-points-per-segment", type=int, default=0)
    parser.add_argument("--rounding-mode", default="round")
    parser.add_argument("--fallback-count-mode", default="")
    parser.add_argument("--count-fusion-mode", default="learned_only")
    parser.add_argument("--fusion-learned-weight", type=float, default=0.7)
    parser.add_argument("--density-gt-col", default="P10")
    parser.add_argument("--model-name", default="poisson")
    parser.add_argument("--model-alpha", type=float, default=0.1)
    parser.add_argument("--model-max-iter", type=int, default=1000)
    parser.add_argument("--tweedie-power", type=float, default=1.5)
    parser.add_argument("--feature-preset", default="")
    parser.add_argument("--feature-cols", default=",".join(dataset_builder.DEFAULT_FEATURE_COLS))
    parser.add_argument("--xgb-n-estimators", type=int, default=300)
    parser.add_argument("--xgb-max-depth", type=int, default=3)
    parser.add_argument("--xgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--xgb-subsample", type=float, default=0.8)
    parser.add_argument("--xgb-colsample-bytree", type=float, default=0.8)
    parser.add_argument("--hgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--hgb-max-depth", type=int, default=3)
    args = parser.parse_args()

    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / sanitize(args.exp_id))
    save_dir.mkdir(parents=True, exist_ok=True)
    well_names = parse_list_arg(args.well_names) if args.well_names else None
    feature_cols = resolve_feature_cols(str(args.feature_preset), str(args.feature_cols))
    fallback_count_mode = validate_optional_count_mode(str(args.fallback_count_mode))

    dataset_df, per_well, config = dataset_builder.build_segment_dataset(
        exist_exp_dir=str(args.exist_exp_dir),
        well_names=well_names,
        save_dir=str(save_dir / "segment_dataset"),
        prob_weight_gamma=float(args.prob_weight_gamma),
        gt_min_points_per_segment=int(args.gt_min_points_per_segment),
        gt_rounding_mode=str(args.gt_rounding_mode),
        density_gt_col=str(args.density_gt_col),
    )
    if dataset_df.empty:
        raise ValueError("Segment dataset is empty")

    missing_features = [feature for feature in feature_cols if feature not in dataset_df.columns]
    if missing_features:
        raise ValueError(f"Missing feature cols: {missing_features}")

    summary_rows = []
    coef_rows = []
    for well_name in sorted(per_well.keys()):
        train_df = dataset_df[dataset_df["WellName"] != well_name].copy()
        val_df = dataset_df[dataset_df["WellName"] == well_name].copy()
        if train_df.empty or val_df.empty:
            continue

        X_train = train_df[feature_cols].fillna(0.0)
        y_train = train_df["GTPointCountInPredSegment"].fillna(0.0).clip(lower=0.0)
        X_val = val_df[feature_cols].fillna(0.0)

        regressor = build_regressor(
            model_name=str(args.model_name),
            alpha=float(args.model_alpha),
            max_iter=int(args.model_max_iter),
            tweedie_power=float(args.tweedie_power),
            xgb_n_estimators=int(args.xgb_n_estimators),
            xgb_max_depth=int(args.xgb_max_depth),
            xgb_learning_rate=float(args.xgb_learning_rate),
            xgb_subsample=float(args.xgb_subsample),
            xgb_colsample_bytree=float(args.xgb_colsample_bytree),
            hgb_learning_rate=float(args.hgb_learning_rate),
            hgb_max_depth=int(args.hgb_max_depth),
        )
        regressor.fit(X_train, y_train)
        learned_pred_count_float = np.clip(regressor.predict(X_val), a_min=0.0, a_max=None)

        fallback_basis_col = ""
        fallback_scale = np.nan
        fallback_pred_count_float = None
        if fallback_count_mode:
            fallback_basis_col, fallback_scale = fit_basis_scale(train_df, fallback_count_mode)
            fallback_pred_count_float = np.clip(
                val_df[fallback_basis_col].fillna(0.0).clip(lower=0.0).to_numpy(dtype=np.float64) * fallback_scale,
                a_min=0.0,
                a_max=None,
            )

        pred_count_float = fuse_count_predictions(
            learned_pred=learned_pred_count_float,
            fallback_pred=fallback_pred_count_float,
            fusion_mode=str(args.count_fusion_mode),
            learned_weight=float(args.fusion_learned_weight),
        )

        coef_rows.extend(build_coef_rows(well_name, regressor, feature_cols))

        pred_segment_df = val_df.copy()
        pred_segment_df["LearnedPredPointCountFloat"] = learned_pred_count_float
        pred_segment_df["FallbackCountBasisCol"] = fallback_basis_col
        pred_segment_df["FallbackCountScale"] = fallback_scale
        pred_segment_df["FallbackPredPointCountFloat"] = (
            fallback_pred_count_float if fallback_pred_count_float is not None else np.nan
        )
        pred_segment_df["PredPointCountFloat"] = pred_count_float
        pred_segment_df["PredPointCount"] = pred_segment_df["PredPointCountFloat"].apply(
            lambda value: compute_pred_count(
                raw_value=float(value),
                pred_min_points_per_segment=int(args.pred_min_points_per_segment),
                rounding_mode=str(args.rounding_mode),
            )
        )
        pred_segment_df["SegmentCountAbsError"] = (
            pred_segment_df["PredPointCount"] - pred_segment_df["GTPointCountInPredSegment"]
        ).abs()

        well_data = per_well[well_name]
        pred_points = build_pred_points_from_segment_df(
            df_sorted=well_data["df_sorted"],
            payloads=well_data["payloads"],
            pred_segment_df=pred_segment_df,
            config=config,
        )
        gt_points = well_data["gt_points"]
        gt_depths = well_data["gt_depths"]
        pred_depths = (
            pred_points[config.depth_col].to_numpy(dtype=np.float64)
            if not pred_points.empty
            else np.array([], dtype=np.float64)
        )
        pred_to_gt = dataset_builder.POINT.nearest_distance_stats(pred_depths, gt_depths)
        gt_to_pred = dataset_builder.POINT.nearest_distance_stats(gt_depths, pred_depths)

        well_dir = save_dir / well_name
        well_dir.mkdir(parents=True, exist_ok=True)
        gt_points.to_csv(well_dir / "gt_fracture_points.csv", index=False, encoding="utf-8-sig")
        pred_points.to_csv(well_dir / "pred_fracture_points.csv", index=False, encoding="utf-8-sig")
        pred_segment_df.to_csv(well_dir / "pred_segment_summary.csv", index=False, encoding="utf-8-sig")
        train_df.to_csv(well_dir / "train_segment_dataset.csv", index=False, encoding="utf-8-sig")

        metrics = {
            "well": well_name,
            "exist_csv": well_data["exist_csv"],
            "model_name": args.model_name,
            "model_alpha": args.model_alpha,
            "model_max_iter": args.model_max_iter,
            "tweedie_power": args.tweedie_power,
            "hgb_learning_rate": args.hgb_learning_rate,
            "hgb_max_depth": args.hgb_max_depth,
            "feature_preset": args.feature_preset,
            "feature_cols": feature_cols,
            "prob_weight_gamma": args.prob_weight_gamma,
            "gt_min_points_per_segment": args.gt_min_points_per_segment,
            "gt_rounding_mode": args.gt_rounding_mode,
            "fallback_count_mode": fallback_count_mode,
            "count_fusion_mode": args.count_fusion_mode,
            "fusion_learned_weight": args.fusion_learned_weight,
            "pred_min_points_per_segment": args.pred_min_points_per_segment,
            "rounding_mode": args.rounding_mode,
            "num_train_segments": int(len(train_df)),
            "num_val_segments": int(len(val_df)),
            "segment_count_mae": float(pred_segment_df["SegmentCountAbsError"].mean()),
            "segment_count_rmse": float(
                np.sqrt(
                    np.mean(
                        np.square(
                            pred_segment_df["PredPointCount"].to_numpy(dtype=np.float64)
                            - pred_segment_df["GTPointCountInPredSegment"].to_numpy(dtype=np.float64)
                        )
                    )
                )
            ),
            "N_gt_points": int(len(gt_depths)),
            "N_pred_points": int(len(pred_depths)),
            "count_diff": int(len(pred_depths) - len(gt_depths)),
            "pred_to_gt_mean_dist": pred_to_gt["mean"],
            "pred_to_gt_median_dist": pred_to_gt["median"],
            "pred_to_gt_p90_dist": pred_to_gt["p90"],
            "gt_to_pred_mean_dist": gt_to_pred["mean"],
            "gt_to_pred_median_dist": gt_to_pred["median"],
            "gt_to_pred_p90_dist": gt_to_pred["p90"],
        }
        with open(well_dir / "metrics.json", "w", encoding="utf-8") as file_obj:
            json.dump(metrics, file_obj, ensure_ascii=False, indent=2)
        summary_rows.append(metrics)

    summary_df = pd.DataFrame(summary_rows)
    summary_csv = save_dir / "segment_count_refine_summary.csv"
    summary_df.to_csv(summary_csv, index=False, encoding="utf-8-sig")
    coef_df = pd.DataFrame(coef_rows)
    coef_df.to_csv(save_dir / "loo_model_coefficients.csv", index=False, encoding="utf-8-sig")
    with open(save_dir / "config.json", "w", encoding="utf-8") as file_obj:
        json.dump(
            {
                "exp_id": args.exp_id,
                "exist_exp_dir": str(args.exist_exp_dir),
                "feature_cols": feature_cols,
                "feature_preset": args.feature_preset,
                "prob_weight_gamma": args.prob_weight_gamma,
                "gt_min_points_per_segment": args.gt_min_points_per_segment,
                "gt_rounding_mode": args.gt_rounding_mode,
                "fallback_count_mode": fallback_count_mode,
                "count_fusion_mode": args.count_fusion_mode,
                "fusion_learned_weight": args.fusion_learned_weight,
                "pred_min_points_per_segment": args.pred_min_points_per_segment,
                "rounding_mode": args.rounding_mode,
                "density_gt_col": args.density_gt_col,
                "model_name": args.model_name,
                "model_alpha": args.model_alpha,
                "model_max_iter": args.model_max_iter,
                "tweedie_power": args.tweedie_power,
                "hgb_learning_rate": args.hgb_learning_rate,
                "hgb_max_depth": args.hgb_max_depth,
                "xgb_n_estimators": args.xgb_n_estimators,
                "xgb_max_depth": args.xgb_max_depth,
                "xgb_learning_rate": args.xgb_learning_rate,
                "xgb_subsample": args.xgb_subsample,
                "xgb_colsample_bytree": args.xgb_colsample_bytree,
                "save_dir": str(save_dir),
            },
            file_obj,
            ensure_ascii=False,
            indent=2,
        )

    with summary_csv.open("r", encoding="utf-8-sig", newline="") as file_obj:
        results = list(csv.DictReader(file_obj))
    append_result_to_docx(
        Path(args.docx_path),
        f"实验 {args.exp_id}",
        {
            "exist_exp_dir": str(args.exist_exp_dir),
            "model_name": args.model_name,
            "model_alpha": args.model_alpha,
            "model_max_iter": args.model_max_iter,
            "tweedie_power": args.tweedie_power,
            "hgb_learning_rate": args.hgb_learning_rate,
            "hgb_max_depth": args.hgb_max_depth,
            "feature_preset": args.feature_preset,
            "feature_cols": feature_cols,
            "prob_weight_gamma": args.prob_weight_gamma,
            "gt_min_points_per_segment": args.gt_min_points_per_segment,
            "gt_rounding_mode": args.gt_rounding_mode,
            "fallback_count_mode": fallback_count_mode,
            "count_fusion_mode": args.count_fusion_mode,
            "fusion_learned_weight": args.fusion_learned_weight,
            "pred_min_points_per_segment": args.pred_min_points_per_segment,
            "rounding_mode": args.rounding_mode,
            "save_dir": str(save_dir),
            "summary_csv": str(summary_csv),
            "xgb_n_estimators": args.xgb_n_estimators,
            "xgb_max_depth": args.xgb_max_depth,
            "xgb_learning_rate": args.xgb_learning_rate,
            "xgb_subsample": args.xgb_subsample,
            "xgb_colsample_bytree": args.xgb_colsample_bytree,
        },
        results,
    )

    print(
        json.dumps(
            {
                "exp_id": args.exp_id,
                "save_dir": str(save_dir),
                "summary_csv": str(summary_csv),
                "feature_cols": feature_cols,
                "results": results,
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
