import argparse
import csv
import json
import os
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from workflow_paths import FRACTURE_POINT_REFINE_SCRIPT_PATH, WORKFLOW_ROOT

ROOT = WORKFLOW_ROOT
SCRIPT_PATH = FRACTURE_POINT_REFINE_SCRIPT_PATH
BASE_SAVE_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/裂缝位置预测/基于密度的裂缝点位分析/裂缝点位精细化"
)
DEFAULT_EXIST_EXP_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/裂缝存在性预测/LSTM/单井验证/成像测井裂缝预测/cnn+lstm/exp40_missing_drop_ac_gr_3x3_seq5_AC_GR"
)
DEFAULT_DOCX_PATH = Path(r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/实验记录/实验记录20260319.docx")


def sanitize(text: str) -> str:
    text = text.replace('"', "")
    text = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]+", "_", text)
    return text.strip("_")


def open_or_create_doc(docx_path: Path) -> Document:
    if docx_path.exists():
        try:
            return Document(str(docx_path))
        except PackageNotFoundError:
            pass

    doc = Document()
    doc.add_heading("实验记录 20260319", level=1)
    return doc


def format_float(value: object) -> str:
    try:
        return f"{float(value):.4f}"
    except (TypeError, ValueError):
        return str(value)


def mean_of_numeric(rows: list[dict], key: str) -> float | None:
    values = []
    for row in rows:
        raw_value = row.get(key, "")
        if raw_value in {"", None}:
            continue
        try:
            numeric_value = float(raw_value)
        except (TypeError, ValueError):
            continue
        if numeric_value != numeric_value:
            continue
        values.append(numeric_value)
    if not values:
        return None
    return sum(values) / len(values)


def append_result_to_docx(docx_path: Path, title: str, config: dict, results: list[dict]) -> None:
    doc = open_or_create_doc(docx_path)
    doc.add_heading(title, level=2)
    doc.add_paragraph(f"记录时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    config_lines = [
        "task: 裂缝点位精细化",
        f"script: {SCRIPT_PATH}",
        f"exist_exp_dir: {config['exist_exp_dir']}",
        f"density_exp_dir: {config['density_exp_dir']}",
        f"use_pred_density: {config['use_pred_density']}",
        f"use_prob_weighting: {config['use_prob_weighting']}",
        f"segment_count_mode: {config['segment_count_mode']}",
        f"shape_mode: {config['shape_mode']}",
        f"prob_weight_gamma: {config['prob_weight_gamma']}",
        f"min_points_per_segment: {config['min_points_per_segment']}",
        f"rounding_mode: {config['rounding_mode']}",
        f"density_clip_max: {config['density_clip_max']}",
        f"density_gt_col: {config['density_gt_col']}",
        f"density_pred_col: {config['density_pred_col']}",
        f"well_names: {config['well_names']}",
        f"save_dir: {config['save_dir']}",
        f"summary_csv: {config['summary_csv']}",
    ]
    doc.add_paragraph("\n".join(config_lines))

    headers = [
        "井名",
        "GT点数",
        "Pred点数",
        "点数差",
        "pred->gt_mean",
        "pred->gt_p90",
        "gt->pred_mean",
        "gt->pred_p90",
    ]
    key_order = [
        "well",
        "N_gt_points",
        "N_pred_points",
        "count_diff",
        "pred_to_gt_mean_dist",
        "pred_to_gt_p90_dist",
        "gt_to_pred_mean_dist",
        "gt_to_pred_p90_dist",
    ]

    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header

    for row in results:
        cells = table.add_row().cells
        for idx, key in enumerate(key_order):
            cells[idx].text = str(row[key]) if key == "well" else format_float(row[key])

    if results:
        avg_count_diff = mean_of_numeric(results, "count_diff")
        avg_pred_to_gt = mean_of_numeric(results, "pred_to_gt_mean_dist")
        avg_gt_to_pred = mean_of_numeric(results, "gt_to_pred_mean_dist")
        summary_parts = []
        if avg_count_diff is not None:
            summary_parts.append(f"AVG_count_diff={avg_count_diff:.4f}")
        if avg_pred_to_gt is not None:
            summary_parts.append(f"AVG_pred_to_gt_mean={avg_pred_to_gt:.4f}")
        if avg_gt_to_pred is not None:
            summary_parts.append(f"AVG_gt_to_pred_mean={avg_gt_to_pred:.4f}")
        doc.add_paragraph("总结: " + ", ".join(summary_parts) if summary_parts else "总结: 已记录本次点位精细化实验结果。")
    else:
        doc.add_paragraph("总结: 已记录本次点位精细化实验结果。")

    doc.add_paragraph("")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", required=True)
    parser.add_argument("--density-exp-dir", required=True)
    parser.add_argument("--exist-exp-dir", default=str(DEFAULT_EXIST_EXP_DIR))
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument("--result-dir")
    parser.add_argument("--skip-run", action="store_true")
    parser.add_argument("--well-names", default="")
    parser.add_argument("--use-pred-density", default="true")
    parser.add_argument("--use-prob-weighting", default="true")
    parser.add_argument("--segment-count-mode", default="density_integral")
    parser.add_argument("--shape-mode", default="")
    parser.add_argument("--prob-weight-gamma", default="2.0")
    parser.add_argument("--min-points-per-segment", default="1")
    parser.add_argument("--rounding-mode", default="round")
    parser.add_argument("--density-clip-max", default="20.0")
    parser.add_argument("--density-gt-col", default="P10")
    parser.add_argument("--density-pred-col", default="P10_PRED_RAW")
    args = parser.parse_args()

    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / sanitize(args.exp_id))
    save_dir.mkdir(parents=True, exist_ok=True)

    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["EXP_EXIST_EXP_DIR"] = str(args.exist_exp_dir)
    env["EXP_DENSITY_EXP_DIR"] = str(args.density_exp_dir)
    env["EXP_SAVE_DIR"] = str(save_dir)
    env["EXP_USE_PRED_DENSITY"] = str(args.use_pred_density)
    env["EXP_USE_PROB_WEIGHTING"] = str(args.use_prob_weighting)
    env["EXP_SEGMENT_COUNT_MODE"] = str(args.segment_count_mode)
    if args.shape_mode:
        env["EXP_SHAPE_MODE"] = str(args.shape_mode)
    env["EXP_PROB_WEIGHT_GAMMA"] = str(args.prob_weight_gamma)
    env["EXP_MIN_POINTS_PER_SEGMENT"] = str(args.min_points_per_segment)
    env["EXP_ROUNDING_MODE"] = str(args.rounding_mode)
    env["EXP_DENSITY_CLIP_MAX"] = str(args.density_clip_max)
    env["EXP_DENSITY_GT_COL"] = str(args.density_gt_col)
    env["EXP_DENSITY_PRED_COL"] = str(args.density_pred_col)
    if args.well_names:
        env["EXP_WELL_NAMES"] = args.well_names

    if not args.skip_run:
        completed = subprocess.run(
            [sys.executable, str(SCRIPT_PATH)],
            env=env,
            cwd=str(ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
        )
        (save_dir / "run_stdout.log").write_text(completed.stdout, encoding="utf-8")
        (save_dir / "run_stderr.log").write_text(completed.stderr, encoding="utf-8")
        if completed.returncode != 0:
            print(completed.stdout[-4000:], file=sys.stdout)
            print(completed.stderr[-4000:], file=sys.stderr)
            return completed.returncode

    summary_csv = save_dir / "fracture_point_refine_summary.csv"
    if not summary_csv.exists():
        print(f"Missing result file: {summary_csv}", file=sys.stderr)
        return 2

    with summary_csv.open("r", encoding="utf-8-sig", newline="") as file_obj:
        results = list(csv.DictReader(file_obj))

    config = {
        "exist_exp_dir": str(args.exist_exp_dir),
        "density_exp_dir": str(args.density_exp_dir),
        "use_pred_density": args.use_pred_density,
        "use_prob_weighting": args.use_prob_weighting,
        "segment_count_mode": args.segment_count_mode,
        "shape_mode": args.shape_mode or ("density_prob" if args.use_prob_weighting.lower() == "true" else "density_only"),
        "prob_weight_gamma": args.prob_weight_gamma,
        "min_points_per_segment": args.min_points_per_segment,
        "rounding_mode": args.rounding_mode,
        "density_clip_max": args.density_clip_max,
        "density_gt_col": args.density_gt_col,
        "density_pred_col": args.density_pred_col,
        "well_names": args.well_names,
        "save_dir": str(save_dir),
        "summary_csv": str(summary_csv),
    }
    append_result_to_docx(Path(args.docx_path), f"实验 {args.exp_id}", config, results)

    print(
        json.dumps(
            {
                "exp_id": args.exp_id,
                "save_dir": str(save_dir),
                "summary_csv": str(summary_csv),
                "results": results,
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
