import argparse
import json
import os
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from workflow_paths import FRACTURE_EXISTENCE_LSTM_SCRIPT_PATH, WORKFLOW_ROOT

ROOT = WORKFLOW_ROOT
SCRIPT_PATH = FRACTURE_EXISTENCE_LSTM_SCRIPT_PATH
BASE_SAVE_DIR = Path(r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/裂缝存在性预测/LSTM/单井验证/成像测井裂缝预测/cnn+lstm")
FLOW_RESULT_ROOT = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/两阶段裂缝预测流程结果"
)
BASE_SAVE_DIR = FLOW_RESULT_ROOT / "manual_runs" / "run_lstm_experiment"


def sanitize(text: str) -> str:
    text = text.replace('"', "")
    text = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]+", "_", text)
    return text.strip("_")


def safe_console_write(text: str, stream) -> None:
    if not text:
        return
    try:
        stream.write(text)
    except UnicodeEncodeError:
        encoding = getattr(stream, "encoding", None) or "utf-8"
        fallback = text.encode(encoding, errors="replace").decode(encoding, errors="replace")
        stream.write(fallback)


def append_result_to_docx(docx_path: Path, title: str, config: dict, results: list[dict]) -> None:
    if docx_path.exists():
        try:
            doc = Document(str(docx_path))
        except PackageNotFoundError:
            doc = Document()
            doc.add_heading("实验记录 20260319", level=1)
    else:
        doc = Document()
        doc.add_heading("实验记录 20260319", level=1)

    doc.add_heading(title, level=2)
    doc.add_paragraph(f"记录时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    config_lines = [
        f"data_dir: {config['data_dir']}",
        f"SEIS_MODE: {config['SEIS_MODE']}",
        f"SEQ_LEN: {config['SEQ_LEN']}",
        f"imaging_well_features: {config['imaging_well_features']}",
        f"selection_metric: {config['selection_metric']}",
        f"use_dynamic_threshold: {config['use_dynamic_threshold']}",
        f"thresh_search_min: {config['thresh_search_min']}",
        f"thresh_search_max: {config['thresh_search_max']}",
        f"thresh_search_step: {config['thresh_search_step']}",
        f"use_density_regression: {config['use_density_regression']}",
        f"use_domain_adversarial: {config['use_domain_adversarial']}",
        f"use_all_other_wells: {config['use_all_other_wells']}",
        f"use_wellwise_log_scaling: {config['use_wellwise_log_scaling']}",
        f"use_well_balanced_sampling: {config['use_well_balanced_sampling']}",
        f"use_oversample: {config['use_oversample']}",
        f"manual_pos_weight: {config['manual_pos_weight']}",
        f"target_ratio: {config['target_ratio']}",
        f"edge_exclude: {config['edge_exclude']}",
        f"max_pos_weight: {config['max_pos_weight']}",
        f"use_dice_loss: {config['use_dice_loss']}",
        f"dice_loss_weight: {config['dice_loss_weight']}",
        f"fixed_threshold: {config['fixed_threshold']}",
        f"use_selection_accuracy_floor: {config['use_selection_accuracy_floor']}",
        f"MIN_SELECTION_ACCURACY: {config['min_selection_accuracy']}",
        f"use_threshold_constraints: {config['use_threshold_constraints']}",
        f"MIN_SELECTION_RECALL: {config['min_selection_recall']}",
        f"MIN_SELECTION_POS_RATIO: {config['min_selection_pos_ratio']}",
        f"MIN_SELECTION_POS_RATIO_SCALE: {config['min_selection_pos_ratio_scale']}",
        f"use_selection_accuracy_floor_by_well: {config['use_selection_accuracy_floor_by_well_json']}",
        f"min_selection_accuracy_by_well: {config['min_selection_accuracy_by_well_json']}",
        f"use_threshold_constraints_by_well: {config['use_threshold_constraints_by_well_json']}",
        f"min_selection_recall_by_well: {config['min_selection_recall_by_well_json']}",
        f"min_selection_pos_ratio_by_well: {config['min_selection_pos_ratio_by_well_json']}",
        f"min_selection_pos_ratio_scale_by_well: {config['min_selection_pos_ratio_scale_by_well_json']}",
        f"MIN_TRAIN_WELLS: {config['min_train_wells']}",
        f"DIST_THRESHOLD: {config['dist_threshold']}",
        f"only_val_wells: {config['only_val_wells']}",
        f"custom_train_wells_json: {config['custom_train_wells_json']}",
        f"save_dir: {config['save_dir']}",
    ]
    doc.add_paragraph("\n".join(config_lines))

    table = doc.add_table(rows=1, cols=7)
    hdr = table.rows[0].cells
    hdr[0].text = "井名"
    hdr[1].text = "Accuracy"
    hdr[2].text = "AUC"
    hdr[3].text = "Precision"
    hdr[4].text = "Recall"
    hdr[5].text = "F1"
    hdr[6].text = "Threshold"

    for row in results:
        cells = table.add_row().cells
        cells[0].text = str(row["val_well"])
        cells[1].text = f"{float(row['Accuracy']):.4f}"
        cells[2].text = f"{float(row['AUC']):.4f}"
        cells[3].text = f"{float(row['Precision']):.4f}"
        cells[4].text = f"{float(row['Recall']):.4f}"
        cells[5].text = f"{float(row['F1']):.4f}"
        cells[6].text = f"{float(row['Threshold']):.4f}"

    all_over_85 = all(float(row["Accuracy"]) > 0.85 for row in results)
    doc.add_paragraph(f"四口井 accuracy 是否全部 > 85%: {'是' if all_over_85 else '否'}")
    doc.add_paragraph("")
    doc.save(str(docx_path))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", required=True)
    parser.add_argument("--seq-len", type=int, required=True)
    parser.add_argument("--seis-mode", required=True)
    parser.add_argument("--features-json", required=True)
    parser.add_argument("--docx-path", required=True)
    parser.add_argument("--data-dir", default="")
    parser.add_argument("--result-dir")
    parser.add_argument("--skip-run", action="store_true")
    parser.add_argument("--selection-metric", default="accuracy")
    parser.add_argument("--use-dynamic-threshold", default="true")
    parser.add_argument("--use-density-regression", default="false")
    parser.add_argument("--use-domain-adversarial", default="false")
    parser.add_argument("--use-all-other-wells", default="false")
    parser.add_argument("--use-wellwise-log-scaling", default="false")
    parser.add_argument("--use-well-balanced-sampling", default="false")
    parser.add_argument("--use-oversample", default="false")
    parser.add_argument("--manual-pos-weight", default="2.0")
    parser.add_argument("--target-ratio", default="0.35")
    parser.add_argument("--edge-exclude", default="1")
    parser.add_argument("--max-pos-weight", default="3.0")
    parser.add_argument("--use-dice-loss", default="true")
    parser.add_argument("--dice-loss-weight", default="0.5")
    parser.add_argument("--fixed-threshold", default="0.5")
    parser.add_argument("--thresh-search-min", default="0.30")
    parser.add_argument("--thresh-search-max", default="0.90")
    parser.add_argument("--thresh-search-step", default="0.02")
    parser.add_argument("--use-selection-accuracy-floor", default="true")
    parser.add_argument("--min-selection-accuracy", default="0.80")
    parser.add_argument("--use-threshold-constraints", default="true")
    parser.add_argument("--min-selection-recall", default="0.20")
    parser.add_argument("--min-selection-pos-ratio", default="0.03")
    parser.add_argument("--min-selection-pos-ratio-scale", default="0.35")
    parser.add_argument("--use-selection-accuracy-floor-by-well-json", default="")
    parser.add_argument("--min-selection-accuracy-by-well-json", default="")
    parser.add_argument("--use-threshold-constraints-by-well-json", default="")
    parser.add_argument("--min-selection-recall-by-well-json", default="")
    parser.add_argument("--min-selection-pos-ratio-by-well-json", default="")
    parser.add_argument("--min-selection-pos-ratio-scale-by-well-json", default="")
    parser.add_argument("--min-train-wells", default="2")
    parser.add_argument("--dist-threshold", default="0.4")
    parser.add_argument("--only-val-wells", default="")
    parser.add_argument("--custom-train-wells-json", default="")
    args = parser.parse_args()

    try:
        features = json.loads(args.features_json)
    except json.JSONDecodeError:
        features = [item.strip() for item in args.features_json.split(",") if item.strip()]
    feature_tag = sanitize("_".join(features))
    exp_tag = f"{args.exp_id}_{args.seis_mode}_seq{args.seq_len}_{feature_tag}"
    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / exp_tag)
    save_dir.mkdir(parents=True, exist_ok=True)

    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    env["EXP_SEQ_LEN"] = str(args.seq_len)
    env["EXP_SEIS_MODE"] = args.seis_mode
    env["EXP_IMAGING_FEATURES"] = json.dumps(features, ensure_ascii=False)
    if args.data_dir:
        env["EXP_DATA_DIR"] = str(args.data_dir)
    env["EXP_SELECTION_METRIC"] = args.selection_metric
    env["EXP_USE_DYNAMIC_THRESHOLD"] = args.use_dynamic_threshold
    env["EXP_USE_DENSITY_REGRESSION"] = args.use_density_regression
    env["EXP_USE_DOMAIN_ADVERSARIAL"] = args.use_domain_adversarial
    env["EXP_USE_ALL_OTHER_WELLS"] = args.use_all_other_wells
    env["EXP_USE_WELLWISE_LOG_SCALING"] = args.use_wellwise_log_scaling
    env["EXP_USE_WELL_BALANCED_SAMPLING"] = args.use_well_balanced_sampling
    env["EXP_USE_OVERSAMPLE"] = args.use_oversample
    env["EXP_MANUAL_POS_WEIGHT"] = str(args.manual_pos_weight)
    env["EXP_TARGET_RATIO"] = str(args.target_ratio)
    env["EXP_EDGE_EXCLUDE"] = str(args.edge_exclude)
    env["EXP_MAX_POS_WEIGHT"] = str(args.max_pos_weight)
    env["EXP_USE_DICE_LOSS"] = args.use_dice_loss
    env["EXP_DICE_LOSS_WEIGHT"] = str(args.dice_loss_weight)
    env["EXP_FIXED_THRESHOLD"] = str(args.fixed_threshold)
    env["EXP_THRESH_SEARCH_MIN"] = str(args.thresh_search_min)
    env["EXP_THRESH_SEARCH_MAX"] = str(args.thresh_search_max)
    env["EXP_THRESH_SEARCH_STEP"] = str(args.thresh_search_step)
    env["EXP_USE_SELECTION_ACCURACY_FLOOR"] = args.use_selection_accuracy_floor
    env["EXP_MIN_SELECTION_ACCURACY"] = str(args.min_selection_accuracy)
    env["EXP_USE_THRESHOLD_CONSTRAINTS"] = args.use_threshold_constraints
    env["EXP_MIN_SELECTION_RECALL"] = str(args.min_selection_recall)
    env["EXP_MIN_SELECTION_POS_RATIO"] = str(args.min_selection_pos_ratio)
    env["EXP_MIN_SELECTION_POS_RATIO_SCALE"] = str(args.min_selection_pos_ratio_scale)
    if args.use_selection_accuracy_floor_by_well_json:
        env["EXP_USE_SELECTION_ACCURACY_FLOOR_BY_WELL"] = args.use_selection_accuracy_floor_by_well_json
    if args.min_selection_accuracy_by_well_json:
        env["EXP_MIN_SELECTION_ACCURACY_BY_WELL"] = args.min_selection_accuracy_by_well_json
    if args.use_threshold_constraints_by_well_json:
        env["EXP_USE_THRESHOLD_CONSTRAINTS_BY_WELL"] = args.use_threshold_constraints_by_well_json
    if args.min_selection_recall_by_well_json:
        env["EXP_MIN_SELECTION_RECALL_BY_WELL"] = args.min_selection_recall_by_well_json
    if args.min_selection_pos_ratio_by_well_json:
        env["EXP_MIN_SELECTION_POS_RATIO_BY_WELL"] = args.min_selection_pos_ratio_by_well_json
    if args.min_selection_pos_ratio_scale_by_well_json:
        env["EXP_MIN_SELECTION_POS_RATIO_SCALE_BY_WELL"] = args.min_selection_pos_ratio_scale_by_well_json
    env["EXP_MIN_TRAIN_WELLS"] = str(args.min_train_wells)
    env["EXP_DIST_THRESHOLD"] = str(args.dist_threshold)
    if args.only_val_wells:
        only_val_wells = [item.strip() for item in args.only_val_wells.split(",") if item.strip()]
        env["EXP_ONLY_VAL_WELLS"] = json.dumps(only_val_wells, ensure_ascii=False)
    if args.custom_train_wells_json:
        env["EXP_CUSTOM_TRAIN_WELLS_MAP"] = args.custom_train_wells_json
    env["EXP_SAVE_DIR"] = str(save_dir)

    if not args.skip_run:
        cmd = [sys.executable, str(SCRIPT_PATH)]
        completed = subprocess.run(
            cmd,
            env=env,
            cwd=str(ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
        )
        (save_dir / "run_stdout.log").write_text(completed.stdout, encoding="utf-8")
        (save_dir / "run_stderr.log").write_text(completed.stderr, encoding="utf-8")
        if completed.returncode != 0:
            safe_console_write(completed.stdout[-4000:], sys.stdout)
            safe_console_write(completed.stderr[-4000:], sys.stderr)
            return completed.returncode

    result_path = save_dir / "loo_results.csv"
    if not result_path.exists():
        print(f"Missing result file: {result_path}", file=sys.stderr)
        return 2

    import csv

    with result_path.open("r", encoding="utf-8-sig", newline="") as f:
        results = list(csv.DictReader(f))

    docx_path = Path(args.docx_path)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    config = {
        "SEQ_LEN": args.seq_len,
        "SEIS_MODE": args.seis_mode,
        "data_dir": args.data_dir,
        "imaging_well_features": features,
        "selection_metric": args.selection_metric,
        "use_dynamic_threshold": args.use_dynamic_threshold,
        "use_density_regression": args.use_density_regression,
        "use_domain_adversarial": args.use_domain_adversarial,
        "use_all_other_wells": args.use_all_other_wells,
        "use_wellwise_log_scaling": args.use_wellwise_log_scaling,
        "use_well_balanced_sampling": args.use_well_balanced_sampling,
        "use_oversample": args.use_oversample,
        "manual_pos_weight": args.manual_pos_weight,
        "target_ratio": args.target_ratio,
        "edge_exclude": args.edge_exclude,
        "max_pos_weight": args.max_pos_weight,
        "use_dice_loss": args.use_dice_loss,
        "dice_loss_weight": args.dice_loss_weight,
        "fixed_threshold": args.fixed_threshold,
        "thresh_search_min": args.thresh_search_min,
        "thresh_search_max": args.thresh_search_max,
        "thresh_search_step": args.thresh_search_step,
        "use_selection_accuracy_floor": args.use_selection_accuracy_floor,
        "min_selection_accuracy": args.min_selection_accuracy,
        "use_threshold_constraints": args.use_threshold_constraints,
        "min_selection_recall": args.min_selection_recall,
        "min_selection_pos_ratio": args.min_selection_pos_ratio,
        "min_selection_pos_ratio_scale": args.min_selection_pos_ratio_scale,
        "use_selection_accuracy_floor_by_well_json": args.use_selection_accuracy_floor_by_well_json,
        "min_selection_accuracy_by_well_json": args.min_selection_accuracy_by_well_json,
        "use_threshold_constraints_by_well_json": args.use_threshold_constraints_by_well_json,
        "min_selection_recall_by_well_json": args.min_selection_recall_by_well_json,
        "min_selection_pos_ratio_by_well_json": args.min_selection_pos_ratio_by_well_json,
        "min_selection_pos_ratio_scale_by_well_json": args.min_selection_pos_ratio_scale_by_well_json,
        "min_train_wells": args.min_train_wells,
        "dist_threshold": args.dist_threshold,
        "only_val_wells": args.only_val_wells or env.get("EXP_ONLY_VAL_WELLS", ""),
        "custom_train_wells_json": args.custom_train_wells_json or env.get("EXP_CUSTOM_TRAIN_WELLS_MAP", ""),
        "save_dir": str(save_dir),
    }
    append_result_to_docx(docx_path, f"实验 {args.exp_id}", config, results)

    print(json.dumps({
        "exp_id": args.exp_id,
        "save_dir": str(save_dir),
        "results": results,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
