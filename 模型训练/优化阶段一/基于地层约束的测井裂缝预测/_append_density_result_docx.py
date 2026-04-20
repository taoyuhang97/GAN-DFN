from pathlib import Path
from datetime import datetime
import csv
import sys
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

docx_path = Path(sys.argv[1])
result_csv = Path(sys.argv[2])
if docx_path.exists():
    try:
        doc = Document(str(docx_path))
    except PackageNotFoundError:
        doc = Document()
        doc.add_heading('实验记录 20260319', level=1)
else:
    doc = Document()
    doc.add_heading('实验记录 20260319', level=1)

with result_csv.open('r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.DictReader(f))

doc.add_heading('实验 fracture_density_predict_p10_only_3x3_seq5_AC_GR', level=2)
doc.add_paragraph(f'记录时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
doc.add_paragraph('\n'.join([
    'task: 裂缝密度预测',
    'script: 模型训练/优化阶段一/裂缝位置分析/基于密度的裂缝点位分析/fracture_density_predict.py',
    'SEIS_MODE: 3x3',
    'SEQ_LEN: 5',
    "features: ['AC', 'GR']",
    "targets: ['P10']",
    'train_rule: select_train_wells + LOO',
    f'result_csv: {result_csv}',
]))

table = doc.add_table(rows=1, cols=5)
h = table.rows[0].cells
h[0].text = '井名'
h[1].text = 'P10_R2'
h[2].text = 'P10_MAE'
h[3].text = 'P10_RMSE'
h[4].text = 'best_val_loss'
for row in rows:
    c = table.add_row().cells
    c[0].text = str(row['val_well'])
    c[1].text = f"{float(row['P10_R2']):.4f}"
    c[2].text = f"{float(row['P10_MAE']):.4f}"
    c[3].text = f"{float(row['P10_RMSE']):.4f}"
    c[4].text = f"{float(row['best_val_loss']):.4f}"

doc.add_paragraph('结论: 默认 P10 单目标版比三目标版更稳定，但跨井 R2 仍整体偏低，说明后续还需继续优化特征或训练策略。')
doc.add_paragraph('')
doc.save(str(docx_path))
print('docx_appended')
