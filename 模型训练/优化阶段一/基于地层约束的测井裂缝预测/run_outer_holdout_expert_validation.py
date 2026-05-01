from __future__ import annotations

import argparse
import importlib.util
import json
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from manual_strata_workflow.common import DEFAULT_WELL_NAMES as MANUAL_DEFAULT_WELL_NAMES
from workflow_paths import NEAREST_SCRIPT_PATH, RAW_POINT_SCRIPT_PATH, WORKFLOW_ROOT

ROOT = WORKFLOW_ROOT
PYTHON_EXE_DEFAULT = Path(os.environ.get("PYTHON_EXE", sys.executable))
STAGE1_WRAPPER_PATH = ROOT / "run_manual_strata_lstm_experiment.py"
STAGE2_WRAPPER_PATH = ROOT / "run_manual_strata_raw_point_refine.py"

FLOW_RESULT_ROOT = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/两阶段裂缝预测流程结果"
)
DEFAULT_SAVE_ROOT = FLOW_RESULT_ROOT / "outer_holdout"
DEFAULT_SAMPLE_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容一/成像测井/裂缝样本"
)
DEFAULT_RAW_LABEL_DIR = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/研究内容一/成像测井/裂缝标注"
)
DEFAULT_DOCX_PATH = Path(r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/实验记录/实验记录20260323.docx")
DOC_TITLE = "实验记录 20260323"

DEFAULT_WELL_NAMES = list(MANUAL_DEFAULT_WELL_NAMES)
DEFAULT_STAGE1_FEATURES = ["AC", "GR"]
SIMILARITY_GROUP_WEIGHTS = {
    "seismic": 0.6,
    "AC": 0.2,
    "GR": 0.2,
}
EXPERT_SELECTION_STAGE1_F1_FLOOR = 0.25
EXPERT_SELECTION_STAGE2_COUNT_ERROR_CEIL = 75.0
EXPERT_SELECTION_SIMILARITY_WEIGHT = 0.50
EXPERT_SELECTION_STAGE1_WEIGHT = 0.25
EXPERT_SELECTION_STAGE2_WEIGHT = 0.25
_MODULE_CACHE: dict[str, object] = {}


def sanitize(text: str) -> str:
    return "".join(ch if ch.isalnum() or ch in {"_", "-", "."} else "_" for ch in str(text)).strip("_")


def parse_bool(raw: object) -> bool:
    return str(raw).strip().lower() in {"1", "true", "yes", "y"}


def parse_list_arg(raw: object) -> list[str]:
    text = str(raw or "").strip()
    if not text:
        return []
    try:
        value = json.loads(text)
        if isinstance(value, list):
            return [str(item).strip() for item in value if str(item).strip()]
    except Exception:
        pass
    if text.startswith("[") and text.endswith("]"):
        text = text[1:-1].strip()
    return [
        item.strip().strip("'\"")
        for item in text.split(",")
        if item.strip().strip("'\"")
    ]


def to_jsonable(value):
    if isinstance(value, dict):
        return {str(key): to_jsonable(val) for key, val in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [to_jsonable(item) for item in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, pd.DataFrame):
        return value.to_dict(orient="records")
    if isinstance(value, pd.Series):
        return value.to_dict()
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.integer, np.floating, np.bool_)):
        return value.item()
    return value


def load_module(module_path: Path, module_name: str):
    cache_key = f"{module_name}:{module_path}"
    if cache_key in _MODULE_CACHE:
        return _MODULE_CACHE[cache_key]
    spec = importlib.util.spec_from_file_location(module_name, module_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Failed to load module from: {module_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    added_to_syspath = False
    parent_dir = str(module_path.parent)
    if parent_dir not in sys.path:
        sys.path.insert(0, parent_dir)
        added_to_syspath = True
    try:
        spec.loader.exec_module(module)
    finally:
        if added_to_syspath:
            try:
                sys.path.remove(parent_dir)
            except ValueError:
                pass
    _MODULE_CACHE[cache_key] = module
    return module


def get_stage1_module():
    return load_module(STAGE1_WRAPPER_PATH, "outer_holdout_stage1_module")


def get_nearest_module():
    return load_module(NEAREST_SCRIPT_PATH, "outer_holdout_nearest_module")


def get_raw_point_module():
    return load_module(RAW_POINT_SCRIPT_PATH, "outer_holdout_raw_point_module")


def canonicalize_well_name(well_name: str) -> str:
    stage1_module = get_stage1_module()
    if hasattr(stage1_module, "canonicalize_well_name"):
        return str(stage1_module.canonicalize_well_name(well_name)).strip()
    return str(well_name).strip()


def open_or_create_doc(docx_path: Path) -> Document:
    if docx_path.exists():
        try:
            return Document(str(docx_path))
        except PackageNotFoundError:
            pass
    doc = Document()
    doc.add_heading(DOC_TITLE, level=1)
    return doc


def format_float(value: object) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return str(value)
    if not np.isfinite(numeric):
        return "nan"
    return f"{numeric:.4f}"


def append_outer_validation_to_docx(
    docx_path: Path,
    title: str,
    config: dict,
    strata_rows: list[dict],
    whole_well_row: dict,
) -> None:
    doc = open_or_create_doc(docx_path)
    doc.add_heading(title, level=2)
    doc.add_paragraph(
        "\n".join(
            [
                "task: 外层留一整体流程验证",
                f"time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                f"holdout_well: {config['holdout_well']}",
                f"training_wells: {config['training_wells']}",
                f"target_strata: {config['target_strata']}",
                f"stage1_library_dir: {config['stage1_library_dir']}",
                f"stage2_library_dir: {config['stage2_library_dir']}",
                f"result_root: {config['result_root']}",
            ]
        )
    )
    if strata_rows:
        table = doc.add_table(rows=1, cols=9)
        headers = [
            "Strata",
            "Expert",
            "GT points",
            "Pred points",
            "Count err %",
            "pred->gt",
            "gt->pred",
            "SegMAE",
            "Stage1 F1",
        ]
        for idx, header in enumerate(headers):
            table.rows[0].cells[idx].text = header
        for row in strata_rows:
            cells = table.add_row().cells
            cells[0].text = str(row.get("strata_name", ""))
            cells[1].text = str(row.get("selected_expert_well", ""))
            cells[2].text = str(row.get("N_gt_points", ""))
            cells[3].text = str(row.get("N_pred_points", ""))
            cells[4].text = format_float(row.get("count_error_pct", np.nan))
            cells[5].text = format_float(row.get("pred_to_gt_mean_dist", np.nan))
            cells[6].text = format_float(row.get("gt_to_pred_mean_dist", np.nan))
            cells[7].text = format_float(row.get("segment_count_mae", np.nan))
            cells[8].text = format_float(row.get("stage1_f1", np.nan))
    doc.add_paragraph(
        "\n".join(
            [
                f"whole_well_gt_points: {whole_well_row.get('N_gt_points', '')}",
                f"whole_well_pred_points: {whole_well_row.get('N_pred_points', '')}",
                f"whole_well_count_error_pct: {format_float(whole_well_row.get('count_error_pct', np.nan))}",
                f"whole_well_pred_to_gt_mean: {format_float(whole_well_row.get('pred_to_gt_mean_dist', np.nan))}",
                f"whole_well_gt_to_pred_mean: {format_float(whole_well_row.get('gt_to_pred_mean_dist', np.nan))}",
                f"whole_well_segmae: {format_float(whole_well_row.get('segment_count_mae', np.nan))}",
                f"whole_well_stage1_f1: {format_float(whole_well_row.get('stage1_f1', np.nan))}",
            ]
        )
    )
    doc.add_paragraph("")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def run_command(cmd: list[str], workdir: Path, stdout_path: Path, stderr_path: Path) -> None:
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    completed = subprocess.run(
        cmd,
        cwd=str(workdir),
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="ignore",
    )
    stdout_path.parent.mkdir(parents=True, exist_ok=True)
    stdout_path.write_text(completed.stdout, encoding="utf-8")
    stderr_path.write_text(completed.stderr, encoding="utf-8")
    if completed.returncode != 0:
        raise RuntimeError(
            f"Command failed with code={completed.returncode}\n"
            f"cmd={' '.join(cmd)}\n"
            f"stdout_tail={completed.stdout[-4000:]}\n"
            f"stderr_tail={completed.stderr[-4000:]}"
        )


def load_all_labeled_data(
    sample_dir: Path,
    selected_wells: list[str],
    boundary_tolerance: float,
) -> tuple[pd.DataFrame, dict[str, list[str]]]:
    stage1_module = get_stage1_module()
    return stage1_module.build_labeled_dataset(
        data_dir=sample_dir,
        selected_wells=selected_wells,
        boundary_tolerance=boundary_tolerance,
    )


def infer_holdout_target_strata(labeled_df: pd.DataFrame, holdout_well: str, requested: list[str]) -> list[str]:
    if requested:
        return requested
    subset = labeled_df[labeled_df["WellName"] == holdout_well].copy()
    strata = [
        str(item).strip()
        for item in subset["StrataName"].dropna().tolist()
        if str(item).strip()
    ]
    deduped = []
    for name in strata:
        if name not in deduped:
            deduped.append(name)
    if not deduped:
        raise ValueError(f"No target strata found for holdout well: {holdout_well}")
    return deduped


def build_holdout_segment_files(
    labeled_df: pd.DataFrame,
    source_columns: dict[str, list[str]],
    holdout_well: str,
    target_strata: list[str],
    output_dir: Path,
) -> tuple[dict[str, Path], pd.DataFrame]:
    output_dir.mkdir(parents=True, exist_ok=True)
    segment_paths: dict[str, Path] = {}
    range_rows = []
    output_cols = [col for col in source_columns[holdout_well] if col in labeled_df.columns]
    for strata_name in target_strata:
        sub = labeled_df[
            (labeled_df["WellName"] == holdout_well) & (labeled_df["StrataName"] == strata_name)
        ].copy()
        if sub.empty:
            continue
        sub = sub.sort_values("ROW_IN_WELL").reset_index(drop=True)
        strata_dir = output_dir / sanitize(strata_name)
        strata_dir.mkdir(parents=True, exist_ok=True)
        save_path = strata_dir / f"{holdout_well}_sample.csv"
        sub[output_cols].to_csv(save_path, index=False, encoding="utf-8-sig")
        segment_paths[strata_name] = save_path
        depth_col = "TVD" if "TVD" in sub.columns else ("DEPT" if "DEPT" in sub.columns else "MD")
        range_rows.append(
            {
                "WellName": holdout_well,
                "StrataName": strata_name,
                "DepthMin": float(pd.to_numeric(sub[depth_col], errors="coerce").min()),
                "DepthMax": float(pd.to_numeric(sub[depth_col], errors="coerce").max()),
                "StrataTop": pd.to_numeric(sub["StrataTop"], errors="coerce").dropna().iloc[0]
                if "StrataTop" in sub.columns and not pd.to_numeric(sub["StrataTop"], errors="coerce").dropna().empty
                else np.nan,
                "StrataBase": pd.to_numeric(sub["StrataBase"], errors="coerce").dropna().iloc[0]
                if "StrataBase" in sub.columns and not pd.to_numeric(sub["StrataBase"], errors="coerce").dropna().empty
                else np.nan,
            }
        )
    range_df = pd.DataFrame(range_rows)
    range_df.to_csv(output_dir / "holdout_strata_range.csv", index=False, encoding="utf-8-sig")
    return segment_paths, range_df


def build_stage1_command(
    args,
    result_dir: Path,
    internal_docx_path: Path,
    training_wells: list[str],
    target_strata: list[str],
) -> list[str]:
    return [
        str(Path(args.python_exe)),
        "-X",
        "utf8",
        str(STAGE1_WRAPPER_PATH),
        "--exp-id",
        f"{args.exp_id}_inner_stage1_{sanitize(args.holdout_well)}",
        "--data-dir",
        str(args.sample_dir),
        "--docx-path",
        str(internal_docx_path),
        "--result-dir",
        str(result_dir),
        "--well-names",
        ",".join(training_wells),
        "--target-strata",
        ",".join(target_strata),
        "--lstm-features-json",
        json.dumps(args.stage1_lstm_features, ensure_ascii=False),
        "--seq-len",
        str(args.stage1_seq_len),
        "--seis-mode",
        str(args.stage1_seis_mode),
        "--selection-metric",
        str(args.stage1_selection_metric),
        "--use-dynamic-threshold",
        str(bool(args.stage1_use_dynamic_threshold)).lower(),
        "--use-density-regression",
        str(bool(args.stage1_use_density_regression)).lower(),
        "--use-domain-adversarial",
        str(bool(args.stage1_use_domain_adversarial)).lower(),
        "--use-wellwise-log-scaling",
        str(bool(args.stage1_use_wellwise_log_scaling)).lower(),
        "--use-well-balanced-sampling",
        str(bool(args.stage1_use_well_balanced_sampling)).lower(),
        "--use-oversample",
        str(bool(args.stage1_use_oversample)).lower(),
        "--manual-pos-weight",
        str(args.stage1_manual_pos_weight),
        "--target-ratio",
        str(args.stage1_target_ratio),
        "--edge-exclude",
        str(args.stage1_edge_exclude),
        "--max-pos-weight",
        str(args.stage1_max_pos_weight),
        "--use-dice-loss",
        str(bool(args.stage1_use_dice_loss)).lower(),
        "--dice-loss-weight",
        str(args.stage1_dice_loss_weight),
        "--fixed-threshold",
        str(args.stage1_fixed_threshold),
        "--thresh-search-min",
        str(args.stage1_thresh_search_min),
        "--thresh-search-max",
        str(args.stage1_thresh_search_max),
        "--thresh-search-step",
        str(args.stage1_thresh_search_step),
        "--use-selection-accuracy-floor",
        str(bool(args.stage1_use_selection_accuracy_floor)).lower(),
        "--min-selection-accuracy",
        str(args.stage1_min_selection_accuracy),
        "--use-threshold-constraints",
        str(bool(args.stage1_use_threshold_constraints)).lower(),
        "--min-selection-recall",
        str(args.stage1_min_selection_recall),
        "--min-selection-pos-ratio",
        str(args.stage1_min_selection_pos_ratio),
        "--min-selection-pos-ratio-scale",
        str(args.stage1_min_selection_pos_ratio_scale),
        "--use-selection-accuracy-floor-by-well-json",
        str(args.stage1_use_selection_accuracy_floor_by_well_json),
        "--min-selection-accuracy-by-well-json",
        str(args.stage1_min_selection_accuracy_by_well_json),
        "--use-threshold-constraints-by-well-json",
        str(args.stage1_use_threshold_constraints_by_well_json),
        "--min-selection-recall-by-well-json",
        str(args.stage1_min_selection_recall_by_well_json),
        "--min-selection-pos-ratio-by-well-json",
        str(args.stage1_min_selection_pos_ratio_by_well_json),
        "--min-selection-pos-ratio-scale-by-well-json",
        str(args.stage1_min_selection_pos_ratio_scale_by_well_json),
        "--dist-threshold",
        str(args.stage1_dist_threshold),
        "--min-train-wells",
        str(args.stage1_min_train_wells),
        "--min-seq-per-well",
        str(args.stage1_min_seq_per_well),
        "--train-selection-mode",
        str(args.stage1_train_selection_mode),
        "--boundary-tolerance",
        str(args.boundary_tolerance),
    ]


def build_stage2_command(
    args,
    result_dir: Path,
    internal_docx_path: Path,
    training_wells: list[str],
    target_strata: list[str],
    exist_exp_map: dict[str, str],
) -> list[str]:
    cmd = [
        str(Path(args.python_exe)),
        "-X",
        "utf8",
        str(STAGE2_WRAPPER_PATH),
        "--exp-id",
        f"{args.exp_id}_inner_stage2_{sanitize(args.holdout_well)}",
        "--sample-dir",
        str(args.sample_dir),
        "--raw-label-dir",
        str(args.raw_label_dir),
        "--docx-path",
        str(internal_docx_path),
        "--result-dir",
        str(result_dir),
        "--well-names",
        ",".join(training_wells),
        "--target-strata",
        ",".join(target_strata),
        "--exist-exp-map-json",
        json.dumps(exist_exp_map, ensure_ascii=False),
        "--boundary-tolerance",
        str(args.boundary_tolerance),
        "--gt-dev-rule",
        str(args.stage2_gt_dev_rule),
        "--prob-weight-gamma",
        str(args.stage2_prob_weight_gamma),
        "--pred-min-points-per-segment",
        str(args.stage2_pred_min_points_per_segment),
        "--rounding-mode",
        str(args.stage2_rounding_mode),
        "--fallback-count-mode",
        str(args.stage2_fallback_count_mode),
        "--count-train-mode",
        str(args.stage2_count_train_mode),
        "--positive-floor-prob-min",
        str(args.stage2_positive_floor_prob_min),
        "--positive-floor-allowed-labels",
        str(args.stage2_positive_floor_allowed_labels),
        "--positive-floor-target-wells",
        str(args.stage2_positive_floor_target_wells),
        "--type-calibration-mode",
        str(args.stage2_type_calibration_mode),
        "--type-calibration-scale-min",
        str(args.stage2_type_calibration_scale_min),
        "--type-calibration-scale-max",
        str(args.stage2_type_calibration_scale_max),
        "--count-fusion-mode",
        str(args.stage2_count_fusion_mode),
        "--fusion-learned-weight",
        str(args.stage2_fusion_learned_weight),
        "--global-post-scale-mode",
        str(args.stage2_global_post_scale_mode),
        "--global-post-scale-min",
        str(args.stage2_global_post_scale_min),
        "--global-post-scale-max",
        str(args.stage2_global_post_scale_max),
        "--selective-post-mode",
        str(args.stage2_selective_post_mode),
        "--selective-post-target-wells",
        str(args.stage2_selective_post_target_wells),
        "--selective-post-feature",
        str(args.stage2_selective_post_feature),
        "--selective-post-feature-2",
        str(args.stage2_selective_post_feature_2),
        "--selective-post-quantile",
        str(args.stage2_selective_post_quantile),
        "--selective-post-quantile-2",
        str(args.stage2_selective_post_quantile_2),
        "--selective-post-max-pred-float",
        str(args.stage2_selective_post_max_pred_float),
        "--soft-negative-weight",
        str(args.stage2_soft_negative_weight),
        "--train-distance-weight-mode",
        str(args.stage2_train_distance_weight_mode),
        "--train-distance-csv",
        str(args.stage2_train_distance_csv),
        "--train-distance-weight-power",
        str(args.stage2_train_distance_weight_power),
        "--train-distance-weight-min",
        str(args.stage2_train_distance_weight_min),
        "--train-distance-weight-max",
        str(args.stage2_train_distance_weight_max),
        "--boundary-expand-mode",
        str(args.stage2_boundary_expand_mode),
        "--boundary-expand-prob-min",
        str(args.stage2_boundary_expand_prob_min),
        "--boundary-expand-max-steps",
        str(args.stage2_boundary_expand_max_steps),
        "--boundary-expand-max-depth",
        str(args.stage2_boundary_expand_max_depth),
        "--gate-mode",
        str(args.stage2_gate_mode),
        "--gate-prob-threshold",
        str(args.stage2_gate_prob_threshold),
        "--gate-target-wells",
        str(args.stage2_gate_target_wells),
        "--orientation-mode",
        str(args.stage2_orientation_mode),
        "--orientation-num-families",
        str(args.stage2_orientation_num_families),
        "--orientation-min-points-per-segment",
        str(args.stage2_orientation_min_points_per_segment),
        "--model-name",
        str(args.stage2_model_name),
        "--model-alpha",
        str(args.stage2_model_alpha),
        "--model-max-iter",
        str(args.stage2_model_max_iter),
        "--tweedie-power",
        str(args.stage2_tweedie_power),
        "--feature-preset",
        str(args.stage2_feature_preset),
        "--feature-cols",
        str(args.stage2_feature_cols),
        "--xgb-n-estimators",
        str(args.stage2_xgb_n_estimators),
        "--xgb-max-depth",
        str(args.stage2_xgb_max_depth),
        "--xgb-learning-rate",
        str(args.stage2_xgb_learning_rate),
        "--xgb-subsample",
        str(args.stage2_xgb_subsample),
        "--xgb-colsample-bytree",
        str(args.stage2_xgb_colsample_bytree),
        "--hgb-learning-rate",
        str(args.stage2_hgb_learning_rate),
        "--hgb-max-depth",
        str(args.stage2_hgb_max_depth),
        "--save-model-artifacts",
        "1",
    ]
    if str(getattr(args, "stage2_config_profile", "") or "").strip():
        cmd.extend(["--config-profile", str(args.stage2_config_profile)])
    if str(getattr(args, "stage2_config_profile_map_json", "") or "").strip():
        cmd.extend(["--config-profile-map-json", str(args.stage2_config_profile_map_json)])
    return cmd


def train_inner_libraries(
    args,
    result_root: Path,
    training_wells: list[str],
    target_strata: list[str],
) -> dict[str, Path]:
    internal_docx_path = result_root / "internal_library_runs.docx"
    stage1_library_dir = result_root / "inner_stage1_library"
    stage2_library_dir = result_root / "inner_stage2_library"
    stage1_library_dir.mkdir(parents=True, exist_ok=True)
    stage2_library_dir.mkdir(parents=True, exist_ok=True)

    stage1_cmd = build_stage1_command(
        args=args,
        result_dir=stage1_library_dir,
        internal_docx_path=internal_docx_path,
        training_wells=training_wells,
        target_strata=target_strata,
    )
    run_command(
        cmd=stage1_cmd,
        workdir=ROOT,
        stdout_path=result_root / "inner_stage1_run_stdout.log",
        stderr_path=result_root / "inner_stage1_run_stderr.log",
    )

    exist_exp_map = {
        strata_name: str(stage1_library_dir / f"{sanitize(strata_name)}_lstm")
        for strata_name in target_strata
    }
    stage2_cmd = build_stage2_command(
        args=args,
        result_dir=stage2_library_dir,
        internal_docx_path=internal_docx_path,
        training_wells=training_wells,
        target_strata=target_strata,
        exist_exp_map=exist_exp_map,
    )
    run_command(
        cmd=stage2_cmd,
        workdir=ROOT,
        stdout_path=result_root / "inner_stage2_run_stdout.log",
        stderr_path=result_root / "inner_stage2_run_stderr.log",
    )
    return {
        "stage1_library_dir": stage1_library_dir,
        "stage2_library_dir": stage2_library_dir,
    }


def build_expert_registry(
    stage1_library_dir: Path,
    stage2_library_dir: Path,
    target_strata: list[str],
) -> pd.DataFrame:
    stage1_summary_path = stage1_library_dir / "manual_strata_lstm_summary.csv"
    stage2_summary_path = stage2_library_dir / "manual_strata_refine_summary.csv"
    stage1_summary = (
        pd.read_csv(stage1_summary_path, encoding="utf-8-sig")
        if stage1_summary_path.exists()
        else pd.DataFrame()
    )
    stage2_summary = (
        pd.read_csv(stage2_summary_path, encoding="utf-8-sig")
        if stage2_summary_path.exists()
        else pd.DataFrame()
    )

    rows = []
    for strata_name in target_strata:
        strata_tag = sanitize(strata_name)
        stage1_strata_dir = stage1_library_dir / f"{strata_tag}_lstm"
        stage2_strata_dir = stage2_library_dir / f"{strata_tag}_refine"
        reference_sample_dir = stage1_library_dir / "filtered_datasets" / strata_tag
        if not stage1_strata_dir.exists():
            continue

        for verify_dir in sorted(stage1_strata_dir.glob("verify_*")):
            expert_well = verify_dir.name.replace("verify_", "", 1)
            stage1_config_path = verify_dir / "config.json"
            stage1_model_path = verify_dir / "model.pth"
            stage1_scaler_path = verify_dir / "scaler.pkl"
            stage1_val_csv = verify_dir / "val_pred_full_log.csv"
            stage2_model_dir = stage2_strata_dir / verify_dir.name
            stage2_artifact_path = stage2_model_dir / "segment_refine_model.joblib"
            reference_sample_csv = reference_sample_dir / f"{expert_well}_sample.csv"
            if not stage1_config_path.exists() or not stage1_model_path.exists() or not stage1_scaler_path.exists():
                continue
            if not stage2_artifact_path.exists():
                continue

            row = {
                "strata_name": strata_name,
                "expert_well": expert_well,
                "stage1_model_dir": str(verify_dir),
                "stage1_config_path": str(stage1_config_path),
                "stage1_model_path": str(stage1_model_path),
                "stage1_scaler_path": str(stage1_scaler_path),
                "stage1_val_csv": str(stage1_val_csv) if stage1_val_csv.exists() else "",
                "stage2_model_dir": str(stage2_model_dir),
                "stage2_artifact_path": str(stage2_artifact_path),
                "reference_sample_csv": str(reference_sample_csv) if reference_sample_csv.exists() else "",
            }

            if not stage1_summary.empty:
                sub1 = stage1_summary[
                    (stage1_summary["StrataName"].astype(str) == str(strata_name))
                    & (stage1_summary["val_well"].astype(str) == str(expert_well))
                ]
                if not sub1.empty:
                    first = sub1.iloc[0]
                    row.update(
                        {
                            "inner_stage1_accuracy": pd.to_numeric(first.get("Accuracy"), errors="coerce"),
                            "inner_stage1_precision": pd.to_numeric(first.get("Precision"), errors="coerce"),
                            "inner_stage1_recall": pd.to_numeric(first.get("Recall"), errors="coerce"),
                            "inner_stage1_f1": pd.to_numeric(first.get("F1"), errors="coerce"),
                            "inner_stage1_iou": pd.to_numeric(first.get("IoU"), errors="coerce"),
                            "inner_stage1_threshold": pd.to_numeric(first.get("Threshold"), errors="coerce"),
                        }
                    )

            if not stage2_summary.empty:
                sub2 = stage2_summary[
                    (stage2_summary["StrataName"].astype(str) == str(strata_name))
                    & (stage2_summary["well"].astype(str) == str(expert_well))
                ]
                if not sub2.empty:
                    first = sub2.iloc[0]
                    row.update(
                        {
                            "inner_stage2_count_error_pct": pd.to_numeric(first.get("count_error_pct"), errors="coerce"),
                            "inner_stage2_pred_to_gt_mean_dist": pd.to_numeric(first.get("pred_to_gt_mean_dist"), errors="coerce"),
                            "inner_stage2_gt_to_pred_mean_dist": pd.to_numeric(first.get("gt_to_pred_mean_dist"), errors="coerce"),
                            "inner_stage2_segmae": pd.to_numeric(first.get("segment_count_mae"), errors="coerce"),
                        }
                    )
            rows.append(row)

    registry_df = pd.DataFrame(rows)
    if registry_df.empty:
        return registry_df
    return registry_df.sort_values(["strata_name", "expert_well"]).reset_index(drop=True)


def resolve_stage1_feature_groups(stage1_config_path: Path, requested_log_features: list[str]) -> tuple[list[str], dict[str, list[str]]]:
    with stage1_config_path.open("r", encoding="utf-8-sig") as file_obj:
        config = json.load(file_obj)
    features = [str(item).strip() for item in config.get("features", []) if str(item).strip()]
    if not features:
        raise ValueError(f"No features found in stage1 config: {stage1_config_path}")
    log_features = [col for col in requested_log_features if col in features]
    seismic_features = [col for col in features if col not in log_features]
    feature_groups = {"seismic": seismic_features}
    for log_col in log_features:
        feature_groups[log_col] = [log_col]
    return features, feature_groups


def prepare_similarity_df(
    sample_csv: Path,
    feature_cols: list[str],
    log_features: list[str],
) -> pd.DataFrame:
    df = pd.read_csv(sample_csv, encoding="utf-8-sig")
    keep_cols = [col for col in feature_cols if col in df.columns]
    if not keep_cols:
        raise ValueError(f"No requested similarity feature columns found in: {sample_csv}")
    out = df[keep_cols].copy()
    for col in keep_cols:
        series = pd.to_numeric(out[col], errors="coerce")
        if col in log_features:
            invalid_mask = series.isna() | series.isin([-999.25, -9999.0]) | (series <= -999.0)
            if col.upper() == "GR":
                invalid_mask |= series == 0
            out[col] = series.mask(invalid_mask)
        else:
            out[col] = series
    return out


def compute_group_distance(
    target_df: pd.DataFrame,
    candidate_df: pd.DataFrame,
    cols: list[str],
) -> float:
    col_scores = []
    for col in cols:
        if col not in target_df.columns or col not in candidate_df.columns:
            continue
        target_values = pd.to_numeric(target_df[col], errors="coerce").dropna().to_numpy(dtype=np.float64)
        candidate_values = pd.to_numeric(candidate_df[col], errors="coerce").dropna().to_numpy(dtype=np.float64)
        if target_values.size == 0 or candidate_values.size == 0:
            continue
        pooled = np.concatenate([target_values, candidate_values])
        scale = float(np.nanstd(pooled))
        if not np.isfinite(scale) or scale < 1e-8:
            scale = max(abs(float(np.nanmean(pooled))) if pooled.size else 0.0, 1.0)
        target_stats = [
            float(np.nanmean(target_values)),
            float(np.nanstd(target_values)),
            float(np.nanquantile(target_values, 0.25)),
            float(np.nanquantile(target_values, 0.50)),
            float(np.nanquantile(target_values, 0.75)),
        ]
        candidate_stats = [
            float(np.nanmean(candidate_values)),
            float(np.nanstd(candidate_values)),
            float(np.nanquantile(candidate_values, 0.25)),
            float(np.nanquantile(candidate_values, 0.50)),
            float(np.nanquantile(candidate_values, 0.75)),
        ]
        score = float(
            np.mean(
                [
                    abs(t_value - c_value) / scale
                    for t_value, c_value in zip(target_stats, candidate_stats)
                ]
            )
        )
        col_scores.append(score)
    if not col_scores:
        return np.nan
    return float(np.mean(col_scores))


def compute_similarity_scores(
    segment_paths: dict[str, Path],
    registry_df: pd.DataFrame,
    requested_log_features: list[str],
) -> pd.DataFrame:
    rows = []
    cache: dict[tuple[str, str], pd.DataFrame] = {}
    for row in registry_df.to_dict(orient="records"):
        strata_name = str(row["strata_name"])
        holdout_sample_csv = segment_paths.get(strata_name)
        reference_sample_csv = Path(str(row["reference_sample_csv"]))
        if holdout_sample_csv is None or not holdout_sample_csv.exists() or not reference_sample_csv.exists():
            continue

        stage1_config_path = Path(str(row["stage1_config_path"]))
        feature_cols, feature_groups = resolve_stage1_feature_groups(stage1_config_path, requested_log_features)
        log_features = [col for col in requested_log_features if col in feature_cols]

        target_cache_key = (strata_name, "target")
        if target_cache_key not in cache:
            cache[target_cache_key] = prepare_similarity_df(holdout_sample_csv, feature_cols, log_features)
        target_df = cache[target_cache_key]

        candidate_cache_key = (strata_name, str(reference_sample_csv))
        if candidate_cache_key not in cache:
            cache[candidate_cache_key] = prepare_similarity_df(reference_sample_csv, feature_cols, log_features)
        candidate_df = cache[candidate_cache_key]

        group_scores = {
            group_name: compute_group_distance(target_df, candidate_df, cols)
            for group_name, cols in feature_groups.items()
        }
        weighted_sum = 0.0
        weight_sum = 0.0
        for group_name, weight in SIMILARITY_GROUP_WEIGHTS.items():
            group_value = group_scores.get(group_name, np.nan)
            if np.isfinite(group_value):
                weighted_sum += float(weight) * float(group_value)
                weight_sum += float(weight)

        rows.append(
            {
                "strata_name": strata_name,
                "expert_well": row["expert_well"],
                "similarity_score": float(weighted_sum / weight_sum) if weight_sum > 0 else np.nan,
                "similarity_score_seismic": group_scores.get("seismic", np.nan),
                "similarity_score_AC": group_scores.get("AC", np.nan),
                "similarity_score_GR": group_scores.get("GR", np.nan),
                "num_stage1_features": int(len(feature_cols)),
                "stage1_model_dir": row["stage1_model_dir"],
                "stage2_model_dir": row["stage2_model_dir"],
                "inner_stage1_f1": row.get("inner_stage1_f1", np.nan),
                "inner_stage2_count_error_pct": row.get("inner_stage2_count_error_pct", np.nan),
            }
        )
    return pd.DataFrame(rows)


def compute_rank_loss(series: pd.Series, ascending: bool) -> pd.Series:
    numeric = pd.to_numeric(series, errors="coerce")
    out = pd.Series(1.0, index=series.index, dtype=np.float64)
    valid = numeric.notna()
    valid_count = int(valid.sum())
    if valid_count <= 0:
        return out
    if valid_count == 1:
        out.loc[valid] = 0.0
        return out
    ranks = numeric.loc[valid].rank(method="average", ascending=ascending)
    out.loc[valid] = (ranks - 1.0) / float(valid_count - 1)
    return out


def select_expert_for_strata(score_df: pd.DataFrame) -> pd.DataFrame:
    if score_df.empty:
        raise ValueError("No expert similarity rows were produced")
    rows = []
    for strata_name, sub_df in score_df.groupby("strata_name", sort=False):
        work_df = sub_df.copy()
        work_df["__score"] = pd.to_numeric(work_df["similarity_score"], errors="coerce")
        work_df["__f1"] = pd.to_numeric(work_df["inner_stage1_f1"], errors="coerce")
        work_df["__ce"] = pd.to_numeric(work_df["inner_stage2_count_error_pct"], errors="coerce")
        work_df["SimilarityRankLoss"] = compute_rank_loss(work_df["__score"], ascending=True)
        work_df["Stage1RankLoss"] = compute_rank_loss(work_df["__f1"], ascending=False)
        work_df["Stage2RankLoss"] = compute_rank_loss(work_df["__ce"], ascending=True)
        work_df["QualityQualified"] = (
            work_df["__f1"].ge(EXPERT_SELECTION_STAGE1_F1_FLOOR)
            & work_df["__ce"].le(EXPERT_SELECTION_STAGE2_COUNT_ERROR_CEIL)
        )
        qualified_df = work_df[work_df["QualityQualified"]].copy()
        select_df = qualified_df if not qualified_df.empty else work_df.copy()
        select_df["SelectionJointScore"] = (
            EXPERT_SELECTION_SIMILARITY_WEIGHT * select_df["SimilarityRankLoss"]
            + EXPERT_SELECTION_STAGE1_WEIGHT * select_df["Stage1RankLoss"]
            + EXPERT_SELECTION_STAGE2_WEIGHT * select_df["Stage2RankLoss"]
        )
        select_df = select_df.sort_values(
            ["SelectionJointScore", "__score", "__f1", "__ce", "expert_well"],
            ascending=[True, True, False, True, True],
        )
        best = select_df.iloc[0]
        rows.append(
            {
                "strata_name": strata_name,
                "selected_expert_well": best["expert_well"],
                "similarity_score": best["similarity_score"],
                "similarity_score_seismic": best["similarity_score_seismic"],
                "similarity_score_AC": best["similarity_score_AC"],
                "similarity_score_GR": best["similarity_score_GR"],
                "stage1_model_dir": best["stage1_model_dir"],
                "stage2_model_dir": best["stage2_model_dir"],
                "inner_stage1_f1": best["inner_stage1_f1"],
                "inner_stage2_count_error_pct": best["inner_stage2_count_error_pct"],
                "quality_qualified": bool(best["QualityQualified"]),
                "selection_joint_score": best["SelectionJointScore"],
                "selection_similarity_rank_loss": best["SimilarityRankLoss"],
                "selection_stage1_rank_loss": best["Stage1RankLoss"],
                "selection_stage2_rank_loss": best["Stage2RankLoss"],
            }
        )
    return pd.DataFrame(rows)


def resolve_depth_column(df: pd.DataFrame) -> str:
    for col in ["TVD", "DEPT", "MD"]:
        if col in df.columns:
            return col
    raise ValueError("No depth column found. Expected one of TVD/DEPT/MD.")


def parse_stage1_eval_by_rule(df: pd.DataFrame, eval_rule: str) -> tuple[pd.Series, pd.Series]:
    rule = str(eval_rule or "gt_label").strip().lower()
    if rule == "gt_label":
        if "GT_LABEL" not in df.columns:
            raise ValueError("Stage1 eval rule gt_label requires GT_LABEL column")
        gt_numeric = pd.to_numeric(df["GT_LABEL"], errors="coerce")
        return gt_numeric.ge(1.0), gt_numeric.notna()
    if rule == "frac_azimuth_notna":
        if "Frac_Azimuth" not in df.columns:
            raise ValueError("Stage1 eval rule frac_azimuth_notna requires Frac_Azimuth column")
        return df["Frac_Azimuth"].notna(), pd.Series(True, index=df.index)
    if rule == "any_density":
        density_cols = [col for col in ["P10", "P21", "P33"] if col in df.columns]
        if not density_cols:
            raise ValueError("Stage1 eval rule any_density requires one of P10/P21/P33")
        density_df = df[density_cols].apply(pd.to_numeric, errors="coerce").fillna(0.0)
        return density_df.max(axis=1).gt(0.0), pd.Series(True, index=df.index)
    raise ValueError(f"Unsupported stage1 eval rule: {eval_rule}")


def compute_stage1_eval_mask(df: pd.DataFrame, eval_rule: str) -> tuple[pd.Series, pd.Series]:
    y_true, gt_valid_mask = parse_stage1_eval_by_rule(df, eval_rule)
    pred_valid_mask = df["PRED_PROB"].notna() & df["PRED_LABEL"].notna()
    eval_mask = pred_valid_mask & gt_valid_mask
    return eval_mask, y_true


def compute_binary_metrics(y_true: np.ndarray, y_pred: np.ndarray) -> dict[str, float]:
    y_true_int = np.asarray(y_true, dtype=np.int64)
    y_pred_int = np.asarray(y_pred, dtype=np.int64)
    if y_true_int.size == 0:
        return {
            "Accuracy": np.nan,
            "Precision": np.nan,
            "Recall": np.nan,
            "F1": np.nan,
            "IoU": np.nan,
            "PredPositiveRatio": np.nan,
            "TruePositiveRatio": np.nan,
        }
    tp = int(np.sum((y_true_int == 1) & (y_pred_int == 1)))
    tn = int(np.sum((y_true_int == 0) & (y_pred_int == 0)))
    fp = int(np.sum((y_true_int == 0) & (y_pred_int == 1)))
    fn = int(np.sum((y_true_int == 1) & (y_pred_int == 0)))
    precision = float(tp / (tp + fp)) if (tp + fp) > 0 else 0.0
    recall = float(tp / (tp + fn)) if (tp + fn) > 0 else 0.0
    f1 = float((2 * precision * recall) / (precision + recall)) if (precision + recall) > 0 else 0.0
    iou = float(tp / (tp + fp + fn)) if (tp + fp + fn) > 0 else 0.0
    accuracy = float((tp + tn) / y_true_int.size)
    return {
        "Accuracy": accuracy,
        "Precision": precision,
        "Recall": recall,
        "F1": f1,
        "IoU": iou,
        "PredPositiveRatio": float(np.mean(y_pred_int == 1)),
        "TruePositiveRatio": float(np.mean(y_true_int == 1)),
    }


def evaluate_stage1_output(
    strata_name: str,
    holdout_well: str,
    selected_expert_well: str,
    stage1_pred_csv: Path,
    eval_rule: str,
) -> tuple[dict, pd.DataFrame]:
    df = pd.read_csv(stage1_pred_csv, encoding="utf-8-sig")
    eval_mask, y_true_series = compute_stage1_eval_mask(df, eval_rule)
    eval_df = df.loc[eval_mask].copy()
    eval_df["Stage1EvalYTrue"] = y_true_series.loc[eval_mask].astype(np.int64)
    eval_df["Stage1EvalYPred"] = pd.to_numeric(eval_df["PRED_LABEL"], errors="coerce").fillna(0).astype(np.int64)
    eval_df["Stage1EvalYProb"] = pd.to_numeric(eval_df["PRED_PROB"], errors="coerce")
    eval_df["StrataName"] = strata_name
    eval_df["HoldoutWell"] = holdout_well
    eval_df["SelectedExpertWell"] = selected_expert_well

    metrics = compute_binary_metrics(
        y_true=eval_df["Stage1EvalYTrue"].to_numpy(dtype=np.int64),
        y_pred=eval_df["Stage1EvalYPred"].to_numpy(dtype=np.int64),
    )
    result = {
        "holdout_well": holdout_well,
        "strata_name": strata_name,
        "selected_expert_well": selected_expert_well,
        "stage1_pred_csv": str(stage1_pred_csv),
        "stage1_eval_rule": eval_rule,
        "stage1_n_eval_samples": int(len(eval_df)),
        "stage1_threshold": float(pd.to_numeric(df["FIRST_STAGE_THRESHOLD"], errors="coerce").dropna().iloc[0])
        if "FIRST_STAGE_THRESHOLD" in df.columns and not pd.to_numeric(df["FIRST_STAGE_THRESHOLD"], errors="coerce").dropna().empty
        else np.nan,
        "stage1_num_predicted_positive_centers": int(eval_df["Stage1EvalYPred"].sum()),
    }
    result.update(
        {
            "stage1_accuracy": metrics["Accuracy"],
            "stage1_precision": metrics["Precision"],
            "stage1_recall": metrics["Recall"],
            "stage1_f1": metrics["F1"],
            "stage1_iou": metrics["IoU"],
            "stage1_pred_positive_ratio": metrics["PredPositiveRatio"],
            "stage1_true_positive_ratio": metrics["TruePositiveRatio"],
        }
    )
    return result, eval_df


def build_whole_well_stage1_summary(
    holdout_well: str,
    stage1_eval_parts: list[pd.DataFrame],
) -> dict:
    if stage1_eval_parts:
        eval_df = pd.concat(stage1_eval_parts, ignore_index=True)
    else:
        eval_df = pd.DataFrame(columns=["Stage1EvalYTrue", "Stage1EvalYPred"])
    metrics = compute_binary_metrics(
        y_true=eval_df["Stage1EvalYTrue"].to_numpy(dtype=np.int64) if not eval_df.empty else np.array([], dtype=np.int64),
        y_pred=eval_df["Stage1EvalYPred"].to_numpy(dtype=np.int64) if not eval_df.empty else np.array([], dtype=np.int64),
    )
    return {
        "holdout_well": holdout_well,
        "stage1_n_eval_samples": int(len(eval_df)),
        "stage1_accuracy": metrics["Accuracy"],
        "stage1_precision": metrics["Precision"],
        "stage1_recall": metrics["Recall"],
        "stage1_f1": metrics["F1"],
        "stage1_iou": metrics["IoU"],
        "stage1_pred_positive_ratio": metrics["PredPositiveRatio"],
        "stage1_true_positive_ratio": metrics["TruePositiveRatio"],
    }


def evaluate_stage2_prediction(
    args,
    strata_name: str,
    holdout_well: str,
    selected_expert_row: dict,
    stage1_pred_csv: Path,
    strata_range_df: pd.DataFrame,
    output_dir: Path,
) -> tuple[dict, dict]:
    raw_module = get_raw_point_module()
    output_dir.mkdir(parents=True, exist_ok=True)

    artifact = raw_module.load_model_artifact(Path(str(selected_expert_row["stage2_model_dir"])))
    config = raw_module.dataset_builder.make_config(
        exist_exp_dir=str(artifact.get("exist_exp_dir", "")),
        prob_weight_gamma=float(artifact.get("prob_weight_gamma", args.stage2_prob_weight_gamma)),
        gt_min_points_per_segment=1,
        gt_rounding_mode="round",
        density_gt_col="P10",
    )
    well_data = raw_module.build_well_segment_dataset(
        exist_csv=str(stage1_pred_csv),
        well_name=holdout_well,
        config=config,
        sample_dir=Path(args.sample_dir),
        raw_label_dir=Path(args.raw_label_dir),
        gt_dev_rule=str(args.stage2_gt_dev_rule),
        strata_name=strata_name,
        strata_range_df=strata_range_df,
        boundary_expand_mode=str(args.stage2_boundary_expand_mode),
        boundary_expand_prob_min=float(args.stage2_boundary_expand_prob_min),
        boundary_expand_max_steps=int(args.stage2_boundary_expand_max_steps),
        boundary_expand_max_depth=float(args.stage2_boundary_expand_max_depth),
    )

    raw_gt_points = well_data["raw_gt_points"]
    raw_gt_depths = well_data["raw_gt_depths"]
    if well_data["segment_df"].empty:
        pred_segment_df = well_data["segment_df"].copy()
        pred_segment_df["SegmentCountAbsError"] = []
        pred_points = pd.DataFrame()
        lowconf_mask = np.array([], dtype=bool)
    else:
        pred_segment_df, lowconf_mask = raw_module.predict_segments_with_artifact(
            segment_df=well_data["segment_df"].copy(),
            artifact=artifact,
            include_gt_columns=True,
        )
        pred_points = raw_module.build_pred_points_from_segment_df(
            df_sorted=well_data["df_sorted"],
            payloads=well_data["payloads"],
            pred_segment_df=pred_segment_df,
            config=config,
            orientation_strategy=dict(artifact.get("orientation_strategy") or {}),
        )

    raw_gt_points.to_csv(output_dir / "raw_gt_points.csv", index=False, encoding="utf-8-sig")
    well_data["gt_dev_segments"].to_csv(output_dir / "gt_dev_segments.csv", index=False, encoding="utf-8-sig")
    pred_segment_df.to_csv(output_dir / "pred_segment_summary.csv", index=False, encoding="utf-8-sig")
    pred_points.to_csv(output_dir / "pred_fracture_points.csv", index=False, encoding="utf-8-sig")

    pred_depths = (
        pred_points[config.depth_col].to_numpy(dtype=np.float64)
        if not pred_points.empty and config.depth_col in pred_points.columns
        else np.array([], dtype=np.float64)
    )
    pred_to_gt = raw_module.dataset_builder.POINT.nearest_distance_stats(pred_depths, raw_gt_depths)
    gt_to_pred = raw_module.dataset_builder.POINT.nearest_distance_stats(raw_gt_depths, pred_depths)
    segment_count_mae = (
        float(pd.to_numeric(pred_segment_df["SegmentCountAbsError"], errors="coerce").dropna().mean())
        if "SegmentCountAbsError" in pred_segment_df.columns and not pred_segment_df.empty
        else np.nan
    )

    metrics = {
        "holdout_well": holdout_well,
        "strata_name": strata_name,
        "selected_expert_well": selected_expert_row["selected_expert_well"],
        "stage2_model_dir": str(selected_expert_row["stage2_model_dir"]),
        "N_gt_points": int(len(raw_gt_depths)),
        "N_pred_points": int(len(pred_depths)),
        "count_diff": int(len(pred_depths) - len(raw_gt_depths)),
        "count_error_pct": raw_module.compute_count_error_pct(int(len(pred_depths) - len(raw_gt_depths)), int(len(raw_gt_depths))),
        "pred_to_gt_mean_dist": pred_to_gt["mean"],
        "pred_to_gt_p90_dist": pred_to_gt["p90"],
        "gt_to_pred_mean_dist": gt_to_pred["mean"],
        "gt_to_pred_p90_dist": gt_to_pred["p90"],
        "segment_count_mae": segment_count_mae,
        "raw_points_covered_by_pred_segments": int(well_data["raw_points_covered_by_pred_segments"]),
        "raw_points_missed_outside_pred_segments": int(well_data["raw_points_missed_outside_pred_segments"]),
        "num_pred_segments": int(len(pred_segment_df)),
        "num_lowconf_segments": int(np.sum(lowconf_mask.astype(np.int64))),
    }
    return metrics, {
        "config": config,
        "raw_gt_points": raw_gt_points,
        "pred_points": pred_points,
        "pred_segment_df": pred_segment_df,
        "stage1_pred_df": pd.read_csv(stage1_pred_csv, encoding="utf-8-sig"),
    }


def build_whole_well_stage2_summary(
    holdout_well: str,
    stage2_payloads: list[dict],
    raw_module,
) -> tuple[dict, dict[str, pd.DataFrame]]:
    combined_raw_gt_points = pd.concat(
        [item["raw_gt_points"] for item in stage2_payloads if not item["raw_gt_points"].empty],
        ignore_index=True,
    ) if stage2_payloads else pd.DataFrame()
    combined_pred_points = pd.concat(
        [item["pred_points"] for item in stage2_payloads if not item["pred_points"].empty],
        ignore_index=True,
    ) if stage2_payloads else pd.DataFrame()
    combined_pred_segment_df = pd.concat(
        [item["pred_segment_df"] for item in stage2_payloads if not item["pred_segment_df"].empty],
        ignore_index=True,
    ) if stage2_payloads else pd.DataFrame()
    combined_stage1_pred_df = pd.concat(
        [item["stage1_pred_df"] for item in stage2_payloads if not item["stage1_pred_df"].empty],
        ignore_index=True,
    ) if stage2_payloads else pd.DataFrame()

    depth_col = "TVD"
    for df in [combined_pred_points, combined_raw_gt_points, combined_stage1_pred_df]:
        if not df.empty:
            depth_col = resolve_depth_column(df)
            break
    gt_depths = (
        combined_raw_gt_points[depth_col].to_numpy(dtype=np.float64)
        if not combined_raw_gt_points.empty and depth_col in combined_raw_gt_points.columns
        else np.array([], dtype=np.float64)
    )
    pred_depths = (
        combined_pred_points[depth_col].to_numpy(dtype=np.float64)
        if not combined_pred_points.empty and depth_col in combined_pred_points.columns
        else np.array([], dtype=np.float64)
    )
    pred_to_gt = raw_module.dataset_builder.POINT.nearest_distance_stats(pred_depths, gt_depths)
    gt_to_pred = raw_module.dataset_builder.POINT.nearest_distance_stats(gt_depths, pred_depths)
    count_diff = int(len(pred_depths) - len(gt_depths))
    segment_count_mae = (
        float(pd.to_numeric(combined_pred_segment_df["SegmentCountAbsError"], errors="coerce").dropna().mean())
        if "SegmentCountAbsError" in combined_pred_segment_df.columns and not combined_pred_segment_df.empty
        else np.nan
    )

    summary_row = {
        "holdout_well": holdout_well,
        "N_gt_points": int(len(gt_depths)),
        "N_pred_points": int(len(pred_depths)),
        "count_diff": count_diff,
        "count_error_pct": raw_module.compute_count_error_pct(count_diff, int(len(gt_depths))),
        "pred_to_gt_mean_dist": pred_to_gt["mean"],
        "pred_to_gt_p90_dist": pred_to_gt["p90"],
        "gt_to_pred_mean_dist": gt_to_pred["mean"],
        "gt_to_pred_p90_dist": gt_to_pred["p90"],
        "segment_count_mae": segment_count_mae,
    }
    return summary_row, {
        "combined_raw_gt_points": combined_raw_gt_points,
        "combined_pred_points": combined_pred_points,
        "combined_pred_segment_df": combined_pred_segment_df,
        "combined_stage1_pred_df": combined_stage1_pred_df,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", default=f"outer_holdout_expert_validation_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
    parser.add_argument("--python-exe", default=str(PYTHON_EXE_DEFAULT))
    parser.add_argument("--result-dir", default="")
    parser.add_argument("--sample-dir", default=str(DEFAULT_SAMPLE_DIR))
    parser.add_argument("--raw-label-dir", default=str(DEFAULT_RAW_LABEL_DIR))
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument("--well-names", default=",".join(DEFAULT_WELL_NAMES))
    parser.add_argument("--holdout-well", default="车151HF")
    parser.add_argument("--target-strata", default="")
    parser.add_argument("--boundary-tolerance", type=float, default=1.0)
    parser.add_argument("--reuse-inner-libraries", type=parse_bool, default=False)

    parser.add_argument("--stage1-lstm-features-json", default=json.dumps(DEFAULT_STAGE1_FEATURES, ensure_ascii=False))
    parser.add_argument("--stage1-seq-len", type=int, default=5)
    parser.add_argument("--stage1-seis-mode", default="3x3")
    parser.add_argument("--stage1-selection-metric", default="iou")
    parser.add_argument("--stage1-use-dynamic-threshold", type=parse_bool, default=True)
    parser.add_argument("--stage1-use-density-regression", type=parse_bool, default=False)
    parser.add_argument("--stage1-use-domain-adversarial", type=parse_bool, default=False)
    parser.add_argument("--stage1-use-wellwise-log-scaling", type=parse_bool, default=True)
    parser.add_argument("--stage1-use-well-balanced-sampling", type=parse_bool, default=True)
    parser.add_argument("--stage1-use-oversample", type=parse_bool, default=False)
    parser.add_argument("--stage1-manual-pos-weight", type=float, default=2.0)
    parser.add_argument("--stage1-target-ratio", type=float, default=0.35)
    parser.add_argument("--stage1-edge-exclude", type=int, default=1)
    parser.add_argument("--stage1-max-pos-weight", type=float, default=3.0)
    parser.add_argument("--stage1-use-dice-loss", type=parse_bool, default=True)
    parser.add_argument("--stage1-dice-loss-weight", type=float, default=0.5)
    parser.add_argument("--stage1-fixed-threshold", type=float, default=0.5)
    parser.add_argument("--stage1-thresh-search-min", type=float, default=0.30)
    parser.add_argument("--stage1-thresh-search-max", type=float, default=0.90)
    parser.add_argument("--stage1-thresh-search-step", type=float, default=0.02)
    parser.add_argument("--stage1-use-selection-accuracy-floor", type=parse_bool, default=True)
    parser.add_argument("--stage1-min-selection-accuracy", type=float, default=0.80)
    parser.add_argument("--stage1-use-threshold-constraints", type=parse_bool, default=True)
    parser.add_argument("--stage1-min-selection-recall", type=float, default=0.20)
    parser.add_argument("--stage1-min-selection-pos-ratio", type=float, default=0.03)
    parser.add_argument("--stage1-min-selection-pos-ratio-scale", type=float, default=0.35)
    parser.add_argument("--stage1-use-selection-accuracy-floor-by-well-json", default=json.dumps({"车151HF": False, "车663": False}, ensure_ascii=False))
    parser.add_argument("--stage1-min-selection-accuracy-by-well-json", default="")
    parser.add_argument("--stage1-use-threshold-constraints-by-well-json", default="")
    parser.add_argument("--stage1-min-selection-recall-by-well-json", default=json.dumps({"车151HF": 0.65, "车663": 0.25}, ensure_ascii=False))
    parser.add_argument("--stage1-min-selection-pos-ratio-by-well-json", default="")
    parser.add_argument("--stage1-min-selection-pos-ratio-scale-by-well-json", default="")
    parser.add_argument("--stage1-dist-threshold", type=float, default=0.35)
    parser.add_argument("--stage1-min-train-wells", type=int, default=3)
    parser.add_argument("--stage1-min-seq-per-well", type=int, default=30)
    parser.add_argument("--stage1-train-selection-mode", default="distance")
    parser.add_argument("--stage1-eval-rule", default="frac_azimuth_notna")
    parser.add_argument("--stage1-threshold-override", type=float, default=float("nan"))

    parser.add_argument("--stage2-gt-dev-rule", default="any_density")
    parser.add_argument("--stage2-config-profile", default="")
    parser.add_argument("--stage2-config-profile-map-json", default="")
    parser.add_argument("--stage2-prob-weight-gamma", type=float, default=2.0)
    parser.add_argument("--stage2-pred-min-points-per-segment", type=int, default=0)
    parser.add_argument("--stage2-rounding-mode", default="round")
    parser.add_argument("--stage2-fallback-count-mode", default="global_prob_mass_scale")
    parser.add_argument("--stage2-count-train-mode", default="positive_only")
    parser.add_argument("--stage2-positive-floor-prob-min", type=float, default=0.7)
    parser.add_argument("--stage2-positive-floor-allowed-labels", default="other,short_strong")
    parser.add_argument("--stage2-positive-floor-target-wells", default="")
    parser.add_argument("--stage2-type-calibration-mode", default="rule_v1")
    parser.add_argument("--stage2-type-calibration-scale-min", type=float, default=0.25)
    parser.add_argument("--stage2-type-calibration-scale-max", type=float, default=1.15)
    parser.add_argument("--stage2-count-fusion-mode", default="max")
    parser.add_argument("--stage2-fusion-learned-weight", type=float, default=0.5)
    parser.add_argument("--stage2-global-post-scale-mode", default="none")
    parser.add_argument("--stage2-global-post-scale-min", type=float, default=0.5)
    parser.add_argument("--stage2-global-post-scale-max", type=float, default=1.5)
    parser.add_argument("--stage2-selective-post-mode", default="train_zero_dual_quantile_downscale")
    parser.add_argument("--stage2-selective-post-target-wells", default="车660-2")
    parser.add_argument("--stage2-selective-post-feature", default="ProbMassPerLength")
    parser.add_argument("--stage2-selective-post-feature-2", default="ProbMean")
    parser.add_argument("--stage2-selective-post-quantile", type=float, default=0.7)
    parser.add_argument("--stage2-selective-post-quantile-2", type=float, default=0.7)
    parser.add_argument("--stage2-selective-post-max-pred-float", type=float, default=0.6)
    parser.add_argument("--stage2-soft-negative-weight", type=float, default=0.5)
    parser.add_argument("--stage2-train-distance-weight-mode", default="none")
    parser.add_argument("--stage2-train-distance-csv", default="")
    parser.add_argument("--stage2-train-distance-weight-power", type=float, default=1.0)
    parser.add_argument("--stage2-train-distance-weight-min", type=float, default=0.7)
    parser.add_argument("--stage2-train-distance-weight-max", type=float, default=1.3)
    parser.add_argument("--stage2-boundary-expand-mode", default="prob_shoulder")
    parser.add_argument("--stage2-boundary-expand-prob-min", type=float, default=0.4)
    parser.add_argument("--stage2-boundary-expand-max-steps", type=int, default=3)
    parser.add_argument("--stage2-boundary-expand-max-depth", type=float, default=0.0)
    parser.add_argument("--stage2-gate-mode", default="logistic_soft")
    parser.add_argument("--stage2-gate-prob-threshold", type=float, default=0.45)
    parser.add_argument("--stage2-gate-target-wells", default="车151HF,车662")
    parser.add_argument("--stage2-orientation-mode", default="family_classifier")
    parser.add_argument("--stage2-orientation-num-families", type=int, default=3)
    parser.add_argument("--stage2-orientation-min-points-per-segment", type=int, default=1)
    parser.add_argument("--stage2-model-name", default="tweedie")
    parser.add_argument("--stage2-model-alpha", type=float, default=0.03)
    parser.add_argument("--stage2-model-max-iter", type=int, default=3000)
    parser.add_argument("--stage2-tweedie-power", type=float, default=1.5)
    parser.add_argument("--stage2-feature-preset", default="high_density_compact_v1")
    parser.add_argument("--stage2-feature-cols", default="")
    parser.add_argument("--stage2-xgb-n-estimators", type=int, default=300)
    parser.add_argument("--stage2-xgb-max-depth", type=int, default=3)
    parser.add_argument("--stage2-xgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--stage2-xgb-subsample", type=float, default=0.8)
    parser.add_argument("--stage2-xgb-colsample-bytree", type=float, default=0.8)
    parser.add_argument("--stage2-hgb-learning-rate", type=float, default=0.05)
    parser.add_argument("--stage2-hgb-max-depth", type=int, default=3)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    args.stage1_lstm_features = parse_list_arg(args.stage1_lstm_features_json)
    if not args.stage1_lstm_features:
        raise ValueError("stage1_lstm_features is empty")

    selected_wells = []
    for well_name in parse_list_arg(args.well_names):
        canonical_name = canonicalize_well_name(well_name)
        if canonical_name and canonical_name not in selected_wells:
            selected_wells.append(canonical_name)
    holdout_well = canonicalize_well_name(args.holdout_well)
    args.holdout_well = holdout_well
    if holdout_well not in selected_wells:
        selected_wells.append(holdout_well)
    training_wells = [well_name for well_name in selected_wells if well_name != holdout_well]
    if not training_wells:
        raise ValueError("No training wells available after removing holdout well")

    result_root = (
        Path(args.result_dir).resolve()
        if str(args.result_dir).strip()
        else (DEFAULT_SAVE_ROOT / sanitize(args.exp_id)).resolve()
    )
    result_root.mkdir(parents=True, exist_ok=True)

    labeled_df, source_columns = load_all_labeled_data(
        sample_dir=Path(args.sample_dir),
        selected_wells=selected_wells,
        boundary_tolerance=float(args.boundary_tolerance),
    )
    requested_target_strata = parse_list_arg(args.target_strata)
    target_strata = infer_holdout_target_strata(labeled_df, holdout_well, requested_target_strata)

    label_summary_df = (
        labeled_df.groupby(["WellName", "StrataName"], dropna=False)
        .agg(
            SampleCount=("WellName", "size"),
            DepthMin=("DepthForStrata", "min"),
            DepthMax=("DepthForStrata", "max"),
        )
        .reset_index()
    )
    label_summary_df.to_csv(result_root / "manual_strata_label_summary.csv", index=False, encoding="utf-8-sig")

    holdout_segment_dir = result_root / "holdout_segments"
    segment_paths, holdout_strata_range_df = build_holdout_segment_files(
        labeled_df=labeled_df,
        source_columns=source_columns,
        holdout_well=holdout_well,
        target_strata=target_strata,
        output_dir=holdout_segment_dir,
    )
    if not segment_paths:
        raise ValueError(f"No holdout strata segment files were exported for {holdout_well}")

    stage1_library_dir = result_root / "inner_stage1_library"
    stage2_library_dir = result_root / "inner_stage2_library"
    if not bool(args.reuse_inner_libraries) or not (stage1_library_dir.exists() and stage2_library_dir.exists()):
        library_dirs = train_inner_libraries(
            args=args,
            result_root=result_root,
            training_wells=training_wells,
            target_strata=target_strata,
        )
        stage1_library_dir = library_dirs["stage1_library_dir"]
        stage2_library_dir = library_dirs["stage2_library_dir"]

    registry_df = build_expert_registry(
        stage1_library_dir=stage1_library_dir,
        stage2_library_dir=stage2_library_dir,
        target_strata=target_strata,
    )
    if registry_df.empty:
        raise ValueError("Expert registry is empty")
    registry_df.to_csv(result_root / "expert_model_registry.csv", index=False, encoding="utf-8-sig")

    score_df = compute_similarity_scores(
        segment_paths=segment_paths,
        registry_df=registry_df,
        requested_log_features=args.stage1_lstm_features,
    )
    if score_df.empty:
        raise ValueError("No similarity scores were produced")
    score_df.to_csv(result_root / "expert_selection_scores.csv", index=False, encoding="utf-8-sig")

    selected_expert_df = select_expert_for_strata(score_df)
    selected_expert_df.to_csv(result_root / "selected_expert_by_strata.csv", index=False, encoding="utf-8-sig")

    nearest_module = get_nearest_module()
    raw_module = get_raw_point_module()
    stage1_rows = []
    stage1_eval_parts = []
    stage2_rows = []
    stage2_payloads = []

    for strata_name in target_strata:
        selected_match = selected_expert_df[selected_expert_df["strata_name"].astype(str) == str(strata_name)]
        if selected_match.empty:
            raise ValueError(f"Missing selected expert row for strata: {strata_name}")
        selected_row = selected_match.iloc[0].to_dict()

        strata_output_dir = result_root / "prediction_by_strata" / sanitize(strata_name)
        strata_output_dir.mkdir(parents=True, exist_ok=True)

        threshold_override = float(args.stage1_threshold_override) if np.isfinite(args.stage1_threshold_override) else None
        stage1_info = nearest_module.run_first_stage_prediction(
            input_csv=segment_paths[strata_name],
            well_name=holdout_well,
            model_well=str(selected_row["selected_expert_well"]),
            model_dir=Path(str(selected_row["stage1_model_dir"])),
            output_dir=strata_output_dir,
            threshold_override=threshold_override,
        )

        stage1_row, stage1_eval_df = evaluate_stage1_output(
            strata_name=strata_name,
            holdout_well=holdout_well,
            selected_expert_well=str(selected_row["selected_expert_well"]),
            stage1_pred_csv=Path(stage1_info["output_csv"]),
            eval_rule=str(args.stage1_eval_rule),
        )
        stage1_row.update(
            {
                "similarity_score": selected_row["similarity_score"],
                "similarity_score_seismic": selected_row["similarity_score_seismic"],
                "similarity_score_AC": selected_row["similarity_score_AC"],
                "similarity_score_GR": selected_row["similarity_score_GR"],
                "stage1_model_dir": selected_row["stage1_model_dir"],
                "stage1_num_input_rows": stage1_info["num_input_rows"],
                "stage1_num_rows_after_clean": stage1_info["num_rows_after_clean"],
                "stage1_num_removed_rows": stage1_info["num_removed_rows"],
                "stage1_num_sequences": stage1_info["num_sequences"],
            }
        )
        stage1_eval_df.to_csv(strata_output_dir / "stage1_eval_rows.csv", index=False, encoding="utf-8-sig")
        stage1_rows.append(stage1_row)
        stage1_eval_parts.append(stage1_eval_df)

        stage2_row, stage2_payload = evaluate_stage2_prediction(
            args=args,
            strata_name=strata_name,
            holdout_well=holdout_well,
            selected_expert_row=selected_row,
            stage1_pred_csv=Path(stage1_info["output_csv"]),
            strata_range_df=holdout_strata_range_df,
            output_dir=strata_output_dir,
        )
        stage2_row.update(
            {
                "similarity_score": selected_row["similarity_score"],
                "similarity_score_seismic": selected_row["similarity_score_seismic"],
                "similarity_score_AC": selected_row["similarity_score_AC"],
                "similarity_score_GR": selected_row["similarity_score_GR"],
                "stage1_f1": stage1_row["stage1_f1"],
                "stage1_iou": stage1_row["stage1_iou"],
                "stage1_recall": stage1_row["stage1_recall"],
                "stage1_precision": stage1_row["stage1_precision"],
            }
        )
        with (strata_output_dir / "stage2_metrics.json").open("w", encoding="utf-8") as file_obj:
            json.dump(to_jsonable(stage2_row), file_obj, ensure_ascii=False, indent=2)
        stage2_rows.append(stage2_row)
        stage2_payloads.append(stage2_payload)

    stage1_summary_df = pd.DataFrame(stage1_rows)
    stage1_summary_df.to_csv(result_root / "stage1_strata_summary.csv", index=False, encoding="utf-8-sig")
    stage1_whole_row = build_whole_well_stage1_summary(holdout_well=holdout_well, stage1_eval_parts=stage1_eval_parts)
    pd.DataFrame([stage1_whole_row]).to_csv(result_root / "stage1_whole_well_summary.csv", index=False, encoding="utf-8-sig")

    stage2_summary_df = pd.DataFrame(stage2_rows)
    stage2_summary_df.to_csv(result_root / "stage2_strata_summary.csv", index=False, encoding="utf-8-sig")
    whole_well_stage2_row, combined_outputs = build_whole_well_stage2_summary(
        holdout_well=holdout_well,
        stage2_payloads=stage2_payloads,
        raw_module=raw_module,
    )
    whole_well_stage2_row["stage1_f1"] = stage1_whole_row["stage1_f1"]
    whole_well_stage2_row["stage1_iou"] = stage1_whole_row["stage1_iou"]
    whole_well_stage2_row["stage1_recall"] = stage1_whole_row["stage1_recall"]
    whole_well_stage2_row["stage1_precision"] = stage1_whole_row["stage1_precision"]
    pd.DataFrame([whole_well_stage2_row]).to_csv(result_root / "whole_well_final_summary.csv", index=False, encoding="utf-8-sig")

    combined_outputs["combined_raw_gt_points"].to_csv(result_root / "whole_well_raw_gt_points.csv", index=False, encoding="utf-8-sig")
    combined_outputs["combined_pred_points"].to_csv(result_root / "whole_well_pred_fracture_points.csv", index=False, encoding="utf-8-sig")
    combined_outputs["combined_pred_segment_df"].to_csv(result_root / "whole_well_pred_segment_summary.csv", index=False, encoding="utf-8-sig")
    combined_outputs["combined_stage1_pred_df"].to_csv(result_root / "whole_well_stage1_pred_full_log.csv", index=False, encoding="utf-8-sig")

    summary = {
        "exp_id": args.exp_id,
        "holdout_well": holdout_well,
        "training_wells": training_wells,
        "target_strata": target_strata,
        "result_root": str(result_root),
        "stage1_library_dir": str(stage1_library_dir),
        "stage2_library_dir": str(stage2_library_dir),
        "registry_csv": str(result_root / "expert_model_registry.csv"),
        "score_csv": str(result_root / "expert_selection_scores.csv"),
        "selected_expert_csv": str(result_root / "selected_expert_by_strata.csv"),
        "stage1_strata_summary_csv": str(result_root / "stage1_strata_summary.csv"),
        "stage2_strata_summary_csv": str(result_root / "stage2_strata_summary.csv"),
        "whole_well_final_summary_csv": str(result_root / "whole_well_final_summary.csv"),
    }
    with (result_root / "summary.json").open("w", encoding="utf-8") as file_obj:
        json.dump(to_jsonable(summary), file_obj, ensure_ascii=False, indent=2)

    append_outer_validation_to_docx(
        docx_path=Path(args.docx_path),
        title=f"实验 {args.exp_id}",
        config=summary,
        strata_rows=stage2_summary_df.to_dict(orient="records"),
        whole_well_row=whole_well_stage2_row,
    )
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
