from __future__ import annotations

import argparse
import json
import math
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document

from run_conventional_log_strata_validation import (
    CSV_ENCODINGS,
    DEFAULT_DOCX_PATH,
    DEFAULT_LAYER_DIR,
    DEFAULT_STAGE1_LIBRARY_DIR,
    DEFAULT_STAGE2_LIBRARY_DIR,
    FINAL_LOG_FILENAME,
    FINAL_META_FILENAME,
    FINAL_POINTS_FILENAME,
    FINAL_SEGMENTS_FILENAME,
    FINAL_STRATA_SEGMENTATION_FILENAME,
    infer_well_name,
)
from strata_expert_deploy.runtime import DEFAULT_LOG_FEATURES, parse_json_list, sanitize


DEFAULT_INCLINED_DIR = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容一\井斜\测井-地震时窗"
)
DEFAULT_VERTICAL_DIR = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容一\测井\测井-地震时窗"
)
DEFAULT_OUTPUT_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容一\两阶段裂缝预测流程结果\现有常规测井裂缝预测"
)

SUMMARY_CSV_NAME = "batch_prediction_summary.csv"
MISSING_FEATURE_DOCX_NAME = "缺失AC_GR属性情况.docx"
REQUIRED_RESULT_FILES = {
    FINAL_LOG_FILENAME,
    FINAL_POINTS_FILENAME,
    FINAL_SEGMENTS_FILENAME,
    FINAL_STRATA_SEGMENTATION_FILENAME,
    FINAL_META_FILENAME,
}


def read_csv_header_flexible(path: Path) -> tuple[list[str], str]:
    last_error: Exception | None = None
    for encoding in CSV_ENCODINGS:
        try:
            header_df = pd.read_csv(path, encoding=encoding, nrows=0)
            return [str(col).strip() for col in header_df.columns], encoding
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Failed to read CSV header: {path}") from last_error


def discover_target_csvs(
    inclined_dir: Path,
    vertical_dir: Path,
    include_wells: list[str] | None = None,
    limit: int = 0,
) -> list[tuple[str, Path]]:
    include_set = {str(item).strip() for item in (include_wells or []) if str(item).strip()}
    rows: list[tuple[str, Path]] = []
    for source_group, base_dir in [("inclined", inclined_dir), ("vertical", vertical_dir)]:
        if not base_dir.exists():
            raise FileNotFoundError(f"Input directory not found: {base_dir}")
        for csv_path in sorted(base_dir.glob("*.csv"), key=lambda item: item.name):
            well_name = infer_well_name(csv_path)
            if include_set and well_name not in include_set:
                continue
            rows.append((source_group, csv_path))
    rows = sorted(rows, key=lambda item: (item[0], infer_well_name(item[1])))
    if limit and limit > 0:
        rows = rows[:limit]
    return rows


def detect_feature_status(csv_path: Path, requested_candidates: list[str]) -> dict:
    columns, encoding = read_csv_header_flexible(csv_path)
    col_set = set(columns)
    available_requested = [col for col in requested_candidates if col in col_set]
    missing_requested = [col for col in requested_candidates if col not in col_set]
    depth_col = next((col for col in ["TVD", "DEPT", "MD"] if col in col_set), "")
    has_required_basic = all(col in col_set for col in ["TIME", "X", "Y"]) and bool(depth_col)
    has_ac = "AC" in col_set
    has_gr = "GR" in col_set

    if not has_required_basic:
        handling = "skip_missing_basic_columns"
    elif len(missing_requested) == 0:
        handling = "use_all_requested_features"
    else:
        handling = f"skip_missing_log_features_{'_'.join(missing_requested)}"

    return {
        "encoding": encoding,
        "columns": columns,
        "depth_col": depth_col,
        "has_time": int("TIME" in col_set),
        "has_x": int("X" in col_set),
        "has_y": int("Y" in col_set),
        "has_ac": int(has_ac),
        "has_gr": int(has_gr),
        "available_requested": available_requested,
        "missing_requested": missing_requested,
        "handling": handling,
    }


def is_completed_result_dir(result_dir: Path) -> bool:
    if not result_dir.exists() or not result_dir.is_dir():
        return False
    existing_names = {child.name for child in result_dir.iterdir()}
    return REQUIRED_RESULT_FILES.issubset(existing_names)


def write_summary_csv(summary_rows: list[dict], save_path: Path) -> None:
    df = pd.DataFrame(summary_rows)
    if not df.empty:
        preferred_cols = [
            "WellName",
            "SourceGroup",
            "CsvPath",
            "ResultDir",
            "HeaderEncoding",
            "DepthColumn",
            "HasTIME",
            "HasX",
            "HasY",
            "HasAC",
            "HasGR",
            "RequestedFeatureCandidates",
            "UsedLogFeatures",
            "MissingRequestedFeatures",
            "FeatureHandling",
            "RunStatus",
            "ReturnCode",
            "ElapsedSeconds",
            "Note",
        ]
        ordered_cols = [col for col in preferred_cols if col in df.columns] + [
            col for col in df.columns if col not in preferred_cols
        ]
        df = df[ordered_cols]
    last_error: Exception | None = None
    for _ in range(3):
        try:
            save_path.parent.mkdir(parents=True, exist_ok=True)
            df.to_csv(save_path, index=False, encoding="utf-8-sig")
            return
        except Exception as exc:
            last_error = exc
            time.sleep(2.0)
    if last_error is not None:
        raise last_error


def write_missing_feature_docx(summary_rows: list[dict], save_path: Path) -> None:
    doc = Document()
    doc.add_heading("现有常规测井 AC/GR 属性缺失情况", level=1)
    doc.add_paragraph(f"time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    missing_rows = [
        row for row in summary_rows
        if str(row.get("FeatureHandling", "")).startswith("skip_missing_log_features")
        or str(row.get("FeatureHandling", "")).startswith("skip_missing_basic_columns")
    ]
    doc.add_paragraph(f"total_wells: {len(summary_rows)}")
    doc.add_paragraph(f"special_handling_wells: {len(missing_rows)}")

    if not missing_rows:
        doc.add_paragraph("all wells contain AC and GR and passed the basic column check.")
    else:
        table = doc.add_table(rows=1, cols=8)
        headers = [
            "WellName",
            "Source",
            "HasAC",
            "HasGR",
            "UsedFeatures",
            "Handling",
            "RunStatus",
            "Note",
        ]
        for idx, header in enumerate(headers):
            table.rows[0].cells[idx].text = header
        for row in missing_rows:
            cells = table.add_row().cells
            cells[0].text = str(row.get("WellName", ""))
            cells[1].text = str(row.get("SourceGroup", ""))
            cells[2].text = str(row.get("HasAC", ""))
            cells[3].text = str(row.get("HasGR", ""))
            cells[4].text = str(row.get("UsedLogFeatures", ""))
            cells[5].text = str(row.get("FeatureHandling", ""))
            cells[6].text = str(row.get("RunStatus", ""))
            cells[7].text = str(row.get("Note", ""))[:240]
    last_error: Exception | None = None
    for _ in range(3):
        try:
            save_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(save_path))
            return
        except Exception as exc:
            last_error = exc
            time.sleep(2.0)
    if last_error is not None:
        raise last_error


def build_single_run_command(
    python_executable: Path,
    single_script_path: Path,
    csv_path: Path,
    output_root: Path,
    docx_path: Path,
    layer_dir: Path,
    stage1_library_dir: Path,
    stage2_library_dir: Path,
    requested_log_features: list[str],
    strata_determination_mode: str,
    strata_determination_topk: int,
    unknown_strata_policy: str,
    first_stage_threshold: float,
    crossing_tolerance_ms: float,
    min_interval_depth: float,
    save_debug: bool,
) -> tuple[list[str], str, Path]:
    well_name = infer_well_name(csv_path)
    exp_id = sanitize(well_name) or well_name
    result_dir = output_root / exp_id
    command = [
        str(python_executable),
        "-X",
        "utf8",
        str(single_script_path),
        "--exp-id",
        exp_id,
        "--target-sample-csv",
        str(csv_path),
        "--target-well-name",
        well_name,
        "--layer-dir",
        str(layer_dir),
        "--stage1-library-dir",
        str(stage1_library_dir),
        "--stage2-library-dir",
        str(stage2_library_dir),
        "--output-root",
        str(output_root),
        "--docx-path",
        str(docx_path),
        "--requested-log-features-json",
        json.dumps(requested_log_features, ensure_ascii=False),
        "--strata-determination-mode",
        str(strata_determination_mode),
        "--strata-determination-topk",
        str(int(strata_determination_topk)),
        "--unknown-strata-policy",
        str(unknown_strata_policy),
        "--crossing-tolerance-ms",
        str(float(crossing_tolerance_ms)),
        "--min-interval-depth",
        str(float(min_interval_depth)),
        "--save-debug",
        str(int(bool(save_debug))),
    ]
    if math.isfinite(float(first_stage_threshold)):
        command.extend(["--first-stage-threshold", str(float(first_stage_threshold))])
    return command, well_name, result_dir


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--inclined-dir", default=str(DEFAULT_INCLINED_DIR))
    parser.add_argument("--vertical-dir", default=str(DEFAULT_VERTICAL_DIR))
    parser.add_argument("--layer-dir", default=str(DEFAULT_LAYER_DIR))
    parser.add_argument("--stage1-library-dir", default=str(DEFAULT_STAGE1_LIBRARY_DIR))
    parser.add_argument("--stage2-library-dir", default=str(DEFAULT_STAGE2_LIBRARY_DIR))
    parser.add_argument("--output-root", default=str(DEFAULT_OUTPUT_ROOT))
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument(
        "--requested-log-features-json",
        default=json.dumps(DEFAULT_LOG_FEATURES, ensure_ascii=False),
    )
    parser.add_argument(
        "--include-wells-json",
        default="[]",
        help="Optional list of well names for subset run.",
    )
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--skip-existing", type=int, default=1)
    parser.add_argument("--python-executable", default=sys.executable)
    parser.add_argument(
        "--strata-determination-mode",
        default="always_auto",
        choices=["use_csv", "auto_if_missing", "always_auto"],
    )
    parser.add_argument("--strata-determination-topk", type=int, default=2)
    parser.add_argument("--unknown-strata-policy", default="skip", choices=["skip", "error"])
    parser.add_argument("--first-stage-threshold", type=float, default=float("nan"))
    parser.add_argument("--crossing-tolerance-ms", type=float, default=8.0)
    parser.add_argument("--min-interval-depth", type=float, default=0.5)
    parser.add_argument("--save-debug", type=int, default=0)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    inclined_dir = Path(args.inclined_dir)
    vertical_dir = Path(args.vertical_dir)
    layer_dir = Path(args.layer_dir)
    stage1_library_dir = Path(args.stage1_library_dir)
    stage2_library_dir = Path(args.stage2_library_dir)
    output_root = Path(args.output_root)
    docx_path = Path(args.docx_path)
    python_executable = Path(args.python_executable)
    single_script_path = Path(__file__).with_name("run_conventional_log_strata_validation.py")
    output_root.mkdir(parents=True, exist_ok=True)

    requested_candidates = parse_json_list(args.requested_log_features_json)
    if not requested_candidates:
        raise ValueError("requested_log_features cannot be empty")
    include_wells = parse_json_list(args.include_wells_json)
    target_rows = discover_target_csvs(
        inclined_dir=inclined_dir,
        vertical_dir=vertical_dir,
        include_wells=include_wells,
        limit=int(args.limit),
    )
    if not target_rows:
        raise ValueError("No target CSV files were discovered for batch prediction")

    summary_rows: list[dict] = []
    summary_csv_path = output_root / SUMMARY_CSV_NAME
    missing_docx_path = output_root / MISSING_FEATURE_DOCX_NAME

    total = len(target_rows)
    for idx, (source_group, csv_path) in enumerate(target_rows, start=1):
        feature_status = detect_feature_status(csv_path, requested_candidates=requested_candidates)
        well_name = infer_well_name(csv_path)
        exp_id = sanitize(well_name) or well_name
        result_dir = output_root / exp_id
        summary_row = {
            "WellName": well_name,
            "SourceGroup": source_group,
            "CsvPath": str(csv_path),
            "ResultDir": str(result_dir),
            "HeaderEncoding": feature_status["encoding"],
            "DepthColumn": feature_status["depth_col"],
            "HasTIME": feature_status["has_time"],
            "HasX": feature_status["has_x"],
            "HasY": feature_status["has_y"],
            "HasAC": feature_status["has_ac"],
            "HasGR": feature_status["has_gr"],
            "RequestedFeatureCandidates": ",".join(requested_candidates),
            "UsedLogFeatures": ",".join(feature_status["available_requested"]),
            "MissingRequestedFeatures": ",".join(feature_status["missing_requested"]),
            "FeatureHandling": feature_status["handling"],
            "RunStatus": "",
            "ReturnCode": "",
            "ElapsedSeconds": "",
            "Note": "",
        }

        if feature_status["handling"] == "skip_missing_basic_columns":
            summary_row["RunStatus"] = "skipped_missing_basic_columns"
            summary_row["Note"] = "Required columns missing. Need TIME/X/Y and one depth column among TVD/DEPT/MD."
            summary_rows.append(summary_row)
            write_summary_csv(summary_rows, summary_csv_path)
            continue

        if feature_status["handling"].startswith("skip_missing_log_features"):
            removed_existing_result = False
            if result_dir.exists():
                shutil.rmtree(result_dir, ignore_errors=False)
                removed_existing_result = True
            summary_row["RunStatus"] = "skipped_missing_log_features"
            summary_row["Note"] = (
                "Missing AC or GR. Prediction skipped for reliability."
                + (" Existing result directory was removed." if removed_existing_result else "")
            )
            summary_rows.append(summary_row)
            write_summary_csv(summary_rows, summary_csv_path)
            print(
                f"[{idx}/{total}] skip missing log features: {well_name} | missing={feature_status['missing_requested']}",
                flush=True,
            )
            continue

        if bool(int(args.skip_existing)) and is_completed_result_dir(result_dir):
            summary_row["RunStatus"] = "skipped_existing"
            summary_row["Note"] = "Required result files already exist."
            summary_rows.append(summary_row)
            write_summary_csv(summary_rows, summary_csv_path)
            print(
                f"[{idx}/{total}] skip existing: {well_name} -> {result_dir}",
                flush=True,
            )
            continue

        command, _, result_dir = build_single_run_command(
            python_executable=python_executable,
            single_script_path=single_script_path,
            csv_path=csv_path,
            output_root=output_root,
            docx_path=docx_path,
            layer_dir=layer_dir,
            stage1_library_dir=stage1_library_dir,
            stage2_library_dir=stage2_library_dir,
            requested_log_features=requested_candidates,
            strata_determination_mode=str(args.strata_determination_mode),
            strata_determination_topk=int(args.strata_determination_topk),
            unknown_strata_policy=str(args.unknown_strata_policy),
            first_stage_threshold=float(args.first_stage_threshold),
            crossing_tolerance_ms=float(args.crossing_tolerance_ms),
            min_interval_depth=float(args.min_interval_depth),
            save_debug=bool(int(args.save_debug)),
        )
        print(
            f"[{idx}/{total}] run: {well_name} | available={feature_status['available_requested']} | missing={feature_status['missing_requested']} | source={source_group}",
            flush=True,
        )
        start_time = time.time()
        completed = subprocess.run(
            command,
            cwd=str(single_script_path.parent),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
        )
        elapsed = round(time.time() - start_time, 2)
        summary_row["ReturnCode"] = int(completed.returncode)
        summary_row["ElapsedSeconds"] = elapsed
        if completed.returncode == 0 and is_completed_result_dir(result_dir):
            summary_row["RunStatus"] = "success"
            stdout_tail = completed.stdout.strip().splitlines()
            summary_row["Note"] = stdout_tail[-1][:400] if stdout_tail else ""
        elif completed.returncode == 0:
            summary_row["RunStatus"] = "failed_incomplete_outputs"
            summary_row["Note"] = "Process returned 0 but required result files are incomplete."
        else:
            summary_row["RunStatus"] = "failed"
            stderr_text = (completed.stderr or "").strip()
            stdout_text = (completed.stdout or "").strip()
            error_text = stderr_text[-1000:] if stderr_text else stdout_text[-1000:]
            summary_row["Note"] = error_text
        summary_rows.append(summary_row)
        write_summary_csv(summary_rows, summary_csv_path)

    write_summary_csv(summary_rows, summary_csv_path)
    write_missing_feature_docx(summary_rows, missing_docx_path)

    summary_df = pd.DataFrame(summary_rows)
    success_count = int(summary_df["RunStatus"].astype(str).eq("success").sum()) if not summary_df.empty else 0
    skipped_existing_count = int(summary_df["RunStatus"].astype(str).eq("skipped_existing").sum()) if not summary_df.empty else 0
    skipped_missing_feature_count = int(
        summary_df["RunStatus"].astype(str).isin(
            ["skipped_missing_basic_columns", "skipped_missing_log_features"]
        ).sum()
    ) if not summary_df.empty else 0
    failed_count = int(summary_df["RunStatus"].astype(str).isin(["failed", "failed_incomplete_outputs"]).sum()) if not summary_df.empty else 0

    print(
        json.dumps(
            {
                "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "output_root": str(output_root),
                "total_wells": total,
                "success_count": success_count,
                "skipped_existing_count": skipped_existing_count,
                "skipped_missing_feature_count": skipped_missing_feature_count,
                "failed_count": failed_count,
                "summary_csv": str(summary_csv_path),
                "missing_feature_docx": str(missing_docx_path),
            },
            ensure_ascii=False,
        ),
        flush=True,
    )
    return 0 if failed_count == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
