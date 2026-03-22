import argparse
import csv
import json
import os
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError


ROOT = Path(__file__).resolve().parent
SCRIPT_PATH = (
    ROOT
    / "模型训练"
    / "优化阶段一"
    / "裂缝位置分析"
    / "基于密度的裂缝点位分析"
    / "fracture_density_predict.py"
)
BASE_SAVE_DIR = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容一\裂缝位置预测\基于密度的裂缝点位分析\裂缝点位精细化\density_experiments"
)
DEFAULT_DOCX_PATH = Path(r"D:\项目\石油开采\断缝储实验\实验记录20260319.docx")


def sanitize(text: str) -> str:
    text = text.replace('"', "")
    text = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]+", "_", text)
    return text.strip("_")


def load_json_list(raw: str) -> list:
    try:
        value = json.loads(raw)
    except json.JSONDecodeError:
        value = [item.strip() for item in raw.split(",") if item.strip()]
    if not isinstance(value, list):
        raise ValueError(f"Expected a JSON list, got: {raw}")
    return value


def load_json_dict(raw: str) -> dict:
    value = json.loads(raw)
    if not isinstance(value, dict):
        raise ValueError(f"Expected a JSON object, got: {raw}")
    return value


def open_or_create_doc(docx_path: Path) -> Document:
    if docx_path.exists():
        try:
            return Document(str(docx_path))
        except PackageNotFoundError:
            pass

    doc = Document()
    doc.add_heading("实验记录 20260319", level=1)
    return doc


def format_float(value: object) -> str:
    try:
        return f"{float(value):.4f}"
    except (TypeError, ValueError):
        return str(value)


def append_result_to_docx(docx_path: Path, title: str, config: dict, results: list[dict], target_cols: list[str]) -> None:
    doc = open_or_create_doc(docx_path)
    doc.add_heading(title, level=2)
    doc.add_paragraph(f"记录时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    config_lines = [
        "task: 裂缝密度预测",
        f"script: {SCRIPT_PATH}",
        f"SEQ_LEN: {config['SEQ_LEN']}",
        f"SEIS_MODE: {config['SEIS_MODE']}",
        f"imaging_well_features: {config['imaging_well_features']}",
        f"target_cols: {config['target_cols']}",
        f"target_weights: {config['target_weights']}",
        f"target_transform_scales: {config['target_transform_scales']}",
        f"LOSS_NAME: {config['LOSS_NAME']}",
        f"EPOCHS: {config['EPOCHS']}",
        f"PATIENCE: {config['PATIENCE']}",
        f"use_all_other_wells: {config['use_all_other_wells']}",
        f"MIN_TRAIN_WELLS: {config['min_train_wells']}",
        f"DIST_THRESHOLD: {config['dist_threshold']}",
        f"only_val_wells: {config['only_val_wells']}",
        f"custom_train_wells_json: {config['custom_train_wells_json']}",
        f"save_dir: {config['save_dir']}",
        f"result_csv: {config['result_csv']}",
    ]
    doc.add_paragraph("\n".join(config_lines))

    headers = ["井名", "训练井", "TrainSeq", "ValSeq"]
    key_order = ["val_well", "train_wells", "NumTrainDensitySeq", "NumValDensitySeq"]
    for target_col in target_cols:
        headers.extend([f"{target_col}_R2", f"{target_col}_MAE", f"{target_col}_RMSE", f"{target_col}_BIAS"])
        key_order.extend(
            [f"{target_col}_R2", f"{target_col}_MAE", f"{target_col}_RMSE", f"{target_col}_BIAS"]
        )

    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header

    for row in results:
        cells = table.add_row().cells
        for idx, key in enumerate(key_order):
            cells[idx].text = format_float(row[key]) if key not in {"val_well", "train_wells"} else str(row[key])

    summary_parts = []
    for target_col in target_cols:
        r2_values = [float(row[f"{target_col}_R2"]) for row in results if row.get(f"{target_col}_R2", "") not in {"", "nan"}]
        mae_values = [float(row[f"{target_col}_MAE"]) for row in results if row.get(f"{target_col}_MAE", "") not in {"", "nan"}]
        if r2_values and mae_values:
            summary_parts.append(
                f"{target_col}: AVG_R2={sum(r2_values) / len(r2_values):.4f}, AVG_MAE={sum(mae_values) / len(mae_values):.4f}"
            )

    if summary_parts:
        doc.add_paragraph("总结: " + " | ".join(summary_parts))
    else:
        doc.add_paragraph("总结: 已记录本次密度实验结果。")

    doc.add_paragraph("")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", required=True)
    parser.add_argument("--seq-len", type=int, default=5)
    parser.add_argument("--seis-mode", default="3x3")
    parser.add_argument("--features-json", default='["AC", "GR"]')
    parser.add_argument("--target-cols-json", default='["P10"]')
    parser.add_argument("--target-weights-json", default="[1.0]")
    parser.add_argument("--target-transform-scales-json", default='{"P10": 1.0, "P21": 1.0, "P33": 10000.0}')
    parser.add_argument("--loss-name", default="smooth_l1")
    parser.add_argument("--epochs", default="60")
    parser.add_argument("--patience", default="10")
    parser.add_argument("--batch-size", default="32")
    parser.add_argument("--lr", default="1e-4")
    parser.add_argument("--hidden-dim", default="32")
    parser.add_argument("--dist-threshold", default="0.35")
    parser.add_argument("--min-train-wells", default="2")
    parser.add_argument("--use-all-other-wells", default="false")
    parser.add_argument("--use-augmentation", default="false")
    parser.add_argument("--save-pred-for-all-centers", default="false")
    parser.add_argument("--only-val-wells", default="")
    parser.add_argument("--custom-train-wells-json", default="")
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument("--result-dir")
    parser.add_argument("--skip-run", action="store_true")
    args = parser.parse_args()

    features = load_json_list(args.features_json)
    target_cols = load_json_list(args.target_cols_json)
    target_weights = load_json_list(args.target_weights_json)
    target_transform_scales = load_json_dict(args.target_transform_scales_json)

    feature_tag = sanitize("_".join(str(item) for item in features))
    target_tag = sanitize("_".join(str(item) for item in target_cols))
    exp_tag = f"{sanitize(args.exp_id)}_{sanitize(args.seis_mode)}_seq{args.seq_len}_{feature_tag}_{target_tag}"
    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / exp_tag)
    save_dir.mkdir(parents=True, exist_ok=True)

    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["EXP_SEQ_LEN"] = str(args.seq_len)
    env["EXP_SEIS_MODE"] = args.seis_mode
    env["EXP_IMAGING_FEATURES"] = json.dumps(features, ensure_ascii=False)
    env["EXP_DENSITY_TARGET_COLS"] = json.dumps(target_cols, ensure_ascii=False)
    env["EXP_DENSITY_TARGET_WEIGHTS"] = json.dumps(target_weights, ensure_ascii=False)
    env["EXP_DENSITY_TARGET_TRANSFORM_SCALES"] = json.dumps(target_transform_scales, ensure_ascii=False)
    env["EXP_DENSITY_LOSS"] = args.loss_name
    env["EXP_EPOCHS"] = str(args.epochs)
    env["EXP_PATIENCE"] = str(args.patience)
    env["EXP_BATCH_SIZE"] = str(args.batch_size)
    env["EXP_LR"] = str(args.lr)
    env["EXP_HIDDEN_DIM"] = str(args.hidden_dim)
    env["EXP_DIST_THRESHOLD"] = str(args.dist_threshold)
    env["EXP_MIN_TRAIN_WELLS"] = str(args.min_train_wells)
    env["EXP_USE_ALL_OTHER_WELLS"] = args.use_all_other_wells
    env["EXP_USE_AUGMENTATION"] = args.use_augmentation
    env["EXP_SAVE_PRED_FOR_ALL_CENTERS"] = args.save_pred_for_all_centers
    env["EXP_SAVE_DIR"] = str(save_dir)
    if args.only_val_wells:
        only_val_wells = [item.strip() for item in args.only_val_wells.split(",") if item.strip()]
        env["EXP_ONLY_VAL_WELLS"] = json.dumps(only_val_wells, ensure_ascii=False)
    if args.custom_train_wells_json:
        env["EXP_CUSTOM_TRAIN_WELLS_MAP"] = args.custom_train_wells_json

    if not args.skip_run:
        completed = subprocess.run(
            [sys.executable, str(SCRIPT_PATH)],
            env=env,
            cwd=str(ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
        )
        (save_dir / "run_stdout.log").write_text(completed.stdout, encoding="utf-8")
        (save_dir / "run_stderr.log").write_text(completed.stderr, encoding="utf-8")
        if completed.returncode != 0:
            print(completed.stdout[-4000:], file=sys.stdout)
            print(completed.stderr[-4000:], file=sys.stderr)
            return completed.returncode

    result_csv = save_dir / "loo_density_results.csv"
    if not result_csv.exists():
        print(f"Missing result file: {result_csv}", file=sys.stderr)
        return 2

    with result_csv.open("r", encoding="utf-8-sig", newline="") as file_obj:
        results = list(csv.DictReader(file_obj))

    config = {
        "SEQ_LEN": args.seq_len,
        "SEIS_MODE": args.seis_mode,
        "imaging_well_features": features,
        "target_cols": target_cols,
        "target_weights": target_weights,
        "target_transform_scales": target_transform_scales,
        "LOSS_NAME": args.loss_name,
        "EPOCHS": args.epochs,
        "PATIENCE": args.patience,
        "use_all_other_wells": args.use_all_other_wells,
        "min_train_wells": args.min_train_wells,
        "dist_threshold": args.dist_threshold,
        "only_val_wells": args.only_val_wells or env.get("EXP_ONLY_VAL_WELLS", ""),
        "custom_train_wells_json": args.custom_train_wells_json or env.get("EXP_CUSTOM_TRAIN_WELLS_MAP", ""),
        "save_dir": str(save_dir),
        "result_csv": str(result_csv),
    }
    append_result_to_docx(Path(args.docx_path), f"实验 {args.exp_id}", config, results, target_cols)

    print(
        json.dumps(
            {
                "exp_id": args.exp_id,
                "save_dir": str(save_dir),
                "result_csv": str(result_csv),
                "results": results,
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
