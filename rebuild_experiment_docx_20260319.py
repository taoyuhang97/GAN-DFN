import json
from datetime import datetime
from pathlib import Path

from docx import Document


RESULT_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容一\裂缝位置预测\基于密度的裂缝点位分析\裂缝点位精细化"
)
OUTPUT_DOCX = Path(r"D:\项目\石油开采\断缝储实验\实验记录20260319.docx")
TARGET_DATE_PREFIX = "2026-03-19"


def summarize_rows(rows: list[dict]) -> str:
    if not rows:
        return "总结: 无结果行"

    def avg_of(key: str) -> str:
        values = []
        for row in rows:
            raw = row.get(key, "")
            try:
                value = float(raw)
            except (TypeError, ValueError):
                continue
            if value == value:
                values.append(value)
        if not values:
            return ""
        return f"{sum(values) / len(values):.4f}"

    parts = []
    for key, label in (
        ("count_diff", "AVG_count_diff"),
        ("pred_to_gt_mean_dist", "AVG_pred_to_gt_mean"),
        ("gt_to_pred_mean_dist", "AVG_gt_to_pred_mean"),
        ("segment_count_mae", "AVG_segment_count_mae"),
    ):
        avg_text = avg_of(key)
        if avg_text:
            parts.append(f"{label}={avg_text}")
    return "总结: " + (", ".join(parts) if parts else "无可汇总指标")


def load_rows(summary_path: Path) -> list[dict]:
    import csv

    with summary_path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        return list(csv.DictReader(file_obj))


def iter_experiments():
    for exp_dir in sorted(RESULT_ROOT.iterdir(), key=lambda item: item.stat().st_mtime):
        if not exp_dir.is_dir():
            continue
        modified = datetime.fromtimestamp(exp_dir.stat().st_mtime)
        if modified.strftime("%Y-%m-%d") != TARGET_DATE_PREFIX:
            continue

        config_path = exp_dir / "config.json"
        if not config_path.exists():
            continue

        summary_candidates = [
            exp_dir / "segment_count_refine_summary.csv",
            exp_dir / "fracture_point_refine_summary.csv",
        ]
        summary_path = next((path for path in summary_candidates if path.exists()), None)
        if summary_path is None:
            continue

        with config_path.open("r", encoding="utf-8") as file_obj:
            config = json.load(file_obj)
        rows = load_rows(summary_path)
        yield {
            "name": exp_dir.name,
            "modified": modified,
            "config": config,
            "summary_path": summary_path,
            "rows": rows,
        }


def write_table(doc: Document, rows: list[dict]) -> None:
    if not rows:
        doc.add_paragraph("无结果行")
        return

    first_row = rows[0]
    if "segment_count_mae" in first_row:
        headers = [
            ("well", "井名"),
            ("N_gt_points", "GT点数"),
            ("N_pred_points", "Pred点数"),
            ("count_diff", "点数差"),
            ("pred_to_gt_mean_dist", "pred->gt_mean"),
            ("gt_to_pred_mean_dist", "gt->pred_mean"),
            ("segment_count_mae", "SegMAE"),
        ]
    else:
        headers = [
            ("well", "井名"),
            ("N_gt_points", "GT点数"),
            ("N_pred_points", "Pred点数"),
            ("count_diff", "点数差"),
            ("pred_to_gt_mean_dist", "pred->gt_mean"),
            ("gt_to_pred_mean_dist", "gt->pred_mean"),
        ]

    table = doc.add_table(rows=1, cols=len(headers))
    for idx, (_, label) in enumerate(headers):
        table.rows[0].cells[idx].text = label

    for row in rows:
        cells = table.add_row().cells
        for idx, (key, _) in enumerate(headers):
            value = row.get(key, "")
            if key == "well":
                cells[idx].text = str(value)
            else:
                try:
                    cells[idx].text = f"{float(value):.4f}"
                except (TypeError, ValueError):
                    cells[idx].text = str(value)


def main() -> int:
    doc = Document()
    doc.add_heading("实验记录 20260319", level=1)

    experiments = list(iter_experiments())
    for exp in experiments:
        doc.add_heading(f"实验 {exp['name']}", level=2)
        doc.add_paragraph(f"记录时间: {exp['modified'].strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(
            json.dumps(
                exp["config"],
                ensure_ascii=False,
                indent=2,
            )
        )
        doc.add_paragraph(f"summary_csv: {exp['summary_path']}")
        write_table(doc, exp["rows"])
        doc.add_paragraph(summarize_rows(exp["rows"]))
        doc.add_paragraph("")

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_DOCX))
    print(
        json.dumps(
            {
                "output_docx": str(OUTPUT_DOCX),
                "num_experiments": len(experiments),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
