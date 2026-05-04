# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import math
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from scipy.ndimage import generate_binary_structure, label as ndi_label

try:
    from scipy.optimize import linear_sum_assignment
except Exception:  # pragma: no cover
    linear_sum_assignment = None


DEFAULT_UNIT_DFN_ROOT = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容二/单元DFN构建/批量生成"
)
DEFAULT_OUTPUT_ROOT = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/研究内容三/DFN体素互转实验"
)
DEFAULT_DOCX_PATH = Path(r"/data/shared/project-oil/wx数据/砂砾岩/优化阶段一/实验记录/实验记录20260328.docx")
DEFAULT_TRACE_HEADER_CSV = Path(
    r"/data/shared/project-oil/wx数据/砂砾岩/研究内容一/trace_header_xy.csv"
)
DEFAULT_BLOCK_SIZE = 25
LAYER_SURFACE_PAIR_KEY_COL = "LayerSurfacePairKey"
UNIT_LAYER_SEGMENT_KEY_COL = "UnitLayerSegmentKey"

PATCH_VERTEX_COLUMNS = [
    f"V{vertex_idx}{axis}"
    for vertex_idx in range(1, 5)
    for axis in ("X", "Y", "Z")
]

PATCH_OUTPUT_COLUMNS = [
    "PatchIndex",
    "PatchID",
    "UnitID",
    "BlockX",
    "BlockY",
    "GeoIntervalKey",
    LAYER_SURFACE_PAIR_KEY_COL,
    UNIT_LAYER_SEGMENT_KEY_COL,
    "StrataName",
    "TopSurfaceCode",
    "BaseSurfaceCode",
    "SeedID",
    "SourceKind",
    "SourceName",
    "SeedType",
    "ParentSeedID",
    "ParentSourceKind",
    "CenterX",
    "CenterY",
    "CenterTIME",
    "CenterDepth",
    "Azimuth",
    "Dip",
    "DensityWeight",
    "LengthWeight",
    "Confidence",
    "PatchLength",
    "PatchHeight",
    "NormalX",
    "NormalY",
    "NormalZ",
] + PATCH_VERTEX_COLUMNS


@dataclass
class GridSpec:
    unit_id: str
    block_x: int
    block_y: int
    x_min: float
    x_max: float
    y_min: float
    y_max: float
    z_min: float
    z_max: float
    nx: int
    ny: int
    nz: int
    dx: float
    dy: float
    dz: float

    def to_json_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["x_centers"] = [float(v) for v in self.x_centers]
        payload["y_centers"] = [float(v) for v in self.y_centers]
        payload["z_centers"] = [float(v) for v in self.z_centers]
        return payload

    @property
    def x_centers(self) -> np.ndarray:
        return self.x_min + (np.arange(self.nx, dtype=float) + 0.5) * self.dx

    @property
    def y_centers(self) -> np.ndarray:
        return self.y_min + (np.arange(self.ny, dtype=float) + 0.5) * self.dy

    @property
    def z_centers(self) -> np.ndarray:
        return self.z_min + (np.arange(self.nz, dtype=float) + 0.5) * self.dz


@dataclass
class VtkPatchExportConfig:
    display_z_scale: float = 5.0
    invert_time: bool = False


def _normalize_layer_key_text(value: Any) -> str:
    text = str(value or "").strip()
    return "" if (not text or text.lower() == "nan") else text


def build_layer_surface_pair_key(top_surface_code: Any, base_surface_code: Any) -> str:
    return f"{_normalize_layer_key_text(top_surface_code)}->{_normalize_layer_key_text(base_surface_code)}"


def resolve_layer_surface_pair_key_from_row(row: pd.Series | dict[str, Any]) -> str:
    if hasattr(row, "get"):
        existing = _normalize_layer_key_text(row.get(LAYER_SURFACE_PAIR_KEY_COL, ""))
        if existing:
            return existing
        top_surface_code = _normalize_layer_key_text(row.get("TopSurfaceCode", ""))
        base_surface_code = _normalize_layer_key_text(row.get("BaseSurfaceCode", ""))
        if top_surface_code or base_surface_code:
            return build_layer_surface_pair_key(top_surface_code, base_surface_code)
        return _normalize_layer_key_text(row.get("GeoIntervalKey", ""))
    return ""


def resolve_unit_layer_segment_key_from_row(row: pd.Series | dict[str, Any]) -> str:
    if hasattr(row, "get"):
        existing = _normalize_layer_key_text(row.get(UNIT_LAYER_SEGMENT_KEY_COL, ""))
        if existing:
            return existing
        pair_key = resolve_layer_surface_pair_key_from_row(row)
        unit_id = _normalize_layer_key_text(row.get("UnitID", ""))
        interval_key = _normalize_layer_key_text(row.get("GeoIntervalKey", ""))
        if unit_id or interval_key or pair_key:
            return f"{unit_id or 'UNKNOWN_UNIT'}__{interval_key or 'UNKNOWN_INTERVAL'}__{pair_key or 'UNKNOWN_LAYER_PAIR'}"
    return ""


def preferred_layer_group_series(
    patch_df: pd.DataFrame,
    prefer_unit_segment: bool = False,
) -> pd.Series:
    fallback = patch_df.get("GeoIntervalKey", pd.Series([""] * len(patch_df), index=patch_df.index, dtype="object")).fillna("").astype(str)
    pair_series = (
        patch_df.get(LAYER_SURFACE_PAIR_KEY_COL, pd.Series([""] * len(patch_df), index=patch_df.index, dtype="object"))
        .fillna("")
        .astype(str)
        .str.strip()
    )
    if prefer_unit_segment:
        segment_series = (
            patch_df.get(UNIT_LAYER_SEGMENT_KEY_COL, pd.Series([""] * len(patch_df), index=patch_df.index, dtype="object"))
            .fillna("")
            .astype(str)
            .str.strip()
        )
        return segment_series.mask(segment_series.eq(""), pair_series).mask(lambda s: s.eq(""), fallback)
    return pair_series.mask(pair_series.eq(""), fallback)


def read_csv_utf8(path: Path) -> pd.DataFrame:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError:
            continue
    return pd.read_csv(path)


def write_csv_utf8(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8-sig")


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def summarize_patch_statistics(patch_df: pd.DataFrame) -> dict[str, float]:
    if patch_df.empty:
        return {}
    length = pd.to_numeric(patch_df.get("PatchLength"), errors="coerce").dropna()
    height = pd.to_numeric(patch_df.get("PatchHeight"), errors="coerce").dropna()
    summary: dict[str, float] = {}
    if not length.empty:
        summary.update(
            {
                "patch_length_median": float(length.median()),
                "patch_length_q75": float(length.quantile(0.75)),
                "patch_length_q90": float(length.quantile(0.90)),
            }
        )
    if not height.empty:
        summary.update(
            {
                "patch_height_median": float(height.median()),
                "patch_height_q75": float(height.quantile(0.75)),
                "patch_height_q90": float(height.quantile(0.90)),
            }
        )
    return summary


def normalize_vtk_numeric(values: np.ndarray) -> np.ndarray:
    array = np.asarray(values)
    if np.issubdtype(array.dtype, np.integer):
        return array.astype(int)
    numeric = array.astype(float)
    return np.where(np.isfinite(numeric), numeric, -9999.0)


def write_vtk_scalar_block(lines: list[str], name: str, values: np.ndarray) -> None:
    array = normalize_vtk_numeric(values)
    if np.issubdtype(array.dtype, np.integer):
        vtk_type = "int"
        formatter = lambda value: str(int(value))
    else:
        vtk_type = "float"
        formatter = lambda value: f"{float(value):.6f}"
    lines.append(f"SCALARS {name} {vtk_type} 1")
    lines.append("LOOKUP_TABLE default")
    lines.extend(formatter(value) for value in array.tolist())


def write_legacy_vtk_polygons(
    path: Path,
    title: str,
    points: np.ndarray,
    polygons: list[list[int]],
    cell_data: dict[str, np.ndarray],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    total_polygon_size = sum(len(polygon) + 1 for polygon in polygons)
    lines = [
        "# vtk DataFile Version 3.0",
        title,
        "ASCII",
        "DATASET POLYDATA",
        f"POINTS {len(points)} float",
    ]
    lines.extend(f"{float(point[0]):.6f} {float(point[1]):.6f} {float(point[2]):.6f}" for point in points)
    lines.append(f"POLYGONS {len(polygons)} {total_polygon_size}")
    lines.extend(f"{len(polygon)} {' '.join(str(int(idx)) for idx in polygon)}" for polygon in polygons)
    if cell_data:
        lines.append(f"CELL_DATA {len(polygons)}")
        for name, values in cell_data.items():
            write_vtk_scalar_block(lines, name, values)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_category_code_map(series: pd.Series) -> tuple[np.ndarray, dict[int, str]]:
    values = series.fillna("").astype(str)
    unique_values = sorted([value for value in values.unique().tolist() if value != ""])
    code_map = {value: idx for idx, value in enumerate(unique_values, start=1)}
    codes = values.map(code_map).fillna(0).astype(int).to_numpy(dtype=int)
    reverse_map = {code: value for value, code in code_map.items()}
    reverse_map[0] = ""
    return codes, reverse_map


def transform_display_z(values: np.ndarray, config: VtkPatchExportConfig) -> np.ndarray:
    z_values = np.asarray(values, dtype=float) * float(config.display_z_scale)
    if bool(config.invert_time):
        z_values = -z_values
    return z_values


def resolve_patch_csv(unit_id: str | None, patch_csv: Path | None, unit_dfn_root: Path) -> Path:
    if patch_csv is not None:
        return Path(patch_csv)
    if not unit_id:
        raise ValueError("必须提供 --unit-id 或 --patch-csv。")
    candidate = Path(unit_dfn_root) / str(unit_id) / "unit_dfn_patches.csv"
    if not candidate.exists():
        raise FileNotFoundError(f"未找到单元裂缝片文件: {candidate}")
    return candidate


def load_patch_table(path: Path) -> pd.DataFrame:
    df = read_csv_utf8(path)
    numeric_cols = [
        "BlockX",
        "BlockY",
        "CenterX",
        "CenterY",
        "CenterTIME",
        "CenterDepth",
        "Azimuth",
        "Dip",
        "DensityWeight",
        "LengthWeight",
        "Confidence",
        "PatchLength",
        "PatchHeight",
        "NormalX",
        "NormalY",
        "NormalZ",
    ] + PATCH_VERTEX_COLUMNS
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    for col in PATCH_OUTPUT_COLUMNS:
        if col not in df.columns:
            df[col] = np.nan if col not in {
                "PatchID",
                "UnitID",
                "GeoIntervalKey",
                LAYER_SURFACE_PAIR_KEY_COL,
                UNIT_LAYER_SEGMENT_KEY_COL,
                "StrataName",
                "TopSurfaceCode",
                "BaseSurfaceCode",
                "SeedID",
                "SourceKind",
                "SourceName",
                "SeedType",
                "ParentSeedID",
                "ParentSourceKind",
            } else ""
    if len(df) > 0:
        pair_series = df.apply(resolve_layer_surface_pair_key_from_row, axis=1)
        pair_fill_mask = df[LAYER_SURFACE_PAIR_KEY_COL].fillna("").astype(str).str.strip().eq("")
        if pair_fill_mask.any():
            df.loc[pair_fill_mask, LAYER_SURFACE_PAIR_KEY_COL] = pair_series.loc[pair_fill_mask]
        segment_series = df.apply(resolve_unit_layer_segment_key_from_row, axis=1)
        segment_fill_mask = df[UNIT_LAYER_SEGMENT_KEY_COL].fillna("").astype(str).str.strip().eq("")
        if segment_fill_mask.any():
            df.loc[segment_fill_mask, UNIT_LAYER_SEGMENT_KEY_COL] = segment_series.loc[segment_fill_mask]
    return df[PATCH_OUTPUT_COLUMNS].copy()


def load_layer_table(unit_dir: Path) -> pd.DataFrame:
    layer_path = unit_dir / "unit_layers_input.csv"
    if not layer_path.exists():
        return pd.DataFrame()
    df = read_csv_utf8(layer_path)
    for col in ("TopTime", "BaseTime", "TopDepth", "BaseDepth", "BlockX", "BlockY"):
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def load_unit_summary(unit_dir: Path) -> dict[str, Any]:
    summary_path = unit_dir / "unit_dfn_summary.json"
    if not summary_path.exists():
        return {}
    return json.loads(summary_path.read_text(encoding="utf-8"))


def compute_unit_bounds_from_trace_header(
    trace_header_csv: Path,
    block_x: int,
    block_y: int,
    block_size: int = DEFAULT_BLOCK_SIZE,
) -> tuple[float, float, float, float]:
    trace_df = read_csv_utf8(trace_header_csv)
    unique_x = np.sort(trace_df["X"].dropna().unique())
    unique_y = np.sort(trace_df["Y"].dropna().unique())
    start_x = int(block_x) * (int(block_size) - 1)
    start_y = int(block_y) * (int(block_size) - 1)
    end_x = start_x + int(block_size)
    end_y = start_y + int(block_size)
    if end_x > len(unique_x) or end_y > len(unique_y):
        raise ValueError(f"BlockX={block_x}, BlockY={block_y} 超出 trace header 网格范围。")
    target_x = unique_x[start_x:end_x]
    target_y = unique_y[start_y:end_y]
    return float(target_x.min()), float(target_x.max()), float(target_y.min()), float(target_y.max())


def compute_unit_bounds_from_patch_table(patch_df: pd.DataFrame) -> tuple[float, float, float, float]:
    x_values = patch_df[[f"V{i}X" for i in range(1, 5)]].to_numpy(dtype=float).reshape(-1)
    y_values = patch_df[[f"V{i}Y" for i in range(1, 5)]].to_numpy(dtype=float).reshape(-1)
    x_values = x_values[np.isfinite(x_values)]
    y_values = y_values[np.isfinite(y_values)]
    if x_values.size == 0 or y_values.size == 0:
        raise ValueError("裂缝片顶点为空，无法回退计算单元边界。")
    return float(x_values.min()), float(x_values.max()), float(y_values.min()), float(y_values.max())


def compute_time_bounds(
    patch_df: pd.DataFrame,
    layers_df: pd.DataFrame,
    z_padding_ms: float = 0.0,
) -> tuple[float, float]:
    z_values = patch_df[[f"V{i}Z" for i in range(1, 5)]].to_numpy(dtype=float).reshape(-1)
    z_values = z_values[np.isfinite(z_values)]
    if not layers_df.empty and {"TopTime", "BaseTime"}.issubset(layers_df.columns):
        layer_values = layers_df[["TopTime", "BaseTime"]].to_numpy(dtype=float).reshape(-1)
        layer_values = layer_values[np.isfinite(layer_values)]
        if layer_values.size:
            z_values = np.concatenate([z_values, layer_values])
    lower = float(np.nanmin(z_values))
    upper = float(np.nanmax(z_values))
    lower -= float(z_padding_ms)
    upper += float(z_padding_ms)
    if not np.isfinite(lower) or not np.isfinite(upper) or upper <= lower:
        raise ValueError("无法确定有效时间范围。")
    return lower, upper


def build_grid_spec(
    unit_id: str,
    block_x: int,
    block_y: int,
    x_bounds: tuple[float, float],
    y_bounds: tuple[float, float],
    z_bounds: tuple[float, float],
    xy_resolution: int,
    z_step_ms: float,
) -> GridSpec:
    x_min, x_max = x_bounds
    y_min, y_max = y_bounds
    z_min, z_max = z_bounds
    nx = int(xy_resolution)
    ny = int(xy_resolution)
    if nx <= 0 or ny <= 0:
        raise ValueError("XY 分辨率必须大于 0。")
    dz = float(z_step_ms)
    if dz <= 0:
        raise ValueError("Z 采样间隔必须大于 0。")
    dx = float(x_max - x_min) / float(nx)
    dy = float(y_max - y_min) / float(ny)
    nz = max(1, int(math.ceil(float(z_max - z_min) / dz)))
    z_max_aligned = float(z_min + nz * dz)
    return GridSpec(
        unit_id=str(unit_id),
        block_x=int(block_x),
        block_y=int(block_y),
        x_min=float(x_min),
        x_max=float(x_max),
        y_min=float(y_min),
        y_max=float(y_max),
        z_min=float(z_min),
        z_max=float(z_max_aligned),
        nx=nx,
        ny=ny,
        nz=nz,
        dx=dx,
        dy=dy,
        dz=dz,
    )


def physical_to_index(points_xyz: np.ndarray, grid: GridSpec) -> np.ndarray:
    work = np.asarray(points_xyz, dtype=float).copy()
    work[..., 0] = (work[..., 0] - grid.x_min) / grid.dx
    work[..., 1] = (work[..., 1] - grid.y_min) / grid.dy
    work[..., 2] = (work[..., 2] - grid.z_min) / grid.dz
    return work


def index_to_physical(points_idx: np.ndarray, grid: GridSpec) -> np.ndarray:
    work = np.asarray(points_idx, dtype=float).copy()
    work[..., 0] = grid.x_min + work[..., 0] * grid.dx
    work[..., 1] = grid.y_min + work[..., 1] * grid.dy
    work[..., 2] = grid.z_min + work[..., 2] * grid.dz
    return work


def normalize(vector: np.ndarray) -> np.ndarray:
    vector = np.asarray(vector, dtype=float)
    norm = float(np.linalg.norm(vector))
    if norm <= 1e-12:
        return np.zeros_like(vector, dtype=float)
    return vector / norm


def patch_row_to_vertices(row: pd.Series) -> np.ndarray:
    return np.array(
        [[float(row[f"V{i}X"]), float(row[f"V{i}Y"]), float(row[f"V{i}Z"])] for i in range(1, 5)],
        dtype=float,
    )


def build_patch_geometry_in_index_space(row: pd.Series, grid: GridSpec) -> dict[str, np.ndarray | float]:
    vertices_phys = patch_row_to_vertices(row)
    vertices_idx = physical_to_index(vertices_phys, grid)
    center_idx = np.mean(vertices_idx, axis=0)
    edge_u = vertices_idx[1] - vertices_idx[0]
    edge_v = vertices_idx[3] - vertices_idx[0]
    length_idx = float(np.linalg.norm(edge_u))
    height_idx = float(np.linalg.norm(edge_v))
    u_vec = normalize(edge_u)
    v_vec = normalize(edge_v)
    normal_vec = normalize(np.cross(u_vec, v_vec))
    return {
        "vertices_idx": vertices_idx,
        "center_idx": center_idx,
        "u_vec": u_vec,
        "v_vec": v_vec,
        "normal_vec": normal_vec,
        "length_idx": length_idx,
        "height_idx": height_idx,
    }


def rasterize_patches_to_voxel(
    patch_df: pd.DataFrame,
    grid: GridSpec,
    thickness_vox: float = 1.0,
    include_normals: bool = True,
) -> tuple[np.ndarray, np.ndarray | None, dict[str, Any]]:
    occupancy = np.zeros((grid.nx, grid.ny, grid.nz), dtype=np.float32)
    normals_sum = np.zeros((grid.nx, grid.ny, grid.nz, 3), dtype=np.float32) if include_normals else None
    normals_count = np.zeros((grid.nx, grid.ny, grid.nz), dtype=np.int32) if include_normals else None
    half_thickness = max(float(thickness_vox) / 2.0, 1e-6)

    for _, row in patch_df.iterrows():
        geom = build_patch_geometry_in_index_space(row, grid)
        vertices_idx = np.asarray(geom["vertices_idx"], dtype=float)
        center_idx = np.asarray(geom["center_idx"], dtype=float)
        u_vec = np.asarray(geom["u_vec"], dtype=float)
        v_vec = np.asarray(geom["v_vec"], dtype=float)
        normal_vec = np.asarray(geom["normal_vec"], dtype=float)
        length_idx = float(geom["length_idx"])
        height_idx = float(geom["height_idx"])

        if length_idx <= 1e-6 or height_idx <= 1e-6:
            continue

        pad = max(1.0, half_thickness + 1.0)
        mins = np.floor(vertices_idx.min(axis=0) - pad).astype(int)
        maxs = np.ceil(vertices_idx.max(axis=0) + pad).astype(int)
        ix0 = max(0, int(mins[0]))
        iy0 = max(0, int(mins[1]))
        iz0 = max(0, int(mins[2]))
        ix1 = min(grid.nx - 1, int(maxs[0]))
        iy1 = min(grid.ny - 1, int(maxs[1]))
        iz1 = min(grid.nz - 1, int(maxs[2]))
        if ix0 > ix1 or iy0 > iy1 or iz0 > iz1:
            continue

        local_x = np.arange(ix0, ix1 + 1, dtype=float) + 0.5
        local_y = np.arange(iy0, iy1 + 1, dtype=float) + 0.5
        local_z = np.arange(iz0, iz1 + 1, dtype=float) + 0.5
        xx, yy, zz = np.meshgrid(local_x, local_y, local_z, indexing="ij")
        points = np.stack([xx, yy, zz], axis=-1)
        delta = points - center_idx
        du = np.tensordot(delta, u_vec, axes=([3], [0]))
        dv = np.tensordot(delta, v_vec, axes=([3], [0]))
        dn = np.abs(np.tensordot(delta, normal_vec, axes=([3], [0])))
        mask = (
            (np.abs(du) <= (length_idx / 2.0))
            & (np.abs(dv) <= (height_idx / 2.0))
            & (dn <= half_thickness)
        )
        if not np.any(mask):
            continue

        occ_block = occupancy[ix0 : ix1 + 1, iy0 : iy1 + 1, iz0 : iz1 + 1]
        occ_block[mask] = 1.0
        occupancy[ix0 : ix1 + 1, iy0 : iy1 + 1, iz0 : iz1 + 1] = occ_block

        if include_normals and normals_sum is not None and normals_count is not None:
            normal_block = normals_sum[ix0 : ix1 + 1, iy0 : iy1 + 1, iz0 : iz1 + 1]
            count_block = normals_count[ix0 : ix1 + 1, iy0 : iy1 + 1, iz0 : iz1 + 1]
            normal_block[mask] += normal_vec.astype(np.float32)
            count_block[mask] += 1
            normals_sum[ix0 : ix1 + 1, iy0 : iy1 + 1, iz0 : iz1 + 1] = normal_block
            normals_count[ix0 : ix1 + 1, iy0 : iy1 + 1, iz0 : iz1 + 1] = count_block

    normals = None
    if include_normals and normals_sum is not None and normals_count is not None:
        normals = np.zeros_like(normals_sum)
        valid = normals_count > 0
        normals[valid] = normals_sum[valid] / normals_count[valid][..., None]

    summary = {
        "input_patch_count": int(len(patch_df)),
        "occupied_voxel_count": int(np.count_nonzero(occupancy > 0.5)),
        "occupancy_ratio": float(np.mean(occupancy > 0.5)),
        "include_normals": bool(include_normals),
        "thickness_vox": float(thickness_vox),
    }
    return occupancy, normals, summary


def export_voxel_preview_vtk(
    output_path: Path,
    occupancy: np.ndarray,
    grid: GridSpec,
    threshold: float = 0.5,
) -> Path:
    idx = np.argwhere(occupancy >= float(threshold))
    points = np.column_stack(
        [
            grid.x_min + (idx[:, 0].astype(float) + 0.5) * grid.dx,
            grid.y_min + (idx[:, 1].astype(float) + 0.5) * grid.dy,
            grid.z_min + (idx[:, 2].astype(float) + 0.5) * grid.dz,
        ]
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# vtk DataFile Version 3.0",
        "voxel occupancy preview",
        "ASCII",
        "DATASET POLYDATA",
        f"POINTS {len(points)} float",
    ]
    lines.extend(f"{float(p[0]):.6f} {float(p[1]):.6f} {float(p[2]):.6f}" for p in points)
    vertex_size = int(len(points) * 2)
    lines.append(f"VERTICES {len(points)} {vertex_size}")
    lines.extend(f"1 {idx_value}" for idx_value in range(len(points)))
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def build_patch_vtk_cell_data(
    patch_df: pd.DataFrame,
    extra_cell_data: dict[str, np.ndarray] | None = None,
) -> tuple[dict[str, np.ndarray], dict[str, dict[int, str]]]:
    source_codes, source_map = build_category_code_map(patch_df.get("SourceKind", pd.Series(dtype=str)))
    seed_type_codes, seed_type_map = build_category_code_map(patch_df.get("SeedType", pd.Series(dtype=str)))
    layer_codes, layer_map = build_category_code_map(preferred_layer_group_series(patch_df, prefer_unit_segment=False))
    layer_segment_codes, layer_segment_map = build_category_code_map(preferred_layer_group_series(patch_df, prefer_unit_segment=True))
    top_surface_codes, top_surface_map = build_category_code_map(patch_df.get("TopSurfaceCode", pd.Series(dtype=str)))
    base_surface_codes, base_surface_map = build_category_code_map(patch_df.get("BaseSurfaceCode", pd.Series(dtype=str)))

    def series_int(name: str) -> np.ndarray:
        if name not in patch_df.columns:
            return np.full(len(patch_df), -9999, dtype=int)
        return pd.to_numeric(patch_df[name], errors="coerce").fillna(-9999).to_numpy(dtype=int)

    def series_float(name: str) -> np.ndarray:
        if name not in patch_df.columns:
            return np.full(len(patch_df), -9999.0, dtype=float)
        return pd.to_numeric(patch_df[name], errors="coerce").fillna(-9999.0).to_numpy(dtype=float)

    cell_data = {
        "PatchIndex": series_int("PatchIndex"),
        "BlockX": series_int("BlockX"),
        "BlockY": series_int("BlockY"),
        "SourceKindCode": source_codes,
        "SeedTypeCode": seed_type_codes,
        "LayerCode": layer_codes,
        "LayerSegmentCode": layer_segment_codes,
        "TopSurfaceCodeInt": top_surface_codes,
        "BaseSurfaceCodeInt": base_surface_codes,
        "CenterTime": series_float("CenterTIME"),
        "CenterDepth": series_float("CenterDepth"),
        "Azimuth": series_float("Azimuth"),
        "Dip": series_float("Dip"),
        "PatchLength": series_float("PatchLength"),
        "PatchHeight": series_float("PatchHeight"),
        "DensityWeight": series_float("DensityWeight"),
        "LengthWeight": series_float("LengthWeight"),
        "Confidence": series_float("Confidence"),
    }
    if extra_cell_data:
        cell_data.update(extra_cell_data)
    mappings = {
        "SourceKindCode": source_map,
        "SeedTypeCode": seed_type_map,
        "LayerCode": layer_map,
        "LayerSegmentCode": layer_segment_map,
        "TopSurfaceCodeInt": top_surface_map,
        "BaseSurfaceCodeInt": base_surface_map,
    }
    return cell_data, mappings


def export_patch_vtk_files(
    patch_df: pd.DataFrame,
    output_dir: Path,
    base_name: str,
    title_prefix: str,
    config: VtkPatchExportConfig,
    extra_cell_data: dict[str, np.ndarray] | None = None,
) -> dict[str, Any]:
    raw_path = output_dir / f"{base_name}_raw_time.vtk"
    display_path = output_dir / f"{base_name}_display.vtk"
    if patch_df.empty:
        empty_points = np.empty((0, 3), dtype=float)
        empty_polygons: list[list[int]] = []
        empty_cell_data: dict[str, np.ndarray] = {}
        write_legacy_vtk_polygons(raw_path, f"{title_prefix}_raw_time", empty_points, empty_polygons, empty_cell_data)
        write_legacy_vtk_polygons(display_path, f"{title_prefix}_display", empty_points, empty_polygons, empty_cell_data)
        return {
            "raw_vtk": str(raw_path),
            "display_vtk": str(display_path),
            "mappings": {},
            "patch_count": 0,
        }
    cell_data, mappings = build_patch_vtk_cell_data(patch_df, extra_cell_data=extra_cell_data)
    raw_points: list[list[float]] = []
    display_points: list[list[float]] = []
    polygons: list[list[int]] = []
    for _, row in patch_df.iterrows():
        vertices = np.array([[row[f"V{i}X"], row[f"V{i}Y"], row[f"V{i}Z"]] for i in range(1, 5)], dtype=float)
        start_idx = len(raw_points)
        raw_points.extend(vertices.tolist())
        display_vertices = vertices.copy()
        display_vertices[:, 2] = transform_display_z(display_vertices[:, 2], config)
        display_points.extend(display_vertices.tolist())
        polygons.append([start_idx, start_idx + 1, start_idx + 2, start_idx + 3])

    write_legacy_vtk_polygons(raw_path, f"{title_prefix}_raw_time", np.asarray(raw_points, dtype=float), polygons, cell_data)
    write_legacy_vtk_polygons(display_path, f"{title_prefix}_display", np.asarray(display_points, dtype=float), polygons, cell_data)
    return {
        "raw_vtk": str(raw_path),
        "display_vtk": str(display_path),
        "mappings": mappings,
        "patch_count": int(len(patch_df)),
    }


def export_patch_comparison_vtk(
    input_patch_df: pd.DataFrame,
    output_patch_df: pd.DataFrame,
    output_dir: Path,
    config: VtkPatchExportConfig,
) -> dict[str, Any]:
    frames: list[pd.DataFrame] = []
    if not input_patch_df.empty:
        work = input_patch_df.copy()
        work["CompareSourceCode"] = 1
        frames.append(work)
    if not output_patch_df.empty:
        work = output_patch_df.copy()
        work["CompareSourceCode"] = 2
        frames.append(work)
    if not frames:
        return {
            "raw_vtk": "",
            "display_vtk": "",
            "mappings": {"CompareSourceCode": {0: "", 1: "input", 2: "roundtrip"}},
            "patch_count": 0,
        }
    merged = pd.concat(frames, ignore_index=True, sort=False)
    export_result = export_patch_vtk_files(
        patch_df=merged,
        output_dir=output_dir,
        base_name="patch_compare",
        title_prefix="patch_compare",
        config=config,
        extra_cell_data={"CompareSourceCode": merged["CompareSourceCode"].fillna(0).astype(int).to_numpy(dtype=int)},
    )
    mappings = dict(export_result.get("mappings", {}))
    mappings["CompareSourceCode"] = {0: "", 1: "input", 2: "roundtrip"}
    export_result["mappings"] = mappings
    return export_result


def save_voxel_package(
    output_path: Path,
    occupancy: np.ndarray,
    normals: np.ndarray | None,
    grid: GridSpec,
    metadata: dict[str, Any],
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    save_payload: dict[str, Any] = {
        "occupancy": occupancy.astype(np.float32),
        "x_centers": grid.x_centers.astype(np.float32),
        "y_centers": grid.y_centers.astype(np.float32),
        "z_centers": grid.z_centers.astype(np.float32),
    }
    if normals is not None:
        save_payload["normals"] = normals.astype(np.float32)
    np.savez_compressed(output_path, **save_payload)
    meta_path = output_path.with_suffix(".json")
    metadata_payload = dict(metadata)
    metadata_payload["grid"] = grid.to_json_dict()
    write_json(meta_path, metadata_payload)
    return output_path


def load_voxel_package(npz_path: Path) -> tuple[np.ndarray, np.ndarray | None, GridSpec, dict[str, Any]]:
    data = np.load(npz_path, allow_pickle=True)
    occupancy = data["occupancy"].astype(np.float32)
    normals = data["normals"].astype(np.float32) if "normals" in data.files else None
    meta_path = npz_path.with_suffix(".json")
    metadata = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}
    grid_payload = metadata.get("grid")
    if grid_payload is None:
        x_centers = data["x_centers"].astype(float)
        y_centers = data["y_centers"].astype(float)
        z_centers = data["z_centers"].astype(float)
        grid = GridSpec(
            unit_id=str(metadata.get("unit_id", "")),
            block_x=int(metadata.get("block_x", 0)),
            block_y=int(metadata.get("block_y", 0)),
            x_min=float(x_centers[0] - (x_centers[1] - x_centers[0]) / 2.0),
            x_max=float(x_centers[-1] + (x_centers[1] - x_centers[0]) / 2.0),
            y_min=float(y_centers[0] - (y_centers[1] - y_centers[0]) / 2.0),
            y_max=float(y_centers[-1] + (y_centers[1] - y_centers[0]) / 2.0),
            z_min=float(z_centers[0] - (z_centers[1] - z_centers[0]) / 2.0),
            z_max=float(z_centers[-1] + (z_centers[1] - z_centers[0]) / 2.0),
            nx=int(len(x_centers)),
            ny=int(len(y_centers)),
            nz=int(len(z_centers)),
            dx=float(x_centers[1] - x_centers[0]),
            dy=float(y_centers[1] - y_centers[0]),
            dz=float(z_centers[1] - z_centers[0]),
        )
    else:
        grid = GridSpec(**grid_payload)
    return occupancy, normals, grid, metadata


def connectivity_structure(connectivity: int) -> np.ndarray:
    connectivity = int(connectivity)
    if connectivity <= 1:
        return generate_binary_structure(3, 1)
    if connectivity == 2:
        return generate_binary_structure(3, 2)
    return generate_binary_structure(3, 3)


def assign_layer_by_time(center_time: float, layers_df: pd.DataFrame) -> dict[str, Any]:
    if layers_df.empty:
        return {}
    for _, row in layers_df.iterrows():
        top_time = pd.to_numeric(row.get("TopTime"), errors="coerce")
        base_time = pd.to_numeric(row.get("BaseTime"), errors="coerce")
        if pd.isna(top_time) or pd.isna(base_time):
            continue
        lower = float(min(top_time, base_time))
        upper = float(max(top_time, base_time))
        if lower <= float(center_time) <= upper:
            return row.to_dict()
    return {}


def normal_to_azimuth_dip(normal_vec: np.ndarray) -> tuple[float, float]:
    normal_vec = normalize(normal_vec)
    if normal_vec[2] < 0:
        normal_vec = -normal_vec
    dip = math.degrees(math.acos(float(np.clip(normal_vec[2], -1.0, 1.0))))
    azimuth = math.degrees(math.atan2(float(normal_vec[0]), float(normal_vec[1])))
    return float(azimuth % 360.0), float(dip)


def compute_component_geometry(
    points_idx: np.ndarray,
    weights: np.ndarray,
    grid: GridSpec,
    normals_component: np.ndarray | None = None,
) -> dict[str, Any]:
    center_idx = np.average(points_idx, axis=0, weights=weights)
    centered = points_idx - center_idx
    cov = (centered * weights[:, None]).T @ centered / float(weights.sum())
    eigvals, eigvecs = np.linalg.eigh(cov)
    order = np.argsort(eigvals)[::-1]
    u_vec = normalize(eigvecs[:, order[0]])
    v_vec = normalize(eigvecs[:, order[1]])
    normal_vec = normalize(eigvecs[:, order[2]])
    if normals_component is not None and len(normals_component) == len(points_idx):
        normal_mean = normalize(np.mean(normals_component, axis=0))
        if np.linalg.norm(normal_mean) > 1e-6 and float(np.dot(normal_vec, normal_mean)) < 0:
            normal_vec = -normal_vec
    proj_u = centered @ u_vec
    proj_v = centered @ v_vec
    u_min = float(proj_u.min() - 0.5)
    u_max = float(proj_u.max() + 0.5)
    v_min = float(proj_v.min() - 0.5)
    v_max = float(proj_v.max() + 0.5)
    rect_center_idx = center_idx + ((u_min + u_max) / 2.0) * u_vec + ((v_min + v_max) / 2.0) * v_vec
    length_idx = max(1e-6, float(u_max - u_min))
    height_idx = max(1e-6, float(v_max - v_min))
    vertices_idx = np.array(
        [
            rect_center_idx - (length_idx / 2.0) * u_vec - (height_idx / 2.0) * v_vec,
            rect_center_idx + (length_idx / 2.0) * u_vec - (height_idx / 2.0) * v_vec,
            rect_center_idx + (length_idx / 2.0) * u_vec + (height_idx / 2.0) * v_vec,
            rect_center_idx - (length_idx / 2.0) * u_vec + (height_idx / 2.0) * v_vec,
        ],
        dtype=float,
    )
    vertices_phys = index_to_physical(vertices_idx, grid)
    rect_center_phys = index_to_physical(rect_center_idx, grid)
    edge_u_phys = vertices_phys[1] - vertices_phys[0]
    edge_v_phys = vertices_phys[3] - vertices_phys[0]
    normal_phys = normalize(np.cross(edge_u_phys, edge_v_phys))
    return {
        "center_idx": center_idx,
        "u_vec": u_vec,
        "v_vec": v_vec,
        "normal_vec": normal_vec,
        "proj_u": proj_u,
        "proj_v": proj_v,
        "u_min": u_min,
        "u_max": u_max,
        "v_min": v_min,
        "v_max": v_max,
        "vertices_idx": vertices_idx,
        "vertices_phys": vertices_phys,
        "rect_center_phys": rect_center_phys,
        "edge_u_phys": edge_u_phys,
        "edge_v_phys": edge_v_phys,
        "normal_phys": normal_phys,
        "length_phys": float(np.linalg.norm(edge_u_phys)),
        "height_phys": float(np.linalg.norm(edge_v_phys)),
    }


def resolve_split_limits(reference_patch_stats: dict[str, Any] | None) -> tuple[float, float]:
    default_length_limit = 70.0
    default_height_limit = 16.0
    if not reference_patch_stats:
        return default_length_limit, default_height_limit
    length_candidates = [
        float(reference_patch_stats.get("patch_length_q90", np.nan)) * 1.20,
        float(reference_patch_stats.get("patch_length_q75", np.nan)) * 1.45,
        float(reference_patch_stats.get("patch_length_median", np.nan)) * 1.85,
        default_length_limit,
    ]
    height_candidates = [
        float(reference_patch_stats.get("patch_height_q90", np.nan)) * 1.20,
        float(reference_patch_stats.get("patch_height_q75", np.nan)) * 1.45,
        float(reference_patch_stats.get("patch_height_median", np.nan)) * 1.85,
        default_height_limit,
    ]
    length_limit = max(value for value in length_candidates if np.isfinite(value) and value > 0)
    height_limit = max(value for value in height_candidates if np.isfinite(value) and value > 0)
    return float(length_limit), float(height_limit)


def find_projection_split(
    projection: np.ndarray,
    weights: np.ndarray,
    min_points_side: int,
    valley_ratio: float = 0.55,
    bin_width: float = 1.0,
) -> float | None:
    values = np.asarray(projection, dtype=float)
    if values.size < max(6, int(min_points_side) * 2):
        return None
    lower = float(np.min(values))
    upper = float(np.max(values))
    if upper - lower < max(2.5, float(bin_width) * 3.0):
        return None

    edges = np.arange(math.floor(lower), math.ceil(upper) + float(bin_width), float(bin_width), dtype=float)
    if len(edges) < 4:
        return None

    weighted_hist, _ = np.histogram(values, bins=edges, weights=np.asarray(weights, dtype=float))
    count_hist, _ = np.histogram(values, bins=edges)
    if count_hist.size < 3:
        return None

    kernel = np.array([1.0, 2.0, 3.0, 2.0, 1.0], dtype=float)
    kernel /= kernel.sum()
    smooth_hist = np.convolve(weighted_hist, kernel, mode="same")

    zero_bins = np.flatnonzero(count_hist == 0)
    best_zero_cut: tuple[float, float] | None = None
    for bin_idx in zero_bins.tolist():
        left_count = int(count_hist[: bin_idx + 1].sum())
        right_count = int(count_hist[bin_idx + 1 :].sum())
        if left_count < int(min_points_side) or right_count < int(min_points_side):
            continue
        score = abs(left_count - right_count)
        cut_coord = float(edges[bin_idx + 1])
        if best_zero_cut is None or score < best_zero_cut[1]:
            best_zero_cut = (cut_coord, float(score))
    if best_zero_cut is not None:
        return best_zero_cut[0]

    best_cut: tuple[float, float] | None = None
    for bin_idx in range(1, len(smooth_hist) - 1):
        left_count = int(count_hist[: bin_idx + 1].sum())
        right_count = int(count_hist[bin_idx + 1 :].sum())
        if left_count < int(min_points_side) or right_count < int(min_points_side):
            continue
        valley = float(smooth_hist[bin_idx])
        left_peak = float(np.max(smooth_hist[: bin_idx + 1]))
        right_peak = float(np.max(smooth_hist[bin_idx + 1 :]))
        if left_peak <= 1e-8 or right_peak <= 1e-8:
            continue
        ratio = valley / max(min(left_peak, right_peak), 1e-8)
        if ratio > float(valley_ratio):
            continue
        balance_penalty = abs(left_count - right_count) / max(left_count + right_count, 1)
        score = ratio + 0.35 * balance_penalty
        cut_coord = float(edges[bin_idx + 1])
        if best_cut is None or score < best_cut[1]:
            best_cut = (cut_coord, score)
    return best_cut[0] if best_cut is not None else None


def split_component_points_recursive(
    points_idx: np.ndarray,
    weights: np.ndarray,
    grid: GridSpec,
    min_component_voxels: int,
    reference_patch_stats: dict[str, Any] | None,
    normals_component: np.ndarray | None = None,
    depth: int = 0,
    max_depth: int = 3,
) -> tuple[list[tuple[np.ndarray, np.ndarray, np.ndarray | None]], int]:
    if depth >= int(max_depth) or len(points_idx) < max(int(min_component_voxels) * 2, 12):
        return [(points_idx, weights, normals_component)], 0

    geom = compute_component_geometry(points_idx, weights, grid, normals_component=normals_component)
    length_limit, height_limit = resolve_split_limits(reference_patch_stats)
    candidate_axes: list[tuple[str, np.ndarray, float, float]] = []
    if float(geom["length_phys"]) > float(length_limit):
        candidate_axes.append(("u", np.asarray(geom["proj_u"], dtype=float), float(geom["length_phys"]), float(length_limit)))
    if float(geom["height_phys"]) > float(height_limit):
        candidate_axes.append(("v", np.asarray(geom["proj_v"], dtype=float), float(geom["height_phys"]), float(height_limit)))

    candidate_axes.sort(key=lambda item: item[2] / max(item[3], 1e-6), reverse=True)
    min_points_side = max(int(min_component_voxels), min(24, int(len(points_idx) * 0.12)))

    for axis_name, projection, extent_value, limit_value in candidate_axes:
        cut_coord = find_projection_split(
            projection=projection,
            weights=weights,
            min_points_side=min_points_side,
            valley_ratio=0.60 if extent_value / max(limit_value, 1e-6) > 1.8 else 0.50,
            bin_width=1.0,
        )
        if cut_coord is None:
            continue
        left_mask = projection <= float(cut_coord)
        right_mask = ~left_mask
        if int(np.count_nonzero(left_mask)) < int(min_component_voxels) or int(np.count_nonzero(right_mask)) < int(min_component_voxels):
            continue

        left_normals = normals_component[left_mask] if normals_component is not None else None
        right_normals = normals_component[right_mask] if normals_component is not None else None
        left_parts, left_split_count = split_component_points_recursive(
            points_idx=points_idx[left_mask],
            weights=weights[left_mask],
            grid=grid,
            min_component_voxels=min_component_voxels,
            reference_patch_stats=reference_patch_stats,
            normals_component=left_normals,
            depth=depth + 1,
            max_depth=max_depth,
        )
        right_parts, right_split_count = split_component_points_recursive(
            points_idx=points_idx[right_mask],
            weights=weights[right_mask],
            grid=grid,
            min_component_voxels=min_component_voxels,
            reference_patch_stats=reference_patch_stats,
            normals_component=right_normals,
            depth=depth + 1,
            max_depth=max_depth,
        )
        return left_parts + right_parts, 1 + left_split_count + right_split_count

    return [(points_idx, weights, normals_component)], 0


def fit_voxel_components_to_patches(
    occupancy: np.ndarray,
    grid: GridSpec,
    threshold: float,
    min_component_voxels: int,
    connectivity: int,
    layers_df: pd.DataFrame,
    normals: np.ndarray | None = None,
    reference_patch_stats: dict[str, Any] | None = None,
    max_split_depth: int = 3,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    binary = occupancy >= float(threshold)
    labels, component_count = ndi_label(binary, structure=connectivity_structure(connectivity))
    rows: list[dict[str, Any]] = []
    skipped_small = 0
    split_component_count = 0

    for label_idx in range(1, int(component_count) + 1):
        idx = np.argwhere(labels == label_idx)
        if len(idx) < int(min_component_voxels):
            skipped_small += 1
            continue
        weights = occupancy[labels == label_idx].astype(float)
        if float(weights.sum()) <= 0:
            weights = np.ones(len(idx), dtype=float)
        points_idx = idx.astype(float) + 0.5
        normals_component = normals[labels == label_idx].reshape(-1, 3) if normals is not None else None
        subcomponents, split_count = split_component_points_recursive(
            points_idx=points_idx,
            weights=weights,
            grid=grid,
            min_component_voxels=min_component_voxels,
            reference_patch_stats=reference_patch_stats,
            normals_component=normals_component,
            depth=0,
            max_depth=max_split_depth,
        )
        split_component_count += int(split_count)

        for split_idx, (sub_points_idx, sub_weights, sub_normals) in enumerate(subcomponents, start=1):
            if len(sub_points_idx) < int(min_component_voxels):
                skipped_small += 1
                continue
            geom = compute_component_geometry(sub_points_idx, sub_weights, grid, normals_component=sub_normals)
            vertices_phys = np.asarray(geom["vertices_phys"], dtype=float)
            rect_center_phys = np.asarray(geom["rect_center_phys"], dtype=float)
            edge_u_phys = np.asarray(geom["edge_u_phys"], dtype=float)
            edge_v_phys = np.asarray(geom["edge_v_phys"], dtype=float)
            normal_phys = normalize(np.asarray(geom["normal_phys"], dtype=float))
            azimuth, dip = normal_to_azimuth_dip(normal_phys)
            layer_info = assign_layer_by_time(float(rect_center_phys[2]), layers_df)

            row = {
                "PatchIndex": int(len(rows) + 1),
                "PatchID": f"{grid.unit_id}_ROUNDTRIP_{len(rows) + 1:04d}",
                "UnitID": str(grid.unit_id),
                "BlockX": int(grid.block_x),
                "BlockY": int(grid.block_y),
                "GeoIntervalKey": str(layer_info.get("GeoIntervalKey", "")),
                LAYER_SURFACE_PAIR_KEY_COL: resolve_layer_surface_pair_key_from_row(layer_info),
                UNIT_LAYER_SEGMENT_KEY_COL: resolve_unit_layer_segment_key_from_row(layer_info),
                "StrataName": str(layer_info.get("StrataName", "")),
                "TopSurfaceCode": str(layer_info.get("TopSurfaceCode", "")),
                "BaseSurfaceCode": str(layer_info.get("BaseSurfaceCode", "")),
                "SeedID": f"roundtrip_component_{label_idx:04d}_split_{split_idx:02d}",
                "SourceKind": "roundtrip_fit",
                "SourceName": "voxel_roundtrip",
                "SeedType": "roundtrip_component",
                "ParentSeedID": f"roundtrip_component_{label_idx:04d}",
                "ParentSourceKind": "component",
                "CenterX": float(rect_center_phys[0]),
                "CenterY": float(rect_center_phys[1]),
                "CenterTIME": float(rect_center_phys[2]),
                "CenterDepth": float(rect_center_phys[2]),
                "Azimuth": float(azimuth),
                "Dip": float(dip),
                "DensityWeight": float(len(sub_points_idx)),
                "LengthWeight": float(len(sub_points_idx)),
                "Confidence": 1.0,
                "PatchLength": float(np.linalg.norm(edge_u_phys)),
                "PatchHeight": float(np.linalg.norm(edge_v_phys)),
                "NormalX": float(normal_phys[0]),
                "NormalY": float(normal_phys[1]),
                "NormalZ": float(normal_phys[2]),
            }
            for vertex_offset, vertex in enumerate(vertices_phys, start=1):
                row[f"V{vertex_offset}X"] = float(vertex[0])
                row[f"V{vertex_offset}Y"] = float(vertex[1])
                row[f"V{vertex_offset}Z"] = float(vertex[2])
            rows.append(row)

    patch_df = pd.DataFrame(rows, columns=PATCH_OUTPUT_COLUMNS)
    summary = {
        "input_component_count": int(component_count),
        "output_patch_count": int(len(patch_df)),
        "skipped_small_component_count": int(skipped_small),
        "recursive_split_count": int(split_component_count),
        "threshold": float(threshold),
        "min_component_voxels": int(min_component_voxels),
        "connectivity": int(connectivity),
        "max_split_depth": int(max_split_depth),
    }
    return patch_df, summary


def patch_angle_diff_deg(value_a: float, value_b: float) -> float:
    diff = abs(float(value_a) - float(value_b)) % 360.0
    return float(min(diff, 360.0 - diff))


def build_patch_match_metrics(
    input_patches: pd.DataFrame,
    output_patches: pd.DataFrame,
    grid: GridSpec,
) -> dict[str, Any]:
    if input_patches.empty or output_patches.empty:
        return {
            "matched_patch_count": 0,
            "input_patch_count": int(len(input_patches)),
            "output_patch_count": int(len(output_patches)),
        }

    span_x = max(grid.x_max - grid.x_min, 1e-6)
    span_y = max(grid.y_max - grid.y_min, 1e-6)
    span_t = max(grid.z_max - grid.z_min, 1e-6)
    cost = np.zeros((len(input_patches), len(output_patches)), dtype=float)
    for i, (_, src) in enumerate(input_patches.iterrows()):
        for j, (_, dst) in enumerate(output_patches.iterrows()):
            dx = (float(src["CenterX"]) - float(dst["CenterX"])) / span_x
            dy = (float(src["CenterY"]) - float(dst["CenterY"])) / span_y
            dt = (float(src["CenterTIME"]) - float(dst["CenterTIME"])) / span_t
            center_cost = math.sqrt(dx * dx + dy * dy + dt * dt)
            az_cost = patch_angle_diff_deg(float(src["Azimuth"]), float(dst["Azimuth"])) / 180.0
            dip_cost = abs(float(src["Dip"]) - float(dst["Dip"])) / 90.0
            interval_penalty = 0.0
            if resolve_layer_surface_pair_key_from_row(src) != resolve_layer_surface_pair_key_from_row(dst):
                interval_penalty = 0.5
            cost[i, j] = center_cost + 0.25 * az_cost + 0.25 * dip_cost + interval_penalty

    if linear_sum_assignment is not None:
        row_idx, col_idx = linear_sum_assignment(cost)
    else:  # pragma: no cover
        row_idx = np.arange(min(cost.shape[0], cost.shape[1]))
        col_idx = np.argmin(cost[row_idx], axis=1)

    center_offsets = []
    azimuth_diffs = []
    dip_diffs = []
    length_rel_errors = []
    height_rel_errors = []
    accepted_matches = 0
    for src_idx, dst_idx in zip(row_idx, col_idx):
        if cost[int(src_idx), int(dst_idx)] > 2.0:
            continue
        src = input_patches.iloc[int(src_idx)]
        dst = output_patches.iloc[int(dst_idx)]
        dx = float(dst["CenterX"]) - float(src["CenterX"])
        dy = float(dst["CenterY"]) - float(src["CenterY"])
        dt = float(dst["CenterTIME"]) - float(src["CenterTIME"])
        center_offsets.append(math.sqrt(dx * dx + dy * dy + dt * dt))
        azimuth_diffs.append(patch_angle_diff_deg(float(src["Azimuth"]), float(dst["Azimuth"])))
        dip_diffs.append(abs(float(src["Dip"]) - float(dst["Dip"])))
        src_len = max(abs(float(src["PatchLength"])), 1e-6)
        src_height = max(abs(float(src["PatchHeight"])), 1e-6)
        length_rel_errors.append(abs(float(dst["PatchLength"]) - float(src["PatchLength"])) / src_len)
        height_rel_errors.append(abs(float(dst["PatchHeight"]) - float(src["PatchHeight"])) / src_height)
        accepted_matches += 1

    return {
        "matched_patch_count": int(accepted_matches),
        "input_patch_count": int(len(input_patches)),
        "output_patch_count": int(len(output_patches)),
        "center_offset_median": float(np.median(center_offsets)) if center_offsets else None,
        "azimuth_diff_median_deg": float(np.median(azimuth_diffs)) if azimuth_diffs else None,
        "dip_diff_median_deg": float(np.median(dip_diffs)) if dip_diffs else None,
        "length_rel_error_median": float(np.median(length_rel_errors)) if length_rel_errors else None,
        "height_rel_error_median": float(np.median(height_rel_errors)) if height_rel_errors else None,
    }


def build_voxel_overlap_metrics(
    input_occupancy: np.ndarray,
    output_occupancy: np.ndarray,
    threshold: float,
) -> dict[str, Any]:
    input_mask = input_occupancy >= float(threshold)
    output_mask = output_occupancy >= float(threshold)
    intersection = int(np.count_nonzero(input_mask & output_mask))
    union = int(np.count_nonzero(input_mask | output_mask))
    input_count = int(np.count_nonzero(input_mask))
    output_count = int(np.count_nonzero(output_mask))
    precision = float(intersection / output_count) if output_count else None
    recall = float(intersection / input_count) if input_count else None
    iou = float(intersection / union) if union else None
    return {
        "input_voxel_count": input_count,
        "output_voxel_count": output_count,
        "intersection_voxel_count": intersection,
        "union_voxel_count": union,
        "voxel_precision": precision,
        "voxel_recall": recall,
        "voxel_iou": iou,
    }


def open_or_create_doc(docx_path: Path) -> Document:
    if docx_path.exists():
        try:
            return Document(str(docx_path))
        except PackageNotFoundError:
            pass
    doc = Document()
    doc.add_heading("实验记录 20260328", level=1)
    return doc


def append_roundtrip_summary_to_docx(
    docx_path: Path,
    title: str,
    config: dict[str, Any],
    summary: dict[str, Any],
) -> None:
    doc = open_or_create_doc(docx_path)
    doc.add_heading(title, level=2)
    lines = [
        "task: 单元 DFN 与体素互转验证",
        f"time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"unit_id: {config.get('unit_id', '')}",
        f"patch_csv: {config.get('patch_csv', '')}",
        f"output_dir: {config.get('output_dir', '')}",
        f"xy_resolution: {config.get('xy_resolution', '')}",
        f"z_step_ms: {config.get('z_step_ms', '')}",
        f"thickness_vox: {config.get('thickness_vox', '')}",
        f"channels: {config.get('channels', '')}",
        f"occupancy_threshold: {config.get('occupancy_threshold', '')}",
        f"min_component_voxels: {config.get('min_component_voxels', '')}",
        f"input_patch_count: {summary.get('input_patch_count', '')}",
        f"roundtrip_patch_count: {summary.get('roundtrip_patch_count', '')}",
        f"occupied_voxel_count: {summary.get('occupied_voxel_count', '')}",
        f"voxel_iou: {summary.get('voxel_iou', '')}",
        f"matched_patch_count: {summary.get('matched_patch_count', '')}",
        f"center_offset_median: {summary.get('center_offset_median', '')}",
        f"azimuth_diff_median_deg: {summary.get('azimuth_diff_median_deg', '')}",
        f"dip_diff_median_deg: {summary.get('dip_diff_median_deg', '')}",
        f"length_rel_error_median: {summary.get('length_rel_error_median', '')}",
        f"height_rel_error_median: {summary.get('height_rel_error_median', '')}",
        f"input_patches_display_vtk: {summary.get('input_patches_display_vtk', '')}",
        f"roundtrip_patches_display_vtk: {summary.get('roundtrip_patches_display_vtk', '')}",
        f"patch_compare_display_vtk: {summary.get('patch_compare_display_vtk', '')}",
        f"summary_json: {summary.get('summary_json', '')}",
    ]
    doc.add_paragraph("\n".join(lines))
    doc.add_paragraph("")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
