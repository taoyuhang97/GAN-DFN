import argparse
import json
import os
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError


ROOT = Path(__file__).resolve().parent
SCRIPT_PATH = ROOT / "模型训练" / "优化阶段一" / "裂缝存在性分析" / "LSTM" / "imaging_well_to_fracture_cnn_lstm_test_1.py"
BASE_SAVE_DIR = Path(r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容一\裂缝存在性预测\LSTM\单井验证\成像测井裂缝预测\cnn+lstm")


def sanitize(text: str) -> str:
    text = text.replace('"', "")
    text = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]+", "_", text)
    return text.strip("_")


def append_result_to_docx(docx_path: Path, title: str, config: dict, results: list[dict]) -> None:
    if docx_path.exists():
        try:
            doc = Document(str(docx_path))
        except PackageNotFoundError:
            doc = Document()
            doc.add_heading("实验记录 20260319", level=1)
    else:
        doc = Document()
        doc.add_heading("实验记录 20260319", level=1)

    doc.add_heading(title, level=2)
    doc.add_paragraph(f"记录时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    config_lines = [
        f"SEIS_MODE: {config['SEIS_MODE']}",
        f"SEQ_LEN: {config['SEQ_LEN']}",
        f"imaging_well_features: {config['imaging_well_features']}",
        f"selection_metric: {config['selection_metric']}",
        f"use_dynamic_threshold: {config['use_dynamic_threshold']}",
        f"use_density_regression: {config['use_density_regression']}",
        f"use_domain_adversarial: {config['use_domain_adversarial']}",
        f"use_all_other_wells: {config['use_all_other_wells']}",
        f"use_selection_accuracy_floor: {config['use_selection_accuracy_floor']}",
        f"MIN_SELECTION_ACCURACY: {config['min_selection_accuracy']}",
        f"MIN_TRAIN_WELLS: {config['min_train_wells']}",
        f"DIST_THRESHOLD: {config['dist_threshold']}",
        f"only_val_wells: {config['only_val_wells']}",
        f"custom_train_wells_json: {config['custom_train_wells_json']}",
        f"save_dir: {config['save_dir']}",
    ]
    doc.add_paragraph("\n".join(config_lines))

    table = doc.add_table(rows=1, cols=7)
    hdr = table.rows[0].cells
    hdr[0].text = "井名"
    hdr[1].text = "Accuracy"
    hdr[2].text = "AUC"
    hdr[3].text = "Precision"
    hdr[4].text = "Recall"
    hdr[5].text = "F1"
    hdr[6].text = "Threshold"

    for row in results:
        cells = table.add_row().cells
        cells[0].text = str(row["val_well"])
        cells[1].text = f"{float(row['Accuracy']):.4f}"
        cells[2].text = f"{float(row['AUC']):.4f}"
        cells[3].text = f"{float(row['Precision']):.4f}"
        cells[4].text = f"{float(row['Recall']):.4f}"
        cells[5].text = f"{float(row['F1']):.4f}"
        cells[6].text = f"{float(row['Threshold']):.4f}"

    all_over_85 = all(float(row["Accuracy"]) > 0.85 for row in results)
    doc.add_paragraph(f"四口井 accuracy 是否全部 > 85%: {'是' if all_over_85 else '否'}")
    doc.add_paragraph("")
    doc.save(str(docx_path))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", required=True)
    parser.add_argument("--seq-len", type=int, required=True)
    parser.add_argument("--seis-mode", required=True)
    parser.add_argument("--features-json", required=True)
    parser.add_argument("--docx-path", required=True)
    parser.add_argument("--result-dir")
    parser.add_argument("--skip-run", action="store_true")
    parser.add_argument("--selection-metric", default="accuracy")
    parser.add_argument("--use-dynamic-threshold", default="true")
    parser.add_argument("--use-density-regression", default="false")
    parser.add_argument("--use-domain-adversarial", default="false")
    parser.add_argument("--use-all-other-wells", default="false")
    parser.add_argument("--use-selection-accuracy-floor", default="true")
    parser.add_argument("--min-selection-accuracy", default="0.80")
    parser.add_argument("--min-train-wells", default="2")
    parser.add_argument("--dist-threshold", default="0.4")
    parser.add_argument("--only-val-wells", default="")
    parser.add_argument("--custom-train-wells-json", default="")
    args = parser.parse_args()

    try:
        features = json.loads(args.features_json)
    except json.JSONDecodeError:
        features = [item.strip() for item in args.features_json.split(",") if item.strip()]
    feature_tag = sanitize("_".join(features))
    exp_tag = f"{args.exp_id}_{args.seis_mode}_seq{args.seq_len}_{feature_tag}"
    save_dir = Path(args.result_dir) if args.result_dir else (BASE_SAVE_DIR / exp_tag)
    save_dir.mkdir(parents=True, exist_ok=True)

    env = os.environ.copy()
    env["EXP_SEQ_LEN"] = str(args.seq_len)
    env["EXP_SEIS_MODE"] = args.seis_mode
    env["EXP_IMAGING_FEATURES"] = json.dumps(features, ensure_ascii=False)
    env["EXP_SELECTION_METRIC"] = args.selection_metric
    env["EXP_USE_DYNAMIC_THRESHOLD"] = args.use_dynamic_threshold
    env["EXP_USE_DENSITY_REGRESSION"] = args.use_density_regression
    env["EXP_USE_DOMAIN_ADVERSARIAL"] = args.use_domain_adversarial
    env["EXP_USE_ALL_OTHER_WELLS"] = args.use_all_other_wells
    env["EXP_USE_SELECTION_ACCURACY_FLOOR"] = args.use_selection_accuracy_floor
    env["EXP_MIN_SELECTION_ACCURACY"] = str(args.min_selection_accuracy)
    env["EXP_MIN_TRAIN_WELLS"] = str(args.min_train_wells)
    env["EXP_DIST_THRESHOLD"] = str(args.dist_threshold)
    if args.only_val_wells:
        only_val_wells = [item.strip() for item in args.only_val_wells.split(",") if item.strip()]
        env["EXP_ONLY_VAL_WELLS"] = json.dumps(only_val_wells, ensure_ascii=False)
    if args.custom_train_wells_json:
        env["EXP_CUSTOM_TRAIN_WELLS_MAP"] = args.custom_train_wells_json
    env["EXP_SAVE_DIR"] = str(save_dir)

    if not args.skip_run:
        cmd = [sys.executable, str(SCRIPT_PATH)]
        completed = subprocess.run(
            cmd,
            env=env,
            cwd=str(ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
        )
        (save_dir / "run_stdout.log").write_text(completed.stdout, encoding="utf-8")
        (save_dir / "run_stderr.log").write_text(completed.stderr, encoding="utf-8")
        if completed.returncode != 0:
            print(completed.stdout[-4000:], file=sys.stdout)
            print(completed.stderr[-4000:], file=sys.stderr)
            return completed.returncode

    result_path = save_dir / "loo_results.csv"
    if not result_path.exists():
        print(f"Missing result file: {result_path}", file=sys.stderr)
        return 2

    import csv

    with result_path.open("r", encoding="utf-8-sig", newline="") as f:
        results = list(csv.DictReader(f))

    docx_path = Path(args.docx_path)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    config = {
        "SEQ_LEN": args.seq_len,
        "SEIS_MODE": args.seis_mode,
        "imaging_well_features": features,
        "selection_metric": args.selection_metric,
        "use_dynamic_threshold": args.use_dynamic_threshold,
        "use_density_regression": args.use_density_regression,
        "use_domain_adversarial": args.use_domain_adversarial,
        "use_all_other_wells": args.use_all_other_wells,
        "use_selection_accuracy_floor": args.use_selection_accuracy_floor,
        "min_selection_accuracy": args.min_selection_accuracy,
        "min_train_wells": args.min_train_wells,
        "dist_threshold": args.dist_threshold,
        "only_val_wells": args.only_val_wells or env.get("EXP_ONLY_VAL_WELLS", ""),
        "custom_train_wells_json": args.custom_train_wells_json or env.get("EXP_CUSTOM_TRAIN_WELLS_MAP", ""),
        "save_dir": str(save_dir),
    }
    append_result_to_docx(docx_path, f"实验 {args.exp_id}", config, results)

    print(json.dumps({
        "exp_id": args.exp_id,
        "save_dir": str(save_dir),
        "results": results,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
