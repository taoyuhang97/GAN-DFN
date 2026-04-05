from __future__ import annotations

import argparse
import importlib.util
import json
import math
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document


def find_repo_root(start: Path) -> Path:
    current = start.resolve()
    if current.is_file():
        current = current.parent
    for candidate in [current, *current.parents]:
        if (candidate / "README.md").exists() and (candidate / "模型训练").exists():
            return candidate
    raise FileNotFoundError(f"Failed to locate repo root from: {start}")


REPO_ROOT = find_repo_root(Path(__file__))
WORKFLOW_ROOT = REPO_ROOT / "模型训练" / "优化阶段一" / "基于地层约束的测井裂缝预测"
DEFAULT_SINGLE_SCRIPT_PATH = WORKFLOW_ROOT / "run_conventional_log_strata_validation.py"
RUNTIME_MODULE_PATH = WORKFLOW_ROOT / "strata_expert_deploy" / "runtime.py"
if str(WORKFLOW_ROOT) not in sys.path:
    sys.path.insert(0, str(WORKFLOW_ROOT))


def load_module_from_file(module_name: str, module_path: Path):
    spec = importlib.util.spec_from_file_location(module_name, str(module_path))
    if spec is None or spec.loader is None:
        raise ImportError(f"Failed to load module from: {module_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


validation_module = load_module_from_file("virtual_fracture_validation_module", DEFAULT_SINGLE_SCRIPT_PATH)
runtime_module = load_module_from_file("virtual_fracture_runtime_module", RUNTIME_MODULE_PATH)

CSV_ENCODINGS = validation_module.CSV_ENCODINGS
DEFAULT_LAYER_DIR = validation_module.DEFAULT_LAYER_DIR
DEFAULT_STAGE1_LIBRARY_DIR = validation_module.DEFAULT_STAGE1_LIBRARY_DIR
DEFAULT_STAGE2_LIBRARY_DIR = validation_module.DEFAULT_STAGE2_LIBRARY_DIR
FINAL_LOG_FILENAME = validation_module.FINAL_LOG_FILENAME
FINAL_META_FILENAME = validation_module.FINAL_META_FILENAME
FINAL_POINTS_FILENAME = validation_module.FINAL_POINTS_FILENAME
FINAL_SEGMENTS_FILENAME = validation_module.FINAL_SEGMENTS_FILENAME
FINAL_STRATA_SEGMENTATION_FILENAME = validation_module.FINAL_STRATA_SEGMENTATION_FILENAME

DEFAULT_LOG_FEATURES = runtime_module.DEFAULT_LOG_FEATURES
parse_json_list = runtime_module.parse_json_list
sanitize = runtime_module.sanitize


DEFAULT_VIRTUAL_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容二\虚拟测井构建\虚拟测井批量生成"
)
DEFAULT_OUTPUT_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容二\虚拟测井构建\虚拟裂缝预测"
)
DEFAULT_DOCX_PATH = Path(r"D:\项目\石油开采\断缝储实验\实验记录20260328.docx")

SUMMARY_CSV_NAME = "batch_prediction_summary.csv"
SUMMARY_JSON_NAME = "batch_prediction_summary.json"
MISSING_FEATURE_DOCX_NAME = "虚拟测井缺失AC_GR属性情况.docx"
INCOMPLETE_INPUT_CSV_NAME = "missing_or_incomplete_virtual_wells.csv"
VIRTUAL_BATCH_META_FILENAME = "virtual_batch_meta.json"

REQUIRED_RESULT_FILES = {
    FINAL_LOG_FILENAME,
    FINAL_POINTS_FILENAME,
    FINAL_SEGMENTS_FILENAME,
    FINAL_STRATA_SEGMENTATION_FILENAME,
    FINAL_META_FILENAME,
}
REQUIRED_INPUT_FILES = {
    "virtual_well_index.csv",
    "virtual_well_around_data.csv",
}
UNIT_ID_PATTERN = re.compile(r"^BX(\d+)_BY(\d+)$", re.IGNORECASE)


def read_csv_flexible(path: Path) -> tuple[pd.DataFrame, str]:
    last_error: Exception | None = None
    for encoding in CSV_ENCODINGS:
        try:
            return pd.read_csv(path, encoding=encoding), encoding
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Failed to read CSV: {path}") from last_error


def read_csv_header_flexible(path: Path) -> tuple[list[str], str]:
    last_error: Exception | None = None
    for encoding in CSV_ENCODINGS:
        try:
            header_df = pd.read_csv(path, encoding=encoding, nrows=0)
            return [str(col).strip() for col in header_df.columns], encoding
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Failed to read CSV header: {path}") from last_error


def parse_unit_sort_key(unit_id: str) -> tuple[int, int, str]:
    match = UNIT_ID_PATTERN.match(str(unit_id).strip())
    if not match:
        return (10**9, 10**9, str(unit_id))
    return (int(match.group(1)), int(match.group(2)), str(unit_id))


def normalize_path_text(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        return str(Path(text).resolve(strict=False))
    except Exception:
        return text


def normalize_text(value: object) -> str:
    return str(value or "").strip()


def normalize_feature_key(features: list[str]) -> str:
    return json.dumps([str(item).strip() for item in features if str(item).strip()], ensure_ascii=False)


def coalesce_record_value(record: dict[str, object], candidates: list[str], default: object = "") -> object:
    for key in candidates:
        if key not in record:
            continue
        value = record.get(key)
        if pd.isna(value):
            continue
        text = str(value).strip()
        if text == "":
            continue
        return value
    return default


def coalesce_int(record: dict[str, object], candidates: list[str], default: int = 0) -> int:
    value = coalesce_record_value(record, candidates, default="")
    try:
        return int(float(value))
    except Exception:
        return int(default)


def load_json_safe(path: Path) -> dict[str, object]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def discover_virtual_well_units(
    virtual_root: Path,
    include_unit_ids: list[str] | None = None,
    limit: int = 0,
) -> list[Path]:
    if not virtual_root.exists():
        raise FileNotFoundError(f"Virtual well root not found: {virtual_root}")
    include_set = {str(item).strip() for item in (include_unit_ids or []) if str(item).strip()}
    unit_dirs: list[Path] = []
    for unit_dir in sorted(
        [path for path in virtual_root.iterdir() if path.is_dir()],
        key=lambda item: parse_unit_sort_key(item.name),
    ):
        if include_set and unit_dir.name not in include_set:
            continue
        unit_dirs.append(unit_dir)
    if limit and limit > 0:
        unit_dirs = unit_dirs[:limit]
    return unit_dirs


def load_virtual_well_context(unit_dir: Path) -> dict[str, object]:
    index_csv = unit_dir / "virtual_well_index.csv"
    around_csv = unit_dir / "virtual_well_around_data.csv"
    strata_csv = unit_dir / "virtual_well_strata_segmentation.csv"
    curve_csv = unit_dir / "virtual_well_curve.csv"
    run_summary_json = unit_dir / "run_summary.json"
    missing_inputs = [name for name in sorted(REQUIRED_INPUT_FILES) if not (unit_dir / name).exists()]

    if index_csv.exists():
        index_df, index_encoding = read_csv_flexible(index_csv)
    else:
        index_df, index_encoding = pd.DataFrame(), ""
    row = index_df.iloc[0].to_dict() if not index_df.empty else {}
    run_summary = load_json_safe(run_summary_json) if run_summary_json.exists() else {}

    unit_id = normalize_text(coalesce_record_value(row, ["unit_id", "UnitID"], default=unit_dir.name)) or unit_dir.name
    virtual_well_name = normalize_text(
        coalesce_record_value(
            row,
            ["VirtualWellName", "virtual_well_name", "WellName"],
            default=run_summary.get("virtual_well_name", unit_id),
        )
    ) or unit_id
    predict_method = normalize_text(
        coalesce_record_value(
            row,
            ["PredictMethod", "predict_method", "Method"],
            default=run_summary.get("predict_method", ""),
        )
    )
    block_x = coalesce_int(row, ["BlockX", "block_x"], default=parse_unit_sort_key(unit_id)[0] if UNIT_ID_PATTERN.match(unit_id) else 0)
    block_y = coalesce_int(row, ["BlockY", "block_y"], default=parse_unit_sort_key(unit_id)[1] if UNIT_ID_PATTERN.match(unit_id) else 0)
    return {
        "UnitID": unit_id,
        "VirtualWellName": virtual_well_name,
        "PredictMethod": predict_method,
        "BlockX": block_x,
        "BlockY": block_y,
        "UnitDir": unit_dir,
        "AroundCsv": around_csv,
        "IndexCsv": index_csv,
        "StrataCsv": strata_csv,
        "CurveCsv": curve_csv,
        "RunSummaryJson": run_summary_json,
        "IndexEncoding": index_encoding,
        "MissingInputs": missing_inputs,
    }


def detect_feature_status(csv_path: Path, requested_candidates: list[str]) -> dict[str, object]:
    columns, encoding = read_csv_header_flexible(csv_path)
    col_set = set(columns)
    available_requested = [col for col in requested_candidates if col in col_set]
    missing_requested = [col for col in requested_candidates if col not in col_set]
    depth_col = next((col for col in ["TVD", "DEPT", "MD"] if col in col_set), "")
    has_required_basic = all(col in col_set for col in ["TIME", "X", "Y"]) and bool(depth_col)
    if not has_required_basic:
        handling = "skip_missing_basic_columns"
    elif len(missing_requested) == 0:
        handling = "use_all_requested_features"
    else:
        handling = f"skip_missing_log_features_{'_'.join(missing_requested)}"
    return {
        "encoding": encoding,
        "columns": columns,
        "depth_col": depth_col,
        "has_time": int("TIME" in col_set),
        "has_x": int("X" in col_set),
        "has_y": int("Y" in col_set),
        "has_ac": int("AC" in col_set),
        "has_gr": int("GR" in col_set),
        "available_requested": available_requested,
        "missing_requested": missing_requested,
        "handling": handling,
    }


def load_batch_meta(result_dir: Path) -> dict[str, object]:
    return load_json_safe(result_dir / VIRTUAL_BATCH_META_FILENAME)


def load_prediction_meta(result_dir: Path) -> dict[str, object]:
    return load_json_safe(result_dir / FINAL_META_FILENAME)


def is_completed_result_dir(
    result_dir: Path,
    input_csv: Path,
    target_well_name: str,
    stage1_library_dir: Path,
    stage2_library_dir: Path,
    requested_log_features: list[str],
    strata_determination_mode: str,
    strata_determination_topk: int,
    unknown_strata_policy: str,
    first_stage_threshold: float,
    crossing_tolerance_ms: float,
    min_interval_depth: float,
    save_debug: bool,
) -> tuple[bool, str]:
    if not result_dir.exists() or not result_dir.is_dir():
        return False, "result_dir_missing"
    existing_names = {child.name for child in result_dir.iterdir()}
    if not REQUIRED_RESULT_FILES.issubset(existing_names):
        return False, "required_result_files_incomplete"

    expected_values = {
        "target_sample_csv": normalize_path_text(input_csv),
        "target_well_name": normalize_text(target_well_name),
        "stage1_library_dir": normalize_path_text(stage1_library_dir),
        "stage2_library_dir": normalize_path_text(stage2_library_dir),
        "requested_log_features": normalize_feature_key(requested_log_features),
        "strata_determination_mode": normalize_text(strata_determination_mode),
        "strata_determination_topk": int(strata_determination_topk),
        "unknown_strata_policy": normalize_text(unknown_strata_policy),
        "first_stage_threshold": "__nan__" if not math.isfinite(float(first_stage_threshold)) else float(first_stage_threshold),
        "crossing_tolerance_ms": float(crossing_tolerance_ms),
        "min_interval_depth": float(min_interval_depth),
        "save_debug": bool(save_debug),
    }

    batch_meta = load_batch_meta(result_dir)
    if batch_meta:
        for key, expected in expected_values.items():
            actual = batch_meta.get(key)
            if key.endswith("_csv") or key.endswith("_dir"):
                if normalize_path_text(actual) != str(expected):
                    return False, f"batch_meta_mismatch_{key}"
            elif isinstance(expected, str):
                if normalize_text(actual) != expected:
                    return False, f"batch_meta_mismatch_{key}"
            else:
                if actual != expected:
                    return False, f"batch_meta_mismatch_{key}"
        return True, "batch_meta_matched"

    prediction_meta = load_prediction_meta(result_dir)
    if not prediction_meta:
        return False, "prediction_meta_missing"
    if normalize_path_text(prediction_meta.get("target_sample_csv", "")) != expected_values["target_sample_csv"]:
        return False, "prediction_meta_mismatch_target_sample_csv"
    if normalize_text(prediction_meta.get("target_well_name", "")) != expected_values["target_well_name"]:
        return False, "prediction_meta_mismatch_target_well_name"
    if normalize_path_text(prediction_meta.get("stage1_library_dir", "")) != expected_values["stage1_library_dir"]:
        return False, "prediction_meta_mismatch_stage1_library_dir"
    if normalize_path_text(prediction_meta.get("stage2_library_dir", "")) != expected_values["stage2_library_dir"]:
        return False, "prediction_meta_mismatch_stage2_library_dir"
    return True, "prediction_meta_matched"


def write_summary_csv(summary_rows: list[dict[str, object]], save_path: Path) -> None:
    df = pd.DataFrame(summary_rows)
    if not df.empty:
        preferred_cols = [
            "UnitID",
            "BlockX",
            "BlockY",
            "WellName",
            "SourceGroup",
            "PredictMethod",
            "VirtualWellDir",
            "CsvPath",
            "IndexCsvPath",
            "StrataCsvPath",
            "ResultDir",
            "HeaderEncoding",
            "DepthColumn",
            "HasTIME",
            "HasX",
            "HasY",
            "HasAC",
            "HasGR",
            "RequestedFeatureCandidates",
            "UsedLogFeatures",
            "MissingRequestedFeatures",
            "MissingInputs",
            "FeatureHandling",
            "RunStatus",
            "CompletionCheck",
            "ReturnCode",
            "ElapsedSeconds",
            "Note",
        ]
        ordered_cols = [col for col in preferred_cols if col in df.columns] + [col for col in df.columns if col not in preferred_cols]
        df = df[ordered_cols]
    last_error: Exception | None = None
    for _ in range(3):
        try:
            save_path.parent.mkdir(parents=True, exist_ok=True)
            df.to_csv(save_path, index=False, encoding="utf-8-sig")
            return
        except Exception as exc:
            last_error = exc
            time.sleep(2.0)
    if last_error is not None:
        raise last_error


def write_incomplete_input_csv(summary_rows: list[dict[str, object]], save_path: Path) -> None:
    df = pd.DataFrame(summary_rows)
    if not df.empty:
        df = df[df["RunStatus"].astype(str).eq("skipped_missing_virtual_inputs")].copy()
    save_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(save_path, index=False, encoding="utf-8-sig")


def write_missing_feature_docx(summary_rows: list[dict[str, object]], save_path: Path) -> None:
    doc = Document()
    doc.add_heading("虚拟测井 AC/GR 属性缺失情况", level=1)
    doc.add_paragraph(f"time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    missing_rows = [
        row
        for row in summary_rows
        if str(row.get("FeatureHandling", "")).startswith("skip_missing_log_features")
        or str(row.get("FeatureHandling", "")).startswith("skip_missing_basic_columns")
    ]
    doc.add_paragraph(f"total_units: {len(summary_rows)}")
    doc.add_paragraph(f"special_handling_units: {len(missing_rows)}")
    if not missing_rows:
        doc.add_paragraph("all virtual well inputs contain AC and GR and passed the basic column check.")
    else:
        table = doc.add_table(rows=1, cols=9)
        headers = [
            "UnitID",
            "WellName",
            "HasAC",
            "HasGR",
            "UsedFeatures",
            "Handling",
            "RunStatus",
            "MissingInputs",
            "Note",
        ]
        for idx, header in enumerate(headers):
            table.rows[0].cells[idx].text = header
        for row in missing_rows:
            cells = table.add_row().cells
            cells[0].text = str(row.get("UnitID", ""))
            cells[1].text = str(row.get("WellName", ""))
            cells[2].text = str(row.get("HasAC", ""))
            cells[3].text = str(row.get("HasGR", ""))
            cells[4].text = str(row.get("UsedLogFeatures", ""))
            cells[5].text = str(row.get("FeatureHandling", ""))
            cells[6].text = str(row.get("RunStatus", ""))
            cells[7].text = str(row.get("MissingInputs", ""))
            cells[8].text = str(row.get("Note", ""))[:240]
    save_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(save_path))


def build_single_run_command(
    python_executable: Path,
    single_script_path: Path,
    unit_context: dict[str, object],
    output_root: Path,
    docx_path: Path,
    layer_dir: Path,
    stage1_library_dir: Path,
    stage2_library_dir: Path,
    requested_log_features: list[str],
    strata_determination_mode: str,
    strata_determination_topk: int,
    unknown_strata_policy: str,
    first_stage_threshold: float,
    crossing_tolerance_ms: float,
    min_interval_depth: float,
    save_debug: bool,
) -> tuple[list[str], str, Path]:
    unit_id = str(unit_context["UnitID"])
    well_name = str(unit_context["VirtualWellName"])
    csv_path = Path(str(unit_context["AroundCsv"]))
    exp_id = sanitize(unit_id) or unit_id
    result_dir = output_root / exp_id
    command = [
        str(python_executable),
        "-X",
        "utf8",
        str(single_script_path),
        "--exp-id",
        exp_id,
        "--target-sample-csv",
        str(csv_path),
        "--target-well-name",
        well_name,
        "--layer-dir",
        str(layer_dir),
        "--stage1-library-dir",
        str(stage1_library_dir),
        "--stage2-library-dir",
        str(stage2_library_dir),
        "--output-root",
        str(output_root),
        "--docx-path",
        str(docx_path),
        "--requested-log-features-json",
        json.dumps(requested_log_features, ensure_ascii=False),
        "--strata-determination-mode",
        str(strata_determination_mode),
        "--strata-determination-topk",
        str(int(strata_determination_topk)),
        "--unknown-strata-policy",
        str(unknown_strata_policy),
        "--crossing-tolerance-ms",
        str(float(crossing_tolerance_ms)),
        "--min-interval-depth",
        str(float(min_interval_depth)),
        "--save-debug",
        str(int(bool(save_debug))),
    ]
    if math.isfinite(float(first_stage_threshold)):
        command.extend(["--first-stage-threshold", str(float(first_stage_threshold))])
    return command, well_name, result_dir


def add_front_columns(df: pd.DataFrame, front_cols: list[str]) -> pd.DataFrame:
    ordered_front = [col for col in front_cols if col in df.columns]
    ordered_rest = [col for col in df.columns if col not in ordered_front]
    return df[ordered_front + ordered_rest]


def enrich_prediction_csv(path: Path, column_values: dict[str, object], front_cols: list[str]) -> None:
    if not path.exists():
        return
    df, _ = read_csv_flexible(path)
    for col, value in column_values.items():
        df[col] = value
    df = add_front_columns(df, front_cols)
    df.to_csv(path, index=False, encoding="utf-8-sig")


def write_virtual_batch_meta(
    result_dir: Path,
    unit_context: dict[str, object],
    requested_log_features: list[str],
    stage1_library_dir: Path,
    stage2_library_dir: Path,
    strata_determination_mode: str,
    strata_determination_topk: int,
    unknown_strata_policy: str,
    first_stage_threshold: float,
    crossing_tolerance_ms: float,
    min_interval_depth: float,
    save_debug: bool,
) -> Path:
    payload = {
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "unit_id": str(unit_context["UnitID"]),
        "virtual_well_name": str(unit_context["VirtualWellName"]),
        "block_x": int(unit_context["BlockX"]),
        "block_y": int(unit_context["BlockY"]),
        "predict_method": str(unit_context["PredictMethod"]),
        "virtual_well_dir": str(unit_context["UnitDir"]),
        "target_sample_csv": normalize_path_text(unit_context["AroundCsv"]),
        "target_well_name": str(unit_context["VirtualWellName"]),
        "stage1_library_dir": normalize_path_text(stage1_library_dir),
        "stage2_library_dir": normalize_path_text(stage2_library_dir),
        "requested_log_features": normalize_feature_key(requested_log_features),
        "strata_determination_mode": str(strata_determination_mode),
        "strata_determination_topk": int(strata_determination_topk),
        "unknown_strata_policy": str(unknown_strata_policy),
        "first_stage_threshold": "__nan__" if not math.isfinite(float(first_stage_threshold)) else float(first_stage_threshold),
        "crossing_tolerance_ms": float(crossing_tolerance_ms),
        "min_interval_depth": float(min_interval_depth),
        "save_debug": bool(save_debug),
    }
    save_path = result_dir / VIRTUAL_BATCH_META_FILENAME
    save_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return save_path


def enrich_virtual_prediction_outputs(
    result_dir: Path,
    unit_context: dict[str, object],
    requested_log_features: list[str],
    stage1_library_dir: Path,
    stage2_library_dir: Path,
    strata_determination_mode: str,
    strata_determination_topk: int,
    unknown_strata_policy: str,
    first_stage_threshold: float,
    crossing_tolerance_ms: float,
    min_interval_depth: float,
    save_debug: bool,
) -> None:
    common_columns = {
        "UnitID": str(unit_context["UnitID"]),
        "BlockX": int(unit_context["BlockX"]),
        "BlockY": int(unit_context["BlockY"]),
        "VirtualWellName": str(unit_context["VirtualWellName"]),
        "SourceKind": "virtual",
        "SourceName": str(unit_context["VirtualWellName"]),
        "PredictMethod": str(unit_context["PredictMethod"]),
    }
    enrich_prediction_csv(
        result_dir / FINAL_LOG_FILENAME,
        common_columns,
        ["UnitID", "BlockX", "BlockY", "VirtualWellName", "SourceKind", "SourceName", "PredictMethod", "WellName"],
    )
    enrich_prediction_csv(
        result_dir / FINAL_POINTS_FILENAME,
        common_columns,
        ["UnitID", "BlockX", "BlockY", "VirtualWellName", "SourceKind", "SourceName", "PredictMethod", "WellName"],
    )
    enrich_prediction_csv(
        result_dir / FINAL_SEGMENTS_FILENAME,
        common_columns,
        ["UnitID", "BlockX", "BlockY", "VirtualWellName", "SourceKind", "SourceName", "PredictMethod", "WellName"],
    )
    enrich_prediction_csv(
        result_dir / FINAL_STRATA_SEGMENTATION_FILENAME,
        {
            "UnitID": str(unit_context["UnitID"]),
            "BlockX": int(unit_context["BlockX"]),
            "BlockY": int(unit_context["BlockY"]),
            "VirtualWellName": str(unit_context["VirtualWellName"]),
            "PredictMethod": str(unit_context["PredictMethod"]),
        },
        ["UnitID", "BlockX", "BlockY", "VirtualWellName", "PredictMethod", "WellName"],
    )

    prediction_meta_path = result_dir / FINAL_META_FILENAME
    prediction_meta = load_prediction_meta(result_dir)
    if prediction_meta_path.exists():
        prediction_meta.update(
            {
                "source_kind": "virtual",
                "unit_id": str(unit_context["UnitID"]),
                "block_x": int(unit_context["BlockX"]),
                "block_y": int(unit_context["BlockY"]),
                "virtual_well_name": str(unit_context["VirtualWellName"]),
                "predict_method": str(unit_context["PredictMethod"]),
                "virtual_well_dir": str(unit_context["UnitDir"]),
                "virtual_well_index_csv": str(unit_context["IndexCsv"]),
                "virtual_well_strata_segmentation_csv": str(unit_context["StrataCsv"]),
                "virtual_well_input_csv": str(unit_context["AroundCsv"]),
                "requested_log_features": [str(item).strip() for item in requested_log_features if str(item).strip()],
            }
        )
        prediction_meta_path.write_text(json.dumps(prediction_meta, ensure_ascii=False, indent=2), encoding="utf-8")

    batch_meta_path = write_virtual_batch_meta(
        result_dir=result_dir,
        unit_context=unit_context,
        requested_log_features=requested_log_features,
        stage1_library_dir=stage1_library_dir,
        stage2_library_dir=stage2_library_dir,
        strata_determination_mode=strata_determination_mode,
        strata_determination_topk=strata_determination_topk,
        unknown_strata_policy=unknown_strata_policy,
        first_stage_threshold=first_stage_threshold,
        crossing_tolerance_ms=crossing_tolerance_ms,
        min_interval_depth=min_interval_depth,
        save_debug=save_debug,
    )
    prediction_meta = load_prediction_meta(result_dir)
    if prediction_meta:
        prediction_meta["virtual_batch_meta_json"] = str(batch_meta_path)
        prediction_meta_path.write_text(json.dumps(prediction_meta, ensure_ascii=False, indent=2), encoding="utf-8")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("--virtual-root", default=str(DEFAULT_VIRTUAL_ROOT))
    parser.add_argument("--single-script-path", default=str(DEFAULT_SINGLE_SCRIPT_PATH))
    parser.add_argument("--layer-dir", default=str(DEFAULT_LAYER_DIR))
    parser.add_argument("--stage1-library-dir", default=str(DEFAULT_STAGE1_LIBRARY_DIR))
    parser.add_argument("--stage2-library-dir", default=str(DEFAULT_STAGE2_LIBRARY_DIR))
    parser.add_argument("--output-root", default=str(DEFAULT_OUTPUT_ROOT))
    parser.add_argument("--docx-path", default=str(DEFAULT_DOCX_PATH))
    parser.add_argument(
        "--requested-log-features-json",
        default=json.dumps(DEFAULT_LOG_FEATURES, ensure_ascii=False),
    )
    parser.add_argument(
        "--include-unit-ids-json",
        default="[]",
        help="Optional list of unit ids for subset run.",
    )
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--skip-existing", type=int, default=1)
    parser.add_argument("--python-executable", default=sys.executable)
    parser.add_argument(
        "--strata-determination-mode",
        default="always_auto",
        choices=["use_csv", "auto_if_missing", "always_auto"],
    )
    parser.add_argument("--strata-determination-topk", type=int, default=2)
    parser.add_argument("--unknown-strata-policy", default="skip", choices=["skip", "error"])
    parser.add_argument("--first-stage-threshold", type=float, default=float("nan"))
    parser.add_argument("--crossing-tolerance-ms", type=float, default=8.0)
    parser.add_argument("--min-interval-depth", type=float, default=0.5)
    parser.add_argument("--save-debug", type=int, default=0)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    virtual_root = Path(args.virtual_root)
    layer_dir = Path(args.layer_dir)
    stage1_library_dir = Path(args.stage1_library_dir)
    stage2_library_dir = Path(args.stage2_library_dir)
    output_root = Path(args.output_root)
    docx_path = Path(args.docx_path)
    python_executable = Path(args.python_executable)
    single_script_path = Path(args.single_script_path)
    if not single_script_path.exists():
        raise FileNotFoundError(f"Single-run prediction script not found: {single_script_path}")
    if not layer_dir.exists():
        raise FileNotFoundError(f"Layer directory not found: {layer_dir}")
    if not stage1_library_dir.exists():
        raise FileNotFoundError(f"Stage1 library dir not found: {stage1_library_dir}")
    if not stage2_library_dir.exists():
        raise FileNotFoundError(f"Stage2 library dir not found: {stage2_library_dir}")
    output_root.mkdir(parents=True, exist_ok=True)

    requested_candidates = parse_json_list(args.requested_log_features_json)
    if not requested_candidates:
        raise ValueError("requested_log_features cannot be empty")
    include_unit_ids = parse_json_list(args.include_unit_ids_json)
    unit_dirs = discover_virtual_well_units(
        virtual_root=virtual_root,
        include_unit_ids=include_unit_ids,
        limit=int(args.limit),
    )
    if not unit_dirs:
        raise ValueError("No virtual well unit directories were discovered for batch prediction")

    summary_rows: list[dict[str, object]] = []
    summary_csv_path = output_root / SUMMARY_CSV_NAME
    summary_json_path = output_root / SUMMARY_JSON_NAME
    missing_docx_path = output_root / MISSING_FEATURE_DOCX_NAME
    incomplete_input_csv_path = output_root / INCOMPLETE_INPUT_CSV_NAME

    total = len(unit_dirs)
    for idx, unit_dir in enumerate(unit_dirs, start=1):
        unit_context = load_virtual_well_context(unit_dir)
        well_name = str(unit_context["VirtualWellName"])
        unit_id = str(unit_context["UnitID"])
        result_dir = output_root / (sanitize(unit_id) or unit_id)
        summary_row: dict[str, object] = {
            "UnitID": unit_id,
            "BlockX": int(unit_context["BlockX"]),
            "BlockY": int(unit_context["BlockY"]),
            "WellName": well_name,
            "SourceGroup": "virtual",
            "PredictMethod": str(unit_context["PredictMethod"]),
            "VirtualWellDir": str(unit_context["UnitDir"]),
            "CsvPath": str(unit_context["AroundCsv"]),
            "IndexCsvPath": str(unit_context["IndexCsv"]),
            "StrataCsvPath": str(unit_context["StrataCsv"]),
            "ResultDir": str(result_dir),
            "HeaderEncoding": "",
            "DepthColumn": "",
            "HasTIME": "",
            "HasX": "",
            "HasY": "",
            "HasAC": "",
            "HasGR": "",
            "RequestedFeatureCandidates": ",".join(requested_candidates),
            "UsedLogFeatures": "",
            "MissingRequestedFeatures": "",
            "MissingInputs": ",".join(unit_context["MissingInputs"]),
            "FeatureHandling": "",
            "RunStatus": "",
            "CompletionCheck": "",
            "ReturnCode": "",
            "ElapsedSeconds": "",
            "Note": "",
        }

        if unit_context["MissingInputs"]:
            summary_row["RunStatus"] = "skipped_missing_virtual_inputs"
            summary_row["Note"] = "Required virtual well input files are missing."
            summary_rows.append(summary_row)
            write_summary_csv(summary_rows, summary_csv_path)
            continue

        feature_status = detect_feature_status(Path(str(unit_context["AroundCsv"])), requested_candidates=requested_candidates)
        summary_row.update(
            {
                "HeaderEncoding": feature_status["encoding"],
                "DepthColumn": feature_status["depth_col"],
                "HasTIME": feature_status["has_time"],
                "HasX": feature_status["has_x"],
                "HasY": feature_status["has_y"],
                "HasAC": feature_status["has_ac"],
                "HasGR": feature_status["has_gr"],
                "UsedLogFeatures": ",".join(feature_status["available_requested"]),
                "MissingRequestedFeatures": ",".join(feature_status["missing_requested"]),
                "FeatureHandling": feature_status["handling"],
            }
        )

        if feature_status["handling"] == "skip_missing_basic_columns":
            summary_row["RunStatus"] = "skipped_missing_basic_columns"
            summary_row["Note"] = "Required columns missing. Need TIME/X/Y and one depth column among TVD/DEPT/MD."
            summary_rows.append(summary_row)
            write_summary_csv(summary_rows, summary_csv_path)
            continue

        if feature_status["handling"].startswith("skip_missing_log_features"):
            summary_row["RunStatus"] = "skipped_missing_log_features"
            summary_row["Note"] = "Missing AC or GR. Prediction skipped for reliability."
            summary_rows.append(summary_row)
            write_summary_csv(summary_rows, summary_csv_path)
            print(
                f"[{idx}/{total}] skip missing log features: {unit_id} | missing={feature_status['missing_requested']}",
                flush=True,
            )
            continue

        if bool(int(args.skip_existing)):
            completed, completion_reason = is_completed_result_dir(
                result_dir=result_dir,
                input_csv=Path(str(unit_context["AroundCsv"])),
                target_well_name=well_name,
                stage1_library_dir=stage1_library_dir,
                stage2_library_dir=stage2_library_dir,
                requested_log_features=requested_candidates,
                strata_determination_mode=str(args.strata_determination_mode),
                strata_determination_topk=int(args.strata_determination_topk),
                unknown_strata_policy=str(args.unknown_strata_policy),
                first_stage_threshold=float(args.first_stage_threshold),
                crossing_tolerance_ms=float(args.crossing_tolerance_ms),
                min_interval_depth=float(args.min_interval_depth),
                save_debug=bool(int(args.save_debug)),
            )
            summary_row["CompletionCheck"] = completion_reason
            if completed:
                summary_row["RunStatus"] = "skipped_existing"
                summary_row["Note"] = "Required result files already exist and match current batch configuration."
                summary_rows.append(summary_row)
                write_summary_csv(summary_rows, summary_csv_path)
                print(f"[{idx}/{total}] skip existing: {unit_id} -> {result_dir}", flush=True)
                continue

        command, _, result_dir = build_single_run_command(
            python_executable=python_executable,
            single_script_path=single_script_path,
            unit_context=unit_context,
            output_root=output_root,
            docx_path=docx_path,
            layer_dir=layer_dir,
            stage1_library_dir=stage1_library_dir,
            stage2_library_dir=stage2_library_dir,
            requested_log_features=requested_candidates,
            strata_determination_mode=str(args.strata_determination_mode),
            strata_determination_topk=int(args.strata_determination_topk),
            unknown_strata_policy=str(args.unknown_strata_policy),
            first_stage_threshold=float(args.first_stage_threshold),
            crossing_tolerance_ms=float(args.crossing_tolerance_ms),
            min_interval_depth=float(args.min_interval_depth),
            save_debug=bool(int(args.save_debug)),
        )
        print(
            f"[{idx}/{total}] run: {unit_id} | well={well_name} | available={feature_status['available_requested']} | missing={feature_status['missing_requested']}",
            flush=True,
        )
        start_time = time.time()
        completed = subprocess.run(
            command,
            cwd=str(single_script_path.parent),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
        )
        elapsed = round(time.time() - start_time, 2)
        summary_row["ReturnCode"] = int(completed.returncode)
        summary_row["ElapsedSeconds"] = elapsed

        if completed.returncode == 0:
            enrich_virtual_prediction_outputs(
                result_dir=result_dir,
                unit_context=unit_context,
                requested_log_features=requested_candidates,
                stage1_library_dir=stage1_library_dir,
                stage2_library_dir=stage2_library_dir,
                strata_determination_mode=str(args.strata_determination_mode),
                strata_determination_topk=int(args.strata_determination_topk),
                unknown_strata_policy=str(args.unknown_strata_policy),
                first_stage_threshold=float(args.first_stage_threshold),
                crossing_tolerance_ms=float(args.crossing_tolerance_ms),
                min_interval_depth=float(args.min_interval_depth),
                save_debug=bool(int(args.save_debug)),
            )
            completed_after_enrich, completion_reason = is_completed_result_dir(
                result_dir=result_dir,
                input_csv=Path(str(unit_context["AroundCsv"])),
                target_well_name=well_name,
                stage1_library_dir=stage1_library_dir,
                stage2_library_dir=stage2_library_dir,
                requested_log_features=requested_candidates,
                strata_determination_mode=str(args.strata_determination_mode),
                strata_determination_topk=int(args.strata_determination_topk),
                unknown_strata_policy=str(args.unknown_strata_policy),
                first_stage_threshold=float(args.first_stage_threshold),
                crossing_tolerance_ms=float(args.crossing_tolerance_ms),
                min_interval_depth=float(args.min_interval_depth),
                save_debug=bool(int(args.save_debug)),
            )
            summary_row["CompletionCheck"] = completion_reason
            if completed_after_enrich:
                summary_row["RunStatus"] = "success"
                stdout_tail = completed.stdout.strip().splitlines()
                summary_row["Note"] = stdout_tail[-1][:400] if stdout_tail else ""
            else:
                summary_row["RunStatus"] = "failed_incomplete_outputs"
                summary_row["Note"] = "Process returned 0 but required result files or batch metadata are incomplete."
        else:
            summary_row["RunStatus"] = "failed"
            stderr_text = (completed.stderr or "").strip()
            stdout_text = (completed.stdout or "").strip()
            summary_row["Note"] = (stderr_text[-1000:] if stderr_text else stdout_text[-1000:])

        summary_rows.append(summary_row)
        write_summary_csv(summary_rows, summary_csv_path)

    write_summary_csv(summary_rows, summary_csv_path)
    write_missing_feature_docx(summary_rows, missing_docx_path)
    write_incomplete_input_csv(summary_rows, incomplete_input_csv_path)

    summary_df = pd.DataFrame(summary_rows)
    success_count = int(summary_df["RunStatus"].astype(str).eq("success").sum()) if not summary_df.empty else 0
    skipped_existing_count = int(summary_df["RunStatus"].astype(str).eq("skipped_existing").sum()) if not summary_df.empty else 0
    skipped_missing_feature_count = int(
        summary_df["RunStatus"].astype(str).isin(
            ["skipped_missing_basic_columns", "skipped_missing_log_features", "skipped_missing_virtual_inputs"]
        ).sum()
    ) if not summary_df.empty else 0
    failed_count = int(summary_df["RunStatus"].astype(str).isin(["failed", "failed_incomplete_outputs"]).sum()) if not summary_df.empty else 0
    run_summary = {
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "workflow_root": str(WORKFLOW_ROOT),
        "single_script_path": str(single_script_path),
        "virtual_root": str(virtual_root),
        "output_root": str(output_root),
        "docx_path": str(docx_path),
        "total_units": total,
        "success_count": success_count,
        "skipped_existing_count": skipped_existing_count,
        "skipped_missing_feature_count": skipped_missing_feature_count,
        "failed_count": failed_count,
        "summary_csv": str(summary_csv_path),
        "summary_json": str(summary_json_path),
        "missing_feature_docx": str(missing_docx_path),
        "incomplete_input_csv": str(incomplete_input_csv_path),
    }
    summary_json_path.write_text(json.dumps(run_summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(run_summary, ensure_ascii=False), flush=True)
    return 0 if failed_count == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
