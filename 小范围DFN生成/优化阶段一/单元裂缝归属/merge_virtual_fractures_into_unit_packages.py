from __future__ import annotations

import argparse
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


DEFAULT_REAL_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容二\虚拟测井构建\单元裂缝归属\smoke_20260328_pkg_min"
)
DEFAULT_VIRTUAL_FRACTURE_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容二\虚拟测井构建\虚拟裂缝预测"
)
DEFAULT_VIRTUAL_WELL_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容二\虚拟测井构建\虚拟测井批量生成"
)
DEFAULT_OUTPUT_ROOT = Path(
    r"E:\项目\石油项目\断缝储\原始数据\wx数据\砂砾岩\优化阶段一\研究内容二\单元DFN构建\单元测井裂缝"
)
DEFAULT_DOCX_PATH = Path(r"D:\项目\石油开采\断缝储实验\实验记录20260328.docx")

LAYER_COLUMNS = [
    "UnitID",
    "BlockX",
    "BlockY",
    "GeoIntervalKey",
    "StrataName",
    "TopSurfaceCode",
    "BaseSurfaceCode",
    "TopDepth",
    "BaseDepth",
    "TopTime",
    "BaseTime",
    "RealPointSeedCount",
    "RealSegmentSeedCount",
    "VirtualPointSeedCount",
    "VirtualSegmentSeedCount",
    "RealSeedCount",
    "VirtualSeedCount",
    "PreferredSource",
]

SEED_COLUMNS = [
    "UnitID",
    "BlockX",
    "BlockY",
    "GeoIntervalKey",
    "StrataName",
    "TopSurfaceCode",
    "BaseSurfaceCode",
    "SeedID",
    "SourceKind",
    "SourceName",
    "SeedType",
    "CenterX",
    "CenterY",
    "CenterDepth",
    "CenterTime",
    "DepthStart",
    "DepthEnd",
    "TimeStart",
    "TimeEnd",
    "Azimuth",
    "Dip",
    "DensityWeight",
    "LengthWeight",
    "Confidence",
]

CATALOG_COLUMNS = [
    "UnitID",
    "BlockX",
    "BlockY",
    "ReliabilityClass",
    "HasRealData",
    "HasVirtualData",
    "DataMode",
    "RealLayerCount",
    "VirtualLayerCount",
    "LayerCount",
    "CoveredIntervalCount",
    "SeedCount",
    "RealSeedCount",
    "VirtualSeedCount",
    "RealPointSeedCount",
    "RealSegmentSeedCount",
    "VirtualPointSeedCount",
    "VirtualSegmentSeedCount",
    "TopDepth",
    "BaseDepth",
    "TopTime",
    "BaseTime",
    "PackageDir",
]

REAL_NUMERIC_COLUMNS = [
    "BlockX",
    "BlockY",
    "CenterX",
    "CenterY",
    "CenterDepth",
    "CenterTime",
    "DepthStart",
    "DepthEnd",
    "TimeStart",
    "TimeEnd",
    "Azimuth",
    "Dip",
    "DensityWeight",
    "LengthWeight",
    "Confidence",
    "TopDepth",
    "BaseDepth",
    "TopTime",
    "BaseTime",
    "RealPointSeedCount",
    "RealSegmentSeedCount",
    "VirtualPointSeedCount",
    "VirtualSegmentSeedCount",
    "RealSeedCount",
    "VirtualSeedCount",
]

VIRTUAL_POINT_NUMERIC_COLUMNS = [
    "UnitID",
    "BlockX",
    "BlockY",
    "GeoRangeOrder",
    "GeoDepthMin",
    "GeoDepthMax",
    "Segment_ID",
    "Point_ID_In_Segment",
    "TVD",
    "TIME",
    "X",
    "Y",
    "SegStartDepth",
    "SegEndDepth",
    "SegLength",
    "PredPointCount",
    "PointDensityMassPerLengthAllocated",
    "PointDensityMassAllocated",
    "PointAzimuth",
    "PointDip",
    "PredOrientationConfidence",
]

VIRTUAL_SEGMENT_NUMERIC_COLUMNS = [
    "UnitID",
    "BlockX",
    "BlockY",
    "GeoRangeOrder",
    "GeoDepthMin",
    "GeoDepthMax",
    "Segment_ID",
    "SegStartDepth",
    "SegEndDepth",
    "SegLength",
    "PredPointCount",
    "PredDensityStrength",
    "PredP10MassPerLength",
    "PredP10Mass",
    "PredOrientationConfidence",
    "PredAzimuth",
    "PredDip",
]

VIRTUAL_INTERVAL_NUMERIC_COLUMNS = [
    "UnitID",
    "BlockX",
    "BlockY",
    "GeoRangeOrder",
    "GeoDepthMin",
    "GeoDepthMax",
]

AROUND_NUMERIC_COLUMNS = ["TVD", "TIME", "X", "Y", "BlockX", "BlockY"]


def read_csv_utf8(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, encoding="utf-8-sig", low_memory=False)


def write_csv_utf8(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8-sig")


def ensure_numeric(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    for column in columns:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")
    return df


def ensure_string(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    for column in columns:
        if column in df.columns:
            df[column] = df[column].fillna("").astype(str)
    return df


def ensure_columns(df: pd.DataFrame, columns: list[str], default: object = np.nan) -> pd.DataFrame:
    work = df.copy()
    for column in columns:
        if column not in work.columns:
            work[column] = default
    return work


def root_unit_dir(root: Path) -> Path:
    if (root / "units").exists():
        return root / "units"
    return root


def discover_unit_dirs(root: Path) -> dict[str, Path]:
    base = root_unit_dir(root)
    if not base.exists():
        return {}
    unit_dirs: dict[str, Path] = {}
    for path in sorted(base.iterdir()):
        if not path.is_dir():
            continue
        if not re.fullmatch(r"BX\d+_BY\d+", path.name):
            continue
        unit_dirs[path.name] = path
    return unit_dirs


def interval_order_key(value: str) -> tuple[int, str]:
    text = str(value).strip()
    match = re.match(r"^(\d+)_", text)
    if match:
        return int(match.group(1)), text
    return 10**9, text


def sort_layers(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    order_values = df["GeoIntervalKey"].astype(str).map(lambda value: interval_order_key(value)[0])
    work = df.assign(_interval_order=order_values)
    work = work.sort_values(["BlockX", "BlockY", "_interval_order", "GeoIntervalKey"]).drop(columns=["_interval_order"])
    return work.reset_index(drop=True)


def sort_seeds(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    order_values = df["GeoIntervalKey"].astype(str).map(lambda value: interval_order_key(value)[0])
    seed_type_rank = df["SeedType"].astype(str).map({"segment": 0, "point": 1}).fillna(9)
    source_rank = df["SourceKind"].astype(str).map({"real": 0, "virtual": 1}).fillna(9)
    work = df.assign(_interval_order=order_values, _seed_type_rank=seed_type_rank, _source_rank=source_rank)
    work = work.sort_values(
        ["BlockX", "BlockY", "_interval_order", "_source_rank", "_seed_type_rank", "SeedID"]
    ).drop(columns=["_interval_order", "_seed_type_rank", "_source_rank"])
    return work.reset_index(drop=True)


def open_or_create_doc(docx_path: Path) -> Document:
    if docx_path.exists():
        try:
            return Document(str(docx_path))
        except PackageNotFoundError:
            pass
    doc = Document()
    doc.add_heading("实验记录 20260328", level=1)
    return doc


def append_merge_summary_to_docx(docx_path: Path, title: str, config: dict[str, object], summary: dict[str, object]) -> None:
    doc = open_or_create_doc(docx_path)
    doc.add_heading(title, level=2)
    lines = [
        "task: 单元测井裂缝 real + virtual 合并",
        f"time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"real_root: {config.get('real_root', '')}",
        f"virtual_fracture_root: {config.get('virtual_fracture_root', '')}",
        f"virtual_well_root: {config.get('virtual_well_root', '')}",
        f"output_root: {config.get('output_root', '')}",
        f"real_unit_count: {summary.get('real_unit_count', 0)}",
        f"virtual_unit_count: {summary.get('virtual_unit_count', 0)}",
        f"merged_unit_count: {summary.get('merged_unit_count', 0)}",
        f"real_only_unit_count: {summary.get('real_only_unit_count', 0)}",
        f"virtual_only_unit_count: {summary.get('virtual_only_unit_count', 0)}",
        f"real_virtual_unit_count: {summary.get('real_virtual_unit_count', 0)}",
        f"merged_seed_count: {summary.get('merged_seed_count', 0)}",
        f"merged_layer_count: {summary.get('merged_layer_count', 0)}",
        f"unit_catalog_csv: {summary.get('unit_catalog_csv', '')}",
        f"run_summary_json: {summary.get('run_summary_json', '')}",
    ]
    doc.add_paragraph("\n".join(lines))
    doc.add_paragraph("")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def load_real_unit_layers(unit_id: str, unit_dir: Path) -> pd.DataFrame:
    path = unit_dir / "unit_layers.csv"
    if not path.exists():
        return pd.DataFrame(columns=LAYER_COLUMNS)
    df = ensure_numeric(read_csv_utf8(path), REAL_NUMERIC_COLUMNS)
    df = ensure_string(df, ["GeoIntervalKey", "StrataName", "TopSurfaceCode", "BaseSurfaceCode", "PreferredSource"])
    df["UnitID"] = unit_id
    df = ensure_columns(df, LAYER_COLUMNS)
    return sort_layers(df[LAYER_COLUMNS])


def load_real_unit_seeds(unit_id: str, unit_dir: Path) -> pd.DataFrame:
    path = unit_dir / "fracture_seeds.csv"
    if not path.exists():
        return pd.DataFrame(columns=SEED_COLUMNS)
    df = ensure_numeric(read_csv_utf8(path), REAL_NUMERIC_COLUMNS)
    df = ensure_string(
        df,
        [
            "GeoIntervalKey",
            "StrataName",
            "TopSurfaceCode",
            "BaseSurfaceCode",
            "SeedID",
            "SourceKind",
            "SourceName",
            "SeedType",
        ],
    )
    df["UnitID"] = unit_id
    df = ensure_columns(df, SEED_COLUMNS)
    return sort_seeds(df[SEED_COLUMNS])


def build_around_lookup(around_csv: Path) -> pd.DataFrame:
    df = ensure_numeric(read_csv_utf8(around_csv), AROUND_NUMERIC_COLUMNS)
    df = df.dropna(subset=["TVD", "TIME", "X", "Y"]).copy()
    if df.empty:
        return pd.DataFrame(columns=["TVD", "TIME", "X", "Y"])
    df = df.sort_values("TVD").drop_duplicates(subset=["TVD"], keep="first").reset_index(drop=True)
    return df[["TVD", "TIME", "X", "Y"]]


def interpolate_curve(curve_df: pd.DataFrame, depths: pd.Series | np.ndarray, column: str) -> np.ndarray:
    depth_values = pd.to_numeric(pd.Series(depths), errors="coerce").to_numpy(dtype=float)
    result = np.full(len(depth_values), np.nan, dtype=float)
    if curve_df.empty or column not in curve_df.columns:
        return result
    curve_depth = curve_df["TVD"].to_numpy(dtype=float)
    curve_values = pd.to_numeric(curve_df[column], errors="coerce").to_numpy(dtype=float)
    valid_curve = np.isfinite(curve_depth) & np.isfinite(curve_values)
    if valid_curve.sum() == 0:
        return result
    curve_depth = curve_depth[valid_curve]
    curve_values = curve_values[valid_curve]
    valid_depth = np.isfinite(depth_values)
    if valid_depth.any():
        result[valid_depth] = np.interp(depth_values[valid_depth], curve_depth, curve_values)
        low = depth_values[valid_depth] < curve_depth.min()
        high = depth_values[valid_depth] > curve_depth.max()
        invalid = low | high
        if invalid.any():
            valid_idx = np.flatnonzero(valid_depth)
            result[valid_idx[invalid]] = np.nan
    return result


def build_virtual_point_seeds(unit_id: str, points_path: Path) -> pd.DataFrame:
    if not points_path.exists():
        return pd.DataFrame(columns=SEED_COLUMNS)
    points = ensure_numeric(read_csv_utf8(points_path), VIRTUAL_POINT_NUMERIC_COLUMNS)
    points = ensure_string(
        points,
        [
            "UnitID",
            "VirtualWellName",
            "SourceKind",
            "SourceName",
            "PredictMethod",
            "WellName",
            "StrataName",
            "GeoSegmentID",
            "GeoIntervalKey",
            "TopSurfaceCode",
            "BaseSurfaceCode",
        ],
    )
    if points.empty:
        return pd.DataFrame(columns=SEED_COLUMNS)

    source_name = points["SourceName"].replace("", np.nan).fillna(points["VirtualWellName"]).fillna(points["WellName"])
    segment_id = points["Segment_ID"].fillna(-1).astype(int).astype(str)
    point_id = points["Point_ID_In_Segment"].fillna(-1).astype(int).astype(str)
    seed_id = source_name.astype(str) + "::" + points["GeoIntervalKey"].astype(str) + "::SEG" + segment_id + "::PT" + point_id

    seeds = pd.DataFrame(
        {
            "UnitID": unit_id,
            "BlockX": points["BlockX"],
            "BlockY": points["BlockY"],
            "GeoIntervalKey": points["GeoIntervalKey"],
            "StrataName": points["StrataName"],
            "TopSurfaceCode": points["TopSurfaceCode"],
            "BaseSurfaceCode": points["BaseSurfaceCode"],
            "SeedID": seed_id,
            "SourceKind": "virtual",
            "SourceName": source_name.astype(str),
            "SeedType": "point",
            "CenterX": points["X"],
            "CenterY": points["Y"],
            "CenterDepth": points["TVD"],
            "CenterTime": points["TIME"],
            "DepthStart": points["SegStartDepth"],
            "DepthEnd": points["SegEndDepth"],
            "TimeStart": points["TIME"],
            "TimeEnd": points["TIME"],
            "Azimuth": points["PointAzimuth"],
            "Dip": points["PointDip"],
            "DensityWeight": points["PointDensityMassAllocated"],
            "LengthWeight": points["PointDensityMassPerLengthAllocated"],
            "Confidence": points["PredOrientationConfidence"],
        }
    )
    seeds = ensure_columns(seeds, SEED_COLUMNS)
    return sort_seeds(seeds[SEED_COLUMNS])


def build_virtual_segment_seeds(unit_id: str, segments_path: Path, curve_df: pd.DataFrame) -> pd.DataFrame:
    if not segments_path.exists():
        return pd.DataFrame(columns=SEED_COLUMNS)
    segments = ensure_numeric(read_csv_utf8(segments_path), VIRTUAL_SEGMENT_NUMERIC_COLUMNS)
    segments = ensure_string(
        segments,
        [
            "UnitID",
            "VirtualWellName",
            "SourceKind",
            "SourceName",
            "PredictMethod",
            "WellName",
            "StrataName",
            "GeoSegmentID",
            "GeoIntervalKey",
            "TopSurfaceCode",
            "BaseSurfaceCode",
        ],
    )
    if segments.empty:
        return pd.DataFrame(columns=SEED_COLUMNS)

    source_name = segments["SourceName"].replace("", np.nan).fillna(segments["VirtualWellName"]).fillna(segments["WellName"])
    segment_id = segments["Segment_ID"].fillna(-1).astype(int).astype(str)
    mid_depth = (segments["SegStartDepth"] + segments["SegEndDepth"]) / 2.0
    center_time = interpolate_curve(curve_df, mid_depth, "TIME")
    center_x = interpolate_curve(curve_df, mid_depth, "X")
    center_y = interpolate_curve(curve_df, mid_depth, "Y")
    time_start = interpolate_curve(curve_df, segments["SegStartDepth"], "TIME")
    time_end = interpolate_curve(curve_df, segments["SegEndDepth"], "TIME")
    seed_id = source_name.astype(str) + "::" + segments["GeoIntervalKey"].astype(str) + "::SEG" + segment_id

    seeds = pd.DataFrame(
        {
            "UnitID": unit_id,
            "BlockX": segments["BlockX"],
            "BlockY": segments["BlockY"],
            "GeoIntervalKey": segments["GeoIntervalKey"],
            "StrataName": segments["StrataName"],
            "TopSurfaceCode": segments["TopSurfaceCode"],
            "BaseSurfaceCode": segments["BaseSurfaceCode"],
            "SeedID": seed_id,
            "SourceKind": "virtual",
            "SourceName": source_name.astype(str),
            "SeedType": "segment",
            "CenterX": center_x,
            "CenterY": center_y,
            "CenterDepth": mid_depth,
            "CenterTime": center_time,
            "DepthStart": segments["SegStartDepth"],
            "DepthEnd": segments["SegEndDepth"],
            "TimeStart": time_start,
            "TimeEnd": time_end,
            "Azimuth": segments["PredAzimuth"],
            "Dip": segments["PredDip"],
            "DensityWeight": segments["PredP10Mass"],
            "LengthWeight": segments["SegLength"],
            "Confidence": segments["PredOrientationConfidence"],
        }
    )
    seeds = ensure_columns(seeds, SEED_COLUMNS)
    return sort_seeds(seeds[SEED_COLUMNS])


def build_virtual_layers_from_strata(
    unit_id: str,
    strata_path: Path,
    curve_df: pd.DataFrame,
    block_x: int | None,
    block_y: int | None,
) -> pd.DataFrame:
    if not strata_path.exists():
        return pd.DataFrame(columns=LAYER_COLUMNS)
    strata = ensure_numeric(read_csv_utf8(strata_path), VIRTUAL_INTERVAL_NUMERIC_COLUMNS)
    strata = ensure_string(
        strata,
        ["UnitID", "VirtualWellName", "GeoSegmentID", "GeoIntervalKey", "StrataName", "TopSurfaceCode", "BaseSurfaceCode"],
    )
    if strata.empty:
        return pd.DataFrame(columns=LAYER_COLUMNS)

    top_time = interpolate_curve(curve_df, strata["GeoDepthMin"], "TIME")
    base_time = interpolate_curve(curve_df, strata["GeoDepthMax"], "TIME")

    layers = pd.DataFrame(
        {
            "UnitID": unit_id,
            "BlockX": strata["BlockX"].fillna(block_x),
            "BlockY": strata["BlockY"].fillna(block_y),
            "GeoIntervalKey": strata["GeoIntervalKey"],
            "StrataName": strata["StrataName"],
            "TopSurfaceCode": strata["TopSurfaceCode"],
            "BaseSurfaceCode": strata["BaseSurfaceCode"],
            "TopDepth": strata["GeoDepthMin"],
            "BaseDepth": strata["GeoDepthMax"],
            "TopTime": top_time,
            "BaseTime": base_time,
            "RealPointSeedCount": 0,
            "RealSegmentSeedCount": 0,
            "VirtualPointSeedCount": 0,
            "VirtualSegmentSeedCount": 0,
            "RealSeedCount": 0,
            "VirtualSeedCount": 0,
            "PreferredSource": "seismic_fill",
        }
    )
    layers = ensure_columns(layers, LAYER_COLUMNS)
    return sort_layers(layers[LAYER_COLUMNS])


def build_virtual_layers_from_seeds(unit_id: str, virtual_seeds_df: pd.DataFrame) -> pd.DataFrame:
    if virtual_seeds_df.empty:
        return pd.DataFrame(columns=LAYER_COLUMNS)
    group_cols = ["UnitID", "BlockX", "BlockY", "GeoIntervalKey", "StrataName", "TopSurfaceCode", "BaseSurfaceCode"]
    grouped = (
        virtual_seeds_df.groupby(group_cols, as_index=False)
        .agg(
            TopDepth=("DepthStart", "min"),
            BaseDepth=("DepthEnd", "max"),
            TopTime=("TimeStart", "min"),
            BaseTime=("TimeEnd", "max"),
        )
    )
    grouped["RealPointSeedCount"] = 0
    grouped["RealSegmentSeedCount"] = 0
    grouped["VirtualPointSeedCount"] = 0
    grouped["VirtualSegmentSeedCount"] = 0
    grouped["RealSeedCount"] = 0
    grouped["VirtualSeedCount"] = 0
    grouped["PreferredSource"] = "seismic_fill"
    grouped["UnitID"] = unit_id
    grouped = ensure_columns(grouped, LAYER_COLUMNS)
    return sort_layers(grouped[LAYER_COLUMNS])


def merge_layer_rows(real_layers_df: pd.DataFrame, virtual_layers_df: pd.DataFrame) -> pd.DataFrame:
    if real_layers_df.empty and virtual_layers_df.empty:
        return pd.DataFrame(columns=LAYER_COLUMNS)
    if real_layers_df.empty:
        return sort_layers(virtual_layers_df[LAYER_COLUMNS].copy())
    if virtual_layers_df.empty:
        return sort_layers(real_layers_df[LAYER_COLUMNS].copy())

    merged = real_layers_df.copy().reset_index(drop=True)
    merged = ensure_columns(merged, LAYER_COLUMNS)
    virtual_lookup = {str(row["GeoIntervalKey"]): row for row in virtual_layers_df.to_dict(orient="records")}

    for index, row in merged.iterrows():
        virtual_row = virtual_lookup.get(str(row["GeoIntervalKey"]))
        if virtual_row is None:
            continue
        for column in ["StrataName", "TopSurfaceCode", "BaseSurfaceCode", "TopDepth", "BaseDepth", "TopTime", "BaseTime"]:
            current = merged.at[index, column]
            candidate = virtual_row.get(column)
            current_blank = (isinstance(current, str) and current.strip() == "") or pd.isna(current)
            candidate_blank = (isinstance(candidate, str) and str(candidate).strip() == "") or pd.isna(candidate)
            if current_blank and not candidate_blank:
                merged.at[index, column] = candidate

    existing_keys = set(merged["GeoIntervalKey"].astype(str).tolist())
    append_rows = virtual_layers_df[~virtual_layers_df["GeoIntervalKey"].astype(str).isin(existing_keys)].copy()
    if not append_rows.empty:
        merged = pd.concat([merged, append_rows[LAYER_COLUMNS]], ignore_index=True)
    return sort_layers(merged[LAYER_COLUMNS])


def apply_seed_counts(layer_df: pd.DataFrame, seed_df: pd.DataFrame) -> pd.DataFrame:
    layers = ensure_columns(layer_df.copy(), LAYER_COLUMNS)
    for column in [
        "RealPointSeedCount",
        "RealSegmentSeedCount",
        "VirtualPointSeedCount",
        "VirtualSegmentSeedCount",
        "RealSeedCount",
        "VirtualSeedCount",
    ]:
        layers[column] = 0

    if seed_df.empty:
        layers["PreferredSource"] = "seismic_fill"
        return sort_layers(layers[LAYER_COLUMNS])

    stats = (
        seed_df.assign(
            RealPointSeedCount=((seed_df["SourceKind"] == "real") & (seed_df["SeedType"] == "point")).astype(int),
            RealSegmentSeedCount=((seed_df["SourceKind"] == "real") & (seed_df["SeedType"] == "segment")).astype(int),
            VirtualPointSeedCount=((seed_df["SourceKind"] == "virtual") & (seed_df["SeedType"] == "point")).astype(int),
            VirtualSegmentSeedCount=((seed_df["SourceKind"] == "virtual") & (seed_df["SeedType"] == "segment")).astype(int),
        )
        .groupby("GeoIntervalKey", as_index=False)
        .agg(
            RealPointSeedCount=("RealPointSeedCount", "sum"),
            RealSegmentSeedCount=("RealSegmentSeedCount", "sum"),
            VirtualPointSeedCount=("VirtualPointSeedCount", "sum"),
            VirtualSegmentSeedCount=("VirtualSegmentSeedCount", "sum"),
        )
    )
    stats["RealSeedCount"] = stats["RealPointSeedCount"] + stats["RealSegmentSeedCount"]
    stats["VirtualSeedCount"] = stats["VirtualPointSeedCount"] + stats["VirtualSegmentSeedCount"]
    layers = layers.drop(
        columns=[
            "RealPointSeedCount",
            "RealSegmentSeedCount",
            "VirtualPointSeedCount",
            "VirtualSegmentSeedCount",
            "RealSeedCount",
            "VirtualSeedCount",
            "PreferredSource",
        ],
        errors="ignore",
    )
    layers = layers.merge(stats, on="GeoIntervalKey", how="left")
    for column in [
        "RealPointSeedCount",
        "RealSegmentSeedCount",
        "VirtualPointSeedCount",
        "VirtualSegmentSeedCount",
        "RealSeedCount",
        "VirtualSeedCount",
    ]:
        layers[column] = pd.to_numeric(layers[column], errors="coerce").fillna(0).astype(int)
    layers["PreferredSource"] = "seismic_fill"
    layers.loc[layers["VirtualSeedCount"] > 0, "PreferredSource"] = "virtual"
    layers.loc[layers["RealSeedCount"] > 0, "PreferredSource"] = "real"
    return sort_layers(layers[LAYER_COLUMNS])


def build_unit_catalog_row(
    unit_id: str,
    layer_df: pd.DataFrame,
    seed_df: pd.DataFrame,
    has_real_layer: bool,
    has_virtual_layer: bool,
) -> dict[str, object]:
    block_x = pd.to_numeric(layer_df.get("BlockX"), errors="coerce") if not layer_df.empty else pd.Series(dtype=float)
    block_y = pd.to_numeric(layer_df.get("BlockY"), errors="coerce") if not layer_df.empty else pd.Series(dtype=float)
    if layer_df.empty and not seed_df.empty:
        block_x = pd.to_numeric(seed_df.get("BlockX"), errors="coerce")
        block_y = pd.to_numeric(seed_df.get("BlockY"), errors="coerce")

    real_seed_count = int(seed_df["SourceKind"].eq("real").sum()) if not seed_df.empty else 0
    virtual_seed_count = int(seed_df["SourceKind"].eq("virtual").sum()) if not seed_df.empty else 0
    real_point_seed_count = int(((seed_df["SourceKind"] == "real") & (seed_df["SeedType"] == "point")).sum()) if not seed_df.empty else 0
    real_segment_seed_count = int(((seed_df["SourceKind"] == "real") & (seed_df["SeedType"] == "segment")).sum()) if not seed_df.empty else 0
    virtual_point_seed_count = int(((seed_df["SourceKind"] == "virtual") & (seed_df["SeedType"] == "point")).sum()) if not seed_df.empty else 0
    virtual_segment_seed_count = int(((seed_df["SourceKind"] == "virtual") & (seed_df["SeedType"] == "segment")).sum()) if not seed_df.empty else 0
    real_layer_count = int(len(layer_df[layer_df["RealSeedCount"] > 0])) if not layer_df.empty else 0
    virtual_layer_count = int(len(layer_df[layer_df["VirtualSeedCount"] > 0])) if not layer_df.empty else 0

    has_real_data = bool(has_real_layer or real_seed_count > 0)
    has_virtual_data = bool(has_virtual_layer or virtual_seed_count > 0)
    data_mode = "empty"
    if has_real_data and not has_virtual_data:
        data_mode = "real_only"
    elif not has_real_data and has_virtual_data:
        data_mode = "virtual_only"
    elif has_real_data and has_virtual_data:
        data_mode = "real_virtual"

    reliability_class = "empty"
    if has_real_data:
        reliability_class = "real_controlled"
    elif has_virtual_data:
        reliability_class = "virtual_controlled"

    top_depth = pd.to_numeric(layer_df.get("TopDepth"), errors="coerce").min() if not layer_df.empty else np.nan
    base_depth = pd.to_numeric(layer_df.get("BaseDepth"), errors="coerce").max() if not layer_df.empty else np.nan
    top_time = pd.to_numeric(layer_df.get("TopTime"), errors="coerce").min() if not layer_df.empty else np.nan
    base_time = pd.to_numeric(layer_df.get("BaseTime"), errors="coerce").max() if not layer_df.empty else np.nan

    return {
        "UnitID": unit_id,
        "BlockX": int(block_x.dropna().iloc[0]) if not block_x.dropna().empty else np.nan,
        "BlockY": int(block_y.dropna().iloc[0]) if not block_y.dropna().empty else np.nan,
        "ReliabilityClass": reliability_class,
        "HasRealData": has_real_data,
        "HasVirtualData": has_virtual_data,
        "DataMode": data_mode,
        "RealLayerCount": real_layer_count,
        "VirtualLayerCount": virtual_layer_count,
        "LayerCount": int(len(layer_df)),
        "CoveredIntervalCount": int(layer_df["GeoIntervalKey"].nunique()) if not layer_df.empty else 0,
        "SeedCount": int(len(seed_df)),
        "RealSeedCount": real_seed_count,
        "VirtualSeedCount": virtual_seed_count,
        "RealPointSeedCount": real_point_seed_count,
        "RealSegmentSeedCount": real_segment_seed_count,
        "VirtualPointSeedCount": virtual_point_seed_count,
        "VirtualSegmentSeedCount": virtual_segment_seed_count,
        "TopDepth": None if pd.isna(top_depth) else float(top_depth),
        "BaseDepth": None if pd.isna(base_depth) else float(base_depth),
        "TopTime": None if pd.isna(top_time) else float(top_time),
        "BaseTime": None if pd.isna(base_time) else float(base_time),
        "PackageDir": f"units/{unit_id}",
    }


def materialize_unit_package(output_root: Path, unit_id: str, layer_df: pd.DataFrame, seed_df: pd.DataFrame, meta: dict[str, object]) -> None:
    unit_dir = output_root / "units" / unit_id
    unit_dir.mkdir(parents=True, exist_ok=True)
    layer_to_write = sort_layers(layer_df.copy()).drop(columns=["UnitID"], errors="ignore")
    seed_to_write = sort_seeds(seed_df.copy()).drop(columns=["UnitID"], errors="ignore")
    write_csv_utf8(layer_to_write, unit_dir / "unit_layers.csv")
    write_csv_utf8(seed_to_write, unit_dir / "fracture_seeds.csv")
    (unit_dir / "unit_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def load_virtual_unit_bundle(unit_id: str, virtual_dir: Path, virtual_well_root: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    points_path = virtual_dir / "final_fracture_points.csv"
    segments_path = virtual_dir / "final_fracture_segments.csv"
    strata_path = virtual_dir / "final_strata_segmentation.csv"
    around_path = virtual_well_root / unit_id / "virtual_well_around_data.csv"
    if not around_path.exists():
        raise FileNotFoundError(f"Missing virtual well around data for {unit_id}: {around_path}")

    curve_df = build_around_lookup(around_path)
    point_seeds = build_virtual_point_seeds(unit_id, points_path)
    segment_seeds = build_virtual_segment_seeds(unit_id, segments_path, curve_df)
    if point_seeds.empty and segment_seeds.empty:
        virtual_seeds = pd.DataFrame(columns=SEED_COLUMNS)
    else:
        virtual_seeds = pd.concat([point_seeds, segment_seeds], ignore_index=True)
    virtual_seeds = ensure_columns(virtual_seeds, SEED_COLUMNS)
    virtual_seeds = sort_seeds(virtual_seeds[SEED_COLUMNS])

    block_x = None
    block_y = None
    if not virtual_seeds.empty:
        bx = pd.to_numeric(virtual_seeds["BlockX"], errors="coerce").dropna()
        by = pd.to_numeric(virtual_seeds["BlockY"], errors="coerce").dropna()
        if not bx.empty:
            block_x = int(bx.iloc[0])
        if not by.empty:
            block_y = int(by.iloc[0])

    virtual_layers = build_virtual_layers_from_strata(unit_id, strata_path, curve_df, block_x, block_y)
    if virtual_layers.empty:
        virtual_layers = build_virtual_layers_from_seeds(unit_id, virtual_seeds)
    virtual_layers = apply_seed_counts(virtual_layers, virtual_seeds)
    return virtual_layers, virtual_seeds


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Merge real and virtual fracture seeds into unit-scale DFN seed packages.")
    parser.add_argument("--real-root", type=Path, default=DEFAULT_REAL_ROOT)
    parser.add_argument("--virtual-fracture-root", type=Path, default=DEFAULT_VIRTUAL_FRACTURE_ROOT)
    parser.add_argument("--virtual-well-root", type=Path, default=DEFAULT_VIRTUAL_WELL_ROOT)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--docx-path", type=Path, default=DEFAULT_DOCX_PATH)
    parser.add_argument("--include-unit-ids-json", default="[]")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--clean-output", type=int, default=0)
    return parser.parse_args()


def parse_unit_id_filter(raw: str) -> set[str]:
    text = str(raw or "").strip()
    if not text:
        return set()
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return {str(item).strip() for item in parsed if str(item).strip()}
    except json.JSONDecodeError:
        pass
    return {item.strip() for item in text.split(",") if item.strip()}


def main() -> int:
    args = parse_args()

    real_root = Path(args.real_root)
    virtual_fracture_root = Path(args.virtual_fracture_root)
    virtual_well_root = Path(args.virtual_well_root)
    output_root = Path(args.output_root)
    docx_path = Path(args.docx_path)

    include_unit_ids = parse_unit_id_filter(str(args.include_unit_ids_json or ""))

    if int(bool(args.clean_output)):
        units_root = output_root / "units"
        if units_root.exists():
            shutil.rmtree(units_root)

    real_units = discover_unit_dirs(real_root)
    virtual_units = discover_unit_dirs(virtual_fracture_root)

    selected_unit_ids = sorted(set(real_units) | set(virtual_units))
    if include_unit_ids:
        selected_unit_ids = [unit_id for unit_id in selected_unit_ids if unit_id in include_unit_ids]
    if int(args.limit) > 0:
        selected_unit_ids = selected_unit_ids[: int(args.limit)]

    catalog_rows: list[dict[str, object]] = []
    merged_layer_count = 0
    merged_seed_count = 0
    real_only_unit_count = 0
    virtual_only_unit_count = 0
    real_virtual_unit_count = 0

    for unit_id in selected_unit_ids:
        real_layer_df = pd.DataFrame(columns=LAYER_COLUMNS)
        real_seed_df = pd.DataFrame(columns=SEED_COLUMNS)
        if unit_id in real_units:
            real_layer_df = load_real_unit_layers(unit_id, real_units[unit_id])
            real_seed_df = load_real_unit_seeds(unit_id, real_units[unit_id])

        virtual_layer_df = pd.DataFrame(columns=LAYER_COLUMNS)
        virtual_seed_df = pd.DataFrame(columns=SEED_COLUMNS)
        if unit_id in virtual_units:
            virtual_layer_df, virtual_seed_df = load_virtual_unit_bundle(unit_id, virtual_units[unit_id], virtual_well_root)

        if real_seed_df.empty and virtual_seed_df.empty:
            merged_seed_df = pd.DataFrame(columns=SEED_COLUMNS)
        else:
            merged_seed_df = pd.concat([real_seed_df, virtual_seed_df], ignore_index=True)
        merged_seed_df = ensure_columns(merged_seed_df, SEED_COLUMNS)
        merged_seed_df = sort_seeds(merged_seed_df[SEED_COLUMNS])

        merged_layer_df = merge_layer_rows(real_layer_df, virtual_layer_df)
        merged_layer_df = apply_seed_counts(merged_layer_df, merged_seed_df)

        meta = build_unit_catalog_row(
            unit_id=unit_id,
            layer_df=merged_layer_df,
            seed_df=merged_seed_df,
            has_real_layer=not real_layer_df.empty,
            has_virtual_layer=not virtual_layer_df.empty,
        )
        materialize_unit_package(output_root, unit_id, merged_layer_df, merged_seed_df, meta)
        catalog_rows.append(meta)

        merged_layer_count += int(len(merged_layer_df))
        merged_seed_count += int(len(merged_seed_df))
        if meta["DataMode"] == "real_only":
            real_only_unit_count += 1
        elif meta["DataMode"] == "virtual_only":
            virtual_only_unit_count += 1
        elif meta["DataMode"] == "real_virtual":
            real_virtual_unit_count += 1

    catalog_df = pd.DataFrame(catalog_rows)
    if catalog_df.empty:
        catalog_df = pd.DataFrame(columns=CATALOG_COLUMNS)
    else:
        catalog_df = ensure_columns(catalog_df, CATALOG_COLUMNS)
        catalog_df = catalog_df[CATALOG_COLUMNS].sort_values(["BlockX", "BlockY", "UnitID"]).reset_index(drop=True)

    output_root.mkdir(parents=True, exist_ok=True)
    write_csv_utf8(catalog_df, output_root / "unit_catalog.csv")

    summary = {
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "real_root": str(real_root),
        "virtual_fracture_root": str(virtual_fracture_root),
        "virtual_well_root": str(virtual_well_root),
        "output_root": str(output_root),
        "docx_path": str(docx_path),
        "real_unit_count": int(len(real_units)),
        "virtual_unit_count": int(len(virtual_units)),
        "merged_unit_count": int(len(catalog_df)),
        "real_only_unit_count": int(real_only_unit_count),
        "virtual_only_unit_count": int(virtual_only_unit_count),
        "real_virtual_unit_count": int(real_virtual_unit_count),
        "merged_layer_count": int(merged_layer_count),
        "merged_seed_count": int(merged_seed_count),
        "unit_catalog_csv": str(output_root / "unit_catalog.csv"),
        "run_summary_json": str(output_root / "run_summary.json"),
    }
    (output_root / "run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    append_merge_summary_to_docx(
        docx_path=docx_path,
        title=f"实验 unit_fracture_merge_{datetime.now().strftime('%Y%m%d_%H%M%S')}",
        config={
            "real_root": str(real_root),
            "virtual_fracture_root": str(virtual_fracture_root),
            "virtual_well_root": str(virtual_well_root),
            "output_root": str(output_root),
        },
        summary=summary,
    )

    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
